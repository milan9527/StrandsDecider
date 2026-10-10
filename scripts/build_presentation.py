"""Create the v2 deck from the original PPTX, supplement, and dated production evidence."""

import argparse
import copy
import hashlib
import io
import json
from datetime import UTC, datetime
from pathlib import Path

import presentation_layout as L
from PIL import Image
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
BLOG = "https://strandsagents.com/blog/introducing-strands-decider/"
MODEL = "https://huggingface.co/StrandsAgents/strands-decider-2B-hobson-v19"
REPO = "https://github.com/milan9527/StrandsDecider"
DEMO = "https://d3i7ogfx5ml7dc.cloudfront.net/#crypto"
CODE = "https://github.com/strands-labs/strands-decider/tree/63d24ae286e50105fba0bd1db15b0e243aea4650/src/strands_decider"
CASES = {
    "pending": "USDC 未到账",
    "underpaid": "USDT 少付",
    "refund": "稳定币退款",
    "settlement": "商户结算",
}
STAGES = {
    "intake": "客服受理",
    "verify": "交易核验",
    "approve": "风控审批",
    "execute": "业务处置",
    "close": "回复结案",
}
STATUS = {"complete": "业务完成", "handoff": "转人工", "waiting": "等待事件", "running": "尚未完成"}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--base",
        type=Path,
        default=ROOT / "artifacts/presentation/Strands-Decider-模型原理与支付业务演示.pptx",
    )
    parser.add_argument(
        "--supplement",
        type=Path,
        default=ROOT / "artifacts/presentation/Strands-Decider-模型原理与支付业务演示-扩展版.pptx",
    )
    parser.add_argument("--assets", type=Path, default=ROOT / "docs/presentations/v2")
    args = parser.parse_args()
    out, images = args.assets, args.assets / "screenshots"
    bench = json.loads((ROOT / "docs/results/benchmark-2026-10-10.json").read_text())
    loops = json.loads((out / "loop-observations.json").read_text())
    manifest = json.loads((out / "capture-manifest.json").read_text())
    # Crop only for readability; keep the full source screenshot and record exact bounds.
    crops = {
        "customer-reply-focus.png": (15, 310, 450, 540),
        "customer-events-focus.png": (15, 562, 450, 1020),
    }
    source_image = Image.open(images / "customer-conversation.png")
    shots = {row["file"]: row for row in manifest["screenshots"]}
    for filename, bounds in crops.items():
        source_image.crop(bounds).save(images / filename)
        shots[filename] = {
            "file": filename,
            "derived_from": "customer-conversation.png",
            "crop_pixels": list(bounds),
            "data_replaced": False,
            "description": "Crop only for slide readability; original screenshot preserved.",
        }
    for row in shots.values():
        path = images / row["file"]
        row["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        with Image.open(path) as image:
            row["pixels"] = list(image.size)
    manifest["screenshots"] = list(shots.values())
    (out / "capture-manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n"
    )
    prs, supplement = Presentation(args.base), Presentation(args.supplement)
    original = list(prs.slides)
    final, metadata = [], []
    totals = bench["totals"]
    paired = totals["paired_steps"]
    ratio = totals["gpt_over_decider"]
    common_sources = (
        f"\n\n来源：\n官方博客：{BLOG}\n模型卡：{MODEL}\n锁定源码：{CODE}\n项目：{REPO}"
        "\n固定输入实测：docs/results/benchmark-2026-10-10.json"
        "\n闭环实测：docs/presentations/v2/loop-observations.json"
        "\n页面截图：docs/presentations/v2/capture-manifest.json"
        "\nBERT 原始论文：https://aclanthology.org/N19-1423/"
        "\nBERT 分类实现：https://github.com/google-research/bert/blob/master/run_classifier.py"
        "\nNLI 与零样本分类：https://aclanthology.org/D19-1404/"
        "\n版本：前端 crypto10 / API 20261010-loop1；2026-10-10 UTC。"
    )

    def strip_footer(slide):
        for shape in list(slide.shapes):
            if shape.top >= Inches(7.05):
                shape._element.getparent().remove(shape._element)

    def keep(index, technical=False):
        source = supplement.slides[index - 1] if technical else original[index - 1]
        if technical:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = L.rgb(L.BG)
            remap = {}
            for rel in source.part.rels.values():
                if rel.reltype == RT.IMAGE:
                    _, rid = slide.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
                    remap[rel.rId] = rid
            for shape in source.shapes:
                element = copy.deepcopy(shape._element)
                for node in element.iter():
                    for key, value in list(node.attrib.items()):
                        if value in remap and key.startswith(
                            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
                        ):
                            node.set(key, remap[value])
                slide.shapes._spTree.insert_element_before(element, "p:extLst")
        else:
            slide = source
        strip_footer(slide)
        notes = source.notes_slide.notes_text_frame.text.split("\n\n资料来源：")[0]
        slide.notes_slide.notes_text_frame.text = notes + common_sources
        texts = [s.text for s in slide.shapes if s.has_text_frame and s.text.strip()]
        title = texts[1] if len(texts) > 1 else texts[0]
        final.append(slide)
        metadata.append(
            {
                "title": title,
                "source": "原版 / 扩展版沿用的原理与架构；已核对当前实现",
                "dark": False,
            }
        )
        return slide

    def new(
        title, section, subtitle="", notes="", source="当前项目实现 · crypto10 / 20261010-loop1"
    ):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = L.rgb(L.BG)
        L.text(slide, section, 0.58, 0.22, 11.9, 0.26, 10, L.TEAL, True)
        L.text(slide, title, 0.57, 0.67, 12.15, 0.59, 28, L.INK, True)
        if subtitle:
            L.text(slide, subtitle, 0.59, 1.34, 12.0, 0.46, 14, L.MUTED)
        slide.notes_slide.notes_text_frame.text = notes + common_sources
        final.append(slide)
        metadata.append({"title": title, "source": source, "dark": False})
        return slide

    def screenshot(title, filename, subtitle, notes, footer="", sidebar=None):
        slide = new(
            title,
            "03 / 当前演示 · 线上真实截图",
            subtitle,
            notes,
            "2026-10-10 UTC · 线上原始界面 · 无业务数据替换",
        )
        if sidebar:
            L.picture(slide, images / filename, 0.6, 1.98, 9.05, 4.70, frame=True)
            L.bullets(slide, sidebar, 9.95, 2.30, 2.75, size=17, gap=1.1)
        else:
            L.picture(slide, images / filename, 0.6, 1.97, 12.10, 4.72, frame=True)
        if footer:
            L.note(slide, footer, y=6.72, color=L.TEAL)
        return slide

    def table(slide, headers, rows, widths, y=2.05, row_height=0.66, size=15):
        x = 0.6
        for heading, width in zip(headers, widths, strict=True):
            L.rect(slide, x, y, width - 0.03, 0.48, L.NAVY)
            L.text(slide, heading, x + 0.10, y + 0.10, width - 0.2, 0.34, 13, L.WHITE, True)
            x += width
        for i, row in enumerate(rows):
            x = 0.6
            for value, width in zip(row, widths, strict=True):
                yy = y + 0.54 + i * row_height
                L.rect(
                    slide,
                    x,
                    yy,
                    width - 0.03,
                    row_height - 0.03,
                    L.WHITE if i % 2 == 0 else "ECF1F6",
                )
                L.text(
                    slide, str(value), x + 0.10, yy + 0.09, width - 0.20, row_height - 0.12, size
                )
                x += width

    slide = new(
        "Strands Decider\n模型原理与支付业务演示",
        "STRANDS DECIDER / v2 · 2026.10.10",
        notes="基于用户指定的25页原版创建新版，保留模型基础与AWS架构，合入BERT区别、演示参考、当前闭环与逐轮输入输出。原文件未覆盖。全部界面截图重新来自crypto10生产页面。",
    )
    # The two-line cover title gets a larger frame than regular slide headings.
    title_shape = slide.shapes[1]
    title_shape.height = Inches(1.35)
    L.text(
        slide,
        "从模型选择，到实际执行，再到可观察的反馈闭环",
        0.61,
        2.20,
        11.6,
        0.65,
        23,
        L.TEAL,
        True,
    )
    for x, heading, body, color in (
        (0.6, "模型", "Pointer Head / 动态候选\nDecider、GPT 与 BERT 的区别", L.TEAL),
        (4.73, "项目", "私有 S3 + CloudFront OAC\nECS + SageMaker + Bedrock", L.BLUE),
        (8.86, "新版演示", "逐轮输入输出 / 公共 GPU 窗口\n真实固定对比 + 自主闭环", L.ORANGE),
    ):
        L.card(slide, x, 3.55, 3.86, 2.03, heading, body, color, size=18)
    L.note(slide, "基于原版新建 · 原始 PPT 保留 · 前端 crypto10 · 真实截图与观测均标注来源")
    keep(2)
    keep(3)
    keep(4, technical=True)
    keep(4)
    keep(5)
    keep(8, technical=True)
    keep(9, technical=True)
    keep(6)

    screenshot(
        "当前演示入口：先选测试方式，再开始对比",
        "overview.png",
        "固定流程 / 自主闭环 / 仅比较决策，使用同一个「开始对比」入口。",
        "原版介绍了基础工作台和支付工作区。新版支付入口突出速度、逐轮输入输出、实际执行及客服对话；不是另起一个自主Agent产品。",
        sidebar=[
            "顶部选择场景与测试方式",
            "耗时卡片下方直接看输入 / 输出",
            "客户对话与模拟事件默认可见",
        ],
    )
    keep(8)
    screenshot(
        "GPU 管理：所有模块共用一个弹窗",
        "gpu-dialog.png",
        "本次采集时 GPU 已就绪；只复用现有实例，没有发起启停。",
        "所有模块顶部都可打开GPU管理；真实运行前读取状态。未启动、启动中、停止中或状态未知时提示并阻止请求。管理员启动后自动刷新进度，就绪后返回原模块重新运行。非管理员查看状态并联系管理员。",
        sidebar=["真实推理前检查状态", "唯一一组启动 / 停止按钮", "闲置 30 分钟删除 Endpoint"],
    )
    keep(10)

    slide = new(
        "两种对比方式，回答不同的问题",
        "03 / 测试设计",
        "固定流程比较同一输入；自主闭环比较各自实际完成业务的过程。",
        "固定模式先冻结输入并收集五步建议，再执行两套方案。闭环则每轮判断、执行、读取实际结果，再提交带反馈的新请求。两侧分叉后输入可能不同，不能继续将全部轮次当作同输入准确率或延迟比较。",
    )
    table(
        slide,
        ["维度", "固定输入流程", "自主闭环"],
        [
            ["模型输入", "每步使用冻结业务证据", "使用本侧最新状态 + 执行反馈"],
            ["执行时机", "收齐建议后按序执行", "每次判断后立即执行并读取结果"],
            ["外部事件", "推进证据后重新开始", "等待事件后恢复本次工作流"],
            ["对比指标", "成功且同 input_hash 的配对决策", "完成状态、调用数、总耗时与重判"],
            ["业务结果", "处理五步也可能没有资金动作", "可完成、等待、失败或转人工"],
        ],
        [1.8, 5.15, 5.15],
    )
    L.note(slide, "闭环是一个对比测试选项；候选、执行保护与调用预算在两侧相同。")
    keep(16, technical=True)

    differences = []
    for case in bench["cases"]:
        for step in case["steps"]:
            result = step["decider"]
            if result["source"] == "live" and not result["matches_reference"]:
                applied = result["execution"] or {}
                differences.append(
                    [
                        CASES[case["summary"]["case"]] + "\n" + STAGES[step["stage"]],
                        result["reference_label"],
                        result["label"],
                        applied.get("label", "未执行")
                        + ("\n保护调整" if applied.get("safety_adjusted") else "\n保留原建议"),
                    ]
                )
    slide = new(
        "Decider 执行差异：原始选择与实际动作分别保留",
        "03 / 本次真实观测",
        f"本次固定流程中有 {len(differences)} 次 Decider 参考差异；保护调整不计为模型判断正确。",
        "下表直接从本次导出读取。reference()由开发者独立计算，不是模型自己生成的规则。保守选择可以保留并继续，但未满足资金条件的动作仍会被调整。",
    )
    table(
        slide,
        ["场景 / 步骤", "演示参考", "Decider 建议", "实际采用"],
        differences,
        [3.0, 2.9, 2.9, 3.3],
        row_height=min(0.67, 3.9 / max(1, len(differences))),
        size=14,
    )
    L.note(slide, "能观察到条件遵循偏差；仅凭 choice / 概率，不能断言某次偏差的内部根因。")
    slide = new(
        "为什么会不一致？模型评分不是执行规则代码",
        "03 / 解释边界",
        "Decider 输出有限候选中的选择；符合输出结构，不等于正确遵守全部业务条件。",
        "候选语义、字段干扰、条件组合、训练分布与业务校准都可能影响结果，但本项目没有进行归因消融。不要编造模型内部推理过程。BERT也会误分类，GPT也可能选择错误或请求失败。",
    )
    for x, heading, body, color in (
        (
            0.6,
            "可核对的事实",
            "记录模型实际 choice / 分布\n记录原始输入与参考差异\n执行器另行检查实际状态",
            L.TEAL,
        ),
        (
            4.73,
            "不能据此推断",
            "一次差异的内部思维过程\n金融行业的通用准确率\nDecider 比 BERT 更快或更准",
            L.ORANGE,
        ),
        (
            8.86,
            "可以改进的环节",
            "设计清楚的候选与条件\n收集失败和重判样本\n按业务数据评估与校准",
            L.BLUE,
        ),
    ):
        L.card(slide, x, 2.15, 3.86, 3.25, heading, body, color, size=19)
    L.note(slide, "金额计算、地址校验和资金权限等确定性条件，仍由业务代码执行。")

    screenshot(
        "每轮输入 / 输出：默认展开，直接并排对照",
        "fixed-turn-1.png",
        "当前页面实际截取：客户诉求、交易证据、两侧输出、实际执行与时间。",
        "输入保存的是当时业务请求的快照，不会被后来的确认数、金额或消息修改覆盖。模型建议、实际执行和模板客服回复分别标记；原始业务请求与响应可展开。",
    )
    slide = new(
        "固定输入：同一步的选择与速度一起看",
        "03 / 真实模型截图",
        "USDC 未到账：确认数 3/12；表中显示原始建议、实际动作与各步耗时。",
        "截图来自本次生产页面真实调用，没有用历史结果替换UI。模型每步后端耗时与整轮浏览器时间口径不同。",
        source="2026-10-10 UTC · pending-fixed · 实际生产截图",
    )
    L.picture(slide, images / "fixed-summary.png", 0.6, 1.94, 12.1, 1.34, frame=True)
    L.picture(slide, images / "fixed-table.png", 0.6, 3.43, 12.1, 3.25, frame=True)

    slide = new(
        "已处理 5/5 步，不等于业务成功",
        "03 / 业务完成判定",
        "等待、挂起和继续跟进，也属于已处理步骤。",
        "业务完成必须同时确认目标订单状态、对应模拟账本、工单结案，并且无待复核状态。初始USDC确认数不足，因此两侧可处理五步，但没有资金记录。",
    )
    L.picture(slide, images / "fixed-final.png", 0.6, 1.98, 12.1, 2.45, frame=True)
    for x, heading, body in [
        (0.6, "订单达到目标状态", "已收款 / 已退款 / 已对账"),
        (4.73, "存在对应模拟账本", "金额与资金动作可核对"),
        (8.86, "客服工单已结案", "不能仅依据模型建议判断"),
    ]:
        L.card(slide, x, 4.80, 3.86, 1.52, heading, body, size=17)
    screenshot(
        "退款处置：建议与真正执行的结果分开显示",
        "refund-execute.png",
        "地址未验证、审批未完成时，不能因为模型选择退款就认定资金已执行。",
        "保留本次真实Decider与GPT原始建议。执行器检查实际条件，必要时采用hold；继续给出与真实状态一致的模板客服回复。",
    )
    slide = new(
        "客户对话与模拟事件，放在主区域",
        "03 / 当前演示 · 线上真实截图",
        "客服回复由模板按实际订单 / 工单状态生成，不是 GPT 长文本质量测试。",
        "左右均裁自同一张最新生产截图，只改变取景，没有替换业务数据。客服回复对应实际的3/12确认状态。点击模拟按钮后，固定模式重新开始；闭环等待时保留历史并续跑。",
        source="2026-10-10 UTC · customer-conversation.png 局部取景",
    )
    L.label(slide, "模板回复：如实说明待确认状态", 0.6, 1.97, 5.86)
    L.label(slide, "模拟事件与客户输入", 6.85, 1.97, 5.86, L.PALE_BLUE, L.BLUE)
    L.picture(slide, images / "customer-reply-focus.png", 0.6, 2.50, 5.86, 2.60, frame=True)
    L.text(
        slide,
        "未满足确认数 → 持续跟进\n更新证据后重新对比\n闭环等待时则继续当前运行",
        0.84,
        5.37,
        5.32,
        1.06,
        19,
        L.TEAL,
    )
    L.picture(slide, images / "customer-events-focus.png", 6.85, 2.50, 5.86, 4.0, frame=True)

    slide = new(
        "自主闭环：判断 → 执行 → 读取结果 → 再判断",
        "03 / 反馈编排",
        "每次回到模型前，把实际状态和最近反馈写入下一轮输入。",
        "POST /api/crypto/loop/tick每次处理一个受限步骤。不能通过反复请求模型来补齐外部证据。每侧独立执行，模拟器事件不是模型执行真实审批或资金交易。",
    )
    for i, (heading, body) in enumerate(
        [
            ("当前输入", "诉求、证据、执行状态\n上轮真实反馈"),
            ("模型判断", "从给定候选中选择\n保存原始结果与时间"),
            ("应用执行", "校验条件、执行或挂起\n记录实际动作"),
            ("结果反馈", "订单 / 工单 / 账本\n确认是否真的推进"),
        ]
    ):
        x = 0.6 + i * 3.08
        L.step(slide, i + 1, heading, body, x, 2.17, 2.86)
        if i < 3:
            L.line(slide, [(x + 2.88, 2.85), (x + 3.05, 2.85)], arrow=True)
    L.line(slide, [(11.3, 3.62), (11.3, 4.03), (1.9, 4.03), (1.9, 3.62)], color=L.TEAL, arrow=True)
    for x, heading, body, color in (
        (0.6, "条件未齐 → 等待", "等待确认、补款、地址验证等\n事件更新后再核验", L.ORANGE),
        (4.73, "无效动作 → 重判", "把未执行原因传回模型\n同一步两次无进展转人工", L.BLUE),
        (8.86, "业务完成 → 结束", "订单 + 账本 + 结案均确认\n已记账动作不再次提交", L.TEAL),
    ):
        L.card(slide, x, 4.45, 3.86, 1.83, heading, body, color, size=17)
    L.note(slide, "预算：最多 20 次决策 / 32 次 tick；页面连续运行约 5 分钟后暂停。")
    screenshot(
        "闭环输入：带着本侧执行反馈进入下一轮",
        "pending-loop-feedback.png",
        "同样的初始证据，两侧执行后可能进入不同状态。",
        "截图显示闭环请求中的当前订单、工单状态与上一轮反馈。轮次编号仅表示各侧处理顺序，不代表后续输入相同。",
    )

    for case in ("pending", "refund"):
        observation = loops[case]
        a, b = (observation["providers"][p] for p in ("decider", "bedrock"))
        slide = new(
            f"{CASES[case]}闭环：如实展示完成或转人工",
            "03 / 本次真实闭环",
            f"Decider：{STATUS[a['status']]} · {len(a['decisions'])} 次决策；GPT：{STATUS[b['status']]} · {len(b['decisions'])} 次决策。",
            "本页状态与结果直接读取本次生产导出。没有替换模型选择或强制业务完成。闭环包含模拟外部事件，但不是现实中的人工审批或真实交易。",
            source="2026-10-10 UTC · " + case + "-loop · 实测原始结果",
        )
        L.picture(slide, images / f"{case}-loop-final.png", 0.6, 2.05, 12.1, 2.70, frame=True)
        for x, name, result, color in [
            (0.6, "Strands Decider", a, L.TEAL),
            (6.85, "GPT-6.1 Sol", b, L.BLUE),
        ]:
            total = sum(
                d["timings"]["backend_total_ms"]
                for d in result["decisions"]
                if d["source"] == "live"
            )
            L.card(
                slide,
                x,
                5.05,
                5.86,
                1.30,
                name,
                f"决策累计 {total / 1000:.2f} s · 模拟事件 {len(result['events'])} 次",
                color,
                size=17,
            )
        L.note(slide, "两侧后续输入和调用次数可能不同，此处不能解释为固定同输入的推理速度排名。")

    slide = new(
        "闭环的停止原因，也属于演示结果",
        "03 / 真实执行反馈",
        "继续、等待、完成、调用失败和转人工都应可见，不能只展示成功路径。",
        "展示各侧最后一条控制反馈，用于解释为什么结束或暂停。连续无进展会触发转人工；这不是系统把参考答案补给Decider。",
    )
    for i, case in enumerate(("pending", "refund")):
        for j, provider in enumerate(("decider", "bedrock")):
            result = loops[case]["providers"][provider]
            L.card(
                slide,
                0.6 + j * 6.25,
                2.02 + i * 2.10,
                5.86,
                1.85,
                CASES[case] + " · " + ("Decider" if provider == "decider" else "GPT"),
                result["error"] or result["reason"],
                L.TEAL if j == 0 else L.BLUE,
                size=16,
            )
    L.note(slide, "前端可回看每次输入、输出、执行反馈和失败记录；继续运行追加新记录。")

    slide = new(
        f"本次实测：{paired} 组成功配对，累计耗时比 {ratio:.2f}×",
        "04 / 最新真实数据",
        "四个初始场景，每侧尝试 20 次决策；只汇总成功且输入相同的配对。",
        "完整数据包含一次Bedrock ReadTimeoutError。失败请求不进入速度合计，超时对用户等待的影响不能被这个成功请求统计替代。商户结算GPT侧因缺少完整方案未执行沙盒。",
        source="docs/results/benchmark-2026-10-10.json · 单次观测，不是稳定性能保证",
    )
    rows = []
    for case in bench["cases"]:
        r = case["summary"]
        rows.append(
            [
                CASES[r["case"]],
                str(r["paired_steps"]),
                f"{r['decider_total_ms']:.1f} ms",
                f"{r['bedrock_total_ms']:.1f} ms",
                f"{r['decider_reference_matches']}/{r['paired_steps']}",
                f"{r['bedrock_reference_matches']}/{r['paired_steps']}",
            ]
        )
    rows.append(
        [
            "合计",
            str(paired),
            f"{totals['decider']['backend_ms']:.1f} ms",
            f"{totals['bedrock']['backend_ms']:.1f} ms",
            f"{totals['decider']['reference_matches']}/{paired}",
            f"{totals['bedrock']['reference_matches']}/{paired}",
        ]
    )
    table(
        slide,
        ["初始案例", "配对", "Decider 合计", "GPT 合计", "D 参考一致", "G 参考一致"],
        rows,
        [2.15, 0.9, 2.65, 2.65, 1.9, 1.85],
        row_height=0.65,
        size=14,
    )
    L.note(
        slide, "Bedrock 有 1 次读取超时，单独保留错误记录；未使用回放补结果。", fill=L.PALE_ORANGE
    )

    slide = new(
        "普通大模型也可能失败：结果与可用性都要记录",
        "04 / 失败路径",
        "本次商户结算 / 业务处置：Bedrock ReadTimeoutError。",
        "不要因为GPT在历史固定案例中符合参考，就宣称它在每次调用中都能完整跑完。本次5步中返回4个GPT决策，1步失败；没有完整方案，因此该侧模拟执行没有启动。后续重试应保留此前错误。",
        source="2026-10-10 真实采集 · settlement-fixed 导出",
    )
    L.card(
        slide,
        0.6,
        2.1,
        5.85,
        3.08,
        "观察到的事实",
        "Bedrock 返回 4/5 个决策\n业务处置请求读取超时\n该侧没有执行完整沙盒方案\n错误与本轮输入保留",
        L.ORANGE,
        size=21,
    )
    L.card(
        slide,
        6.85,
        2.1,
        5.85,
        3.08,
        "比较时如何处理",
        "速度仅统计成功同输入配对\n失败数量单独披露\n不能把成功请求延迟当总体体验\n不能用规则结果替代失败输出",
        L.BLUE,
        size=21,
    )
    L.note(slide, "这是单次实际故障记录，不足以推断任何模型的长期可用性。")
    keep(22)

    slide = new(
        "原始输入输出可追溯，生产状态还需持久化",
        "05 / 当前实现与边界",
        "当前演示已实现受限闭环；金融生产系统还需要服务端工作流与幂等。",
        "对话记录保存最近80轮，客户会话保存最近80条，刷新清除。每次输入做快照，不会随证据更新被改写。后台目前仍接收浏览器传回的沙盒状态，不能用它作为生产金融系统的状态或资金权限依据。",
    )
    table(
        slide,
        ["当前已实现", "生产化还需补齐"],
        [
            ["判断 → 执行 → 反馈 → 重判 / 等待", "服务端持久化状态、可靠队列与恢复调度"],
            ["当前运行内避免重复资金步骤", "数据库级唯一约束、状态版本与条件写"],
            ["逐轮输入输出、错误历史、导出", "持久审计、权限隔离与敏感数据治理"],
            ["明确标记的模拟确认 / 补款 / 审批", "可验证的链上事件、支付账务和授权审批"],
            ["预算限制、连续无进展转人工", "持续评测、费用预算与人工接管流程"],
        ],
        [6.05, 6.05],
        size=16,
    )
    L.note(slide, "自动按置信度切换到 GPT 属于后续方案；当前不会静默替换 Decider 输出。")
    slide = new(
        "现场演示：先看同输入，再看执行反馈",
        "05 / 操作路线",
        "演示入口：支付流程对比 → 测试方式 → 开始对比。",
        "可先用规则回放讲解交互，明确回放不产生模型速度。真实模式在GPU已就绪时运行。先固定初始USDC显示资金未落实，再选择闭环展示事件推进、重判和各自最终状态。最后展示历史输入不会被改写并导出记录。",
    )
    for i, (stamp, title, body) in enumerate(
        [
            ("00:00", "检查 GPU", "顶部公共入口；未就绪会弹窗，等就绪后返回运行。"),
            ("00:45", "固定流程对比", "看每轮输入与两模型输出，再看实际动作和业务是否完成。"),
            ("02:00", "切换自主闭环", "开启自动模拟证据；逐轮观察反馈、事件、重判与耗时。"),
            ("03:30", "解释分歧与停止", "完成、等待或转人工均保留；不要只用 5/5 判断成功。"),
            ("04:30", "回看历史并导出", "切换对比记录，展开原始请求与返回，保存 JSON。"),
        ]
    ):
        y = 2.0 + i * 0.83
        L.label(slide, stamp, 0.65, y + 0.07, 1.0)
        L.text(slide, title, 1.95, y, 10.0, 0.36, 20, L.INK, True)
        L.text(slide, body, 1.95, y + 0.4, 10.4, 0.36, 14, L.MUTED)
    L.note(slide, DEMO, color=L.BLUE)
    slide = new(
        "结论：快的决策，需要可观察、可约束的执行",
        "05 / 结论",
        "Decider 适合高频有限候选判断；能否完成业务，要看实际状态和反馈。",
        "本次数据只能支持这些案例上的观察。与BERT的比较是结构和任务适配方式，不是性能实测。普通大模型也存在超时或错误，闭环亦不保证成功。确定性规则由代码执行，模型处理语义选择；生产金融操作仍需外部授权与持久账务。",
    )
    for x, heading, body, color in (
        (
            0.6,
            "模型层",
            "候选与条件来自应用\nDecider 直接评分，GPT 生成 JSON\nBERT 固定分类头有不同适配方式",
            L.TEAL,
        ),
        (
            4.73,
            "执行层",
            "原始选择与实际动作分开\n业务完成用订单、账本、结案判定\n反馈闭环可改进，也可能转人工",
            L.BLUE,
        ),
        (
            8.86,
            "评测层",
            "同输入统计与闭环统计分开\n成功延迟与调用失败同时披露\n用真实业务分布继续验证",
            L.ORANGE,
        ),
    ):
        L.card(slide, x, 2.13, 3.86, 3.45, heading, body, color, size=18)
    L.note(
        slide, "推荐组合：代码做确定性校验，模型做语义选择，高影响操作由明确权限与人工流程控制。"
    )
    slide = new(
        "资料来源、截图与本次数据",
        "附录 / 复核",
        "原版文件不覆盖；所有新版截图来自当前生产界面，数据来源与采集时间保留。",
        f"原版SHA256：{hashlib.sha256(args.base.read_bytes()).hexdigest()}。补充判断机制与BERT内容来自此前扩展版，并再次核对锁定源码。最新统计只读取2026-10-10新采集结果，未混入旧PPT的7.49倍或README旧批次8.87倍。",
    )
    for i, (title, desc, url) in enumerate(
        [
            ("Strands 官方博客 / v19 模型卡", "模型目的、Pointer Head、问题类型与限制", BLOG),
            ("Decider 锁定源码", "判断计算、confidence 与实际请求协议", CODE),
            ("BERT 论文与分类器", "BERT 对比未进行性能实测", "https://aclanthology.org/N19-1423/"),
            ("项目与在线演示", "代码、架构、闭环实现及用户操作", REPO),
            (
                "本次证据",
                "benchmark-2026-10-10.json / loop-observations.json / capture-manifest.json",
                REPO + "/tree/main/docs/presentations/v2",
            ),
        ]
    ):
        y = 1.98 + i * 0.94
        L.text(slide, title, 0.64, y, 11.9, 0.34, 17, L.TEAL, True)
        L.text(slide, desc, 0.64, y + 0.36, 11.9, 0.29, 12.5)
        shape = L.text(slide, url, 0.64, y + 0.66, 11.9, 0.23, 10.2, L.BLUE)
        shape.text_frame.paragraphs[0].runs[0].hyperlink.address = url

    # Preserve only selected original slides and newly built slides, in narrative order.
    order = {s.part: s for s in final}
    ids = {prs.part.related_slide(sid.rId).part: sid for sid in list(prs.slides._sldIdLst)}
    for sid in list(prs.slides._sldIdLst):
        slide = prs.part.related_slide(sid.rId)
        prs.slides._sldIdLst.remove(sid)
        if slide.part not in order:
            prs.part.drop_rel(sid.rId)
    for slide in final:
        prs.slides._sldIdLst.append(ids[slide.part])
    for i, (slide, meta) in enumerate(zip(final, metadata, strict=True), 1):
        L.line(slide, [(0.6, 7.08), (12.71, 7.08)], L.LINE, 0.7)
        L.text(slide, meta["source"], 0.62, 7.17, 11.15, 0.22, 8.1, L.MUTED)
        L.text(slide, f"{i:02} / {len(final):02}", 11.98, 7.16, 0.72, 0.22, 9, L.MUTED)
        for effect in slide._element.xpath(".//a:effectRef"):
            effect.set("idx", "0")
        for shape in slide.shapes:
            assert shape.left >= -10 and shape.top >= -10, (i, shape.name)
            assert shape.left + shape.width <= prs.slide_width + 10, (i, shape.name, "width")
            assert shape.top + shape.height <= prs.slide_height + 10, (i, shape.name, "height")
    prs.core_properties.title = "Strands Decider 模型原理与支付业务演示 · v2 · 2026-10-10"
    prs.core_properties.subject = "crypto10；真实逐轮输入输出、执行反馈闭环、模型区别与业务完成"
    prs.core_properties.comments = "Based on the original 25-slide presentation. Fresh production screenshots; failed requests disclosed. BERT comparison is conceptual."
    target = out / "Strands-Decider-模型原理与支付业务演示-v2-20261010.pptx"
    prs.save(target)
    (out / "讲稿与来源.md").write_text(
        "# Strands Decider v2 演示讲稿\n\n"
        + "\n\n".join(
            f"## {i:02}. {m['title']}\n\n{s.notes_slide.notes_text_frame.text}"
            for i, (s, m) in enumerate(zip(final, metadata, strict=True), 1)
        )
        + "\n"
    )
    (out / "build-manifest.json").write_text(
        json.dumps(
            {
                "generated_at": datetime.now(UTC).isoformat(),
                "slides": metadata,
                "base_file": args.base.name,
                "base_sha256": hashlib.sha256(args.base.read_bytes()).hexdigest(),
                "supplement_file": args.supplement.name,
                "supplement_sha256": hashlib.sha256(args.supplement.read_bytes()).hexdigest(),
                "pptx_sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
                "capture_start": manifest["captured_at"],
                "successful_pairs": paired,
                "failed_requests": totals["failed_requests"],
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n"
    )
    print(f"Created {target.relative_to(ROOT)} ({len(final)} slides)")


if __name__ == "__main__":
    main()
