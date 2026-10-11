"""Update the latest 25-slide deck for EC2, preserving the historical measurements."""

import copy
import hashlib
import io
import json
import shutil
from datetime import UTC, datetime
from pathlib import Path

import cairosvg
import presentation_layout as L
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/presentations/v4"
OUT = ROOT / "docs/presentations/v5"
BASE = SOURCE / "Strands-Decider-模型原理与支付业务演示-v4-决策模型与支付对比-25页-20261010.pptx"
NAME = "Strands-Decider-模型原理与支付业务演示-v5-EC2部署版-25页-20261011"
DEMO = "https://d3i7ogfx5ml7dc.cloudfront.net/#crypto"
REPO = "https://github.com/milan9527/StrandsDecider"
HISTORY = "历史实测：2026-10-10 · 迁移前的托管部署；当前 EC2 架构见第 9 页"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def replace_text(shape, old, new, size=None):
    """Preserve run formatting; these deck labels each have a single run per line."""
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.text = run.text.replace(old, new)
            if size is not None:
                run.font.size = Pt(size)
    shape.name = "TEXT: " + shape.text.replace("\n", " ")[:65]


def architecture(slide):
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    for rel in list(slide.part.rels.values()):
        if rel.reltype == RT.IMAGE:
            slide.part.drop_rel(rel.rId)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = L.rgb(L.BG)
    L.text(slide, "STRANDS DECIDER   /   当前项目架构", 0.6, 0.27, 12, 0.25, 10, L.TEAL, True)
    L.text(slide, "一套业务后端，对接 EC2 与 Bedrock", 0.6, 0.79, 12.1, 0.69, 30, L.INK, True)
    L.text(
        slide,
        "2026-10-11 当前部署 · EC2 g6.xlarge / L4 · On-Demand 常驻运行",
        0.62,
        1.56,
        12,
        0.53,
        15,
        L.MUTED,
    )

    def aws_icon(name, x, y, size=0.42):
        path = ROOT / "docs/diagrams/aws-icons" / f"{name}.svg"
        bitmap = cairosvg.svg2png(url=str(path), output_width=256, output_height=256)
        shape = slide.shapes.add_picture(
            io.BytesIO(bitmap), Inches(x), Inches(y), Inches(size), Inches(size)
        )
        shape.name = "ICON: AWS " + name

    def service(name, title, body, x, y, w, h=1.14, fill=L.WHITE, size=16):
        L.rect(slide, x, y, w, h, fill, L.LINE, True)
        aws_icon(name, x + 0.14, y + 0.15)
        L.text(slide, title, x + 0.68, y + 0.17, w - 0.81, 0.32, size, L.INK, True)
        L.text(slide, body, x + 0.14, y + 0.68, w - 0.28, h - 0.73, 12, L.MUTED)

    # CloudFront remains outside the regional boundary. EC2 is inside the private VPC.
    # Include regional S3 in the lower-left branch, while keeping global
    # CloudFront above that boundary and outside us-east-1.
    boundary = slide.shapes.build_freeform(Inches(3.25), Inches(2.18))
    boundary.add_line_segments(
        [
            (Inches(x), Inches(y))
            for x, y in [(12.70, 2.18), (12.70, 6.37), (0.50, 6.37), (0.50, 3.98), (3.25, 3.98)]
        ],
        close=True,
    )
    region = boundary.convert_to_shape()
    region.name = "AWS Region boundary: includes S3, excludes CloudFront"
    region.fill.solid()
    region.fill.fore_color.rgb = L.rgb(L.BG)
    region.line.color.rgb = L.rgb("00A4A6")
    region.line.width = Pt(0.8)
    L.text(slide, "AWS Region · us-east-1", 3.42, 2.32, 8.9, 0.3, 12, L.TEAL, True)
    L.rect(slide, 3.42, 2.86, 6.62, 3.23, "F0F6F6", "8C4FFF")
    L.text(slide, "VPC · 私有子网 / 无公网 IP", 3.57, 2.98, 6.1, 0.27, 11, L.MUTED, True)
    service("cloudfront", "CloudFront", "HTTPS 页面 / API 入口", 0.65, 2.52, 2.34)
    service("s3", "Amazon S3", "私有源站 · OAC", 0.65, 4.15, 2.34)
    L.text(slide, "Block Public Access 全部开启", 0.67, 5.45, 2.4, 0.55, 11, L.MUTED)
    service("alb", "内网 ALB", "CloudFront VPC Origin", 3.62, 3.4, 2.53, 1.02)
    service("ecs", "ECS Fargate", "FastAPI · 业务与模型调用", 3.62, 4.84, 2.53, 1.07)
    service(
        "ec2",
        "Amazon EC2",
        "g6.xlarge · NVIDIA L4 / 24 GB\nECS 模型服务 · Decider v19\n私网 HTTP :8099 · 持续运行",
        6.7,
        3.4,
        3.08,
        2.04,
        L.PALE_TEAL,
        17,
    )
    aws_icon("ebs", 6.75, 5.64, 0.24)
    L.text(slide, "80 GB 加密 EBS · 模型缓存", 7.06, 5.66, 2.79, 0.25, 10.5, L.MUTED)
    service("bedrock", "Bedrock", "GPT-6.1 Sol", 10.32, 3.26, 2.19, 1.12, L.PALE_BLUE)
    service("dynamodb", "DynamoDB", "用户 / 会话 / 推理租约", 10.32, 4.95, 2.19, 1.12, size=14)

    L.line(slide, [(1.82, 3.70), (1.82, 4.09)], L.MUTED, arrow=True)
    L.text(slide, "OAC", 1.96, 3.80, 0.7, 0.25, 10, L.MUTED)
    L.line(slide, [(3.01, 3.1), (3.22, 3.1), (3.22, 3.91), (3.56, 3.91)], L.MUTED, arrow=True)
    L.line(slide, [(4.87, 4.45), (4.87, 4.79)], L.MUTED, arrow=True)
    L.line(slide, [(6.18, 5.15), (6.42, 5.15), (6.42, 4.5), (6.65, 4.5)], L.TEAL, arrow=True)
    L.line(
        slide,
        [(4.87, 5.94), (4.87, 6.23), (10.16, 6.23), (10.16, 3.83), (10.27, 3.83)],
        L.BLUE,
        arrow=True,
    )
    L.line(slide, [(10.16, 5.52), (10.27, 5.52)], L.MUTED, arrow=True)
    L.text(
        slide,
        "API 直接调用常驻模型；后端保留并发与执行校验。历史测速数据保留原采集环境。",
        0.62,
        6.65,
        12.05,
        0.31,
        11.5,
        L.MUTED,
    )
    L.line(slide, [(0.62, 7.04), (12.68, 7.04)], L.LINE, 0.6)
    L.text(slide, "当前部署：EC2 On-Demand · 2026-10-11", 0.64, 7.16, 9.9, 0.21, 8.5, L.MUTED)
    L.text(slide, "09 / 25", 11.76, 7.14, 0.92, 0.23, 10, L.MUTED, align=PP_ALIGN.RIGHT)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "assets").mkdir(exist_ok=True)
    base_hash = digest(BASE)
    prs = Presentation(BASE)
    metadata = copy.deepcopy(json.loads((SOURCE / "build-manifest.json").read_text())["slides"])
    metadata[0]["section"] = "决策模型与支付业务 / v5 · EC2 部署版 · 25 页"
    changed = []
    for number, slide in enumerate(prs.slides, 1):
        for shape in list(slide.shapes):
            if not shape.has_text_frame:
                continue
            replacements = {
                "/ v4 · 25 页": "/ v5 · EC2 部署版 · 25 页",
                "2026.10.10": "2026.10.11",
                "Decider / SageMaker": "Decider / EC2",
            }
            for old, new in replacements.items():
                if old in shape.text:
                    replace_text(shape, old, new)
                    changed.append(number)
            if (
                number == 11
                and shape.text == "返回选择、分布与置信信息\n记录后端决策耗时，独立执行"
            ):
                # Use new editable text while retaining the original card's geometry.
                x, y, w, h = [
                    v / 914400 for v in (shape.left, shape.top, shape.width, shape.height)
                ]
                shape._element.getparent().remove(shape._element)
                L.text(
                    slide,
                    "私网 HTTP → EC2 模型服务\n返回选择、分布与置信信息\n记录耗时，两侧独立执行",
                    x,
                    y,
                    w,
                    h,
                    17,
                    L.INK,
                )
            if number in [*range(12, 23), 24] and shape.text == "决策模型 → 支付对比 → 应用结论":
                replace_text(shape, shape.text, HISTORY)
                changed.append(number)
            if number == 24 and shape.text == "本次实测":
                replace_text(shape, "本次实测", "历史实测")
            if number == 23 and shape.text == "以本次观测为起点，继续按真实业务分布评测。":
                replace_text(shape, "本次观测", "历史观测")
                changed.append(number)
        notes = slide.notes_slide.notes_text_frame.text
        notes = notes.replace(
            "Decider：官方 v19 / SageMaker ml.g6.xlarge / L4；",
            "历史数据的采集环境：Decider 官方 v19 / SageMaker ml.g6.xlarge / L4；",
        ).replace(
            "当前最新基准19组成功配对，另有1次GPT超时。",
            "本稿保留2026-10-10历史基准：19组成功配对，另有1次GPT超时；不标为EC2实测。",
        )
        notes += (
            "\n\nv5 当前部署更新（2026-10-11）：\n"
            "前端仍为私有 S3 + CloudFront OAC；API 仍为 ECS Fargate，经内网 ALB 提供服务。\n"
            "Decider 改为 ECS on EC2，us-east-1d / g6.xlarge / NVIDIA L4 24GB 标称显存。"
            "On-Demand 常驻，不因闲置自动停止，模型权重和缓存位于80GB加密gp3磁盘。"
            "API 经私网 HTTP :8099 调用模型，正常推理不等待前端GPU预检查；"
            "后端继续校验readiness、并发与推理租约。"
            "SageMaker Endpoint 已释放、旧控制器禁用，Model/EndpointConfig保留作回退。\n"
            "EC2计算费参考0.8048美元/小时，730小时约587.50美元；磁盘与其他服务另计。"
            "旧可用区停止实例与其磁盘另行保留，不参与当前推理。\n"
            "本版只更新文稿，没有重新采集模型性能；原速度图、失败样本和截图"
            "保留2026-10-10采集环境，不作为当前EC2性能保证。\n"
            "当前部署与验证依据：deployment/DEPLOYED.md、deployment/README.md、"
            "docs/DEVELOPMENT.md。"
        )
        if number == 9:
            architecture(slide)
            notes = (
                "当前架构已重绘为EC2私网部署：CloudFront访问私有S3，API经VPC Origin、"
                "内网ALB到ECS Fargate；API经私网HTTP直接调用VPC内ECS on EC2模型服务，"
                "另调用Bedrock和DynamoDB。图中使用AWS官方图标，不单列Runtime或数据库"
                "Endpoint节点。GPU常驻，保留磁盘缓存，用户无需逐次启动。\n\n来源与口径：\n"
                + notes.split("来源与口径：", 1)[1]
            )
            metadata[number - 1]["title"] = "一套业务后端，对接 EC2 与 Bedrock"
            metadata[number - 1]["section"] = "当前项目架构"
            changed.append(number)
        slide.notes_slide.notes_text_frame.text = notes
        for effect in slide._element.xpath(".//a:effectRef"):
            effect.set("idx", "0")
        visible = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
        assert "SageMaker" not in visible and "ml.g6" not in visible, (number, visible)
        assert "启动 GPU" not in visible and "停止 GPU" not in visible, number
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0, (number, shape.name)
            assert shape.left + shape.width <= prs.slide_width + 10, (number, shape.name)
            assert shape.top + shape.height <= prs.slide_height + 10, (number, shape.name)
    assert len(prs.slides) == 25
    prs.core_properties.title = "Strands Decider & Jev · EC2 部署版 · 25页"
    prs.core_properties.subject = "EC2 On-Demand架构、Decider与GPT决策对比、历史实测"
    prs.core_properties.comments = (
        "v5 updates current hosting to ECS on EC2 g6.xlarge. "
        "Measurements and screenshots remain dated SageMaker observations from 2026-10-10."
    )
    target = OUT / f"{NAME}.pptx"
    prs.save(target)
    for path in (SOURCE / "assets").glob("*-qr.png"):
        shutil.copyfile(path, OUT / "assets" / path.name)
    shutil.copyfile(SOURCE / "source-manifest.json", OUT / "source-manifest.json")
    (OUT / "讲稿与来源.md").write_text(
        "# 决策模型与支付业务演示 v5 · EC2 部署版 · 25 页讲稿\n\n"
        + "\n\n".join(
            f"## {i:02}. {meta['title']}\n\n{slide.notes_slide.notes_text_frame.text}"
            for i, (slide, meta) in enumerate(zip(prs.slides, metadata, strict=True), 1)
        )
        + "\n"
    )
    (OUT / "build-manifest.json").write_text(
        json.dumps(
            {
                "generated_at": datetime.now(UTC).isoformat(),
                "slides": metadata,
                "base_file": BASE.relative_to(ROOT).as_posix(),
                "base_sha256": base_hash,
                "pptx_sha256": digest(target),
                "changed_visible_slides": sorted(set(changed)),
                "notes_updated": 25,
                "current_deployment_date": "2026-10-11",
                "current_deployment": "ECS on EC2 g6.xlarge / L4 / On-Demand / always on / private HTTP",
                "historical_measurement_date": "2026-10-10",
                "historical_measurement_hosting": "SageMaker ml.g6.xlarge / L4",
                "benchmark_file": "docs/results/benchmark-2026-10-10.json",
                "benchmark_sha256": digest(ROOT / "docs/results/benchmark-2026-10-10.json"),
                "loop_file": "docs/presentations/v2/loop-observations.json",
                "loop_sha256": digest(ROOT / "docs/presentations/v2/loop-observations.json"),
                "source_capture_manifest": "docs/presentations/v2/capture-manifest.json",
                "new_model_calls": False,
                "historical_charts_and_screenshots_unchanged": True,
                "demo_url": DEMO,
                "repo_url": REPO,
                "documentation": [
                    "deployment/DEPLOYED.md",
                    "deployment/README.md",
                    "docs/DEVELOPMENT.md",
                ],
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n"
    )
    assert digest(BASE) == base_hash
    print(
        f"Created {target.relative_to(ROOT)} — 25 slides; EC2 architecture, historical evidence retained."
    )


if __name__ == "__main__":
    main()
