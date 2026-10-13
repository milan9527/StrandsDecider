"""Build v8 with warmed three-model observations, retaining the v7 story and design."""

import copy
import hashlib
import io
import json
import shutil
import statistics
from datetime import UTC, datetime
from pathlib import Path

import cairosvg
import matplotlib
import presentation_layout as L
from build_presentation_v5 import replace_text
from fontTools.ttLib import TTFont
from matplotlib import font_manager
from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/presentations/v8"
BASE = next((ROOT / "docs/presentations/v7").glob("*.pptx"))
NAME = "Strands-Decider-模型原理与支付业务演示-v8-三模型预热性能对比-30页-20261012"
DEMO = "https://d3i7ogfx5ml7dc.cloudfront.net/#crypto"
REPO = "https://github.com/milan9527/StrandsDecider"
BENCH_URL = "https://benchmarkheaven.com/jev-models/api"
BLOG = "https://strandsagents.com/blog/introducing-strands-decider/"
PURPLE, PALE_PURPLE = "8055A4", "F3EDFA"
PROVIDERS = ("decider", "jev", "bedrock")
NAMES = {"decider": "Strands Decider", "jev": "Jev 1.13.0", "bedrock": "GPT-6.1 Sol"}
COLORS = {"decider": L.TEAL, "jev": PURPLE, "bedrock": L.BLUE}
STAGES = ("intake", "verify", "approve", "execute", "close")
STAGE_LABELS = ("客服受理", "交易核验", "风控审批", "业务处置", "回复结案")
CASES = ("pending", "underpaid", "refund", "settlement")
CASE_LABELS = ("USDC 未到账", "USDT 少付", "稳定币退款", "商户结算")


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    assets = OUT / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    base_hash = digest(BASE)
    source = Presentation(BASE)
    prs = Presentation(BASE)
    for sid in list(prs.slides._sldIdLst):
        prs.part.drop_rel(sid.rId)
        prs.slides._sldIdLst.remove(sid)
    observed = json.loads((OUT / "live-observations.json").read_text())
    benchmark = json.loads((OUT / "sources/jevbench-api-20261012.json").read_text())
    bench = {r["key"]: r for r in benchmark["rows"]}
    old = json.loads((ROOT / "docs/results/optimized-2026-10-11.json").read_text())
    slides, crops = [], []
    cases = observed["cases"]
    samples = {p: [r for c in cases.values() for r in c["results"][p].values()] for p in PROVIDERS}
    times = {p: [r["timings"]["backend_total_ms"] for r in samples[p]] for p in PROVIDERS}
    means = {p: statistics.mean(times[p]) for p in PROVIDERS}
    medians = {p: statistics.median(times[p]) for p in PROVIDERS}
    totals = {p: sum(times[p]) for p in PROVIDERS}
    matches = {p: sum(r["matches_reference"] for r in samples[p]) for p in PROVIDERS}
    lambda_evidence = json.loads((OUT / "lambda-warm-verification.json").read_text())
    assert lambda_evidence["verified_all_measured_warm"]
    assert lambda_evidence["measured_calls"] == 20
    assert matches == {"decider": 18, "jev": 20, "bedrock": 20}
    assert all(len(samples[p]) == 20 for p in PROVIDERS)
    assert all(r["source"] == "live" for rows in samples.values() for r in rows)
    common_notes = (
        "\n\n项目实测来源：docs/presentations/v8/live-observations.json；"
        f"采集开始 {observed['captured_at']}。"
        "API 20261012-triple1 / 前端 triple1，us-east-1。"
        "优化流程4场景×5步×3模型=60次真实调用，全部成功。"
        "先各调用三模型一次预热，3次预热不计入60次正式测试；"
        "CloudWatch逐请求时间窗匹配正式20次Jev调用，均无Init Duration，"
        "且复用预热请求的Lambda环境。未从原始计时中减去任何估算值。"
        "这只验证自有Lambda代理未冷启动，不验证上游Jev基础设施状态。"
        "独立USDC闭环复用v7保存的18次真实决策与截图，不混入固定输入60次统计。"
        "Decider v19 / EC2 g6.xlarge NVIDIA L4常驻；Jev实际版本1.13.0；"
        "Bedrock us.openai.gpt-6.1-sol，reasoning low，仅返回候选ID JSON。"
        "每步三侧业务payload完全一致；GPT适配另加输出格式要求。"
        "每步计时backend_total_ms含通信、后端处理与排队；三模型并发调用，"
        "各列累计时间不能相加作为浏览器总经过时间。"
        "演示参考由开发者定义；参考一致不等于通用准确率。"
        "业务为模拟，客服回复来自应用模板；未连接链上或执行真实资金交易。"
        "当前README与PPT采用同一预热测试数据，原始批次在历史目录保留。"
        f"\n公开基准：{BENCH_URL}，"
        f"快照 {benchmark['captured_at']}，board {benchmark['board_version']}。"
        "这是第三方已发布结果，本项目未重跑JevBench。"
        f"\n模型原理：{BLOG}"
        "\nDecider模型卡：https://huggingface.co/StrandsAgents/strands-decider-2B-hobson-v19"
        "\nJev官方：https://typesafe.ai/blog/introducing-system-one-models-and-jev"
        f"\nDemo：{DEMO}\nRepo：{REPO}"
    )

    def text(s, value, x, y, w, h=0.4, size=18, color=L.INK, bold=False, align=None):
        return L.text(s, str(value), x, y, w, h, size, color, bold, align)

    def icon(s, name, x, y, size=0.4, color=L.TEAL, aws=False):
        path = ROOT / (
            f"docs/diagrams/aws-icons/{name}.svg"
            if aws
            else f"docs/presentations/v3/icons/{name}.svg"
        )
        svg = path.read_text().replace("currentColor", f"#{color}")
        blob = cairosvg.svg2png(bytestring=svg.encode(), output_width=160, output_height=160)
        sh = s.shapes.add_picture(
            io.BytesIO(blob), Inches(x), Inches(y), Inches(size), Inches(size)
        )
        sh.name = "ICON: " + ("AWS " if aws else "Lucide ") + name

    def new(title, section, subtitle="", notes="", dark=False, reused=()):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = L.rgb(L.NAVY if dark else L.BG)
        text(
            s,
            "STRANDS DECIDER  /  " + section,
            0.6,
            0.27,
            12,
            0.25,
            10,
            "86D7C6" if dark else L.TEAL,
            True,
        )
        text(s, title, 0.6, 0.79, 12.1, 0.69, 29.5, L.WHITE if dark else L.INK, True)
        if subtitle:
            text(s, subtitle, 0.62, 1.56, 12, 0.46, 15.5, "CBD9E6" if dark else L.MUTED)
        s.notes_slide.notes_text_frame.text = notes + common_notes
        slides.append(
            {"title": title, "section": section, "dark": dark, "reused_v6_pages": list(reused)}
        )
        return s

    def keep(number, notes, replacements=None):
        src = source.slides[number - 1]
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = L.rgb(L.BG)
        remap = {}
        for rel in src.part.rels.values():
            if rel.reltype == RT.IMAGE:
                _, rid = s.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
                remap[rel.rId] = rid
            elif rel.reltype == RT.HYPERLINK:
                remap[rel.rId] = s.part.relate_to(rel.target_ref, RT.HYPERLINK, True)
        for sh in src.shapes:
            if sh.top >= Inches(7.03):
                continue
            node = copy.deepcopy(sh._element)
            for child in node.iter():
                for key, value in list(child.attrib.items()):
                    if value in remap and key.startswith(
                        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
                    ):
                        child.set(key, remap[value])
            s.shapes._spTree.insert_element_before(node, "p:extLst")
        for sh in s.shapes:
            if sh.has_text_frame:
                for before, after in (replacements or {}).items():
                    replace_text(sh, before, after)
        title = next(
            sh.text for sh in s.shapes if sh.has_text_frame and Inches(0.7) < sh.top < Inches(1.0)
        )
        slides.append(
            {"title": title, "section": "复用 v7", "dark": False, "reused_v7_pages": [number]}
        )
        s.notes_slide.notes_text_frame.text = notes + common_notes
        return s

    def note(s, value, dark=False):
        text(s, value, 0.63, 6.69, 12.03, 0.29, 11.3, "CBD9E6" if dark else L.MUTED)

    def source_link(s, label, url, x=0.65, y=6.35, w=12.0, size=11):
        shape = text(s, label, x, y, w, 0.26, size, L.BLUE)
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.hyperlink.address = url
                run.font.underline = True

    def card(s, x, y, w, h, title, body, color=L.TEAL, glyph=None, size=17):
        L.rect(s, x, y, w, h, L.WHITE, L.LINE, True)
        tx = x + 0.22
        if glyph:
            icon(s, glyph, tx, y + 0.19, color=color)
            tx += 0.57
        text(s, title, tx, y + 0.17, w - (tx - x) - 0.18, 0.43, 19, color, True)
        text(s, body, x + 0.22, y + 0.82, w - 0.44, h - 0.9, size)

    def flow(s, items, y=2.4, height=1.5, x=0.65, w=12.0, color=L.TEAL, connect=True):
        gap = 0.3
        cw = (w - gap * (len(items) - 1)) / len(items)
        for i, (title, body, glyph) in enumerate(items):
            xx = x + i * (cw + gap)
            card(s, xx, y, cw, height, title, body, color, glyph, size=15)
            if connect and i < len(items) - 1:
                L.line(
                    s,
                    [(xx + cw + 0.04, y + height / 2), (xx + cw + gap - 0.04, y + height / 2)],
                    color,
                    1.5,
                    arrow=True,
                )

    def table(s, headers, rows, widths, y=2.23, rh=0.54, size=16, highlights=()):
        x = 0.65
        for j, row in enumerate([headers, *rows]):
            xx = x
            for value, width in zip(row, widths, strict=True):
                fill = (
                    L.NAVY
                    if j == 0
                    else L.PALE_ORANGE
                    if j - 1 in highlights
                    else (L.WHITE if j % 2 else "EFF3F7")
                )
                L.rect(s, xx, y + rh * j, width, rh, fill, L.LINE)
                text(
                    s,
                    value,
                    xx + 0.10,
                    y + rh * j + 0.085,
                    width - 0.20,
                    rh - 0.11,
                    size if j else size - 0.5,
                    L.WHITE if j == 0 else L.INK,
                    j == 0,
                )
                xx += width

    def picture(s, filename, x, y, w, h, folder="screenshots"):
        return L.picture(s, OUT / folder / filename, x, y, w, h, frame=True)

    def crop(filename, box, output):
        path = OUT / "screenshots" / filename
        with Image.open(path) as image:
            image.crop(box).save(assets / output)
        crops.append(
            {
                "source": str(path.relative_to(OUT)),
                "source_sha256": digest(path),
                "crop": list(box),
                "file": f"assets/{output}",
                "sha256": digest(assets / output),
            }
        )

    # 01 — The v6 purpose, expanded to three measured providers.
    s = new(
        "让小模型做好高频决策",
        "v8 · 三模型与 JevBench · 30 页",
        "Strands Decider × Jev × GPT：从公开基准，到支付流程中的真实表现",
        "沿用v6模型目的、原理、应用、实测、优化、结论主线。"
        "新增Jev实测与JevBench对照，以最新同批次截图替换双模型演示主体。",
        dark=True,
        reused=(1,),
    )
    for x, num, title, body, glyph in [
        (0.72, "01", "理解模型", "为什么输出候选答案？\n三种实现有何差异？", "brain-circuit"),
        (
            4.9,
            "02",
            "观察证据",
            "公开基准 + 四个支付场景\n速度、选择、实际执行",
            "chart-no-axes-combined",
        ),
        (9.0, "03", "决定下一步", "如何优化 Decider？\n哪些任务交给哪类模型？", "target"),
    ]:
        text(s, num, x, 2.52, 1, 0.7, 36, "567B99", True)
        icon(s, glyph, x, 3.63, color="86D7C6")
        text(s, title, x + 0.57, 3.6, 3.2, 0.5, 23, L.WHITE, True)
        text(s, body, x, 4.48, 3.65, 1.2, 19, "DCE6F0")
    note(s, "模型原理 → 公开基准 → 项目实测 → 差异解释 → 优化路径  ·  2026.10.12", True)

    # 02–03 — Reuse the original concise motivation and typed primitives.
    keep(2, "高频决策具有有限候选、短输出、由代码消费的共同点。模型处理语义，代码核验资金与权限。")
    keep(
        3,
        "System One接口提供choice、noul和score。本支付对比只测choice，"
        "没有把本项目结果推广到另两种类型。共享上下文用于多个独立问题，"
        "相互依赖且需执行反馈的步骤仍应分轮判断。",
    )

    # 04 — Correct the obsolete “Jev not integrated” statement.
    s = new(
        "三个模型，同一个候选动作任务",
        "01 / 模型定位",
        "比较的是软件判断能力；本演示不测自由长文本生成质量。",
        "Jev强调RLCD校准决策训练，但内部结构未在本项目得到验证。"
        "同一System One接口不等于和Decider共享架构。GPT当前只输出候选ID。",
        reused=(4, 25),
    )
    for x, title, body, color, glyph in [
        (
            0.65,
            "Strands Decider",
            "开源 2B · 自行部署\nQwen 主体 + LoRA\nPointer Head 对候选评分\n\n本项目：EC2 L4 常驻",
            L.TEAL,
            "code-xml",
        ),
        (
            4.79,
            "Jev / TypeSafe",
            "托管 System One API\n类型化答案与概率\n官方强调 RLCD 校准训练\n\n本项目：实际版本 1.13.0",
            PURPLE,
            "list-filter",
        ),
        (
            8.93,
            "GPT / Bedrock",
            "通用生成式模型\n根据候选条件生成 JSON\n本项目仅输出 choice ID\n\nGPT-6.1 Sol · 推理 low",
            L.BLUE,
            "brain-circuit",
        ),
    ]:
        card(s, x, 2.45, 3.75, 3.6, title, body, color, glyph, 17.5)
    note(
        s,
        "Decider 与 Jev 的概率有助于观察选择偏向；不同模型的 confidence 不能直接视为同一种正确率。",
    )

    # 05–07 — Preserve the mechanism/BERT content, update paths to three models.
    keep(
        5,
        "官方v19用Qwen3.5-2B-Base主体、rank-16 LoRA与Pointer Head。"
        "候选位置和answer位置隐藏表示投影后评分，避免逐token生成答案。"
        "不能由这个结构保证业务判断正确。",
    )
    s = new(
        "相同业务问题，三条调用路径",
        "02 / 决策如何产生",
        "应用提供事实、问题与候选条件；模型返回建议，执行器再检查。",
        "Decider内部候选评分可从公开代码核实。Jev作为托管服务只描述已公开接口与返回值；"
        "不推断它采用同样Pointer Head。GPT输出格式限制降低了本演示生成长度。",
        reused=(6, 11),
    )
    for y, provider, steps in [
        (2.42, "decider", ["编码事实与候选", "Pointer Head 评分", "选择 + 分布"]),
        (3.60, "jev", ["提交业务事实、问题与候选项", "托管决策模型", "选择 + 分布"]),
        (4.78, "bedrock", ["提示事实与候选", "生成候选 ID JSON", "解析与格式校验"]),
    ]:
        text(s, NAMES[provider], 0.68, y + 0.24, 2.4, 0.4, 18, COLORS[provider], True)
        for j, value in enumerate(steps):
            xx = 3.05 + j * 3.2
            L.rect(s, xx, y, 2.88, 0.9, L.WHITE, L.LINE, True)
            if provider == "jev" and j == 0:
                text(
                    s,
                    value.replace("、", "、\n"),
                    xx + 0.15,
                    y + 0.12,
                    2.58,
                    0.66,
                    15,
                    COLORS[provider],
                    True,
                )
            else:
                text(s, value, xx + 0.15, y + 0.28, 2.58, 0.4, 17, COLORS[provider], True)
            if j < 2:
                L.line(
                    s, [(xx + 2.92, y + 0.45), (xx + 3.13, y + 0.45)], COLORS[provider], arrow=True
                )
    note(s, "输出受约束，只保证形状与候选范围；条件是否理解正确，需要业务评测。")
    keep(
        7,
        "此页沿用v6的一页BERT比较。标准BERT分类头通常绑定固定标签集合；"
        "NLI等方法也可把动态候选描述转成文本对进行评分。BERT并非只能固定标签，"
        "本项目没有BERT对照实测。Jev不在此页假定与Decider架构相同。",
    )

    # 08 — Source screenshot and a compact explanation of the benchmark axes.
    s = new(
        "JevBench：同时衡量能力、校准、速度与成本",
        "03 / 公开基准",
        "Benchmark Heaven 第三方评测：综合分需要结合能力、速度与成本一起看。",
        "榜单独立于TypeSafe。API榜单包含闭源和开源权重的托管端点，"
        "并不是统一GPU的自部署榜单。全量1500题或600题API重测经等值调整；"
        "由历史日期和不同题池组成。Intelligence为归一化、分层及机会校正等处理后的能力分，"
        "Calibration衡量概率表现。两者都不是本项目参考一致次数。",
    )
    picture(s, "jevbench-leaderboard.png", 0.68, 2.15, 7.25, 4.18)
    for y, title, body, glyph in [
        (2.24, "Intelligence", "类型化决策任务的能力分", "brain-circuit"),
        (3.22, "Calibration", "概率有多可信", "target"),
        (4.20, "Speed", "测量或注明调整的响应耗时", "timer"),
        (5.18, "Cost", "标价建模成本；部分为估算", "banknote"),
    ]:
        icon(s, glyph, 8.30, y, color=PURPLE)
        text(s, title, 8.90, y - 0.02, 3.5, 0.37, 19, PURPLE, True)
        text(s, body, 8.90, y + 0.39, 3.5, 0.4, 15)
    note(s, "综合分采用四轴等权调和平均，并对部分低于 50 的轴加惩罚；高综合分不等于最高答题能力。")
    source_link(s, "评测与方法：benchmarkheaven.com/jev-models/api", BENCH_URL, y=6.36)

    # 09 — Exact public numbers, no invented GPT-Sol or Strands board rows.
    s = new(
        "公开结果：Jev 与主流开源、闭源模型",
        "03 / JevBench API 实际结果",
        "能力、校准、综合分越高越好；延迟、成本越低越好。全部是该榜单配置。",
        "选取决策模型与通用LLM代表，不声称覆盖所有主流模型。"
        "OpenAI GPT-6 Luna和GPT-5.6 Luna均为low配置，Gemini为3.1 Flash-Lite，"
        "DeepSeek V4.1 Flash为默认thinking；Qwen3.8 27B由Chutes TEE托管。"
        "Qwen成本是榜单估算，因此加*，不作为本项目AWS账单。"
        "Sage的成本也标估算。Jev综合第4，但能力分63.6低于选列强LLM。"
        "准确模型ID、provider、价格依据、题池说明保存在sources JSON中。",
    )
    selected = [
        ("sage-1.3.0", "Sage 1.3.0", "闭源"),
        ("liquid-d1", "Liquid d1", "闭源"),
        ("jev-1.13.0", "Jev 1.13.0", "闭源"),
        ("gpt-6-luna-low", "GPT-6 Luna · low", "闭源"),
        ("gpt-5.6-luna", "GPT-5.6 Luna · low", "闭源"),
        ("gemini-3.1-flash-lite", "Gemini 3.1 Flash-Lite", "闭源"),
        ("deepseek-flash", "DeepSeek V4.1 Flash", "开放权重"),
        ("qwen3.8-27b", "Qwen3.8 27B", "开放权重"),
    ]
    rows = []
    for key, label, kind in selected:
        r = bench[key]
        rows.append(
            [
                label,
                kind,
                f"{r['jevbench_score']:.1f}",
                f"{r['axes']['intelligence']:.1f}",
                f"{r['axes']['calibration']:.1f}",
                f"{r['speed']['p50_s_adjusted']:.3f} s",
                f"${r['cost']['usd_per_1000']:.3f}"
                + ("*" if r["cost"]["kind"] == "estimate" else ""),
            ]
        )
    table(
        s,
        ["模型 / 配置", "权重", "综合", "能力 I", "校准 C", "延迟 p50", "$/千决策"],
        rows,
        [3.20, 1.25, 1.1, 1.15, 1.15, 1.65, 1.6],
        rh=0.455,
        size=15,
        highlights=(2,),
    )
    note(
        s,
        "* 为榜单估算成本；分数不是准确率百分比。GPT-6.1 Sol、Claude 与 Strands v19 未列入此 API 表。",
    )
    source_link(s, "评测来源：benchmarkheaven.com/jev-models/api", BENCH_URL, y=6.43)

    # 10 — Read the trade-off, explicitly separate Decider's public-subset report.
    # Matplotlib sees only the first face in a TTC; register its Simplified Chinese face.
    ttc = Path("/usr/share/fonts/google-noto-cjk/NotoSansCJK-Regular.ttc")
    if ttc.exists():
        font_file = ROOT / ".deployment/presentation-fonts/NotoSansCJKSC-Regular.otf"
        font_file.parent.mkdir(parents=True, exist_ok=True)
        if not font_file.exists():
            TTFont(ttc, fontNumber=2).save(font_file)
        font_manager.fontManager.addfont(str(font_file))
    plt.rcParams.update({"font.family": "Noto Sans CJK SC", "font.size": 10})
    fig, ax = plt.subplots(figsize=(7.5, 4.6), layout="constrained")
    scatter_keys = [
        ("jev-1.13.0", "Jev", (-5, 12)),
        ("liquid-d1", "d1", (12, -18)),
        ("gpt-6-luna-low", "GPT-6 Luna low", (-85, 10)),
        ("gemini-3.1-flash-lite", "Gemini Flash-Lite", (-40, -20)),
        ("deepseek-flash", "DeepSeek Flash", (10, -21)),
        ("qwen3.8-27b", "Qwen3.8 27B", (-80, -5)),
    ]
    for key, label, offset in scatter_keys:
        r = bench[key]
        x, y = r["speed"]["p50_s_adjusted"], r["axes"]["intelligence"]
        color = PURPLE if key == "jev-1.13.0" else L.TEAL if r["openSource"] else L.BLUE
        ax.scatter(x, y, s=110, color="#" + color, zorder=3, edgecolors="white")
        ax.annotate(label, (x, y), xytext=offset, textcoords="offset points", fontsize=10)
    ax.set(
        xscale="log",
        xlim=(0.16, 10),
        ylim=(48, 106),
        xlabel="API 延迟 p50（秒，对数轴）",
        ylabel="Intelligence 能力分",
    )
    ax.grid(alpha=0.18)
    ax.spines[["top", "right"]].set_visible(False)
    fig.savefig(assets / "jevbench-capability-latency.png", dpi=220)
    fig.savefig(assets / "jevbench-capability-latency.svg")
    plt.close(fig)
    s = new(
        "综合排名高，不等于复杂问题答得更准",
        "03 / 如何解读基准",
        "Jev 的优势来自低延迟和低成本；强通用模型在该基准的能力分更高。",
        "图为同一API榜单的代表行，横轴用榜单p50，纵轴Intelligence。"
        "并未用项目20步参考一致去替代基准分，也没有用延迟推算模型参数量。"
        "Strands官方模型卡单独报告JevBench public 167/231=72.29%，"
        "serving窗口4096或3072；这是公开子集准确率，既不是封闭题结果，"
        "也不是本榜单Intelligence/Composite。当前API榜单没有Strands v19或本项目Sol配置，"
        "不要把Mapika decider-2b等同名条目当作Strands模型。",
    )
    picture(s, "jevbench-capability-latency.png", 0.7, 2.12, 7.2, 4.24, folder="assets")
    card(
        s,
        8.18,
        2.30,
        4.43,
        1.65,
        "先按任务选能力",
        "榜单中 GPT / DeepSeek / Qwen\n能力分高，但速度与成本不同。",
        L.BLUE,
        "target",
        16,
    )
    card(
        s,
        8.18,
        4.20,
        4.43,
        1.93,
        "Decider 的另一份证据",
        "官方公开子集：167 / 231\n准确率 72.29%\n不可与本页归一化分数混排。",
        L.TEAL,
        "scan-eye",
        16,
    )
    note(s, "公开基准用于筛选候选模型；是否适合你的支付策略，需要在真实业务输入上重新验证。")
    source_link(s, "评测：benchmarkheaven.com/jev-models/api", BENCH_URL, w=7.3, y=6.40)
    source_link(
        s,
        "Decider 来源：官方 v19 模型卡",
        "https://huggingface.co/StrandsAgents/strands-decider-2B-hobson-v19",
        x=8.18,
        w=4.43,
        y=6.40,
    )

    # 11 — Rebuild v6 AWS architecture with the actual Jev branch.
    s = new(
        "一套业务后端，连接三个模型",
        "04 / 当前项目架构",
        "us-east-1 · 私有 S3 + CloudFront OAC · ECS Fargate · EC2 g6.xlarge / L4",
        "沿用v6 AWS官方图标与网络边界。S3四项公共访问阻断均开启。"
        "CloudFront API经VPC Origin到内网ALB，再到私有Fargate。"
        "后端直接调用EC2 Decider、Bedrock和DynamoDB。"
        "Jev经IAM Invoke Lambda，Lambda通过HTTPS访问TypeSafe；密钥在Secrets Manager，"
        "服务图省略网络端点，不增加Bedrock Runtime等独立节点。"
        "图中的Jev在AWS外；Lambda和Bedrock在业务VPC边界外但区域AWS内。",
        reused=(9,),
    )
    L.rect(s, 0.65, 2.18, 10.55, 4.22, L.WHITE, "A7B0B8", True)
    text(
        s,
        "AWS Cloud · 区域服务 us-east-1 · CloudFront 全球入口",
        0.87,
        2.27,
        9.7,
        0.3,
        12,
        L.MUTED,
        True,
    )
    L.rect(s, 3.35, 2.77, 4.05, 3.25, "F3F8FC", "89AABD", True)
    text(s, "VPC · 私有子网", 3.55, 2.87, 3.0, 0.3, 12, L.BLUE, True)

    def service(name, title, body, x, y, w=2.1, h=0.95):
        L.rect(s, x, y, w, h, L.WHITE, L.LINE, True)
        icon(s, name, x + 0.12, y + 0.16, size=0.38, aws=True)
        text(s, title, x + 0.59, y + 0.15, w - 0.7, 0.3, 14, L.INK, True)
        text(s, body, x + 0.12, y + 0.59, w - 0.24, h - 0.61, 11.2)

    service("cloudfront", "CloudFront", "页面 / API 入口", 0.95, 3.1)
    service("s3", "Amazon S3", "私有源站 · OAC", 0.95, 4.56)
    service("alb", "内网 ALB", "CloudFront VPC Origin", 3.64, 3.14, 3.46, 0.80)
    service("fargate", "ECS Fargate", "FastAPI · 决策 / 执行 / 登录", 3.64, 4.12, 3.46, 0.80)
    service("ec2", "EC2 · Decider", "g6.xlarge · L4 24GB · 常驻", 3.64, 5.10, 3.46, 0.80)
    service("dynamodb", "DynamoDB", "用户与会话", 7.82, 2.87, 2.99)
    service("bedrock", "Bedrock · GPT", "GPT-6.1 Sol · low", 7.82, 4.03, 2.99)
    service("lambda", "Lambda · Jev 转发", "IAM 调用 · Secrets Manager 密钥", 7.82, 5.18, 2.99)
    L.rect(s, 11.5, 5.18, 1.15, 0.95, PALE_PURPLE, PURPLE, True)
    text(s, "TypeSafe", 11.56, 5.34, 1.03, 0.3, 13, PURPLE, True)
    text(s, "Jev API", 11.56, 5.75, 1.03, 0.25, 11, PURPLE)
    L.line(s, [(2.0, 4.06), (2.0, 4.50)], L.MUTED, arrow=True)
    L.line(s, [(3.09, 3.58), (3.59, 3.58)], L.MUTED, arrow=True)
    L.line(s, [(5.35, 3.95), (5.35, 4.09)], L.MUTED, arrow=True)
    L.line(s, [(5.35, 4.93), (5.35, 5.07)], L.MUTED, arrow=True)
    for yy in (3.32, 4.49, 5.64):
        L.line(s, [(7.13, 4.52), (7.55, 4.52), (7.55, yy), (7.77, yy)], L.MUTED, arrow=True)
    L.line(s, [(10.85, 5.65), (11.44, 5.65)], PURPLE, arrow=True)
    note(s, "自建用户名登录；业务后端统一校验候选和执行条件。图省略网络端点与非主路径的运维组件。")

    # 12 — Reuse the scenario / workflow map.
    keep(
        12,
        "四场景初始事实各有未满足条件。测试是模拟资金业务，不广播真实交易。"
        "五步依次是受理、核验、审批、处置、回复；5/5不是完成收款或退款。",
        {"同一笔业务，分别交给 Decider 和 GPT。": "同一笔业务，分别交给 Decider、Jev 和 GPT。"},
    )

    # 13 — The actual production comparison screen.
    s = new(
        "真实界面：同一输入，三模型输出并排看",
        "04 / 演示方法",
        "相同事实、问题和候选条件；保留各自的原始选择、耗时与实际执行。",
        "本次统一使用优化流程，冻结每场景初始快照；每步三模型并发。"
        "收到5步方案后分别执行，三侧不会共享已修改状态。"
        "画面是现网运行原始截图，只进行矩形裁剪。",
        reused=(11, 12),
    )
    picture(s, "pending-decision.png", 0.68, 2.10, 12.0, 4.38)
    note(s, "先看原始选择，再看实际动作；模型给出建议，业务系统负责执行。")

    # 14 — Main 4-case table; clear numerator, denominator, and metrics.
    s = new(
        "四场景总览：Decider 最快，两处选择仍有差异",
        "05 / 项目实测",
        "每格：该场景五步累计后端耗时 · 与演示参考一致次数",
        "四个案例的全部阶段均成功，所以本页三个提供方都使用20个同输入任务。"
        "含后端通信与排队，不是纯GPU推理或浏览器整轮延迟。"
        "本页使用全新预热测试，日志确认20次Jev正式调用均无Lambda冷启动。"
        "这些是不同任务的单轮观测，不是固定prompt重复延迟分布。",
        reused=(19, 20, 22),
    )
    rows = []
    for case, label in zip(CASES, CASE_LABELS, strict=True):
        row = [label]
        for p in PROVIDERS:
            results = list(cases[case]["results"][p].values())
            row.append(
                f"{sum(r['timings']['backend_total_ms'] for r in results) / 1000:.3f} s · "
                f"{sum(r['matches_reference'] for r in results)}/5"
            )
        rows.append(row)
    rows.append(["20 步合计"] + [f"{totals[p] / 1000:.3f} s · {matches[p]}/20" for p in PROVIDERS])
    table(
        s,
        ["业务场景", "Strands Decider", "Jev 1.13.0", "GPT-6.1 Sol"],
        rows,
        [2.52, 3.19, 3.19, 3.2],
        rh=0.55,
        size=17,
        highlights=(1, 2),
    )
    for i, p in enumerate(PROVIDERS):
        xx = 0.8 + i * 4.14
        text(s, f"{means[p]:.1f} ms", xx, 5.74, 2.6, 0.52, 27, COLORS[p], True)
        text(s, "平均每次后端耗时", xx, 6.28, 3.5, 0.28, 12.5, L.MUTED)
    note(
        s,
        f"预热后：Jev / Decider = {totals['jev'] / totals['decider']:.2f}×；"
        f"GPT / Decider = {totals['bedrock'] / totals['decider']:.2f}×；"
        f"GPT / Jev = {totals['bedrock'] / totals['jev']:.2f}×。耗时含通信与后端处理。",
    )

    # 15 — Per-stage timing, plotted from the same live exports.
    fig, ax = plt.subplots(figsize=(8.8, 4.8), layout="constrained")
    centers = list(range(5))
    stage_max = 0
    for i, p in enumerate(PROVIDERS):
        values = [
            statistics.mean(
                cases[c]["results"][p][stage]["timings"]["backend_total_ms"] for c in CASES
            )
            for stage in STAGES
        ]
        stage_max = max(stage_max, *values)
        bars = ax.barh(
            [y + (i - 1) * 0.23 for y in centers],
            values,
            height=0.21,
            color="#" + COLORS[p],
            label=NAMES[p],
        )
        ax.bar_label(bars, labels=[f"{v:.0f}" for v in values], padding=3, fontsize=10)
    ax.set_yticks(centers, STAGE_LABELS)
    ax.invert_yaxis()
    ax.set_xlabel("平均 backend_total_ms（每节点 n=4）")
    ax.set_xlim(0, stage_max * 1.35)
    ax.legend(loc="upper right", frameon=False)
    ax.grid(axis="x", alpha=0.15)
    ax.spines[["top", "right", "left"]].set_visible(False)
    fig.savefig(assets / "stage-latency.png", dpi=220)
    fig.savefig(assets / "stage-latency.svg")
    plt.close(fig)
    s = new(
        "逐节点看速度：受理、核验、审批、处置、回复",
        "05 / 性能细节",
        "左侧比较各节点平均耗时；右侧查看每个模型的请求中位数。",
        "每个节点平均四个场景的正式测试值，均在预热后采集。"
        "图从0开始，未隐藏慢请求；四场景各运行一次，没有P95/P99统计。",
        reused=(21,),
    )
    picture(s, "stage-latency.png", 0.66, 2.12, 8.7, 4.28, folder="assets")
    for y, p in zip((2.53, 3.68, 4.83), PROVIDERS, strict=True):
        text(s, NAMES[p], 9.65, y, 2.8, 0.3, 16, COLORS[p], True)
        text(s, f"{medians[p]:.1f} ms", 9.65, y + 0.40, 2.8, 0.43, 24, COLORS[p], True)
    text(s, "右侧：20 次请求中位数", 9.55, 6.12, 3.0, 0.3, 12, L.MUTED)
    note(
        s,
        "每节点各 4 次；每模型共 20 次。比较的是应用调用耗时，不能当作模型纯计算时间。",
    )

    # 16–19 — Every scenario has all five steps plus actual output screenshots.
    case_specs = [
        (
            "pending",
            "USDC 未到账：三模型都等待确认",
            "129 USDC 已足额；确认 3/12。审批可通过，资金仍不能入账。",
            "pending-outputs.png",
            "五个节点三个模型的选择完全一致。审批approved只代表当前审批条件通过，"
            "不等于到账；execute=hold，close=follow_up，资金和工单均未完成。",
            "一致的判断：等待确认 → 检查通过 → 挂起并跟进。",
        ),
        (
            "underpaid",
            "USDT 少付：确认已够，仍需补款",
            "应付 100，到账 99.80 USDT；确认 20/12。差额 0.20 才是当前阻塞。",
            "underpaid-outputs.png",
            "Decider verify=await_confirmations，与confirmed=true矛盾；"
            "Jev/GPT=request_topup符合演示参考。Decider该选择概率0.8448、confidence0.8138，"
            "展示高confidence仍可能错。执行器记录参考差异，未把核验选择替换为正确答案。"
            "之后三侧hold/follow_up，仍等待补款。",
            "关键差异：Decider 等待确认；Jev / GPT 要求补款。",
        ),
        (
            "refund",
            "退款审批：原始建议与实际采用动作不同",
            "250 USDC 退款；资金已核验，但地址未验证、人工审批未通过。",
            "refund-outputs.png",
            "Decider approve=approved，最高概率0.4393、confidence0.2524；"
            "Jev/GPT=verify_destination。执行器将Decider的approved调整为verify_destination，"
            "这是一次资金保护，不计作模型正确。execute均hold，不生成退款账本。",
            "执行保护：Decider「检查通过」→「验证地址」；三侧均挂起退款。",
        ),
        (
            "settlement",
            "商户结算：先升级工单，再等待到账",
            "4,800 USDC；120 分钟未到账，演示 SLA 为 60 分钟。",
            "settlement-outputs.png",
            "三个模型五个阶段都选择merchant_ops、settlement_delayed、approved、"
            "escalate_settlement、follow_up。升级工单并不会将settled变为true。"
            "升级动作与资金到账分开记录，不把已升级视为结算完成。",
            "一致的判断：识别超时 → 升级结算工单 → 持续跟进；尚未对账结案。",
        ),
    ]
    for case, title, subtitle, image, notes, footer in case_specs:
        s = new(
            title,
            "05 / 四场景 · " + CASE_LABELS[CASES.index(case)],
            subtitle,
            notes,
            reused=(13, 14, 16) if case != "settlement" else (20,),
        )
        rows = []
        for stage, label in zip(STAGES, STAGE_LABELS, strict=True):
            row = [label]
            for p in PROVIDERS:
                r = cases[case]["results"][p][stage]
                row.append(f"{r['label']} · {r['timings']['backend_total_ms']:.0f} ms")
            rows.append(row)
        table(
            s,
            ["业务节点", "Strands Decider", "Jev", "GPT"],
            rows,
            [1.5, 3.52, 3.52, 3.56],
            y=2.10,
            rh=0.37,
            size=13.9,
            highlights=(1,) if case == "underpaid" else (2,) if case == "refund" else (),
        )
        picture(s, image, 0.67, 4.47, 12.0, 2.05)
        note(s, footer)

    # 20 — Observed behavior versus plausible mechanisms.
    s = new(
        "为什么选择不同？先分清观察与解释",
        "06 / 行为差异",
        "自然语言候选条件由模型近似判断；同一输入不保证同一选择。",
        "不同模型训练目标、表示能力、校准和输出机制不同，但不能由这些输出确认"
        "某个模型内部思维过程。少付场景可观察到Decider忽略confirmed=true；"
        "可能与条件组合、范围或语言表达有关，需用最小差异测试验证。"
        "高confidence错误和低confidence错误都存在，单一阈值无法拦住全部错误。"
        "Jev和GPT本样本20/20不证明对所有支付案例更可靠。",
        reused=(14, 22),
    )
    card(
        s,
        0.65,
        2.32,
        3.80,
        3.52,
        "已观察到",
        "少付：确认已达标，仍等待\n退款：条件不足，过早通过\n\nJev / GPT 本次符合参考\nDecider 两处与参考不同",
        L.TEAL,
        "scan-eye",
        17,
    )
    card(
        s,
        4.80,
        2.32,
        3.80,
        3.52,
        "合理解释，待验证",
        "条件组合与适用范围\n训练分布和语言表达\n候选评分与生成机制\n\n不能把输出当内部思维链",
        PURPLE,
        "brain-circuit",
        17,
    )
    card(
        s,
        8.95,
        2.32,
        3.72,
        3.52,
        "概率不能代替检查",
        "少付误选 confidence 0.814\n退款误选 confidence 0.252\n\n高置信也会错\n需要业务校准与执行保护",
        L.ORANGE,
        "shield-check",
        16,
    )
    note(
        s,
        "GPT / Jev 的 20/20 是本批次参考一致；Decider 的概率分布是偏向证据，不是业务正确性的证明。",
    )

    # 21 — Completion criteria remain explicit.
    s = new(
        "五步已处理，为什么三个模型都没有结案？",
        "06 / 从选择到业务结果",
        "初始四场景都有未补齐的外部条件；固定输入运行不会自动改变证据。",
        "固定优化批次中三个模型各4个场景的fundsComplete=false、complete=false，"
        "但每个execution.audit都有5条。USDC等待确认、少付等待0.20补款、"
        "退款等待地址验证与审批、商户结算等待实际到账。"
        "截图是USDC实际最终结果，三侧无账本、未结案。",
        reused=(15, 22),
    )
    for x, title, body, color in [
        (0.68, "模型返回", "三个模型各 20/20 次成功", L.TEAL),
        (4.82, "步骤处理", "四场景均已处理 5/5 步", PURPLE),
        (8.96, "业务完成", "初始证据下各 0/4 完成", L.ORANGE),
    ]:
        card(s, x, 2.18, 3.72, 1.38, title, body, color, size=15.2)
    crop("pending-outcomes.png", (0, 0, 1950, 200), "pending-outcomes-status.png")
    picture(s, "pending-outcomes-status.png", 0.68, 3.94, 12.0, 1.3, folder="assets")
    text(s, "仍缺什么？", 0.83, 5.60, 1.8, 0.4, 19, L.ORANGE, True)
    text(
        s,
        "未到账：确认达标   ·   少付：补足差额   ·   退款：地址与审批   ·   结算：到账证据",
        2.74,
        5.64,
        9.75,
        0.63,
        16,
    )
    note(
        s,
        "完成标准：目标订单状态 + 对应模拟账本 + 工单结案。模型调用成功、选择正确、业务完成是三件事。",
    )

    # 22 — Actual Jev latency, with the published context.
    s = new(
        "Jev 为什么没有宣传中的几十倍加速？",
        "06 / 速度差异",
        f"本批次 GPT / Jev = {totals['bedrock'] / totals['jev']:.2f}×；比较条件决定加速倍数。",
        "官方发布文章中的70–500ms是端到端范围，并非纯推理时间。"
        "193.6倍等宣传值被官方称为real world gains的较高端，不是每个工作流承诺。"
        "本项目GPT low且只输出choice，官方部分对照生成概率，因此任务输出量不同。"
        "Jev本项目包含ECS、Lambda、HTTPS和客户端建立，无法把差值全归给模型。"
        f"本次{sum(t < 500 for t in times['jev'])}/20小于500ms；"
        "正式调用均无自有Lambda代理冷启动。"
        f"公开基准239ms与项目{medians['jev']:.1f}ms中位数的任务、调用路径、测量机位置均不同，"
        "只能并列说明，不能直接相减当Lambda开销。",
    )
    flow(
        s,
        [
            ("Decider", "ECS → 私网 EC2\n本地部署的候选评分", "code-xml"),
            ("Jev", "ECS → Lambda → HTTPS\n包含代理与外部往返", "route"),
            ("GPT", "ECS → Bedrock\nlow 推理，只输出 ID", "brain-circuit"),
        ],
        y=2.32,
        height=1.75,
        connect=False,
    )
    table(
        s,
        ["同一 Jev，两个观察位置", "响应中位数", "正确解读"],
        [
            ["JevBench API 公开结果", "239.1 ms", "第三方任务与测量网络"],
            ["本项目支付 20 次调用", f"{medians['jev']:.1f} ms", "含 Lambda、HTTPS 与后端处理"],
        ],
        [4.45, 2.10, 5.55],
        y=4.53,
        rh=0.53,
        size=16,
    )
    note(s, "即使预热，网络往返与代理处理仍占用时间；GPT 只输出候选 ID，也会缩小速度差距。")

    # 23 — One concise real feedback loop slide, including the wrong intermediate action.
    loop = observed["loops"]["pending"]
    assert all(loop["businessOutcomes"][p]["complete"] for p in PROVIDERS)
    s = new(
        "闭环做了什么：等待证据、重新核验、完成入账",
        "06 / 独立闭环实测",
        "USDC 案例 · 三侧各 6 次决策 + 1 次模拟确认事件，最终都完成。",
        "此页独立反馈批次，18次真实决策；后续状态各自演进，不计入同输入60次速度表。"
        "Decider初次verify误选settlement_verified；执行检查发现confirmed=false，"
        "没有记资金账本，loop反馈waiting_external。模拟器推进confirm从3到12，"
        "之后Decider/Jev/GPT均payment_verified、approved、credit_order、resolve。"
        "循环通过事件与执行反馈持续，不是模型自行增加链上确认。"
        "Decider中间错误完整保留，最终完成不把它改记为正确。"
        "风险解除仍需人工；无进展或达上限转人工。",
        reused=(17, 18),
    )
    flow(
        s,
        [
            ("实际检查", "确认不足 → 暂停资金", "shield-check"),
            ("模拟器事件", "确认数 3/12 → 12/12", "repeat-2"),
            ("携带反馈重判", "重新核验 → 入账 → 结案", "clipboard-check"),
        ],
        y=2.16,
        height=1.52,
    )
    crop("pending-loop-outcomes.png", (0, 0, 1950, 200), "pending-loop-status.png")
    picture(s, "pending-loop-status.png", 0.69, 3.96, 12.0, 1.30, folder="assets")
    text(
        s,
        "三侧实际结果：订单已收款  ·  模拟账本 129 USDC  ·  工单已结案",
        0.80,
        5.74,
        11.8,
        0.48,
        20,
        L.TEAL,
        True,
    )
    note(
        s,
        "Decider 首轮误选「结算已核验」，执行检查仍等待确认；后续重判完成。完成闭环不代表每次选择都正确。",
    )

    # 24 — Actually reuse the latest v6 optimization example, keeping its evidence.
    keep(
        24,
        "复用v6第23页已验证的输入优化案例，证据为2026-10-11独立80次批次。"
        "同一USDC事实，固定输入Decider等待人工审批75.722ms；优化输入通过80.217ms。"
        "输入字段裁剪和候选范围改写同时发生，不能单独归因，也未提升该单步速度。"
        "GPT该旧批次两种配置均通过，Jev未参加这份前后对照。"
        "旧截图仅为优化入口，当前三模型截图见前面的项目实测。"
        "候选ID、权重、参考及执行校验未变。",
        {
            "STRANDS DECIDER   /   05 / 优化流程实测": "STRANDS DECIDER   /   07 / 已验证的输入优化",
            "USDC 收款 · 同一业务证据，调整输入字段与候选条件": "已验证的前后对照 · USDC 收款 · 只调整输入字段与候选条件",
        },
    )

    # 25 — Historical A/B result, explicitly not a synthetic three-model A/B.
    s = new(
        "已验证的优化：Decider 参考一致 14/20 → 18/20",
        "07 / 优化证据",
        "明确候选适用范围、去除无关字段，减少两类业务条件之间的混淆。",
        "左表数据来自2026-10-11 EC2独立批次，每case先固定后优化一次；"
        "Decider固定14/20优化18/20，GPT都20/20，保护调整4→1。"
        "本次2026-10-12优化批次Decider仍18/20、Jev/GPT20/20。"
        "不能把两日期的耗时变化当优化效果，也没有Jev固定vs优化A/B。"
        "当前剩余差异与此前优化后相同：少付核验和退款审批。",
        reused=(24,),
    )
    rows = []
    for case, label in zip(CASES, CASE_LABELS, strict=True):
        counts = []
        for profile in ("baseline", "optimized"):
            data = next(c for c in old["profiles"][profile]["cases"] if c["case"] == case)
            counts.append(sum(r["decider"]["matches_reference"] for r in data["steps"]))
        rows.append([label, f"{counts[0]}/5", f"{counts[1]}/5"])
    rows.append(["合计", "14/20", "18/20"])
    table(
        s,
        ["业务场景", "Decider 固定", "Decider 优化"],
        rows,
        [2.14, 2.0, 2.0],
        y=2.34,
        rh=0.53,
        size=16,
    )
    card(
        s,
        7.20,
        2.35,
        5.4,
        1.57,
        "已验证的变化",
        "参考一致改善；保护调整 4 → 1\n模型权重未变，未进行领域微调。",
        L.TEAL,
        "circle-check",
        17,
    )
    card(
        s,
        7.20,
        4.23,
        5.4,
        2.08,
        "当前三模型再次观察",
        "Decider 18/20 · Jev / GPT 20/20\n剩余：少付核验、退款审批\n仍需下一轮优化与独立测试。",
        PURPLE,
        "target",
        16,
    )
    note(
        s,
        "优化输入已减少选择错误；模型权重未变。少付核验与退款审批仍需改进。",
    )

    # 26 — Concrete next experiments tied to the observed errors.
    s = new(
        "下一步优化：把剩余错误变成可重复的测试",
        "07 / Decider 优化路线",
        "优先改问题表达与任务拆分，再判断是否需要改模型权重。",
        "以下为建议实验，尚未实施或测得效果。对少付问题，验证confirmed=true的"
        "状态表达、候选顺序、数字边界和优先级是否影响选择；对退款问题，"
        "先判场景再核验address_verified与approved，避免把收款充分条件套到退款。"
        "可按业务类型过滤不适用候选；这会改变任务，要对三模型同样应用，并另记结果。"
        "低置信路由需校准，少付高置信误选证明阈值不能替代确定性保护。",
    )
    flow(
        s,
        [
            ("收集错误", "事实 + 候选 + 原始选择\n实际执行 + 人工正确动作", "receipt-text"),
            ("构造最小对照", "只改一个条件或边界\n加入近似反例与语言变体", "list-filter"),
            ("验证改动", "冻结测试集与版本\n比较错误、延迟和完成率", "clipboard-check"),
        ],
        y=2.34,
        height=1.78,
    )
    card(
        s,
        0.66,
        4.52,
        5.86,
        1.67,
        "少付：把确认与金额分开",
        "确认达到阈值后，明确跳过等待确认；\n用差额 0 / 0.01 / 0.20 测补款分支。",
        L.TEAL,
        "wallet",
        16.5,
    )
    card(
        s,
        6.80,
        4.52,
        5.86,
        1.67,
        "退款：先限定场景，再检查前提",
        "地址验证 → 人工审批 → 执行资格；\n覆盖缺失、冲突、风险优先等反例。",
        L.ORANGE,
        "shield-check",
        16.5,
    )
    note(s, "这是后续实验方案，未宣称已提升效果；参考与真实执行保护必须独立于模型输出。")

    # 27 — Adapt the actual architecture, not a fictitious chat fine-tune.
    s = new(
        "需要微调时：训练决策能力，不只训练 JSON 外形",
        "07 / 领域微调方案",
        "沿用官方训练链路，适配 LoRA 与 Pointer Head；先准备独立领域评测。",
        "这是计划而非已完成训练。官方模型卡与仓库公开训练recipe、数据清单与评测。"
        "基于v19的业务数据应包含state/questions/criteria及人工审核正确候选，"
        "并覆盖noul、score回归。按案件或模板簇与时间切分，防止同义改写泄漏。"
        "优先使用官方decision training接口与目标，训练rank16 LoRA和评分头的适配方案，"
        "是否冻结主体或扩展rank由资源与验证结果决定，不给未经执行的训练CLI。"
        "GPT/Jev可帮助生成候选标注，但必须人工复核，不能把本次20答案直接当足够训练集。"
        "温度或概率校准在独立验证集完成；最终测试集不参与选择阈值或提示调优。",
    )
    flow(
        s,
        [
            ("领域样本", "真实案件去标识\n覆盖正常与困难反例", "receipt-text"),
            ("人工标注", "正确候选 / 判断依据\n教师模型只辅助", "headset"),
            ("参数适配", "LoRA + 候选评分头\n沿用类型化决策目标", "brain-circuit"),
            ("独立评测", "保留集 + 校准 + 延迟\n灰度验证与可回滚", "clipboard-check"),
        ],
        y=2.46,
        height=1.98,
    )
    table(
        s,
        ["必须守住的边界", "目的"],
        [
            ["按案件、模板簇、时间切分", "避免同义样本跨训练与测试集泄漏"],
            ["保留通用任务与所有问题类型的回归集", "防止领域改善伴随能力退化"],
            ["执行保护独立；上线前核验实际业务结果", "不把微调后的高置信选择直接当资金授权"],
        ],
        [6.03, 6.07],
        y=4.79,
        rh=0.43,
        size=15.4,
    )
    note(s, "当前演示没有修改模型权重；本页是下一阶段实施方法，尚无微调后的成绩或成本承诺。")

    # 28 — Production evaluation and routing, compressed from v6's recommendations.
    s = new(
        "落地方式：常见判断用小模型，复杂异常升级",
        "08 / 应用建议",
        "先按业务效果建立分流，再用真实用量评估总体成本。",
        "此页路由为建议，当前demo仍并排测试，并未实现生产级自动择模。"
        "保留金额权限状态机、幂等、人工审核等执行保护。"
        "Decider常驻GPU费用与托管token/调用费用不能直接用榜单每千次代替。"
        "需分段测量应用排队、GPU推理、Lambda和HTTPS，重复随机顺序运行，"
        "记录失败、超时、p50/p95/p99、吞吐、风险错误率、覆盖率和全成本。",
        reused=(25,),
    )
    flow(
        s,
        [
            ("业务请求", "消息 + 经过核验的事实", "receipt-text"),
            ("任务分流", "明确、常见 → 小模型\n复杂、冲突 → GPT / 人工", "route"),
            ("执行系统", "检查条件、权限与账本\n记录结果，反馈再判断", "shield-check"),
        ],
        y=2.4,
        height=1.88,
    )
    for x, title, body, color in [
        (0.7, "质量", "按场景错误率\n高风险误放行\n校准与自动化覆盖率", L.TEAL),
        (4.85, "性能", "同任务重复测量\np50 / p95 / p99\n吞吐、失败与超时", PURPLE),
        (9.0, "业务与成本", "实际完成率与人工介入\n常驻 GPU 利用率\n托管用量与基础设施", L.BLUE),
    ]:
        card(s, x, 4.64, 3.64, 1.95, title, body, color, size=15.5)
    note(s, "置信度阈值应在独立业务集上校准；高置信也可能出错，执行保护始终独立。")

    # 29 — The takeaways are grounded in the observed numbers.
    s = new(
        "总结：同时看速度、选择和真实业务结果",
        "08 / 结论",
        "公开基准帮助选型；项目实测解释行为；独立评测验证优化。",
        "Decider当前批次最快但有两处参考差异。Jev托管接入，当前与参考一致且"
        "较GPT更快；GPT同样20/20，开放式能力在本demo未测试。"
        "JevBench显示高综合分包含成本速度因素，不应解释为Jev在能力轴超过强LLM。"
        "已有输入优化改善Decider选择，后续应针对剩余错误设计实验，必要时领域微调。",
        reused=(26,),
    )
    table(
        s,
        ["模型", "当前观察", "适合继续验证的方向"],
        [
            [
                "Strands Decider",
                f"{means['decider']:.1f} ms / 次 · 18/20",
                "可自行部署的高频有限判断",
            ],
            ["Jev 1.13.0", f"{means['jev']:.1f} ms / 次 · 20/20", "托管决策、概率与接入成本"],
            [
                "GPT-6.1 Sol",
                f"{means['bedrock']:.1f} ms / 次 · 20/20",
                "复杂语义、异常处理和开放式生成",
            ],
        ],
        [2.80, 4.05, 5.25],
        y=2.4,
        rh=0.71,
        size=18,
    )
    text(s, "已验证：输入优化改善选择", 0.82, 5.62, 5.9, 0.48, 22, L.TEAL, True)
    text(s, "待验证：领域微调与生产收益", 6.88, 5.62, 5.8, 0.48, 22, PURPLE, True)
    note(
        s,
        "本批次 20 个不同任务，不代表通用准确率或稳定延迟；四个初始场景均需外部条件才能完成资金业务。",
    )

    # 30 — Keep the original demo / repo ending, with working links and source URL.
    s = new(
        "打开 Demo，查看三模型过程与原始记录",
        "08 / 演示与复现",
        "从相同业务输入开始，核对每一步的建议、实际动作、耗时与最终结果。",
        "仓库已公开。Demo仍需要现有用户名密码。演示附件包含PPT、PDF、"
        "实测导出、截图来源、公开基准快照与逐页讲稿。",
        reused=(27,),
    )
    for x, title, url, qr, body in [
        (0.72, "在线 Demo", DEMO, "demo-qr.png", "加密支付业务台 · 使用现有账号登录"),
        (6.96, "GitHub Repo", REPO, "repo-qr.png", "公开仓库 · 架构、代码、资料与实测记录"),
    ]:
        shutil.copyfile(ROOT / "docs/presentations/v6/assets" / qr, assets / qr)
        icon(s, "workflow" if title == "在线 Demo" else "code-xml", x, 2.57)
        text(s, title, x + 0.6, 2.54, 4.5, 0.5, 23, L.TEAL, True)
        L.picture(s, assets / qr, x, 3.32, 1.80, 1.80)
        sh = text(s, url, x, 5.50, 5.68, 0.66, 13.6, L.BLUE, True)
        sh.text_frame.paragraphs[0].runs[0].hyperlink.address = url
        text(s, body, x + 2.10, 3.75, 3.46, 1.0, 17)
    sh = text(
        s, "公开基准  benchmarkheaven.com/jev-models/api", 0.75, 6.31, 11.8, 0.35, 14, PURPLE, True
    )
    sh.text_frame.paragraphs[0].runs[0].hyperlink.address = BENCH_URL
    note(
        s,
        "原理来源：Strands 官方博客、官方模型卡、TypeSafe System One；完整链接与测量口径见逐页讲稿。",
    )

    assert len(prs.slides) == 30, len(prs.slides)
    for i, (s, meta) in enumerate(zip(prs.slides, slides, strict=True), 1):
        L.line(s, [(0.62, 7.04), (12.68, 7.04)], "425769" if meta["dark"] else L.LINE, 0.6)
        text(
            s,
            "模型原理 → 公开基准 → 项目实测 → 优化路径",
            0.63,
            7.15,
            10.4,
            0.21,
            9,
            "B8C9DB" if meta["dark"] else L.MUTED,
        )
        text(
            s,
            f"{i:02} / 30",
            11.76,
            7.14,
            0.92,
            0.23,
            10,
            "B8C9DB" if meta["dark"] else L.MUTED,
            align=PP_ALIGN.RIGHT,
        )
        for effect in s._element.xpath(".//a:effectRef"):
            effect.set("idx", "0")
        for sh in s.shapes:
            assert min(sh.left, sh.top) >= 0, (i, sh.name)
            assert sh.left + sh.width <= prs.slide_width + 12, (i, sh.name)
            assert sh.top + sh.height <= prs.slide_height + 12, (i, sh.name)
    prs.core_properties.title = "Strands Decider × Jev × GPT · 三模型与JevBench对比"
    prs.core_properties.subject = "决策模型原理、公开基准、四场景实测、行为差异与优化"
    prs.core_properties.comments = (
        "v8 reuses v7 with new prewarmed live measurements and linked benchmark sources."
    )
    target = OUT / f"{NAME}.pptx"
    prs.save(target)
    (OUT / "讲稿与来源.md").write_text(
        "# 三模型与 JevBench 对比 · v8 · 30 页\n\n"
        + "\n\n".join(
            f"## {i:02}. {meta['title']}\n\n{s.notes_slide.notes_text_frame.text}"
            for i, (s, meta) in enumerate(zip(prs.slides, slides, strict=True), 1)
        )
        + "\n"
    )
    manifest = {
        "generated_at": datetime.now(UTC).isoformat(),
        "base_file": str(BASE.relative_to(ROOT)),
        "base_sha256": base_hash,
        "pptx_file": str(target.relative_to(ROOT)),
        "pptx_sha256": digest(target),
        "slides": slides,
        "live_data_sha256": digest(OUT / "live-observations.json"),
        "jevbench_snapshot_sha256": digest(OUT / "sources/jevbench-api-20261012.json"),
        "same_input_live_decisions": 60,
        "excluded_warmup_calls": 3,
        "lambda_verification_sha256": digest(OUT / "lambda-warm-verification.json"),
        "feedback_evidence_reused_from": "docs/presentations/v7/live-observations.json",
        "separate_feedback_live_decisions": 18,
        "models_fine_tuned": False,
        "crops": crops,
        "metrics": {
            "means_ms": means,
            "medians_ms": medians,
            "totals_ms": totals,
            "reference_matches": matches,
        },
        "demo_url": DEMO,
        "repo_url": REPO,
    }
    (OUT / "build-manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n"
    )
    assert digest(BASE) == base_hash
    print(f"Created {target.relative_to(ROOT)} — {len(prs.slides)} slides.")


if __name__ == "__main__":
    main()
