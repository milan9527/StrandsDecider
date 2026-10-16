"""Explain routing scenarios and merge a real support comparison into one slide."""

import json
from datetime import UTC, datetime

import build_presentation_v10 as V
import presentation_layout as L
from build_presentation_v5 import replace_text
from build_presentation_v13 import request_box
from pptx import Presentation
from pptx.util import Inches

ROOT = V.ROOT
OUT = ROOT / "docs/presentations/v14"
BASE = next((ROOT / "docs/presentations/v13").glob("*.pptx"))
NAME = "Strands-Decider-v14-模型路由与高频客服实测-30页-20261015"
PROVIDERS = ("decider", "jev", "bedrock")
LABELS = ("Decider", "Jev", "GPT")
COLORS = (L.TEAL, V.PURPLE, L.BLUE)


def save(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def main():
    (OUT / "assets").mkdir(exist_ok=True)
    V.OUT, V.ASSETS, V.SOURCES = OUT, OUT / "assets", OUT / "sources"
    V.EVIDENCE.clear()
    V.IMAGES.clear()
    V.SLIDES.clear()
    before_hash = V.digest(BASE)
    old = V.load(V.record(ROOT / "docs/presentations/v13/build-manifest.json"))
    live = V.load(V.record(OUT / "sources/support-live.json"))
    review = V.load(V.record(OUT / "sources/support-review.json"))
    V.record(OUT / "capture-manifest.json")
    assert live["complete"]
    sides = live["executed_export"]["cases"]["support"]["sides"]
    comparison = {
        "task": live["task"], "captured_at": live["finished_at"],
        "method": "One support task, six excluded warmups, six routing calls, three target attempts. "
                  "A frequent-support use case, not a high-concurrency or throughput test.",
        "configuration": "Router GPT low; targets use their mapped benchmark effort. "
                         "Maximum 4096 output tokens, 40-second response wait; no retries or substitutions.",
        "routes": {},
    }
    for p in PROVIDERS:
        side = sides[p]
        steps, target = side["steps"], side["execution"]
        http = next(c for c in live["calls"] if c["phase"] == f"target-{p}")
        row = {
            "workflow_choice": steps[0]["choice"], "model_choice": steps[1]["choice"],
            "target_model": steps[1]["selected_model"]["name"],
            "target_model_id": steps[1]["selected_model"]["id"],
            "workflow_ms": steps[0]["timings"]["backend_total_ms"],
            "model_ms": steps[1]["timings"]["backend_total_ms"],
            "routing_ms": sum(s["timings"]["backend_total_ms"] for s in steps),
            "target_status": http["status"], "target_complete": target.get("complete", False),
            "target_has_final_text": target.get("has_final_text", False),
            "target_ms": target.get("backend_total_ms"), "target_usage": target.get("usage"),
            "target_request_fields": target.get("request_fields"),
            "target_text": target.get("text"), "target_error": target.get("error"),
            "review": review["routes"][p],
        }
        if row["review"].get("quote"):
            assert row["review"]["quote"] in target["text"].replace("**", "")
        comparison["routes"][p] = row
    save(OUT / "sources/support-comparison.json", comparison)
    V.record(OUT / "sources/support-comparison.json")
    prs = Presentation(BASE)
    originals, ids = list(prs.slides), list(prs.slides._sldIdLst)
    order, metas = [], []

    def keep(page):
        order.append(ids[page - 1])
        meta = dict(old["slides"][page - 1])
        meta.update(page=len(order), source="v13", v13_page=page, body_modified=False)
        metas.append(meta)
        return originals[page - 1]

    def new(section, title, subtitle, notes):
        slide = V.base(prs, section, title, subtitle, "", notes)
        order.append(prs.slides._sldIdLst[-1])
        meta = dict(V.SLIDES[-1])
        meta.update(page=len(order), source="new")
        metas.append(meta)
        return slide

    for page in range(1, 15):
        keep(page)
    originals[5].notes_slide.notes_text_frame.text = (
        originals[5].notes_slide.notes_text_frame.text.replace("独立 vLLM API 见第 23 页", "独立 vLLM API 见第 22 页")
    )

    s = new("04  模型路由 · 场景与实现",
            "模型路由：同一个入口，按任务选择处理流程与模型",
            "客服、代码、核算和索引需要不同能力；比较三个决策模型如何选型，以及目标模型能否完成任务。",
            "本页新增业务场景解释。四个任务的示例与处理流程来自现有app/routing.py及"
            "docs/MODEL_ROUTING.md。用户只提交任务原文；这些场景说明和评价目标不是新增输入优先级。"
            "第一步从四类处理流程中选择，第二步根据任务、所选流程、Bedrock可用候选与第三方证据选择型号。"
            "程序不按排名替换模型选择；候选合法不等于选型适合或目标完成。"
            "四任务汇总保留原批次。v14第18页单独补采高频客服的路由和三侧目标执行，"
            "原v13结算分析两页移出当前故事线，旧版和记录保留。")
    scenarios = [
        ("高频客服", "headset", "解释 USDC 待支付。\n3/12 次确认意味着什么？", L.TEAL),
        ("交易代码", "code-xml", "生成支付回调函数。\n去重、验签与幂等入账。", L.BLUE),
        ("结算分析", "banknote", "核算到账差额。\n拆解手续费与暂扣款。", V.PURPLE),
        ("中英文索引", "search", "工单转成检索向量。\n选择适合中英文的模型。", L.ORANGE),
    ]
    for i, (title, icon, content, color) in enumerate(scenarios):
        x = .68 + i * 3.06
        L.rect(s, x, 2.19, 2.80, 1.73, L.WHITE, L.LINE, True)
        V.icon(s, icon, x + .17, 2.40, color, .32)
        L.text(s, title, x + .65, 2.38, 1.96, .35, 17, color, True)
        L.text(s, content, x + .17, 2.97, 2.46, .79, 14.2)
    V.flow(s, [
        ("用户给任务", "只提供业务需求", "receipt-text"),
        ("第 1 次决策", "选处理流程", "route"),
        ("第 2 次决策", "按候选证据选型号", "list-filter"),
        ("执行与验收", "看答复、耗时与完成情况", "clipboard-check"),
    ], x=.75, y=4.22, total=11.84, h=1.68)
    V.callout(s, "业务侧不指定品牌或优先级；路由服务补充可用模型与评测事实，再保留三个决策模型的真实选择。",
              y=6.16, size=13.8)

    keep(16)
    keep(17)
    notes = (
        "v14新增高频客服任务实测，替代v13第18–19页结算分析并合并成一页。"
        "实际任务原文直接展示，不加入速度/成本优先级或品牌提示。"
        "6次预热单独保存；正式6次路由调用和3次目标尝试。"
        "同阶段三路获得相同model_input与input_hash，所有实际选择保留。"
        "路由GPT为GPT-6 Sol low；目标模型采用评测映射档位，最多4096输出token，响应等待上限40秒。"
        "路由与生成分别计时；超时没有精确目标耗时或usage时不伪造数据。"
        "任务为高频客服的一个代表样例，未做高并发测试或给出稳定延迟结论。"
        "仅生成答复文本，未向真实客户发送信息，也未执行资金操作。"
        "\n完整输入、HTTP返回、错误、目标答复与实际截图在sources/support-live.json和capture-manifest.json。\n"
        + json.dumps(comparison, ensure_ascii=False, indent=2)
    )
    s = new("04  模型路由 · 高频客服实测",
            "高频客服：一页对比三路选型、实际答复与耗时",
            review["slide_subtitle"], notes)
    request_box(s, live["task"])
    for i, (p, label, color) in enumerate(zip(PROVIDERS, LABELS, COLORS, strict=True)):
        row = comparison["routes"][p]
        x = .68 + i * 4.07
        name = row["target_model"].replace("Claude ", "")
        V.card(s, x, 3.18, 3.84, 3.16, f"{label} → {name}", "", color)
        L.text(s, f'两步路由 {row["routing_ms"]:.1f} ms', x + .18, 3.77, 3.48, .34, 15, color)
        L.text(s, row["review"]["display_status"], x + .18, 4.23, 3.48, .44, 19, color, True)
        if row["review"].get("quote"):
            L.text(s, "“" + row["review"]["quote"] + "”", x + .18, 4.88, 3.48, 1.20, 14.5)
        else:
            shot = V.image(f"docs/presentations/v14/screenshots/support-target-{p}.png",
                           f"support-result-{p}.png")
            L.picture(s, shot, x + .18, 4.84, 3.48, 1.27, frame=False)

    for page in range(20, 32):
        keep(page)
    originals[30].notes_slide.notes_text_frame.text += (
        "\nv14共30页：第15页补充路由四类业务场景；第18页将高频客服的路由与目标结果合并呈现。"
        "其他场景最终对比、vLLM、Jev/EC2计费及交叉图沿用v13。")
    used = {n.rId for n in order}
    for node in list(prs.slides._sldIdLst):
        prs.slides._sldIdLst.remove(node)
        if node.rId not in used:
            prs.part.drop_rel(node.rId)
    for node in order:
        prs.slides._sldIdLst.append(node)
    assert len(prs.slides) == 30
    for i, (slide, meta) in enumerate(zip(prs.slides, metas, strict=True), 1):
        for shape in list(slide.shapes):
            if shape.has_text_frame and Inches(6.5) <= shape.top < Inches(7.03):
                assert not shape.text.strip()
                shape._element.getparent().remove(shape._element)
            elif shape.has_text_frame and shape.top > Inches(7) and (
                shape.text.endswith(" / 31") or shape.text.endswith(" / 32")
            ):
                replace_text(shape, shape.text, f"{i:02} / 30")
        meta["notes"] = slide.notes_slide.notes_text_frame.text
    prs.core_properties.title = "Strands Decider × Jev × GPT：模型路由与高频客服实测"
    prs.core_properties.subject = "v14 / 30 pages / routing scenarios and a one-page real support comparison"
    output = OUT / f"{NAME}.pptx"
    prs.save(output)
    assert V.digest(BASE) == before_hash
    successes = sum(r["target_complete"] for r in comparison["routes"].values())
    save(OUT / "build-manifest.json", {
        "created_at": datetime.now(UTC).isoformat(), "base_file": str(BASE.relative_to(ROOT)),
        "base_sha256": before_hash, "pptx_file": str(output.relative_to(ROOT)),
        "pptx_sha256": V.digest(output), "pages": 30, "slides": metas,
        "evidence": V.EVIDENCE, "images": V.IMAGES, "new_pages": [15, 18],
        "new_warmup_calls": 6, "new_formal_routing_calls": 6, "new_target_attempts": 3,
        "target_successes": successes, "target_incomplete": 3 - successes,
        "aws_resource_mutations": 0, "external_customer_messages": 0,
        "lower_explanations_removed_from_all_slides": True,
    })
    rows = ["# v14 讲稿与来源", "", "30 页；高频客服的路由与目标执行合并在第18页。", ""]
    for meta in metas:
        rows += [f'## {meta["page"]:02} · {meta["title"]}', "", meta["notes"], ""]
    (OUT / "讲稿与来源.md").write_text("\n".join(rows))
    (OUT / "客服实际请求与结果.md").write_text(
        "# 高频客服：实际请求与结果\n\n" + notes + "\n\n## 目标接口完整请求与返回\n\n```json\n"
        + json.dumps([c for c in live["calls"] if c["phase"].startswith("target-")], ensure_ascii=False, indent=2)
        + "\n```\n")
    print(json.dumps({"pptx": str(output), "pages": 30, "new_pages": [15, 18],
                      "target_successes": successes, "target_incomplete": 3 - successes}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
