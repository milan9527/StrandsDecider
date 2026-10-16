"""Audit the rendered v9 deck and create review images without model calls."""

import hashlib
import json
import re
import unicodedata
from datetime import UTC, datetime
from pathlib import Path

import pymupdf as fitz
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/presentations/v9"


def normal(text):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", text)).replace("\u00ad", "")


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def body(slide):
    return [sh.text for sh in slide.shapes if sh.has_text_frame and sh.top < Inches(7.03)]


def main():
    manifest = json.loads((OUT / "build-manifest.json").read_text())
    pptx = ROOT / manifest["pptx_file"]
    prs = Presentation(pptx)
    pdf = fitz.open(pptx.with_suffix(".pdf"))
    base = Presentation(ROOT / manifest["base_file"])
    assert digest(ROOT / manifest["base_file"]) == manifest["base_sha256"]
    assert digest(pptx) == manifest["pptx_sha256"]
    assert len(prs.slides) == len(pdf) == 52
    previews = OUT / "previews"
    previews.mkdir(exist_ok=True)
    missing, bounds, preserved = [], [], []
    thumbnails = []
    for index, (slide, page, meta) in enumerate(zip(prs.slides, pdf, manifest["slides"], strict=True), 1):
        rendered = normal(page.get_text())
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text.strip() and normal(sh.text) not in rendered:
                missing.append({"page": index, "text": sh.text})
            if min(sh.left, sh.top) < 0 or sh.left + sh.width > prs.slide_width + 12 or sh.top + sh.height > prs.slide_height + 12:
                bounds.append({"page": index, "shape": sh.name})
        if meta["source"] == "v8" and meta["v8_page"] not in manifest["preserved_body_except_pages"]:
            before = base.slides[meta["v8_page"] - 1]
            assert body(slide) == body(before), (index, "v8 body text changed")
            images = lambda s: [hashlib.sha256(sh.image.blob).hexdigest() for sh in s.shapes if sh.shape_type == 13]  # noqa: E731
            assert images(slide) == images(before), (index, "v8 image changed")
            preserved.append(index)
        pix = page.get_pixmap(matrix=fitz.Matrix(1.35, 1.35), alpha=False)
        path = previews / f"page-{index:02}.png"
        pix.save(path)
        with Image.open(path) as im:
            thumb = im.copy()
            thumb.thumbnail((384, 216))
            tile = Image.new("RGB", (404, 248), "#e7edf2")
            tile.paste(thumb, ((404 - thumb.width) // 2, 6))
            ImageDraw.Draw(tile).text((12, 226), f"{index:02}", fill="#203448")
            thumbnails.append(tile)
    for suffix, ids in [("all", list(range(52))), ("new", [i for i, s in enumerate(manifest["slides"]) if s["source"] == "new"])]:
        cols = 4
        rows = (len(ids) + cols - 1) // cols
        canvas = Image.new("RGB", (cols * 404, rows * 248), "#e7edf2")
        for i, index in enumerate(ids):
            canvas.paste(thumbnails[index], ((i % cols) * 404, (i // cols) * 248))
        canvas.save(OUT / ("slides-overview.png" if suffix == "all" else "new-slides-overview.png"))
    audit = {
        "audited_at": datetime.now(UTC).isoformat(), "pages": len(pdf),
        "base_unchanged": True, "preserved_v8_body_text_and_images_pages": preserved,
        "shape_bounds_issues": bounds, "rendered_text_missing": missing,
        "pptx_sha256": digest(pptx), "pdf_sha256": digest(pptx.with_suffix(".pdf")),
        "new_model_calls": 0,
    }
    (OUT / "render-audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"pages": len(pdf), "v8_pages_body_preserved": len(preserved), "bounds_issues": bounds, "rendered_text_missing": missing}, ensure_ascii=False, indent=2))
    assert not bounds
    assert not missing


if __name__ == "__main__":
    main()
