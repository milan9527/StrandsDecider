"""Export the blog HTML to PDF and check the rendered text and figures."""

import argparse
import asyncio
import hashlib
import json
import re
from datetime import UTC, datetime
from pathlib import Path

import pymupdf
from bs4 import BeautifulSoup
from docx import Document
from PIL import Image, ImageDraw
from playwright.async_api import async_playwright

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/blogs/strands-decider"
NAME = "Strands-Decider-AWS实践"


def normalize(text):
    return re.sub(r"\s+", "", text)


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


async def render(browser_path):
    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch(
            executable_path=browser_path, headless=True,
            args=["--no-sandbox", "--disable-dev-shm-usage"],
        )
        page = await browser.new_page(viewport={"width": 1400, "height": 1100})
        await page.goto((OUT / f"{NAME}.html").as_uri())
        await page.evaluate("document.fonts.ready")
        images = await page.evaluate("""() => Array.from(document.images).map(img => ({
            loaded: img.complete && img.naturalWidth > 0,
            width: img.naturalWidth, height: img.naturalHeight
        }))""")
        overflow = await page.evaluate(
            "document.documentElement.scrollWidth > window.innerWidth"
        )
        await page.screenshot(path=str(OUT / "article-preview.png"))
        await page.pdf(
            path=str(OUT / f"{NAME}.pdf"), format="A4",
            print_background=True, prefer_css_page_size=True,
        )
        await page.set_viewport_size({"width": 390, "height": 844})
        mobile_overflow = await page.evaluate(
            "document.documentElement.scrollWidth > window.innerWidth"
        )
        await browser.close()

    assert len(images) == 5 and all(image["loaded"] for image in images)
    assert not overflow, "Desktop has horizontal overflow"
    assert not mobile_overflow, "Mobile has horizontal overflow"

    pdf_path = OUT / f"{NAME}.pdf"
    pdf = pymupdf.open(pdf_path)
    pdf_text = normalize("".join(page.get_text() for page in pdf))
    soup = BeautifulSoup((OUT / f"{NAME}.html").read_text(), "html.parser")
    blocks = soup.article.select("h1,h2,h3,p,li,th,td,pre,figcaption")
    missing = [
        block.get_text() for block in blocks
        if normalize(block.get_text()) not in pdf_text
    ]
    assert not missing, f"PDF text missing: {missing}"
    document = Document(OUT / f"{NAME}.docx")
    assert len(document.inline_shapes) == 5
    assert len(document.tables) == 3
    pdf[0].get_pixmap(matrix=pymupdf.Matrix(1.4, 1.4)).save(OUT / "pdf-first-page.png")
    thumbs = []
    for page in pdf:
        pix = page.get_pixmap(matrix=pymupdf.Matrix(.65, .65), alpha=False)
        thumbs.append(Image.frombytes("RGB", [pix.width, pix.height], pix.samples))
    width, height = thumbs[0].size
    canvas = Image.new("RGB", (width * 2 + 60, (height + 36) * ((len(thumbs) + 1) // 2)), "#e8edf1")
    draw = ImageDraw.Draw(canvas)
    for i, thumb in enumerate(thumbs):
        x, y = 20 + i % 2 * (width + 20), 22 + i // 2 * (height + 36)
        draw.text((x, y - 17), f"Page {i+1}", fill="#243447")
        canvas.paste(thumb, (x, y))
    canvas.save(OUT / "pdf-overview.png")
    audit = {
        "audited_at": datetime.now(UTC).isoformat(),
        "pdf_pages": len(pdf), "images_loaded": images,
        "desktop_overflow": overflow, "mobile_overflow": mobile_overflow,
        "pdf_sha256": sha256(pdf_path), "missing_pdf_text_blocks": missing,
        "word_figures": len(document.inline_shapes), "word_tables": len(document.tables),
    }
    (OUT / "render-audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n"
    )
    manifest_path = OUT / "build-manifest.json"
    manifest = json.loads(manifest_path.read_text())
    manifest["files"][pdf_path.name] = sha256(pdf_path)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(audit, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--browser", help="Installed Chromium executable")
    args = parser.parse_args()
    asyncio.run(render(args.browser))
