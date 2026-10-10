"""Shared editable slide elements for the versioned Decider presentation."""

from pathlib import Path

from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

FONT = "Noto Sans CJK SC"
NAVY = "152536"
INK = "203448"
MUTED = "64788A"
TEAL = "087F72"
BLUE = "4669C4"
ORANGE = "BE721F"
GREEN = "BEDD5C"
WHITE = "FFFFFF"
BG = "F6F8FB"
LINE = "DCE4EA"
PALE_TEAL = "E8F5F1"
PALE_BLUE = "EDF1FC"
PALE_ORANGE = "FFF3E4"


def rgb(value):
    return RGBColor.from_string(value)


def rect(s, x, y, w, h, fill=WHITE, stroke=None, radius=False):
    sh = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if stroke:
        sh.line.color.rgb = rgb(stroke)
        sh.line.width = Pt(0.8)
    else:
        sh.line.fill.background()
    if radius:
        sh.adjustments[0] = 0.13
    sh._element.spPr.append(OxmlElement("a:effectLst"))
    return sh


def text(s, value, x, y, w, h, size=18, color=INK, bold=False, align=None, valign=None):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name = "TEXT: " + value.replace("\n", " ")[:65]
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.015)
    tf.margin_top = tf.margin_bottom = 0
    if valign:
        tf.vertical_anchor = valign
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.line_spacing = 1.12
        p.space_after = Pt(5)
        if align is not None:
            p.alignment = align
        for run in p.runs:
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
            props = run._r.get_or_add_rPr()
            for tag in ("a:ea", "a:cs"):
                child = OxmlElement(tag)
                child.set("typeface", FONT)
                props.append(child)
    return sh


def line(s, points, color=MUTED, width=1.5, arrow=False, dashed=False):
    for i, (start, end) in enumerate(zip(points, points[1:], strict=False)):
        sh = s.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(start[0]),
            Inches(start[1]),
            Inches(end[0]),
            Inches(end[1]),
        )
        sh.line.color.rgb = rgb(color)
        sh.line.width = Pt(width)
        ln = sh._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ln")
        if dashed:
            dash = OxmlElement("a:prstDash")
            dash.set("val", "dash")
            ln.append(dash)
        if arrow and i == len(points) - 2:
            tip = OxmlElement("a:tailEnd")
            tip.set("type", "triangle")
            tip.set("w", "sm")
            tip.set("len", "sm")
            ln.append(tip)


def label(s, value, x, y, w, fill=PALE_TEAL, color=TEAL, size=11):
    rect(s, x, y, w, 0.30, fill, radius=True)
    text(s, value, x + 0.04, y + 0.045, w - 0.08, 0.22, size, color, True, PP_ALIGN.CENTER)


def note(s, value, x=0.6, y=6.53, w=12.1, color=MUTED, fill=None):
    if fill:
        rect(s, x - 0.02, y - 0.10, w + 0.04, 0.52, fill, radius=True)
    text(s, value, x, y, w, 0.38, 12, color)


def bullets(s, items, x, y, w, size=18, color=INK, gap=0.67, dot=TEAL):
    for i, item in enumerate(items):
        rect(s, x, y + i * gap + 0.12, 0.06, 0.06, dot)
        text(s, item, x + 0.2, y + i * gap, w - 0.2, gap - 0.06, size, color)


def picture(s, path, x, y, w, h, frame=False, align="center"):
    path = Path(path)
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px = x if align == "left" else x + (w - pw) / 2
    py = y + (h - ph) / 2
    if frame:
        rect(s, px - 0.03, py - 0.03, pw + 0.06, ph + 0.06, WHITE, LINE, True)
    sh = s.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    sh.name = "IMAGE: " + path.name
    return (px, py, pw, ph)


def card(s, x, y, w, h, title, body, accent=TEAL, fill=WHITE, size=17):
    rect(s, x, y, w, h, fill, LINE, True)
    rect(s, x, y, 0.045, h, accent)
    text(s, title, x + 0.2, y + 0.18, w - 0.4, 0.5, 20, accent, True)
    body_y = 0.65 if h <= 1.5 else 0.72 if h < 2 else 0.83
    text(s, body, x + 0.2, y + body_y, w - 0.4, h - body_y - 0.10, size)


def step(s, number, title, body, x, y, w=2.3, color=TEAL):
    rect(s, x, y, w, 1.36, WHITE, LINE, True)
    text(s, f"{number:02}", x + 0.15, y + 0.14, 0.45, 0.38, 22, color, True)
    text(s, title, x + 0.70, y + 0.18, w - 0.82, 0.36, 17, INK, True)
    text(s, body, x + 0.15, y + 0.70, w - 0.3, 0.52, 14, MUTED)
