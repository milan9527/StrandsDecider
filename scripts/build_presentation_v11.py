"""Revise v10 with Jev conclusions and OD/one-year-RI crossover charts.

Uses saved receipts only. Keeps v10 unchanged and reuses its existing slide
objects, screenshots and formatting wherever the content does not change.
"""

import json
import math
import shutil
import statistics
from datetime import UTC, datetime
from pathlib import Path

import build_presentation_v10 as V
import matplotlib.pyplot as plt
import presentation_layout as L
from build_presentation_v5 import replace_text
from matplotlib import font_manager
from pptx import Presentation
from pptx.util import Inches

ROOT = V.ROOT
OUT = ROOT / "docs/presentations/v11"
BASE = next((ROOT / "docs/presentations/v10").glob("*.pptx"))
NAME = "Strands-Decider-v11-Jev场景结论与一年RI成本交叉图-32页-20261015"
RECEIPTS = {
    "skills": "validation/jev_skills/results/2026-10-12/receipts.jsonl",
    "browser_finance": "validation/browser_finance/runs/2026-10-13-live/receipts.jsonl",
    "routing": "validation/model_routing/runs/2026-10-13-live/receipts.jsonl",
}


def read_jsonl(relative):
    return [json.loads(line) for line in V.record(ROOT / relative).read_text().splitlines() if line.strip()]


def scenario_metrics():
    skills = read_jsonl(RECEIPTS["skills"])
    bf = read_jsonl(RECEIPTS["browser_finance"])
    routing = read_jsonl(RECEIPTS["routing"])
    grouped = {
        "skills": [dict(p, case_id=r["case_id"]) for r in skills for p in r["providers"].values()],
        "browser": [r for r in bf if r["experiment"] == "browser"],
        "finance_original": [r for r in bf if r["experiment"] == "finance" and r["variant"] == "original_order"],
        "policy_routing": routing,
    }
    result = {
        "method": "Recomputed from saved receipts; no new inference. Median RTT and nearest-rank P95."
                  " Finance uses original order only; policy routing includes both orders."
                  " Costs use Jev response usage input_tokens at USD0.042/1M, output free."
                  " Per-1000-request costs repeat the observed input mix, not 1000 independent tasks.",
        "sources": RECEIPTS,
        "groups": {},
        "price_usd_per_million_input_tokens": .042,
    }
    for name, rows in grouped.items():
        result["groups"][name] = {}
        for provider in ("decider", "jev"):
            selected = [r for r in rows if r["provider"] == provider]
            assert selected and all(r["status"] == "ok" for r in selected)
            times = sorted(r.get("elapsed_ms", r.get("client_elapsed_ms")) for r in selected)
            entry = {
                "calls": len(selected), "median_rtt_ms": statistics.median(times),
                "p95_rtt_ms": times[math.ceil(.95 * len(times)) - 1],
                "mean_rtt_ms": statistics.mean(times),
            }
            if provider == "jev":
                tokens = [r["response"]["usage"]["input_tokens"] for r in selected]
                entry.update({
                    "input_tokens": sum(tokens), "mean_input_tokens": statistics.mean(tokens),
                    "batch_list_price_usd": sum(tokens) * .042 / 1e6,
                    "usd_per_1000_api_calls": statistics.mean(tokens) * .042 / 1000,
                })
            result["groups"][name][provider] = entry
    b = V.data(V.BROWSER)
    wall_times = sorted(r["wall_ms"] for t in b["browser"]["jev"]["tasks"].values() for r in t["runs"])
    result["browser_jev_complete_tasks"] = {
        "completed": 10, "attempted": 10,
        "task_wall_median_ms": statistics.median(wall_times),
        "api_calls_per_task_mean": result["groups"]["browser"]["jev"]["calls"] / 10,
        "usd_per_1000_observed_tasks": result["groups"]["browser"]["jev"]["batch_list_price_usd"] * 100,
        "cost_scope": "Model API only; excludes Playwright, compute and tool/network costs.",
    }
    return result


def crossover_chart(name, unit_cost, od, ri, maximum, unit_label):
    """Export both a display image and a standalone vector chart."""
    crossings = {"on_demand_millions": od / unit_cost, "one_year_ri_millions": ri / unit_cost}

    def draw(ax):
        xs = [maximum * i / 200 for i in range(201)]
        ax.plot(xs, [x * unit_cost for x in xs], color="#" + V.PURPLE, lw=2.8, label="Jev · 按实测输入计费")
        ax.axhline(od, color="#" + L.ORANGE, lw=2.5, label=f"EC2 On-Demand · ${od:.2f}/月")
        ax.axhline(ri, color="#" + L.TEAL, lw=2.5, label=f"EC2 1 年 RI（无预付） · ${ri:.2f}/月")
        for x, y, color, dx, dy, label in [
            (crossings["one_year_ri_millions"], ri, L.TEAL, -56, -37, f'RI 交点\n{xformat(crossings["one_year_ri_millions"])}'),
            (crossings["on_demand_millions"], od, L.ORANGE, 12, 20, f'按需交点\n{xformat(crossings["on_demand_millions"])}'),
        ]:
            ax.vlines(x, 0, y, color="#" + color, alpha=.55, linestyle="--", lw=1)
            ax.scatter([x], [y], color="#" + color, s=57, zorder=5, edgecolors="white", linewidth=1)
            ax.annotate(label, (x, y), xytext=(dx, dy), textcoords="offset points", color="#" + color,
                        fontsize=11, fontweight="bold",
                        bbox={"boxstyle": "round,pad=.3", "fc": "#" + L.BG, "ec": "none", "alpha": .93})
        ax.set_xlim(0, maximum)
        ax.set_ylim(0, 1100)
        ax.set_xlabel(f"每月{unit_label}（百万）")
        ax.set_ylabel("美元 / 月")
        ax.legend(frameon=False, loc="upper left", fontsize=10.7)
        ax.grid(axis="y", alpha=.14)

    return V.chart(name, draw, 8.0, 4.1), crossings


def xformat(millions):
    return f"{millions * 100:,.0f} 万"


def update_text(slide, old, new):
    found = False
    for shape in slide.shapes:
        if shape.has_text_frame and old in shape.text:
            if "\n" in old:
                assert len(old.split("\n")) == len(new.split("\n"))
                for old_line, new_line in zip(old.split("\n"), new.split("\n"), strict=True):
                    replace_text(shape, old_line, new_line)
            else:
                replace_text(shape, old, new)
            assert new in shape.text
            found = True
    assert found, old


def main():
    (OUT / "sources").mkdir(parents=True, exist_ok=True)
    (OUT / "assets").mkdir(exist_ok=True)
    V.OUT, V.ASSETS, V.SOURCES = OUT, OUT / "assets", OUT / "sources"
    V.EVIDENCE.clear()
    V.SLIDES.clear()
    font = ROOT / ".deployment/presentation-fonts/NotoSansCJKSC-Regular.otf"
    if not font.exists():
        font.parent.mkdir(parents=True, exist_ok=True)
        V.TTFont("/usr/share/fonts/google-noto-cjk/NotoSansCJK-Regular.ttc", fontNumber=2).save(font)
    font_manager.fontManager.addfont(str(font))
    plt.rcParams.update({"font.family": "Noto Sans CJK SC", "font.size": 12})
    before_hash = V.digest(BASE)
    old_manifest = V.load(V.record(ROOT / "docs/presentations/v10/build-manifest.json"))
    V.record(ROOT / "validation/model_comparison/sources/tetris/docs/results/jev-decider-quality/summary.json")
    prs = Presentation(BASE)
    original = list(prs.slides)
    original_ids = list(prs.slides._sldIdLst)
    order, slides = [], []
    source_files = [
        "cost-calculation.json", "aws-g6.2xlarge-pricing.json",
        "aws-ri-overview.txt", "aws-ri-scope.txt", "jev-pricing.txt", "web-source-manifest.json",
    ]
    snapshots = []
    for name in source_files:
        source = V.record(ROOT / "docs/presentations/v10/sources" / name)
        destination = OUT / "sources" / name
        shutil.copy2(source, destination)
        snapshots.append({"file": str(destination.relative_to(ROOT)), "source": str(source.relative_to(ROOT)),
                          "sha256": V.digest(destination)})
    cost = V.load(OUT / "sources/cost-calculation.json")
    od = next(r for r in cost["rows"] if r["years"] == 0)
    ri = next(r for r in cost["rows"] if r["years"] == 1 and r["payment"] == "No Upfront")
    metrics = scenario_metrics()
    (OUT / "sources/jev-scenario-metrics.json").write_text(json.dumps(metrics, ensure_ascii=False, indent=2) + "\n")

    def keep(page, changed=False):
        order.append(original_ids[page - 1])
        meta = dict(old_manifest["slides"][page - 1])
        meta.update({"page": len(order), "source": "v10", "v10_page": page, "body_modified": changed})
        slides.append(meta)
        return original[page - 1]

    def new(section, title, subtitle, foot, notes):
        slide = V.base(prs, section, title, subtitle, foot, notes)
        order.append(prs.slides._sldIdLst[-1])
        meta = dict(V.SLIDES[-1])
        meta.update({"page": len(order), "source": "new"})
        slides.append(meta)
        return slide

    for page in range(1, 21):
        keep(page, page in (1, 6, 20))
    update_text(original[0], "v10 · 2026-10-15", "v11 · 2026-10-15")
    update_text(original[5], "独立 vLLM API 见第 24 页", "独立 vLLM API 见第 25 页")
    # Preserve the game results, but make its Jev recommendation and missing
    # cost evidence explicit without attributing one API RTT to an entire move.
    update_text(original[19], "问题不只在字段形式", "游戏结论：本局优先 Jev")
    update_text(original[19], "要检查候选之间的多指标比较，\n以及完成条件和顺序敏感性。",
                "自然语言版：251.1 vs 1414.5 ms/落点。\n适合继续验证动作选择；仅测一个种子。")
    update_text(original[19],
                "只有一个种子；走法不同会产生不同棋盘。此处为已有实验引用，不计入 466 次新目录调用。",
                "落点耗时为 Jev / Decider 均值，可含多次请求。游戏未保存 token usage，费用不作估算。")
    original[19].notes_slide.notes_text_frame.text += (
        "\nv11补充结论：本局自然语言Jev50块19行，Decider43块2行；每落点平均251.1/1414.5ms。"
        "这是动作级耗时，不能当作单次API RTT。原始游戏汇总未保存token usage，不套用支付输入估算成本。"
    )

    m = metrics["groups"]
    s = new("05  Jev 场景结论", "Jev：证据核验和多步操作表现更好，复杂规则仍需复核",
            "本批业务效果更好；延迟并非每项领先。按实际输入，每千次 Jev API 约 $0.02–0.08。",
            "API RTT 中位数：Jev 直连，Decider 经应用网关；金融仅首轮。费用按实测 usage，不含工具与运维。",
            "数据重新汇总自原始receipts，未新调用模型。Skills32输入65可评分问题；浏览器10个任务，"
            "Decider22次/Jev36次请求；金融只使用原始顺序各30请求；政策路由32题两种顺序各64请求。"
            "Jev直连API，Decider经现有应用网关访问官方原版服务；延迟是整个HTTP调用，不是纯GPU前向耗时。"
            "费用为Jev input_tokens × 0.042USD/百万，输出免费。千次费用是按本批输入配比放大，不是1000个独立任务。"
            "质量分别为Jev/Decider：Skills65/65 vs50/65、浏览器10/10 vs4/10、金融28/30 vs22/30、政策46/64 vs25/64。"
            "分类与局部选择两者接近；Jev在证据、多步完成中更好。金融共享额度误判、路由能力与预算漏判仍存在。"
            f"\n浏览器Jev每任务平均3.6调用，任务总耗时中位数{metrics['browser_jev_complete_tasks']['task_wall_median_ms']:.1f}ms，"
            f"每千条同配比任务的模型费约${metrics['browser_jev_complete_tasks']['usd_per_1000_observed_tasks']:.4f}，"
            "不含浏览器运行费；其他表格展示每千API请求，不能混成每千任务。"
            "\n原始数据与计算：sources/jev-scenario-metrics.json；validation/*/receipts.jsonl。")
    rows = []
    for key, label, quality, conclusion in [
        ("skills", "Skills", "65/65 · 50/65", "证据核验：Jev 优先"),
        ("browser", "真实浏览器", "10/10 · 4/10", "多步查找：Jev 优先"),
        ("finance_original", "金融初筛", "28/30 · 22/30", "初筛可用；金额用代码"),
        ("policy_routing", "政策路由", "46/64 · 25/64", "任务分流；约束用代码"),
    ]:
        row = m[key]
        rows.append([label, quality, f'{row["jev"]["median_rtt_ms"]:.1f}',
                     f'{row["decider"]["median_rtt_ms"]:.1f}',
                     f'${row["jev"]["usd_per_1000_api_calls"]:.3f}', conclusion])
    V.table(s, ["测试场景", "业务结果\nJev / Decider", "Jev ms", "Decider ms", "Jev $/千次", "适用结论"],
            rows, [1.7, 2.42, 1.33, 1.52, 1.60, 3.38], row_h=.70, size=14)
    V.callout(s, "Jev 可优先试点证据核验与多步操作；金额、资格、预算和最终完成条件，仍由程序独立验收。", y=6.04, size=14)

    for page in range(21, 28):
        keep(page)

    cost_notes = (
        "价格来自2026-10-15 AWS Price List：us-east-1/Linux/Shared/g6.2xlarge。"
        "On-Demand0.9776USD/h，730h=713.648USD/月；Standard RI 1年No Upfront0.63642USD/h，"
        "月均464.5866USD，全年承诺5575.0392USD，无预付。停止实例仍付RI小时承诺，区域RI不预留容量。"
        "三条曲线只比较On-Demand、该1年RI与Jev，不展示3年RI。"
        "Jev0.042USD/百万输入token，输出免费；用真实usage而非Decider tokenizer估计。"
        "EC2图线是单台全年常驻的计算费。仅在单台可满足负载且质量达标时比较交点；"
        "g6.2xlarge未进行本批容量实测。未计EBS、网络、后端、运维、冗余、税、下游模型和业务工具。"
        "费用是公开价估算，不是账单；没有购买RI。"
        "\n来源：sources/cost-calculation.json；sources/aws-g6.2xlarge-pricing.json。"
        "\nAWS RI：https://docs.aws.amazon.com/AWSEC2/latest/UserGuide/ec2-reserved-instances.html"
        "\nJev：https://typesafe.ai/blog/introducing-system-one-models-and-jev"
    )
    plots = {}
    for key, title, subtitle, unit_cost, unit_label, maximum, count_label, token_text in [
        ("payment", "支付单次：1 年 RI 的交点约 1684 万次/月",
         "三条曲线：EC2 On-Demand、1 年 Standard RI（无预付）、Jev 按输入用量计费。",
         cost["jev_payment_usd_per_million_decisions"], "支付决策量", 35, "次",
         "平均 656.95 输入 token / 次"),
        ("routing", "两步模型路由：1 年 RI 的交点约 394 万条/月",
         "一次完整路由包含“选流程 + 选型号”两次判断；输入更长，Jev 费用曲线也更陡。",
         cost["jev_route_usd_per_million_routes"], "完整两步路由量", 8, "条",
         "平均 2,805.75 输入 token / 条"),
    ]:
        plot, crossings = crossover_chart(f"cost-crossover-{key}", unit_cost,
                                           od["monthly_amortized_usd"], ri["monthly_amortized_usd"], maximum, unit_label)
        plots[key] = {
            "unit_cost_usd_per_million": unit_cost, "x_max_millions": maximum,
            "on_demand_monthly_usd": od["monthly_amortized_usd"],
            "one_year_ri_monthly_usd": ri["monthly_amortized_usd"],
            **crossings,
            "png": str(plot.relative_to(ROOT)), "png_sha256": V.digest(plot),
            "svg": str(plot.with_suffix(".svg").relative_to(ROOT)), "svg_sha256": V.digest(plot.with_suffix(".svg")),
        }
        s = new("08  On-Demand · 1 年 RI · Jev", title, subtitle,
                "g6.2xlarge · us-east-1 · 730h/月；只比计算/API费用，容量与质量另验。RI 闲置仍计费，其他资源费用另计。",
                cost_notes + f"\n{token_text}，Jev每百万{count_label}${unit_cost:.4f}。"
                f"费用交点：RI{crossings['one_year_ri_millions']:.6f}百万；按需{crossings['on_demand_millions']:.6f}百万。"
                "\n小于RI交点：Jev调用费低于两种EC2计算费；两交点之间：RI<Jev<On-Demand；"
                "高于按需交点：RI<On-Demand<Jev。各区间均假设一台实例可满足负载。")
        L.picture(s, plot, .65, 2.16, 7.95, 4.07)
        V.card(s, 8.82, 2.29, 3.73, 1.10, "1 年 RI 交点",
               f"约 {crossings['one_year_ri_millions'] * 100:,.0f} 万{count_label}/月", L.TEAL, size=19)
        V.card(s, 8.82, 3.64, 3.73, 1.10, "On-Demand 交点",
               f"约 {crossings['on_demand_millions'] * 100:,.0f} 万{count_label}/月", L.ORANGE, size=19)
        L.text(s, token_text, 8.94, 5.09, 3.48, .60, 14, L.MUTED)
        L.text(s, f"Jev：${unit_cost:.2f} / 百万{count_label}", 8.94, 5.82, 3.48, .37, 16, V.PURPLE, True)
        V.callout(s, "单台可承载时：低于 RI 交点，Jev 最省；两交点之间 RI 最省，Jev 仍比按需低；再增量，两种 EC2 都更低。",
                  y=6.20, size=12.8)

    s = new("09  各自适用场景", "业务质量先过关，再在低延迟、自托管与按量之间选择",
            "同一支付批次：Decider 102.2 ms / 18/20；Jev 416.5 ms / 20/20；GPT 1116.7 ms / 20/20。",
            "支付为平均后端耗时与参考一致次数；小样本。公共 vLLM 延迟与验证场景的 RTT 分开统计。",
            "Decider：本批分类/意图13/13、局部动作10/10，与Jev相同；适合候选清楚的高频局部判断，"
            "希望自托管和低单请求延迟的系统。vLLM单并发55.4ms均值，对比Jev126.5ms；"
            "这些性能是g6.xlarge，不能视作g6.2xlarge容量。GPU固定费与运维成本，流量稳定后再评估1年RI。"
            "Jev：本批Skills65/65、浏览器10/10；推荐证据核验、多步DOM操作、低或波动流量的候选决策。"
            "金融28/30、政策46/64仍出错，金额/资格/预算/最终完成交给代码验收。"
            "API按输入量计费，每千调用在本批0.0213至0.0800USD；不必自管GPU，但并非所有场景最低延迟。"
            "GPT：开放式回复、代码生成和复杂异常处理；此次项目只测候选选择，未对这些开放任务做质量对照。"
            "支付候选20/20但平均1116.7ms；高频且只需短ID的步骤要权衡额外延迟。"
            "GPT记录没有费用值，本版不伪造其账单或引用未经查询的价格。")
    V.panel(s, "Decider：高频选择",
            "优先：分类、工具选择、分流。\n\n低单次延迟、可控自托管。\n稳定流量再评估 1 年 RI。\n证据与完成判断需补强。",
            .76, w=3.79, h=3.85, icon_name="list-filter", color=L.TEAL, size=16)
    V.panel(s, "Jev：证据与多步操作",
            "优先：核验、文字浏览器操作。\n\n低或波动流量，按量付费。\n本批业务效果更好。\n金额与路由硬约束仍复核。",
            4.80, w=3.79, h=3.85, icon_name="clipboard-check", color=V.PURPLE, size=16)
    V.panel(s, "GPT：开放式复杂任务",
            "适合：回复、代码、复杂异常。\n\n本次支付参考一致 20/20。\n高频短选择要权衡延迟。\n开放任务质量需单独实测。",
            8.84, w=3.79, h=3.85, icon_name="brain-circuit", color=L.BLUE, size=16)
    keep(32, True)
    update_text(original[31],
                "PPT 内含逐页讲稿与来源；v9 保留，本版新增验证与 RI 后将 52 页重排为 32 页。",
                "PPT 内含逐页讲稿与来源；v11 聚焦 Jev 场景结论、按需与 1 年 RI 成本交叉图。")

    # Drop unused slide relationships, then apply the intended 32-page order.
    used = {node.rId for node in order}
    for node in list(prs.slides._sldIdLst):
        prs.slides._sldIdLst.remove(node)
        if node.rId not in used:
            prs.part.drop_rel(node.rId)
    for node in order:
        prs.slides._sldIdLst.append(node)
    assert len(prs.slides) == 32
    for i, (slide, meta) in enumerate(zip(prs.slides, slides, strict=True), 1):
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.endswith(" / 32") and shape.top > Inches(7):
                replace_text(shape, shape.text, f"{i:02} / 32")
        meta["notes"] = slide.notes_slide.notes_text_frame.text
    prs.core_properties.title = "Strands Decider × Jev × GPT：Jev场景结论与一年RI成本交叉图"
    prs.core_properties.subject = "v11 / 32 pages / On-Demand vs one-year No-Upfront Standard RI vs Jev"
    output = OUT / f"{NAME}.pptx"
    prs.save(output)
    assert V.digest(BASE) == before_hash
    (OUT / "sources/crossover-calculation.json").write_text(json.dumps({
        "pricing_snapshot": "2026-10-15", "ri_years": 1, "ri_payment": "No Upfront",
        "ri_offering_class": "standard", "plots": plots, "new_inference_calls": 0,
    }, ensure_ascii=False, indent=2) + "\n")
    manifest = {
        "created_at": datetime.now(UTC).isoformat(), "base_file": str(BASE.relative_to(ROOT)),
        "base_sha256": before_hash, "pptx_file": str(output.relative_to(ROOT)),
        "pptx_sha256": V.digest(output), "pages": 32, "slides": slides,
        "preserved_v10_pages_except_footer_and_marked_edits": [s["v10_page"] for s in slides if s["source"] == "v10" and not s["body_modified"]],
        "new_pages": [s["page"] for s in slides if s["source"] == "new"],
        "evidence": V.EVIDENCE, "snapshots": snapshots, "charts": plots,
        "new_model_calls": 0, "aws_resource_mutations": 0,
        "changes": [
            "Jev场景结论独立一页：业务效果、分场景延迟、真实usage成本、适用建议",
            "按需与1年无预付RI各一条水平线，Jev为用量斜线；支付与两步路由分别绘图",
            "删除3年RI展示及重复付款表，将成本由3页压缩为2页，整套仍32页",
            "游戏补充动作级耗时、适用结论和usage缺失说明",
            "最终结论分别明确Decider、Jev、GPT的适用场景及实测边界",
        ],
    }
    (OUT / "build-manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    notes = ["# v11 讲稿与来源", "", "32 页；v10 保留。本版重新统计已保存请求，没有新增模型调用。", ""]
    for meta in slides:
        notes.extend([f'## {meta["page"]:02} · {meta["title"]}', "", meta["notes"], ""])
    (OUT / "讲稿与来源.md").write_text("\n".join(notes))
    print(json.dumps({"pptx": str(output), "pages": 32, "new_pages": manifest["new_pages"], "scenario_metrics": metrics["groups"]},
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
