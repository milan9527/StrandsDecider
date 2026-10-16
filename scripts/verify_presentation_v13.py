"""Verify the reduced deck, real target results and billing arithmetic."""

import hashlib
import json
import math
from collections import Counter
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

import pymupdf as fitz
import verify_presentation_v11 as historical
from PIL import Image, ImageChops, ImageDraw
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/presentations/v13"
PROVIDERS = ("decider", "jev", "bedrock")


def load(path):
    return json.loads(path.read_text())


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normal(text):
    return "".join(text.split())


def body(slide):
    return [s.text for s in slide.shapes if s.has_text_frame and s.top < Inches(7.03)]


def images(slide):
    return [hashlib.sha256(s.image.blob).hexdigest() for s in slide.shapes if s.shape_type == 13]


def check_settlement():
    data = load(OUT / "sources/settlement-live.json")
    comparison = load(OUT / "sources/settlement-comparison.json")
    capture = load(OUT / "capture-manifest.json")
    assert digest(OUT / "sources/settlement-live.json") == capture["evidence_sha256"]
    assert data["complete"] and not data["page_errors"] and not data["gpu_writes"]
    assert Counter(r["phase"] for r in data["calls"]) == {
        "warmup": 6, "routing": 6, "target-decider": 1, "target-jev": 1, "target-bedrock": 1}
    exported = data["executed_export"]
    assert exported["router_model"] == "us.openai.gpt-6-sol" and exported["reasoning_effort"] == "low"
    sides = exported["cases"]["analysis"]["sides"]
    for stage in ("workflow", "model"):
        steps = [next(s for s in sides[p]["steps"] if s["stage"] == stage) for p in PROVIDERS]
        assert len({r["input_hash"] for r in steps}) == 1
        assert all(s["model_input"] == steps[0]["model_input"] for s in steps)
        state = json.loads(steps[0]["model_input"]["state"])
        assert state["request"] == data["task"]
        if stage == "workflow":
            assert set(state) == {"request"}
    for p in PROVIDERS:
        side = sides[p]
        summary = comparison["routes"][p]
        assert side["steps"][0]["choice"] == summary["workflow_choice"] == "reasoning"
        assert math.isclose(summary["routing_ms"], sum(s["timings"]["backend_total_ms"] for s in side["steps"]))
        for step in side["steps"]:
            call = next(r for r in data["calls"] if r["phase"] == "routing"
                        and r["request"]["provider"] == p and r["request"]["stage"] == step["stage"])
            assert call["status"] == 200 and call["response"]["source"] == "live"
            assert call["response"] == {k: v for k, v in step.items() if k != "client_total_ms"}
        target = next(r for r in data["calls"] if r["phase"] == f"target-{p}")
        assert target["request"]["task"]["text"] == data["task"]
        assert target["request"]["model_id"] == side["steps"][1]["selected_model"]["id"] == summary["target_model_id"]
        assert target["status"] == summary["target_status"]
        if p == "jev":
            assert target["status"] == 504 and not summary["target_complete"]
            assert target["response"]["detail"] == side["execution"]["error"] == summary["target_error"]
            assert summary["target_ms"] is None and summary["target_usage"] is None
        else:
            response = target["response"]
            assert target["status"] == 200 and response["complete"] and response["source"] == "live"
            assert response == side["execution"]
            assert response["text"] == summary["target_text"] and response["backend_total_ms"] == summary["target_ms"]
            for value in ("96", "24", "120", "11880"):
                assert value in response["text"].replace(",", "")
            fields = response["request_fields"]
            assert (fields.get("reasoning") or fields.get("output_config"))["effort"] == "max"
    ref = comparison["arithmetic_reference"]
    assert Decimal(str(ref["gross"])) * Decimal(str(ref["fee_rate"])) == Decimal(str(ref["fee"]))
    assert ref["fee"] + ref["hold"] == ref["difference"]
    assert ref["gross"] - ref["difference"] == ref["received"] and ref["unexplained"] == 0
    for image in capture["screenshots"]:
        assert digest(ROOT / image["file"]) == image["sha256"] and image["data_replaced"] is False
    return {"warmups_excluded": 6, "formal_routes": 6, "targets": 3, "target_successes": 2,
            "target_timeouts": 1, "same_stage_inputs_verified": True, "quotes_and_arithmetic": "passed",
            "timeout_preserved_without_fabricated_latency_or_usage": True, "gpu_mutations": 0}


def check_billing():
    costs = load(OUT / "sources/cost-calculation.json")
    b = load(OUT / "sources/billing-comparison.json")
    assert b["region"] == "us-east-1" and b["instance"] == "g6.2xlarge"
    assert b["jev_input_usd_per_million_tokens"] == costs["jev_input_per_million_usd"] == .042
    assert b["mean_jev_input_tokens_per_decision"] == costs["jev_payment_mean_input_tokens"] == 656.95
    assert math.isclose(b["jev_example_month_usd"],
                        b["example_monthly_payment_decisions"] * b["mean_jev_input_tokens_per_decision"] * .042 / 1e6)
    for label, hourly, monthly in (
        ("on_demand", "on_demand_hourly_usd", "on_demand_example_month_usd"),
        ("one_year_ri", "one_year_no_upfront_hourly_usd", "one_year_ri_example_month_usd"),
    ):
        assert math.isclose(b[hourly] * b["hours_per_month"], b[monthly]), label
    assert round(b["jev_example_month_usd"], 2) == 27.59
    assert round(b["on_demand_example_month_usd"], 2) == 713.65
    assert round(b["one_year_ri_example_month_usd"], 2) == 464.59
    return {"billing_formulas": "passed", "crossovers": historical.check_crossovers(),
            "validation_receipts": historical.check_metrics()}


def main():
    manifest = load(OUT / "build-manifest.json")
    pptx = ROOT / manifest["pptx_file"]
    prs = Presentation(pptx)
    base = Presentation(ROOT / manifest["base_file"])
    pdf = fitz.open(pptx.with_suffix(".pdf"))
    assert len(prs.slides) == len(pdf) == manifest["pages"] == 31
    assert digest(ROOT / manifest["base_file"]) == manifest["base_sha256"]
    assert digest(pptx) == manifest["pptx_sha256"]
    for item in manifest["evidence"]:
        assert digest(ROOT / item["path"]) == item["sha256"]
    for item in manifest["snapshots"]:
        assert digest(ROOT / item["source"]) == digest(ROOT / item["file"]) == item["sha256"]
    for item in manifest["images"]:
        with Image.open(ROOT / item["source"]) as original, Image.open(ROOT / item["output"]) as rendered:
            crop = original.crop(item["crop"]) if item["crop"] else original
            assert crop.size == rendered.size and ImageChops.difference(crop, rendered).getbbox() is None
        assert digest(ROOT / item["output"]) == item["sha256"]
    missing, bounds, preserved, thumbs = [], [], [], []
    previews = OUT / "previews"
    previews.mkdir(exist_ok=True)
    for i, (slide, page, meta) in enumerate(zip(prs.slides, pdf, manifest["slides"], strict=True), 1):
        rendered = normal(page.get_text())
        assert not any(s.has_text_frame and s.text.strip() and Inches(6.5) <= s.top < Inches(7.03)
                       for s in slide.shapes), i
        assert any(s.has_text_frame and s.text == f"{i:02} / 31" for s in slide.shapes)
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip() and normal(shape.text) not in rendered:
                missing.append({"page": i, "text": shape.text})
            if (min(shape.left, shape.top) < 0 or shape.left + shape.width > prs.slide_width + 12
                    or shape.top + shape.height > prs.slide_height + 12):
                bounds.append({"page": i, "shape": shape.name})
        assert slide.notes_slide.notes_text_frame.text == meta["notes"]
        if meta["source"] == "v12":
            original = base.slides[meta["v12_page"] - 1]
            assert images(slide) == images(original)
            if not meta["body_modified"]:
                assert body(slide) == body(original)
                preserved.append(i)
        path = previews / f"page-{i:02}.png"
        page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(path)
        with Image.open(path) as im:
            im.thumbnail((384, 216))
            tile = Image.new("RGB", (404, 248), "#e7edf2")
            tile.paste(im, ((404 - im.width) // 2, 6))
            ImageDraw.Draw(tile).text((12, 226), f"{i:02}", fill="#203448")
            thumbs.append(tile)
    for name, indexes in (("slides-overview", range(31)), ("updated-overview", [16, 17, 18, 19, 26])):
        cols = 3 if name == "updated-overview" else 4
        canvas = Image.new("RGB", (cols * 404, math.ceil(len(indexes) / cols) * 248), "#e7edf2")
        for n, index in enumerate(indexes):
            canvas.paste(thumbs[index], ((n % cols) * 404, (n // cols) * 248))
        canvas.save(OUT / f"{name}.png")
    task = load(ROOT / "docs/results/workflow-routing-2026-10-13/results.json")["cases"]["retrieval"]["task"]["text"]
    assert normal(task) in normal(pdf[16].get_text())
    report = {
        "audited_at": datetime.now(UTC).isoformat(), "pages": 31,
        "v12_file_unchanged": True, "preserved_body_pages": preserved,
        "all_retained_images_preserved": True, "lower_explanatory_rows": 0,
        "rendered_text_missing": missing, "shape_bounds_issues": bounds,
        "page17_actual_task_verified": True, "settlement_checks": check_settlement(),
        "billing_checks": check_billing(), "pptx_sha256": digest(pptx),
        "pdf_sha256": digest(pptx.with_suffix(".pdf")),
    }
    (OUT / "render-audit.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    assert not missing and not bounds


if __name__ == "__main__":
    main()
