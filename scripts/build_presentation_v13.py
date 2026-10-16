"""Condense validation, expose routing requests and compare billing models."""

import json
import shutil
from datetime import UTC, datetime
from decimal import Decimal

import build_presentation_v10 as V
import presentation_layout as L
from build_presentation_v5 import replace_text
from build_presentation_v11 import update_text
from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = V.ROOT
OUT = ROOT / "docs/presentations/v13"
BASE = next((ROOT / "docs/presentations/v12").glob("*.pptx"))
NAME = "Strands-Decider-v13-路由任务实测与计费模式对比-31页-20261015"
PROVIDERS = ("decider", "jev", "bedrock")
LABELS = ("Decider", "Jev", "GPT")
COLORS = (L.TEAL, V.PURPLE, L.BLUE)


def save(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def compact(slide):
    for shape in slide.shapes:
        if shape.has_text_frame and "\n" in shape.text:
            for p in shape.text_frame.paragraphs:
                p.space_after = Pt(0)
                p.line_spacing = 1.05


def request_box(slide, task):
    L.rect(slide, .68, 2.13, 11.98, .85, L.PALE_BLUE, radius=True)
    L.text(slide, "实际请求", .88, 2.39, 1.19, .32, 15, L.BLUE, True)
    L.text(slide, task, 2.13, 2.24, 10.28, .65, 15.5)


def main():
    (OUT / "assets").mkdir(exist_ok=True)
    V.OUT, V.ASSETS, V.SOURCES = OUT, OUT / "assets", OUT / "sources"
    V.EVIDENCE.clear()
    V.SLIDES.clear()
    V.IMAGES.clear()
    before_hash = V.digest(BASE)
    old = V.load(V.record(ROOT / "docs/presentations/v12/build-manifest.json"))
    live = V.load(V.record(OUT / "sources/settlement-live.json"))
    capture = V.load(V.record(OUT / "capture-manifest.json"))
    assert live["complete"]
    historical = V.data("docs/results/workflow-routing-2026-10-13/results.json")
    validation = V.data("validation/model_comparison/summary.json")
    scenario = V.data("docs/presentations/v11/sources/jev-scenario-metrics.json")
    V.record(ROOT / "validation/model_comparison/sources/tetris/docs/results/jev-decider-quality/summary.json")
    copies = []
    for name in ("cost-calculation.json", "aws-g6.2xlarge-pricing.json", "aws-ri-overview.txt",
                 "aws-ri-scope.txt", "jev-pricing.txt", "crossover-calculation.json"):
        source = V.record(ROOT / "docs/presentations/v11/sources" / name)
        destination = OUT / "sources" / name
        shutil.copy2(source, destination)
        copies.append({"source": str(source.relative_to(ROOT)), "file": str(destination.relative_to(ROOT)),
                       "sha256": V.digest(destination)})
    costs = V.load(OUT / "sources/cost-calculation.json")
    od = next(r for r in costs["rows"] if r["years"] == 0)
    ri = next(r for r in costs["rows"] if r["years"] == 1 and r["payment"] == "No Upfront")
    sides = live["executed_export"]["cases"]["analysis"]["sides"]
    settlement = {
        "task": live["task"], "captured_at": live["finished_at"],
        "arithmetic_reference": {"gross": 12000, "received": 11880, "fee_rate": .008,
                                 "fee": 96, "hold": 24, "difference": 120, "unexplained": 0},
        "target_configuration": "max effort from benchmark mapping; maxTokens=4096; response wait limit 40s.",
        "routes": {},
    }
    for provider in PROVIDERS:
        side = sides[provider]
        steps = side["steps"]
        target = side["execution"]
        call = next(r for r in live["calls"] if r["phase"] == f"target-{provider}")
        settlement["routes"][provider] = {
            "workflow_choice": steps[0]["choice"], "model_choice": steps[1]["choice"],
            "target_model": steps[1]["selected_model"]["name"],
            "target_model_id": steps[1]["selected_model"]["id"],
            "workflow_ms": steps[0]["timings"]["backend_total_ms"],
            "model_ms": steps[1]["timings"]["backend_total_ms"],
            "routing_ms": sum(r["timings"]["backend_total_ms"] for r in steps),
            "target_status": call["status"], "target_complete": target.get("complete", False),
            "target_ms": target.get("backend_total_ms"), "target_usage": target.get("usage"),
            "target_text": target.get("text"), "target_error": target.get("error"),
        }
    save(OUT / "sources/settlement-comparison.json", settlement)
    V.record(OUT / "sources/settlement-comparison.json")
    # The independently calculable business reference is not supplied as an answer to any model.
    assert Decimal("12000") * Decimal(".008") == Decimal("96")
    assert Decimal("12000") - Decimal("96") - Decimal("24") == Decimal("11880")
    quotes = {
        "decider": "总差额 120 USDC，由 96 USDC 手续费 + 24 USDC 暂扣 100% 构成，账务完全平衡，无异常缺口。",
        "bedrock": "差额完全由 96 USDC 手续费和 24 USDC 暂扣款构成，没有未解释的差额。",
    }
    for provider, quote in quotes.items():
        assert quote in sides[provider]["execution"]["text"].replace("**", "")
    prs = Presentation(BASE)
    originals, ids = list(prs.slides), list(prs.slides._sldIdLst)
    order, metas = [], []

    def keep(page, changed=False):
        order.append(ids[page - 1])
        meta = dict(old["slides"][page - 1])
        meta.update(page=len(order), source="v12", v12_page=page, body_modified=changed)
        metas.append(meta)
        return originals[page - 1]

    def new(section, title, subtitle, notes):
        slide = V.base(prs, section, title, subtitle, "", notes)
        order.append(prs.slides._sldIdLst[-1])
        meta = dict(V.SLIDES[-1])
        meta.update(page=len(order), source="new")
        metas.append(meta)
        return slide

    for page in range(1, 17):
        keep(page, page == 15)
    originals[5].notes_slide.notes_text_frame.text = (
        originals[5].notes_slide.notes_text_frame.text.replace("独立 vLLM API 见第 28 页", "独立 vLLM API 见第 23 页")
    )
    # Clarify the scope after adding actual target execution to this version.
    update_text(originals[14], "本次测到的范围", "本版对比范围")
    update_text(originals[14], "两次路由选择与耗时；尚未执行选中的下游模型。",
                "四任务比较路由；结算分析再实际执行所选模型。")
    originals[14].notes_slide.notes_text_frame.text += (
        "\nv13更新：四任务汇总仍为2026-10-13原批次；第18–19页新增2026-10-15"
        "结算分析配对路由与目标执行，分别保存，不与原批次平均耗时合并。")

    # 17: expose the exact user request, retaining the actual routing screenshot.
    retrieval = historical["cases"]["retrieval"]
    s = new("04  模型路由 · 实际请求",
            "中英文工单索引：同一任务，选中了不同的向量模型",
            "第一步仅提供任务原文；第二步追加已选流程、可用候选与同一份检索证据。",
            old["slides"][16]["notes"] + "\nv13增加的实际请求来自原始export，不追加偏好或预设答案。\n"
            + json.dumps({"task": retrieval["task"], "inputs": {
                p: [r["model_input"] for r in retrieval["sides"][p]["steps"]] for p in PROVIDERS
            }}, ensure_ascii=False, indent=2))
    request_box(s, retrieval["task"]["text"])
    shot = V.image("docs/presentations/v9/assets/routing-retrieval-results.png", "retrieval-routing.png")
    L.picture(s, shot, .68, 3.10, 11.98, 2.83, frame=True)
    V.callout(s, "实际提供的检索均分：English 30.65，Multilingual 59.23；Jev / GPT 的选择更符合任务与证据。",
              y=6.08, size=13.6)

    actual_notes = (
        "真实在线演示采集：2026-10-15，6次预热另存，6次正式路由+3次目标执行。"
        "三个决策器同阶段的model_input及input_hash相同，第一步全部reasoning。"
        "路由器GPT为us.openai.gpt-6-sol low，目标GPT为us.openai.gpt-6.1-sol max，角色不能混淆。"
        "Decider使用现有g6.xlarge原版服务；此次未修改GPU生命周期或部署。"
        "目标模型根据评测配置使用max档位，maxTokens4096，SDK响应等待上限40秒，未自动重试或切换。"
        "按Decider、Jev、GPT顺序各执行所选目标一次，因此不是吞吐/稳定延迟实验。"
        "Jev路由返回成功，504发生于Bedrock目标Opus5.5；超时没有目标usage或完整生成时长，不能填0费用或伪造精确耗时。"
        "结果是文本分析，未执行账务或资金操作。相同算术基准不代表评估了全部会计建议的正确性。"
        "\n完整HTTP请求、模型输入、响应、错误与截图见sources/settlement-live.json和capture-manifest.json。\n"
        + json.dumps(settlement, ensure_ascii=False, indent=2)
    )
    # 18: separate the routing decision from later model execution.
    s = new("04  模型路由 · 结算分析",
            "结算分析补测：三路都选“核算与分析”，目标模型不同",
            "实际路由结果：Decider → Opus 5；Jev → Opus 5.5；GPT → GPT-6.1 Sol。",
            actual_notes)
    request_box(s, live["task"])
    shot = V.image("docs/presentations/v13/screenshots/settlement-routing.png",
                   "settlement-routing-choices.png", (0, 0, 1950, 584))
    L.picture(s, shot, .68, 3.08, 11.98, 3.37, frame=True)

    # 19: actual generated answer excerpts, plus the unaltered timeout screen.
    s = new("04  模型路由 · 目标执行",
            "结算分析：两路核算正确，一路目标模型超时",
            "下方是所选目标模型的执行结果与耗时；均使用评测 max 档位，与前页路由耗时分开。",
            actual_notes + "\n本页引文去除Markdown强调后逐字匹配目标返回。"
            "Opus5扩展了暂扣的可能原因、会计处理和核实事项；这些扩展建议未逐项验证。"
            "GPT明确毛额基数假设，输出更聚焦。一次简单算术任务不能推出目标模型通用优劣。")
    V.callout(s, "验算基准：12,000 × 0.8% = 96；96 + 24 = 120；12,000 − 120 = 11,880 USDC。",
              y=2.16, size=15)
    route_names = ("Decider → Opus 5", "Jev → Opus 5.5", "GPT → GPT-6.1 Sol")
    for i, (provider, name, color) in enumerate(zip(PROVIDERS, route_names, COLORS, strict=True)):
        x = .68 + 4.07 * i
        V.card(s, x, 2.88, 3.84, 2.87, name, "", color)
        row = settlement["routes"][provider]
        if row["target_complete"]:
            L.text(s, f'{row["target_ms"] / 1000:.2f} s · 已完成', x + .18, 3.46, 3.48, .42, 21, color, True)
            L.text(s, "“" + quotes[provider] + "”", x + .18, 4.10, 3.48, 1.34, 15)
        else:
            L.text(s, "超时 · 未获得答复", x + .18, 3.46, 3.48, .42, 21, L.ORANGE, True)
            shot = V.image("docs/presentations/v13/screenshots/settlement-target-jev.png", "settlement-target-timeout.png")
            L.picture(s, shot, x + .18, 4.07, 3.48, 1.43, frame=False)
    V.callout(s, "本次差别出现在目标执行：路由选得快，还要验证所选模型能否在时限内完成任务。",
              y=6.02, size=15.2)

    # 20: replace eight scenario walkthrough slides with a final results table.
    groups = scenario["groups"]
    games = {g["provider"]: g for g in validation["tetris"]["primary_games"] if g["policy"] == "model"}
    s = new("05  其他场景 · 最终对比",
            "其他场景：只看最终业务效果、延迟与适用结论",
            "耗时列顺序为 Decider / Jev：前四项是 API RTT 中位数；游戏是每落点均值。此组未测 GPT。",
            "\n".join(old["slides"][i]["notes"] for i in range(17, 25))
            + "\nv13将原v12第18–25页压缩为一页最终对照，未新增这些实验调用。"
            "Skills业务正确50/65 vs65/65；其中分类意图13/13和局部动作10/10两者相同。"
            "Jev千API费用按实测input_tokens与$0.042/MTok计算，不是每千完整任务费用。"
            "游戏仅seed42自然语言版本，无token usage，不估算费用。")
    rows = []
    for key, label, d_score, j_score, conclusion in [
        ("skills", "Skills", "50/65", "65/65", "简单分类接近\n证据核验 Jev 更好"),
        ("browser", "浏览器任务", "4/10 完成", "10/10 完成", "多步操作 Jev 更好"),
        ("finance_original", "金融初筛", "22/30", "28/30", "Jev 更好\n金额仍需程序验算"),
        ("policy_routing", "政策约束选择", "25/64", "46/64", "Jev 更好\n两者仍会漏约束"),
    ]:
        m = groups[key]
        rows.append([label, d_score, j_score,
                     f'{m["decider"]["median_rtt_ms"]:.1f} / {m["jev"]["median_rtt_ms"]:.1f}',
                     f'${m["jev"]["usd_per_1000_api_calls"]:.3f}', conclusion])
    rows.append(["俄罗斯方块", f'{games["decider"]["pieces"]} 块 / {games["decider"]["lines"]} 行',
                 f'{games["jev"]["pieces"]} 块 / {games["jev"]["lines"]} 行',
                 f'{games["decider"]["stats"]["mean_decision_ms"]:.1f} / {games["jev"]["stats"]["mean_decision_ms"]:.1f}',
                 "未记录", "本局 Jev 更好\n仅测一个种子"])
    V.table(s, ["场景", "Decider 结果", "Jev 结果", "耗时 D / J · ms", "Jev $/千请求", "最终结论"],
            rows, [1.60, 1.95, 1.90, 2.00, 1.68, 2.82], row_h=.63, size=12.6)
    compact(s)
    V.callout(s, "分类意图 13/13、局部动作 10/10：两者都答对；复杂证据和多步完成，本批 Jev 表现更好。",
              y=6.18, size=14)

    for page in range(26, 32):
        keep(page)
    # 27: distinguish a variable API bill from provisioned GPU capacity.
    billing = {
        "price_snapshot": costs["queried_at"], "region": costs["region"], "instance": costs["instance"],
        "jev_input_usd_per_million_tokens": costs["jev_input_per_million_usd"],
        "jev_output_usd": 0, "on_demand_hourly_usd": od["recurring_hourly_usd"],
        "one_year_no_upfront_hourly_usd": ri["recurring_hourly_usd"],
        "hours_per_month": 730, "example_monthly_payment_decisions": 1000000,
        "mean_jev_input_tokens_per_decision": costs["jev_payment_mean_input_tokens"],
        "jev_example_month_usd": costs["jev_payment_usd_per_million_decisions"],
        "on_demand_example_month_usd": od["monthly_amortized_usd"],
        "one_year_ri_example_month_usd": ri["monthly_amortized_usd"],
        "assumption": "One instance can serve the load; compute/API only. Target-model execution, EBS, network and application costs extra.",
        "sources": copies,
    }
    save(OUT / "sources/billing-comparison.json", billing)
    V.record(OUT / "sources/billing-comparison.json")
    s = new("08  成本 · 计费模式",
            "Jev 按输入 token 计费；EC2 按运行容量计费",
            "Decider on g6.2xlarge · us-east-1 / Linux；下方按单台可承载、每月常驻 730 小时估算。",
            "公开价格快照2026-10-15。Jev每百万输入token0.042USD，结构化输出免费；"
            "实际输入token包含state、问题与候选，不能只按用户任务的字数估算。"
            "EC2 On-Demand实例运行时间产生计算费，没有请求但实例仍运行也计费；停止后计算费停止，磁盘等另计。"
            "一年Standard RI No Upfront是按承诺换折扣的计费安排，非另一种模型部署；"
            "停止实例仍付承诺，不预留区域容量。月均730小时，全年8760小时，年承诺5575.0392USD。"
            "EC2价格不是按token或请求计费；平均每次费用是计算费除以实际服务请求数，利用率影响分摊。"
            "算例使用此前实测Jev支付输入均值656.95token，不套用本次目标模型生成token。"
            "每月100万决策的Jev API费27.5919USD；EC2按需713.648USD、RI464.5866USD。"
            "尚未测量g6.2xlarge容量，单台承载仅为价格比较前提。下游生成模型与应用/EBS/网络/高可用/运维费用另计。"
            "\n" + json.dumps(billing, ensure_ascii=False, indent=2))
    rows = [
        ["计费对象", "实际输入 token\n结构化输出免费", "实例运行时长", "一年用量承诺\n无预付 Standard RI"],
        ["公开单价", "$0.042 / 百万输入 token", "$0.9776 / 小时", "$0.63642 / 小时"],
        ["没有请求时", "没有模型调用费", "实例仍运行就计费", "停止实例仍付承诺"],
        ["月费如何算", "总输入 token × 单价", "运行小时 × 小时单价", "承诺小时 × 折扣单价"],
        ["每月 100 万次\n支付决策算例", f'${billing["jev_example_month_usd"]:.2f}', f'${od["monthly_amortized_usd"]:.2f}',
         f'${ri["monthly_amortized_usd"]:.2f}'],
    ]
    V.table(s, ["项目", "Jev API", "Decider · On-Demand", "Decider · 1 年 RI"], rows,
            [2.00, 3.32, 3.32, 3.31], row_h=.63, size=13.5)
    compact(s)
    V.callout(s, "算例输入均值 656.95 token/次。EC2 利用率越高，单次分摊越低；存储、网络和目标模型费用另计。",
              y=6.18, size=13.3)
    keep(32)
    keep(33)
    keep(34)
    keep(35)
    originals[34].notes_slide.notes_text_frame.text += (
        "\nv13共31页：其他场景只展示最终对照；原第17页增加任务原文；"
        "新增结算分析真实路由与目标执行、API token与EC2容量计费对照。")

    used = {n.rId for n in order}
    for node in list(prs.slides._sldIdLst):
        prs.slides._sldIdLst.remove(node)
        if node.rId not in used:
            prs.part.drop_rel(node.rId)
    for node in order:
        prs.slides._sldIdLst.append(node)
    assert len(prs.slides) == 31
    for i, (slide, meta) in enumerate(zip(prs.slides, metas, strict=True), 1):
        for shape in list(slide.shapes):
            if shape.has_text_frame and Inches(6.5) <= shape.top < Inches(7.03):
                assert not shape.text.strip()  # No explanatory footer is reintroduced.
                shape._element.getparent().remove(shape._element)
            elif shape.has_text_frame and shape.top > Inches(7) and (
                shape.text.endswith(" / 35") or shape.text.endswith(" / 32")
            ):
                replace_text(shape, shape.text, f"{i:02} / 31")
        meta["notes"] = slide.notes_slide.notes_text_frame.text
    prs.core_properties.title = "Strands Decider × Jev × GPT：路由任务实测与计费模式对比"
    prs.core_properties.subject = "v13 / 31 pages / final scenario results, real routing tasks, API versus EC2 billing"
    output = OUT / f"{NAME}.pptx"
    prs.save(output)
    assert V.digest(BASE) == before_hash
    save(OUT / "build-manifest.json", {
        "created_at": datetime.now(UTC).isoformat(), "base_file": str(BASE.relative_to(ROOT)),
        "base_sha256": before_hash, "pptx_file": str(output.relative_to(ROOT)),
        "pptx_sha256": V.digest(output), "pages": 31, "slides": metas,
        "evidence": V.EVIDENCE, "images": V.IMAGES, "snapshots": copies,
        "new_pages": [m["page"] for m in metas if m["source"] == "new"],
        "new_warmup_calls": 6, "new_formal_routing_calls": 6, "new_target_attempts": 3,
        "target_successes": 2, "target_timeouts": 1,
        "aws_resource_mutations": 0, "external_financial_actions": 0,
        "lower_explanations_removed_from_all_slides": True, "capture": capture,
    })
    notes = ["# v13 讲稿与来源", "", "31 页。保留支付过程与模型路由；其他场景压缩为最终对照。", ""]
    for meta in metas:
        notes += [f'## {meta["page"]:02} · {meta["title"]}', "", meta["notes"], ""]
    (OUT / "讲稿与来源.md").write_text("\n".join(notes))
    (OUT / "结算分析实际请求与结果.md").write_text(
        "# 结算分析：实际请求与结果\n\n" + actual_notes + "\n\n## 三个目标接口的完整请求与返回\n\n```json\n"
        + json.dumps([r for r in live["calls"] if r["phase"].startswith("target-")], ensure_ascii=False, indent=2)
        + "\n```\n")
    print(json.dumps({"pptx": str(output), "pages": 31, "new_pages": [17, 18, 19, 20, 27],
                      "target_successes": 2, "target_timeouts": 1}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
