"""Build the concise v10 story from recorded experiments and current RI prices.

No model calls, deployments or purchases. Every data source and screenshot is
recorded in the manifest; the v9 presentation is kept unchanged.
"""

import hashlib
import json
import shutil
from datetime import UTC, datetime
from pathlib import Path

import cairosvg
import matplotlib
import presentation_layout as L
from fontTools.ttLib import TTFont
from matplotlib import font_manager
from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/presentations/v10"
ASSETS = OUT / "assets"
SOURCES = OUT / "sources"
NAME = "Strands-Decider-v10-场景验证-vLLM与RI成本对比-32页-20261015"
BASE = next((ROOT / "docs/presentations/v9").glob("*.pptx"))
REPO = "https://github.com/milan9527/StrandsDecider"
VREPO = "https://github.com/milan9527/StrandsDeciderVLLM"
DEMO = "https://d3i7ogfx5ml7dc.cloudfront.net"
API = "https://d1e9udi8st2p3g.cloudfront.net"
PURPLE = "8055A4"
COLORS = (L.TEAL, PURPLE, L.BLUE)
PROVIDERS = ("decider", "jev", "bedrock")
PAY = "docs/presentations/v8/live-observations.json"
ROUTE = "docs/results/workflow-routing-2026-10-13/summary.json"
VALID = "validation/model_comparison/summary.json"
BROWSER = "validation/browser_finance/runs/2026-10-13-live/scored-summary.json"
POLICY = "validation/model_routing/runs/2026-10-13-live/summary.json"
NATURAL = "validation/model_routing/runs/2026-10-13-natural/summary.json"
BENCH = "docs/presentations/v8/sources/jevbench-api-20261012.json"
VLLM = "docs/presentations/v9/sources/vllm"
SLIDES, EVIDENCE, IMAGES = [], [], []


def load(path):
    return json.loads(Path(path).read_text())


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def record(path):
    path = Path(path)
    rel = str(path.relative_to(ROOT))
    if not any(x["path"] == rel for x in EVIDENCE):
        EVIDENCE.append({"path": rel, "sha256": digest(path)})
    return path


def data(relative):
    return load(record(ROOT / relative))


def image(relative, name=None, crop=None):
    source = record(ROOT / relative)
    dest = ASSETS / (name or source.name)
    if crop:
        with Image.open(source) as im:
            assert 0 <= crop[0] < crop[2] <= im.width
            assert 0 <= crop[1] < crop[3] <= im.height
            im.crop(crop).save(dest)
    elif source.suffix == ".svg":
        dest = dest.with_suffix(".png")
        cairosvg.svg2png(url=str(source), write_to=str(dest), output_width=2300)
    else:
        shutil.copy2(source, dest)
    IMAGES.append({"source": relative, "output": str(dest.relative_to(ROOT)),
                   "crop": crop, "sha256": digest(dest)})
    return dest


def icon(s, name, x, y, color=L.TEAL, w=.4, aws=False):
    directory = "docs/diagrams/aws-icons" if aws else "docs/presentations/v3/icons"
    source = record(ROOT / directory / f"{name}.svg")
    dest = ASSETS / f"{'aws' if aws else 'icon'}-{name}-{color}.png"
    if not dest.exists():
        svg = source.read_text().replace("currentColor", "#" + color)
        cairosvg.svg2png(bytestring=svg.encode(), write_to=str(dest), output_width=220)
    L.picture(s, dest, x, y, w, w)


def base(prs, section, title, subtitle, foot, notes, old_pages=()):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = L.rgb(L.BG)
    n = len(prs.slides)
    L.rect(s, .65, .39, .13, .25, L.TEAL)
    L.text(s, section, .9, .35, 11.7, .35, 11.5, L.TEAL, True)
    L.text(s, title, .65, .89, 12.0, .56, 28, L.NAVY, True)
    L.text(s, subtitle, .65, 1.56, 12.0, .46, 14.5, L.MUTED)
    L.text(s, foot, .65, 6.67, 12.02, .29, 9.5, L.MUTED)
    L.line(s, [(.65, 7.03), (12.67, 7.03)], L.LINE, .8)
    L.text(s, "Strands Decider × Jev × GPT  |  业务验证 · 推理部署 · 成本", .65, 7.16, 10.8, .21, 9.2, L.MUTED)
    L.text(s, f"{n:02} / 32", 11.65, 7.12, 1.0, .29, 11, L.TEAL, True, PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    SLIDES.append({"page": n, "title": title, "section": section,
                   "v9_pages_reworked": list(old_pages), "notes": notes})
    return s


def table(s, headers, rows, widths, x=.68, y=2.2, row_h=.54, size=16, highlights=()):
    assert all(len(r) == len(widths) for r in rows)
    for i, row in enumerate([headers, *rows]):
        xx = x
        yy = y + i * row_h
        fill = L.NAVY if i == 0 else L.PALE_TEAL if i - 1 in highlights else (L.WHITE if i % 2 else "EEF2F6")
        for j, (value, width) in enumerate(zip(row, widths, strict=True)):
            L.rect(s, xx, yy, width - .018, row_h - .025, fill)
            color = L.WHITE if i == 0 else L.INK
            if i and j in (1, 2, 3) and headers[j] in ("Decider", "Jev", "GPT"):
                color = {"Decider": L.TEAL, "Jev": PURPLE, "GPT": L.BLUE}[headers[j]]
            L.text(s, str(value), xx + .12, yy + .105, width - .23, row_h - .16,
                   size if i else min(size, 15), color, i == 0)
            xx += width


def callout(s, value, y=6.03, color=L.TEAL, fill=L.PALE_TEAL, size=16):
    L.rect(s, .68, y, 11.98, .48, fill, radius=True)
    L.text(s, value, .84, y + .08, 11.66, .33, size, color, True)


def card(s, x, y, w, h, title, body, accent=L.TEAL, size=16):
    """Compact cards with separate title/body baselines even below one inch."""
    L.rect(s, x, y, w, h, L.WHITE, L.LINE, True)
    L.rect(s, x, y, .045, h, accent)
    L.text(s, title, x + .18, y + .10, w - .36, .34, 17.5, accent, True)
    L.text(s, body, x + .18, y + .53, w - .36, h - .56, size)


def panel(s, title, body, x, y=2.3, w=3.8, h=3.6, color=L.TEAL, icon_name=None, size=18):
    L.rect(s, x, y, w, h, L.WHITE, L.LINE, True)
    if icon_name:
        icon(s, icon_name, x + .23, y + .23, color, .42)
    L.text(s, title, x + .23, y + .91, w - .46, .45, 22, color, True)
    L.text(s, body, x + .23, y + 1.52, w - .46, h - 1.72, size)


def flow(s, items, y=2.65, x=.75, total=11.85, h=1.62):
    gap = .30
    w = (total - (len(items) - 1) * gap) / len(items)
    for i, (title, body, ico) in enumerate(items):
        xx = x + i * (w + gap)
        L.rect(s, xx, y, w, h, L.WHITE, L.LINE, True)
        icon(s, ico, xx + .15, y + .17, L.TEAL, .31)
        L.text(s, title, xx + .15, y + .64, w - .3, .32, 17.5, L.TEAL, True)
        L.text(s, body, xx + .15, y + 1.10, w - .3, h - 1.12, 13.5)
        if i + 1 < len(items):
            L.line(s, [(xx + w + .03, y + .78), (xx + w + gap - .03, y + .78)], L.MUTED, 1.5, True)


def chart(name, draw, width=10.5, height=3.8):
    fig, ax = plt.subplots(figsize=(width, height))
    fig.set_facecolor("#" + L.BG)
    ax.set_facecolor("#" + L.BG)
    draw(ax)
    ax.spines[["top", "right"]].set_visible(False)
    ax.spines[["left", "bottom"]].set_color("#" + L.LINE)
    ax.set_axisbelow(True)
    fig.tight_layout(pad=.9)
    for ext in ("png", "svg"):
        fig.savefig(ASSETS / f"{name}.{ext}", dpi=190, facecolor=fig.get_facecolor())
    plt.close(fig)
    return ASSETS / f"{name}.png"


def hyperlink(s, label, url, x, y, w, size=15):
    sh = L.text(s, label, x, y, w, .36, size, L.BLUE)
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            r.hyperlink.address = url


def main():
    ASSETS.mkdir(parents=True, exist_ok=True)
    font = ROOT / ".deployment/presentation-fonts/NotoSansCJKSC-Regular.otf"
    if not font.exists():
        font.parent.mkdir(parents=True, exist_ok=True)
        TTFont("/usr/share/fonts/google-noto-cjk/NotoSansCJK-Regular.ttc", fontNumber=2).save(font)
    font_manager.fontManager.addfont(str(font))
    plt.rcParams.update({"font.family": "Noto Sans CJK SC", "font.size": 12})
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = "Strands Decider：从决策原理到场景验证、vLLM与RI成本"
    prs.core_properties.subject = "Recorded Decider/Jev/GPT comparisons; EC2 Standard RI list prices"
    prs.core_properties.author = "milan9527"
    base_hash = digest(BASE)
    pay, route, valid, bf = data(PAY), data(ROUTE), data(VALID), data(BROWSER)
    policy, natural, bench = data(POLICY), data(NATURAL), data(BENCH)
    sweep = data(f"{VLLM}/sweep-1-128.json")["cells"]
    batch = data(f"{VLLM}/batch-summary.json")
    jev = data(f"{VLLM}/jev-summary.json")
    cost = data("docs/presentations/v10/sources/cost-calculation.json")
    for f in SOURCES.iterdir():
        if f.is_file():
            record(f)
    for rel in ("validation/README.md", "validation/publication-audit.json",
                "validation/browser_finance/evaluation-erratum.json",
                "docs/MODEL_ROUTING_COMPARISON.md", f"{VLLM}/provenance.json",
                f"{VLLM}/original-benchmark.json", f"{VLLM}/vllm-recheck.json",
                f"{VLLM}/requests-usage.json", "docs/presentations/v9/sources/cost-calculation.json"):
        record(ROOT / rel)

    payment_notes = (
        "支付数据：2026-10-12 presentation-v8-warm，3次预热另存，4场景×5节点×3模型=60次正式调用。"
        "Decider v19官方服务在g6.xlarge；Jev 1.13.0经应用的Lambda适配；Bedrock GPT-6.1 Sol low。"
        "计时为后端请求总耗时，预热不计入。与演示参考一致只是这20个节点的参考匹配，不是通用准确率。"
        "初始四场景都缺少资金完成条件，三模型各为0/4资金业务完成。执行保护不改写模型原始成绩。"
        f"\n证据：{PAY}\n截图保持实际输入输出，未修改内容。"
    )
    routing_notes = (
        "2026-10-13 workflow-model-v1：四个任务，每模型先选流程、再选型号，共24次正式请求，另有6次预热。"
        "用户只输入任务，服务提供区域模型目录与第三方能力资料；没有用户指定优先级，也不强制榜单第一。"
        "目录124项，按流程形成22个文本或3个embedding候选。三侧输入指纹相同。"
        "Decider v19原版g6.xlarge、Jev 1.13.0、Bedrock GPT-6 Sol low。"
        "表中时间是两次backend_total_ms相加；没有执行所选下游模型。"
        "首次GPT分析40秒超时保存在initial-run，正式表使用随后完整批次；不能据此声称无失败或稳定SLA。"
        f"\n证据：{ROUTE}\n报告：{REPO}/blob/main/docs/MODEL_ROUTING_COMPARISON.md"
    )
    validation_notes = (
        "2026-10-12—13保存的真实调用，本次仅复核汇总，无新增模型调用。"
        "直接实验为Skills64、浏览器+金融178、两批政策路由224，总计466次。"
        "模型为Decider 2B hobson v19原版服务和Jev1.13.0；这批没有GPT，也没有vLLM。"
        "问题、任务、请求的计数单位不同，不汇总跨场景准确率；候选顺序反转和重复运行不算新增独立案例。"
        "Jev在独立验证中直连，支付经过Lambda；计时边界和输入不同，不能混算加速倍数。"
        f"\n报告：{REPO}/blob/main/validation/README.md\n原始索引：{VALID}"
    )
    perf_notes = (
        "性能证据来自独立StrandsDeciderVLLM项目；v9/sources/vllm/provenance.json保存来源提交与哈希。"
        "v19、vLLM0.31.0，BF16 torso/FP32 Pointer Head，单L4。"
        "原版对齐比较在g6.4xlarge同机，容器限制4vCPU/14GiB，回环HTTP；后续独立公网压测在g6.xlarge。"
        "20个固定支付输入，每次一个choice，固定worker收到完整响应后发下一个请求；失败不重试，预热另计。"
        "吞吐=成功数/正式总时长，失败占用总时长；分位数仅统计成功响应。"
        "相同输入反复调用只测性能，不扩充独立质量样本。数值对齐是对原版，不证明选择正确。"
        f"\n性能报告：{VREPO}/blob/main/docs/VLLM_PERFORMANCE.zh-CN.md"
        f"\nJev对比：{VREPO}/blob/main/docs/JEV_CONCURRENCY.zh-CN.md"
    )
    cost_notes = (
        "2026-10-15 AWS Price List实时查询，us-east-1/Linux/Shared/g6.2xlarge，"
        "On-Demand0.9776USD/h；展示Standard RI的1/3年六种付款方式，非Savings Plans或Convertible RI。"
        "折算月费=预付额/(12×年数)+小时费×730，AWS一年365天=8760h，三年1095天=26280h。"
        "全预付折算月费不是每月账单；无预付也承诺整个期限、停止实例仍付费。"
        "区域RI只提供匹配用量的计费优惠，不保证容量。未购买RI。"
        "Jev官网0.042USD/百万输入token，输出免费；实际支付10240次usage平均656.95输入token；"
        "四任务路由平均2805.75token/完整两步。以这些实测用量估算成本，不是账单或未来价格保证。"
        "仅EC2计算费与Jev调用费，未计EBS、网络、应用后端、运维、冗余、税、所选下游模型执行。"
        "g6.2xlarge仅为指定成本场景，未实测其容量；不把g6.xlarge吞吐当成它的实测。"
        "\nAWS定价源：sources/aws-g6.2xlarge-pricing.json\n公式与交点：sources/cost-calculation.json"
        "\nRI条款：https://docs.aws.amazon.com/AWSEC2/latest/UserGuide/ec2-reserved-instances.html"
        "\n范围：https://docs.aws.amazon.com/AWSEC2/latest/UserGuide/reserved-instances-scope.html"
        "\nJev：https://typesafe.ai/blog/introducing-system-one-models-and-jev"
    )

    # 1–5: introduce one coherent evaluation frame, mechanism and external context.
    s = base(prs, "01  从软件中的判断开始", "让决策模型接得住真实业务",
             "Strands Decider、Jev 与 GPT：从原理、业务验证，到 vLLM 部署与成本选择",
             "v10 · 2026-10-15  |  实验来自保存的真实请求；价格按当日 AWS / Jev 公开口径。",
             "基于v9的52页重排为32页。合并机制、支付重复截图、路由解释、优化和选型结论；新增场景验证与RI。", (1,))
    for i, (title, body, ico) in enumerate([
        ("选得对", "业务事实 → 候选动作\n核对证据和完成条件", "list-filter"),
        ("跑得动", "单次延迟 → 并发吞吐\n看完整请求和尾延迟", "timer"),
        ("算得清", "自托管 → 按量 API\n加入 EC2 RI 长期承诺", "banknote"),
    ]):
        panel(s, title, body, .72 + i * 4.05, w=3.79, icon_name=ico, color=COLORS[i])

    s = base(prs, "01  问题与方法", "业务流程里，很多步骤只需要一个选择",
             "客服分流、证据核验、工具选择、模型路由：输出很短，判断会高频重复。",
             "适用前提：可定义候选、可取得事实、可验证执行结果。开放式写作仍需要生成模型。",
             "合并原适用场景和评估边界。选择的语义正确性、执行完成状态、耗时费用分别记录。", (2, 8, 9, 51))
    flow(s, [("业务状态", "订单、页面或任务", "receipt-text"), ("候选判断", "选择 / 真值 / 评分", "brain-circuit"),
             ("工具执行", "调用受控业务动作", "workflow"), ("结果验收", "事实是否真正改变", "clipboard-check")])
    for x, title, body, color in [(.75, "选择", "与预先定义的参考比较", L.TEAL),
                                 (4.8, "结果", "查看页面、账本或产物", PURPLE),
                                 (8.85, "性能", "记录延迟、吞吐和成本", L.BLUE)]:
        L.text(s, title, x, 4.95, 3.7, .4, 21, color, True)
        L.text(s, body, x, 5.54, 3.7, .4, 17)

    s = base(prs, "02  模型原理", "Decider：读完上下文，直接给候选评分",
             "公开 v19：Qwen3.5-2B 底座 + LoRA + Pointer Head；不逐 token 生成答案。",
             "来源：Strands 官方介绍与 v19 实现。API 的 probability / confidence 不是业务正确率。",
             "输入按官方模板编码，模型得到隐藏表示，Pointer Head比较候选与答案位置后评分，"
             "用校准参数转成输出。去除LM Head的生成路径；约1M参数的头不等于整个模型只有1M。"
             "\nhttps://strandsagents.com/blog/introducing-strands-decider/", (3, 5))
    flow(s, [("提交业务事实", "state + questions", "receipt-text"),
             ("单次前向计算", "2B 底座 + rank-16 LoRA", "layers"),
             ("Pointer Head", "对当前候选评分", "list-filter"),
             ("类型化结果", "choice / noul / score", "check-check")], h=1.73)
    card(s, .75, 4.88, 5.75, 1.15, "共享上下文", "同一请求可含多个独立判断，减少重复输入。", L.TEAL, size=15)
    card(s, 6.75, 4.88, 5.75, 1.15, "依赖执行结果的下一步", "先执行并更新状态，再发起新判断。", L.BLUE, size=15)

    s = base(prs, "02  模型对照", "相同候选问题，可以通过不同模型回答",
             "应用先定义问题和可执行选项；模型负责选择，执行器负责真正改变业务状态。",
             "BERT 的 NLI 用法也可比较文本候选；此处比较典型分类头，没有进行 BERT 实测。",
             "BERT固定标签分类头适合标签稳定、有领域标注的数据。新增类别通常需调整分类头和训练。"
             "Decider接受请求中的动态候选；Jev提供托管决策接口，内部结构未公开，不能推断与Decider相同。"
             "GPT按应用要求生成候选ID，即便输出短仍使用生成式推理路径。", (4, 6, 7))
    table(s, ["维度", "BERT 分类器", "Decider / Jev", "GPT"],
          [["典型输出", "固定标签", "候选 ID、真值、评分", "文本或结构化 JSON"],
           ["新增任务", "标签与领域训练适配", "提交事实、问题、候选", "提交指令与上下文"],
           ["实现", "编码器 + 分类头", "Decider 评分头；Jev API", "生成式模型"],
           ["适合", "稳定标签的专用分类", "高频、有限动作的判断", "开放式分析、写作、代码"]],
          [1.35, 3.10, 4.10, 3.40], row_h=.70, size=16)
    callout(s, "同一个 approved，可能来自分类头、候选评分或文本生成；都需要核对业务条件。")

    s = base(prs, "02  外部参考", "JevBench：同时看能力、校准、速度与成本",
             "公开 API 榜单节选；综合分低，不代表回答正确率同样低。",
             "榜单 v1.7.46 快照：2026-10-12。第三方公布结果，本项目未重跑，也不是 AWS 官方评分。",
             "合并原三页榜单介绍。综合分基于能力、校准、速度、成本四轴的调和平均与低分惩罚，"
             "不是简单准确率。这里无Strands v19和GPT-6.1 Sol同口径条目，不插入自测排名。"
             f"\n原始快照：{BENCH}\n评测链接：https://benchmarkheaven.com/jev-models/api", (10, 11, 12))
    selected = []
    for key, label in [("jev-1.13.0", "Jev 1.13"), ("gpt-6-luna-low", "GPT-6 Luna low"),
                       ("gemini-3.1-flash-lite", "Gemini 3.1 Flash Lite"),
                       ("deepseek-flash", "DeepSeek V4.1 Flash"), ("qwen3.8-27b", "Qwen3.8 27B")]:
        matches = [r for r in bench["rows"] if key == r["key"]]
        assert matches, key
        r = matches[0]
        selected.append([label, f'{r["jevbench_score"]:.1f}', f'{r["axes"]["intelligence"]:.1f}',
                         f'{r["speed"]["p50_s_raw"]:.3f} s'])
    table(s, ["模型", "综合分", "能力轴分", "延迟 P50"], selected, [4.6, 2.3, 2.5, 2.55], row_h=.52)
    L.text(s, "例如：Qwen 的能力轴高，但速度等轴拉低综合分。选型时应分别查看各轴。", .8, 5.57, 11.8, .4, 17)
    hyperlink(s, "评测与方法：benchmarkheaven.com/jev-models/api", "https://benchmarkheaven.com/jev-models/api", .8, 6.05, 11.6)

    # 6–13: retain the concrete application and its genuine screenshots.
    s = base(prs, "03  演示架构", "同一业务后端，连接三个决策模型",
             "前端 S3 + CloudFront OAC；用户名登录；ECS 执行支付、路由与结果记录。",
             "主演示使用原版 Decider g6.xlarge 常驻服务；独立 vLLM API 见第 24 页。",
             "前后端分离，无Cognito，S3 Block Public Access。后端直接连接模型与DynamoDB。"
             "图中Lambda是Jev实际调用路径；没有额外画服务runtime或VPC endpoint节点。", (13,))
    L.rect(s, .76, 2.2, 11.83, 4.10, None, L.LINE, True)
    L.text(s, "AWS Cloud", .95, 2.30, 8, .3, 13, L.MUTED, True)
    L.rect(s, 3.43, 2.72, 4.98, 3.42, None, PURPLE)
    L.text(s, "VPC · us-east-1", 3.62, 2.87, 4.4, .27, 11, PURPLE, True)
    L.rect(s, 5.59, 3.20, 2.63, 2.82, "F0F8F7", "70ABA5")
    L.text(s, "私有子网", 5.77, 3.31, 2.2, .25, 10.5, L.TEAL, True)
    def aws_node(name, label, x, y, w=1.55):
        L.rect(s, x, y, w, .99, L.WHITE, L.LINE, True)
        icon(s, name, x + .1, y + .14, w=.37, aws=True)
        L.text(s, label, x + .54, y + .18, w - .64, .69, 13)
    aws_node("cloudfront", "CloudFront", 1.08, 3.65, 2.1)
    aws_node("s3", "S3 私有桶\nOAC 访问", 1.08, 5.0, 2.1)
    aws_node("alb", "ALB", 3.65, 3.65, 1.65)
    aws_node("fargate", "ECS Fargate\n业务后端", 5.73, 3.65, 2.35)
    aws_node("ec2", "EC2 g6.xlarge\nDecider v19", 5.73, 5.0, 2.35)
    aws_node("dynamodb", "DynamoDB · 状态 / 记录", 9.0, 2.97, 3.14)
    aws_node("bedrock", "Amazon Bedrock · GPT", 9.0, 4.11, 3.14)
    aws_node("lambda", "Lambda → Jev 公网 API", 9.0, 5.25, 3.14)
    for points in [[(2.12,4.66),(2.12,4.97)],[(3.2,4.14),(3.62,4.14)],
                   [(5.32,4.14),(5.70,4.14)],[(6.90,4.66),(6.90,4.97)],
                   [(8.10,4.14),(8.62,4.14),(8.62,3.46),(8.96,3.46)],
                   [(8.10,4.14),(8.62,4.14),(8.62,4.60),(8.96,4.60)],
                   [(8.10,4.14),(8.62,4.14),(8.62,5.74),(8.96,5.74)]]:
        L.line(s, points, L.MUTED, 1.3, True)
    L.text(s, "浏览器 → HTTPS", 1.08, 3.03, 2.1, .3, 12.5, L.TEAL, True)

    s = base(prs, "03  支付流程", "四个场景，同样经过五个业务节点",
             "提交业务事实、问题与候选项；三个模型使用相同输入，分别推进各自的业务状态。",
             "模型建议、实际执行动作、对话输入输出和最终业务结果分别记录。",
             payment_notes, (14, 15))
    flow(s, [("客服受理", "识别用户诉求", "headset"), ("交易核验", "确认数 / 金额", "search"),
             ("风控审批", "风险 / 退款地址", "shield-check"), ("交易处置", "等待 / 补款 / 入账", "wallet"),
             ("客服回复", "跟进 / 结案", "receipt-text")], total=11.85, h=1.63)
    for i, (title, body) in enumerate([("USDC 未到账", "129 USDC；确认数 3/12"),
                                     ("USDT 少付", "应付 100；实收 99.80"),
                                     ("稳定币退款", "退款地址尚未验证"),
                                     ("商户结算", "已超过 60 分钟 SLA")]):
        card(s, .76 + i * 3.01, 4.85, 2.83, 1.24, title, body, COLORS[i % 3], size=13)

    totals = {p: 0 for p in PROVIDERS}
    matches = {p: 0 for p in PROVIDERS}
    payment_rows = []
    for case, name in [("pending", "USDC 未到账"), ("underpaid", "USDT 少付"), ("refund", "稳定币退款"), ("settlement", "商户结算")]:
        row = [name]
        for p in PROVIDERS:
            results = list(pay["cases"][case]["results"][p].values())
            ms = sum(r["timings"]["backend_total_ms"] for r in results)
            correct = sum(r["matches_reference"] for r in results)
            totals[p] += ms
            matches[p] += correct
            row.append(f"{ms / 1000:.3f} s  ·  {correct}/5")
        payment_rows.append(row)
    assert [matches[p] for p in PROVIDERS] == [18, 20, 20]
    s = base(prs, "03  支付实测", "Decider 响应最快，两个节点仍选错动作",
             "每格为五个节点的请求耗时之和 · 与演示参考一致次数。",
             "2026-10-12 预热后批次，60 次正式调用。GPT 为 GPT-6.1 Sol low；不是最终资金完成率。",
             payment_notes, (16, 17, 18, 21, 24, 50))
    table(s, ["场景", "Decider", "Jev", "GPT"], payment_rows, [2.7, 3.08, 3.08, 3.09], row_h=.62)
    for i, p in enumerate(PROVIDERS):
        x = .8 + i * 4.05
        L.text(s, f"{totals[p] / 20:.1f} ms", x, 5.44, 3.6, .57, 29, COLORS[i], True)
        L.text(s, f"单节点平均 · {matches[p]}/20 与参考一致", x, 6.09, 3.8, .3, 13)

    s = base(prs, "03  观察选择与执行", "退款：模型批准了，执行器仍要求验证地址",
             "真实界面保留三模型的原始输出；保护逻辑修正执行，不替模型“改成绩”。",
             "实际截图：refund-decision.png；少付与退款都计作 Decider 选择偏差。",
             payment_notes + "\n退款审批：address_verified=false，Decider approved，Jev/GPT verify_destination。"
             "少付：100应付、99.80实收，确认20/12已经足够，Decider仍选await_confirmations。", (19, 20, 22))
    shot = image("docs/presentations/v8/screenshots/refund-decision.png")
    L.picture(s, shot, .76, 2.12, 11.82, 3.38, frame=True)
    card(s, .76, 5.65, 5.75, .89, "退款偏差", "地址未核验，却选择 approved。", L.ORANGE, size=13)
    card(s, 6.75, 5.65, 5.80, .89, "少付偏差", "确认数已够，却等待更多确认。", L.ORANGE, size=13)

    s = base(prs, "03  业务闭环", "闭环由新事实推动：确认到账，再入账和结案",
             "USDC 沙盒：模拟器推进确认数，执行器更新订单；模型读取执行结果后重新选择。",
             "保存的闭环批次：各模型 6 次决策 + 1 次模拟推进；不是链上真实资金交易。",
             payment_notes + "\n闭环来自v8保留的v7 pending记录，不将其耗时加入本页支付预热批次。"
             "模拟确认3→12，重新核验，入账129USDC并关闭工单；Decider首次选择偏差保留。", (23, 25))
    flow(s, [("等待确认", "当前 3 / 12", "timer"), ("推进模拟交易", "确认数变为 12 / 12", "repeat-2"),
             ("重新核验", "金额足额且已确认", "clipboard-check"), ("入账并结案", "129 USDC 已入账", "circle-check")])
    L.picture(s, image("docs/presentations/v8/assets/pending-loop-status.png"), .8, 4.62, 11.75, 1.36, frame=True)
    callout(s, "五步都返回答案 ≠ 业务已完成。以入账记录、订单状态与工单状态验收。")

    s = base(prs, "04  主演示：模型路由", "用户只给任务，决策模型先选流程、再选型号",
             "服务端补充可用模型与能力资料；三个决策模型比较同一份候选证据。",
             "区域目录快照 124 项；文本流程 22 个候选，向量流程 3 个候选。",
             routing_notes, (26, 27))
    flow(s, [("任务原文", "客服 / 代码 / 分析 / 索引", "receipt-text"),
             ("第 1 次判断", "直接答复、代码、分析、向量", "route"),
             ("补充模型资料", "区域可用性 + 第三方排名", "layers"),
             ("第 2 次判断", "从流程候选中选型号", "list-filter")], h=1.77)
    card(s, .75, 5.02, 5.76, 1.06, "候选事实", "LiveBench 等文本证据；MTEB / MIRACL 向量证据。", L.TEAL, size=13.5)
    card(s, 6.75, 5.02, 5.77, 1.06, "本次测到的范围", "两次路由选择与耗时；尚未执行选中的下游模型。", L.BLUE, size=13.5)

    s = base(prs, "04  主演示：路由结果", "流程判断一致，具体型号选择出现差异",
             "每格为所选型号与两次决策总耗时；GPT 决策器为 GPT-6 Sol low。",
             "2026-10-13 · 四任务各测一次；型号选择差异不能直接当作下游任务准确率。",
             routing_notes, (28, 29, 30))
    model_short = {"Claude Opus 5": "Opus 5", "Claude Opus 5.5": "Opus 5.5",
                   "Claude Sonnet 5.5": "Sonnet 5.5", "GPT-6 Luna": "GPT-6 Luna",
                   "GPT-6.1 Sol": "GPT-6.1 Sol", "Embed English": "Embed English",
                   "Embed Multilingual": "Embed Multilingual"}
    rows = []
    for case, name in [("support", "高频客服"), ("code", "交易代码"), ("analysis", "结算分析"), ("retrieval", "中英文索引")]:
        row = [name]
        for p in PROVIDERS:
            r = route["cases"][case][p]
            row.append(f'{model_short.get(r["model"], r["model"])}\n{r["total_ms"]:.1f} ms')
        rows.append(row)
    table(s, ["任务", "Decider", "Jev", "GPT"], rows, [1.85, 3.36, 3.36, 3.38], row_h=.76, size=15.5)
    callout(s, "代码任务：Jev / GPT 选 Sonnet 5.5；中英文索引：Decider 选了英文向量模型。", y=6.12, size=14)

    s = base(prs, "04  主演示：证据使用", "拿到了同一上下文，仍可能没有选对证据",
             "中英文工单索引的真实截图：三侧都选向量流程，Decider 选择 Embed English。",
             "输入提供均值、型号名与配置信息；下方中文分项用于事后诊断，当时未逐项提供给模型。",
             routing_notes + "\nMIRACL中文分项English4.18/Multilingual60.65，均值30.65/59.23。"
             "代码Coding分项Sonnet5.5=91.37，Opus5=81.45。索引选择与多语言证据不符，但不能据四任务推出通用排名。", (31, 32))
    L.picture(s, image("docs/presentations/v9/assets/routing-retrieval-results.png"), .76, 2.25, 11.85, 3.26, frame=True)
    card(s, .76, 5.61, 5.74, .9, "中文检索证据", "MIRACL 中文：English 4.18；Multilingual 60.65。", L.TEAL, size=12.8)
    card(s, 6.74, 5.61, 5.83, .9, "代码能力证据", "Coding：Sonnet 5.5 为 91.37；Opus 5 为 81.45。", L.BLUE, size=12.8)

    # 14–20: new validation section, distinct counting units and concrete outcomes.
    s = base(prs, "05  新增：场景验证", "从演示样例，扩展到 466 次有记录的调用",
             "Decider v19 原版服务 vs Jev 1.13.0；这轮验证没有 GPT，也没有使用 vLLM。",
             "问题数、任务数与请求数不同，不计算跨场景总准确率；来源：validation/README.md。",
             validation_notes, ())
    table(s, ["实验", "评分单位", "Decider", "Jev", "真实调用"],
          [["Skills", "65 个可判定问题", "50/65", "65/65", "64"],
           ["真实浏览器", "5 任务 × 2 轮", "4/10 完成", "10/10 完成", "58"],
           ["金融业务", "30 题，原始顺序", "22/30", "28/30", "120"],
           ["政策路由", "32 题 × 2 顺序", "25/64", "46/64", "224*"]],
          [2.15, 3.35, 2.2, 2.2, 1.05], row_h=.63, size=15.5)
    L.text(s, "* 路由 224 次包含自然语言候选诊断；金融 120 次包含两种候选顺序。", .82, 5.69, 11.7, .4, 15)
    callout(s, "有效返回 466/466，说明接口能完成请求；业务选择和任务完成需要另外评分。")

    s = base(prs, "05  场景验证：Skills", "分类与局部选择接近，证据核验拉开差距",
             "32 个输入包含多个独立问题；68 个返回中，3 个歧义问题只验协议、不计语义成绩。",
             "参考在调用前冻结；按业务用途事后归类。量表采用固定 ±0.5 容差。",
             validation_notes + "\n证据：validation/jev_skills/scenario_summary.json。"
             "Decider queued任务尚无产物却finish，测试替换成assert True仍通过，缺发票地址时选一般咨询地址。"
             "一个评分1.4936与参考2相差0.5064，刚超容差；不能与错误放行等量解释。"
             "客户端RTT中位数Decider164.6ms/Jev119.9ms；P95911.1/163.8ms。没有正式预热阶段，"
             "Decider首次六问题请求13.511秒原样保留，具体原因未确定。")
    groups = valid["skills"]["by_business"]
    labels = ["分类 / 意图", "工具 / 局部动作", "上下文筛选", "恢复 / 完成检查", "证据核验", "等级评分"]
    table(s, ["业务用途", "Decider", "Jev"], [[name, f'{g["correct"]["decider"]}/{g["total"]}',
          f'{g["correct"]["jev"]}/{g["total"]}'] for name, g in zip(labels, groups, strict=True)],
          [6.1, 2.9, 2.95], row_h=.48, size=15.5)
    L.text(s, "本批客户端 RTT 中位数：Decider 164.6 ms；Jev 119.9 ms。", .84, 5.68, 11.65, .3, 14, L.MUTED)
    callout(s, "典型偏差：任务还在排队却选择结束；断言换成 assert True 仍支持完成。", y=6.07, size=15)

    s = base(prs, "05  场景验证：真实浏览器", "多步操作要验收目标页面，不能只看 finish",
             "模型读取页面文字与 DOM，Playwright 执行真实点击、填写与搜索。",
             "每任务最多 12 步、150 秒；截图留证但不作为视觉输入。五任务各独立运行两次。",
             validation_notes + "\n使用scored-summary修订后结果：CFPB官方页面为60 calendar days，"
             "原匹配60 days漏判。根据保存页面将Jev8/10更正10/10，无模型重跑；修订与原记录均保留。")
    task_labels = [("wikipedia-search", "Wikipedia：Basel III 搜索"),
                   ("cfpb-dispute", "CFPB：信用卡账单争议"),
                   ("ecb-usd", "ECB：美元汇率详情"),
                   ("fdic-faq", "FDIC：保险 FAQ"),
                   ("fdic-bank-search", "FDIC：银行机构查询")]
    table(s, ["任务", "Decider", "Jev"], [[name, f'{bf["browser"]["decider"]["tasks"][k]["success"]}/2',
          f'{bf["browser"]["jev"]["tasks"][k]["success"]}/2'] for k, name in task_labels],
          [6.9, 2.5, 2.55], row_h=.56, size=16)
    callout(s, "Decider 6 次过早结束；没有 API 或浏览器动作错误来解释这些失败。", y=6.05, size=16)

    s = base(prs, "05  真实过程截图", "银行查询：停在入口页，与查到机构是两种结果",
             "同一任务：查找 JPMorgan Chase Bank；截图来自第二轮真实浏览器终点。",
             "原始页面截图仅裁切可读区域，完整图片与每步 URL / DOM / 动作保存在 validation/browser_finance。",
             validation_notes + "\nDecider第二轮2步，在BankFind入口结束；Jev第二轮6步返回机构结果，Cert628。"
             "Jev第一轮10步，包含重复操作；完成率高不代表路径最优。")
    for x, p, title, body, crop in [
        (.76, "decider", "Decider · 2 步后结束", "返回 BankFind 入口，没有提交查询", (0, 0, 1526, 1000)),
        (6.75, "jev", "Jev · 6 步完成", "返回目标机构；FDIC Cert 628", (0, 130, 1440, 1130)),
    ]:
        color = L.TEAL if p == "decider" else PURPLE
        L.text(s, title, x, 2.18, 5.75, .4, 20, color, True)
        img = image(f"validation/browser_finance/runs/2026-10-13-live/browser/fdic-bank-search-{p}-2/final.png",
                    f"bank-search-{p}.png", crop)
        L.picture(s, img, x, 2.77, 5.73, 3.47, frame=True)
        L.text(s, body, x, 6.31, 5.75, .27, 12.5, color)

    s = base(prs, "05  场景验证：金融业务", "会分类，不代表能正确核验金额和条件",
             "新闻、财报、存款保险、汇率业务；30 道独立题，按原始候选顺序统计。",
             "反转顺序后 Decider 24/30、Jev 28/30；仍是同一批 30 道题。",
             validation_notes + "\n新闻人工标签，其他业务输入由测试构造，无真实客户资金。"
             "财报正确收入FY2025=281724百万美元，FY2024=245122百万美元；Decider误认年份/单位。"
             "给定1EUR=1.1181USD，历史报表10000EUR=11181USD，Decider混为真实付款要求可执行报价。"
             "同一人、同一家受保银行、同一所有权类别合并300000美元，250000受保50000超限；两者均误判。")
    fin_labels = [("financial_news", "财经新闻分类"), ("annual_report_evidence", "财报证据核验"),
                  ("deposit_insurance_triage", "存款保险初筛"), ("fx_reporting_and_payment_routing", "汇率报表 / 付款分流")]
    table(s, ["业务", "Decider", "Jev"], [[name, *[
        f'{bf["finance"][p]["categories"][k]["original_order_correct"]}/{bf["finance"][p]["categories"][k]["original_order_total"]}'
        for p in ("decider", "jev")]] for k, name in fin_labels], [6.2, 2.85, 2.9], row_h=.57)
    L.rect(s, .78, 5.29, 11.77, 1.20, L.PALE_ORANGE, radius=True)
    L.text(s, "共同错误：同银行、同类别的 18 万 + 12 万美元账户", 1.0, 5.46, 11.32, .35, 19, L.ORANGE, True)
    L.text(s, "两个模型都选“全部受保”；应合并为 30 万，25 万受保、5 万超限。", 1.0, 6.0, 11.32, .34, 16)

    s = base(prs, "05  场景验证：政策路由", "任务分流能做好，复合约束仍容易漏判",
             "独立于网页四任务实验：32 题 × 2 种顺序；自然语言对照只复测后 24 题。",
             "能力 / 价格来自 OpenRouter 目录，健康 / 延迟等为题设；未调用 openrouter/auto 或下游模型。",
             validation_notes + "\nDecider首轮40次存在合格候选的机会中37次false no_route；"
             "自然语言48/48全拒绝，其中仅8次符合参考。Jev也漏能力与总价条件。"
             "输入最大1485token<4096，不能归因上下文超限。建议硬约束代码核验，但该组合方案尚未完成端到端复测。")
    categories = [("semantic_task_routing", "任务语义分流"),
                  ("capability_constraints", "能力 / 上下文"),
                  ("cost_and_budget", "成本 / 预算"),
                  ("availability_fallback_and_preferences", "可用性 / 回退")]
    def cat_score(d, p, k):
        return sum(v["correct"] for v in d[p]["by_category"][k].values())
    policy_rows = []
    for k, name in categories:
        row = [name, f'{cat_score(policy,"decider",k)}/16', f'{cat_score(policy,"jev",k)}/16']
        row += (["—", "—"] if k == "semantic_task_routing" else [
            f'{cat_score(natural,"decider",k)}/16', f'{cat_score(natural,"jev",k)}/16'])
        policy_rows.append(row)
    table(s, ["题型", "Decider 原始", "Jev 原始", "Decider 自然语言", "Jev 自然语言"],
          policy_rows, [2.7, 2.15, 2.05, 2.82, 2.23], row_h=.64, size=15)
    L.text(s, "复合政策合计", .9, 5.64, 2.5, .4, 18, L.TEAL, True)
    L.text(s, "Decider 9/48 → 8/48     Jev 31/48 → 31/48", 3.55, 5.65, 8.7, .4, 20, L.NAVY, True)
    callout(s, "把候选改成自然语言，没有修复多约束比较；Decider 在自然语言批次全部选择 no_route。", y=6.14, size=14.2)

    s = base(prs, "05  引用验证：动作规划", "俄罗斯方块：改写候选，也未消除动作差距",
             "独立项目记录：同一 seed 42，每局最多 50 块；主比较全部由模型选择落点。",
             "只有一个种子；走法不同会产生不同棋盘。此处为已有实验引用，不计入 466 次新目录调用。",
             validation_notes + "\n游戏原始来源已随validation/model_comparison/sources/tetris打包。"
             "固定3失败棋盘：Decider数字3/3次受支配动作，自然语言2/3；Jev两者0/3。"
             "二选一自然语言两者6/6，不能推广完整候选集合。历史规则版430/550程序代选，排除。")
    games = valid["tetris"]["primary_games"]
    game_rows = []
    for policy_id, label in [("original", "数字字段"), ("model", "自然语言")]:
        row = [label]
        for p in ("decider", "jev"):
            g = next(g for g in games if g["policy"] == policy_id and g["provider"] == p)
            row.append(f'{g["pieces"]} 块 / {g["lines"]} 行')
        game_rows.append(row)
    table(s, ["候选表达", "Decider", "Jev"], game_rows, [3.95, 4.0, 4.0], row_h=.70, size=20)
    card(s, .78, 4.71, 5.73, 1.5, "固定棋盘诊断", "Decider：数字版 3/3 次选较差动作。\nJev：两种表达均 0/3。", L.TEAL, size=16)
    card(s, 6.75, 4.71, 5.80, 1.5, "问题不只在字段形式", "要检查候选之间的多指标比较，\n以及完成条件和顺序敏感性。", PURPLE, size=16)

    # 21–22: one measured intervention, then one next-step plan.
    s = base(prs, "06  已验证的输入优化", "先讲清条件适用范围，20 个节点改善了 4 个",
             "支付实验：权重未变，只改输入与候选说明；参考一致次数 14/20 → 18/20。",
             "2026-10-11 优化对照；20 个节点使用相同演示参考，未进行模型微调。",
             "证据：validation/model_comparison/summary.json的historical_decider_input_optimization。"
             "原始提示退款地址等条件未限定业务类型，非退款也被视为缺少验证。优化用case_type限定条款，"
             "明确省略字段不适用，审批与到账分开，给出代码计算事实。没有微调权重；剩下少付和退款错误保留。", (33, 34))
    panel(s, "优化前", "地址是否验证？\n审批是否通过？\n\n条件混在一起，\n非退款也可能被拦截。", .78, w=5.71, h=3.77, color=L.ORANGE, icon_name="list-filter", size=18)
    panel(s, "优化后", "先确定 case_type。\n仅 refund 检查退款地址。\n\n省略字段表示“不适用”。\n审批通过不等于到账。", 6.75, w=5.80, h=3.77, color=L.TEAL, icon_name="check-check", size=18)
    callout(s, "程序先计算金额差、确认是否足够等事实；模型在适用的业务条件中选择下一步。", y=6.20, size=15)

    s = base(prs, "06  后续优化路线", "把失败样本留下，再决定是否需要微调",
             "按业务拆开语义判断与确定性检查，让模型负责它能可靠完成的部分。",
             "下列为后续方案；尚未完成领域微调，也未宣称新的准确率或生产收益。",
             validation_notes + "\n微调应保留官方模板、底座/LoRA/Pointer Head路径，使用领域标签、无匹配、"
             "缺证据、金额/年份反例，候选数量与顺序扰动，训练/验证/测试按模板或业务来源拆分防泄漏。"
             "独立评估准确、误放行、完成度、置信度校准与拒绝回退，再灰度；不只训练输出JSON形式。", (35, 36, 49))
    flow(s, [("冻结回归集", "缺证据 / 早结束 / 错算", "clipboard-check"),
             ("拆清业务条件", "金额、预算、能力用代码校验", "code-xml"),
             ("再做领域训练", "LoRA / 决策头，加入困难反例", "brain-circuit"),
             ("独立验收上线", "完成度、校准、拒绝与回退", "shield-check")], h=1.85)
    card(s, .78, 5.01, 5.74, 1.12, "浏览器 / 支付", "finish 前核验页面、产物或账本；缺证据继续收集。", L.TEAL, size=13.8)
    card(s, 6.75, 5.01, 5.80, 1.12, "模型路由", "语义分流交给模型；硬约束筛选后再排序和回退。", L.BLUE, size=13.8)

    # 23–27: adapter and performance, with three distinct timing boundaries.
    s = base(prs, "07  vLLM 适配", "保留决策头，把单请求执行接到并发调度",
             "这不是把 Decider 当聊天模型生成 JSON；适配使用官方模板、位置元数据与评分逻辑。",
             "vLLM 0.31.0；BF16 底座 / FP32 Pointer Head；未开启前缀缓存、分块预填充或 CUDA Graph。",
             perf_notes + "\n导出合并LoRA并携带head/calibration/tokenizer；自定义pooler取候选/答案位置，"
             "全prefill后应用官方answer转换。严格比较choice与概率/置信度；容差0.02异常保留。", (37, 38))
    flow(s, [("导出权重", "合并 LoRA，保存决策头", "layers"),
             ("校验产物", "模板 / 分词器 / 校准参数", "clipboard-check"),
             ("启动 API", "定制 Pooler + 并发调度", "workflow"),
             ("对齐原版", "choice 与数值分别验收", "check-check")], h=1.78)
    L.rect(s, .78, 4.96, 11.78, 1.32, L.NAVY, radius=True)
    L.text(s, "decider-vllm export  →  verify-artifacts  →  serve", 1.05, 5.15, 11.25, .42, 20, L.WHITE, True)
    L.text(s, "serve --max-num-seqs 128 --max-num-batched-tokens 8192", 1.05, 5.75, 11.25, .32, 16, "B8D9D3")

    s = base(prs, "07  独立推理服务", "公网 API，私网 EC2 源站，模型持续驻留",
             "独立 StrandsDeciderVLLM 项目：g6.xlarge / 1 × NVIDIA L4；与主演示的原版实例分开。",
             "API 使用访问密钥；CloudFront VPC Origin。图中省略网络端点，后端直接调用相关 AWS 服务。",
             perf_notes + "\n独立API公开入口https://d1e9udi8st2p3g.cloudfront.net，"
             "API key认证，EC2上systemd管理，ECR/私有S3模型产物/Secrets Manager。"
             "本页不是将主演示迁移到vLLM。", (39, 40))
    L.rect(s, 3.02, 2.25, 9.53, 3.41, None, L.LINE)
    L.text(s, "AWS Cloud", 3.19, 2.37, 3.1, .28, 12, L.MUTED, True)
    L.rect(s, 5.65, 2.73, 6.64, 2.69, None, PURPLE)
    L.text(s, "VPC · us-east-1", 5.83, 2.85, 5.8, .26, 11.5, PURPLE, True)
    L.rect(s, 5.88, 3.25, 6.18, 1.96, "F0F8F7", "70ABA5")
    L.text(s, "私有源站", 6.02, 3.35, 5.7, .26, 10.5, L.TEAL, True)
    L.rect(s, .8, 3.49, 1.81, 1.52, L.WHITE, L.LINE, True)
    icon(s, "code-xml", 1.48, 3.64, w=.4)
    L.text(s, "API 客户端", 1.0, 4.2, 1.4, .35, 16, L.TEAL, True)
    L.rect(s, 3.25, 3.49, 1.95, 1.52, L.WHITE, L.LINE, True)
    icon(s, "cloudfront", 3.98, 3.62, w=.45, aws=True)
    L.text(s, "CloudFront", 3.43, 4.19, 1.61, .3, 14, L.NAVY, True)
    L.text(s, "VPC Origin", 3.46, 4.65, 1.6, .25, 12, L.MUTED)
    L.rect(s, 6.16, 3.71, 3.81, 1.32, L.WHITE, L.LINE, True)
    icon(s, "ec2", 6.34, 3.87, w=.42, aws=True)
    L.text(s, "g6.xlarge · 1 × L4", 6.9, 3.91, 2.87, .32, 16, L.NAVY, True)
    L.text(s, "API 鉴权 → vLLM → 决策头", 6.36, 4.51, 3.4, .36, 14, L.TEAL)
    L.rect(s, 10.59, 3.71, 1.2, 1.32, L.WHITE, L.LINE, True)
    icon(s, "ebs", 10.94, 3.85, w=.41, aws=True)
    L.text(s, "EBS gp3\n160 GB", 10.73, 4.38, .95, .57, 12.5, L.NAVY, False, PP_ALIGN.CENTER)
    for a, b in [((2.65,4.27),(3.22,4.27)),((5.24,4.27),(6.13,4.27)),((10.0,4.27),(10.56,4.27))]:
        L.line(s, [a, b], L.MUTED, 1.4, True)
    callout(s, "压测：固定 worker，完整响应后补发；每档预热另计、失败不重试。吞吐按成功请求 / 总时长计算。", y=5.92, size=14)

    s = base(prs, "07  原版与 vLLM", "同一块 L4，8 并发吞吐提升约 2.36 倍",
             "同机实验：g6.4xlarge，两个容器均限制 4 vCPU / 14 GiB；回环 HTTP。",
             "每档 80 请求，预热另计。提升同时包含推理路径和调度变化，不能全部归因于批处理。",
             perf_notes + "\n8并发原版14.76RPS/P95542.8ms；vLLM复测34.91RPS/P95237.6ms。"
             "更早测试的5次超时保留在原始报告，本图为标明的复测。", (41,))
    L.picture(s, image("docs/presentations/v9/assets/original-vllm-rps.png"), .76, 2.18, 8.24, 4.10)
    L.text(s, "14.76 → 34.91", 9.22, 2.77, 3.28, .55, 25, L.TEAL, True)
    L.text(s, "成功请求 / 秒", 9.22, 3.46, 3.2, .35, 16, L.MUTED)
    L.text(s, "542.8 → 237.6", 9.22, 4.42, 3.3, .55, 24, L.BLUE, True)
    L.text(s, "P95 毫秒 · 8 并发", 9.22, 5.11, 3.3, .4, 16, L.MUTED)

    s = base(prs, "07  并发与批次", "超过 8 并发后，这批短输入主要在增加等待",
             "g6.xlarge 公网 API；1–128 并发每档 1,280 请求。输入 304–445 token，均值 373.9。",
             "batch 参数是调度 token 预算，不是请求数。8192 / 16384 / 32768 合计 23,040 次均成功。",
             perf_notes + "\n单次sweep10240成功；另一个重复批次12800成功但1次confidence差.0202超.02容差。"
             "batch三预算短输入总体无稳定吞吐收益，不能推广到长输入。", (42, 43))
    def sweep_plot(ax):
        ax.plot([c["concurrency"] for c in sweep], [c["successful_rps"] for c in sweep], "o-", color="#" + L.TEAL, lw=2.4)
        ax.set_xscale("log", base=2)
        ax.set_xticks([1, 2, 4, 8, 16, 32, 64, 128], labels=[1, 2, 4, 8, 16, 32, 64, 128])
        ax.set_xlabel("并发请求数")
        ax.set_ylabel("成功请求 / 秒")
        ax.set_ylim(0, 40)
        ax.grid(alpha=.15)
    L.picture(s, chart("concurrency-sweep", sweep_plot, 7.8, 3.8), .7, 2.22, 7.5, 3.89)
    c8 = next(c for c in sweep if c["concurrency"] == 8)
    c128 = next(c for c in sweep if c["concurrency"] == 128)
    table(s, ["并发", "吞吐 / s", "P95 ms"],
          [[8, f'{c8["successful_rps"]:.2f}', f'{c8["p95_ms"]:.1f}'],
           [128, f'{c128["successful_rps"]:.2f}', f'{c128["p95_ms"]:.1f}']],
          [1.05, 1.55, 1.75], x=8.28, y=2.54, row_h=.62, size=14)
    L.text(s, "更大的 token 预算\n未提高这批短输入吞吐。\n先找服务容量，再设排队上限。", 8.44, 4.79, 4.0, 1.24, 17, L.MUTED)

    s = base(prs, "07  vLLM vs Jev", "低并发 Decider 延迟低，Jev 在中并发吞吐更高",
             "相同 20 个支付输入循环提交；Decider 8192 token 预算；Jev 直连 API。",
             "2026-10-13，每模型 10,240 次正式尝试；P95 为成功响应端到端延迟，单位毫秒。",
             perf_notes + "\nDecider10239/10240，1次HTTP502，10次confidence差.0202超.02容差；"
             "Jev10240/10240。每侧560预热另计。128并发出现尾延迟，不用短测替代长期SLA。", (44, 45))
    rows = []
    for concurrency in (1, 8, 32, 128):
        cells = {p: next(c for c in jev["cells"] if c["provider"] == p and c["concurrency"] == concurrency) for p in ("decider", "jev")}
        rows.append([concurrency, f'{cells["decider"]["successful_rps"]:.2f}', f'{cells["jev"]["successful_rps"]:.2f}',
                     f'{cells["decider"]["latency_ms"]["p95"]:.1f}', f'{cells["jev"]["latency_ms"]["p95"]:.1f}'])
    table(s, ["并发", "Decider 请求/s", "Jev 请求/s", "Decider P95", "Jev P95"], rows,
          [1.05, 2.85, 2.75, 2.65, 2.65], row_h=.62, size=16)
    L.text(s, "Decider：10,239/10,240 成功，1 次 502；10 次置信度差异略超数值对齐容差。", .82, 5.62, 11.7, .39, 15)
    L.text(s, "Jev：10,240/10,240 成功；128 并发的吞吐回落，尾延迟也明显增加。", .82, 6.16, 11.7, .35, 15)

    # 28–30: new exact RI terms, cost intersection, and utilization/commitment.
    s = base(prs, "08  新增：EC2 RI 成本", "g6.2xlarge：RI 折算月费约降低 35%–60%",
             "us-east-1 · Linux · Shared · Standard RI；按 730 小时/月，美元，单台计算费。",
             "2026-10-15 AWS Price List；折算月费 = 预付额 / 期限月数 + 小时费 × 730。不是实际月账单。",
             cost_notes, (46,))
    rows = [[r["label"], f'${r["upfront_usd"]:,.0f}', f'${r["recurring_hourly_usd"]:.5f}',
             f'${r["monthly_amortized_usd"]:,.2f}', f'{r["savings"] * 100:.1f}%' if r["years"] else "—"] for r in cost["rows"]]
    table(s, ["购买方式", "一次预付", "后续小时费", "折算每月", "相对按需节省"], rows,
          [3.3, 2.0, 2.2, 2.2, 2.25], row_h=.46, size=14.5, highlights=(1, 6))
    callout(s, "Jev 按量：$0.042 / 百万输入 token，输出免费。RI 承诺整个期限，闲置也继续付费。", y=6.19, size=15)

    chosen_cost = [cost["rows"][i] for i in (0, 1, 3, 6)]
    s = base(prs, "08  Decider vs Jev", "3 年全预付 RI：与 Jev 的交点约 1033 万次/月",
             "单台能够承载时：低于交点，Jev 调用费更低；高于交点，EC2 计算费更低。",
             "交点仅表示费用相等，不代表该实例实测容量，也不代表结果质量相同。",
             cost_notes, (47, 48))
    rows = [[r["label"], f'${r["monthly_amortized_usd"]:.2f}',
             f'{r["break_even_monthly_payment_decisions"] / 1e4:,.0f} 万次',
             f'{r["break_even_monthly_two_step_routes"] / 1e4:,.0f} 万条'] for r in chosen_cost]
    table(s, ["计算费方案", "折算每月", "支付决策 / 月", "完整两步路由 / 月"], rows,
          [3.3, 2.2, 3.15, 3.30], row_h=.57, size=15.5)
    card(s, .78, 5.32, 5.76, 1.08, "支付单次", "平均 656.95 输入 token → $27.59 / 百万次", L.TEAL, size=14)
    card(s, 6.76, 5.32, 5.80, 1.08, "两步路由", "平均 2,805.75 token → $117.84 / 百万条流程", PURPLE, size=14)

    s = base(prs, "08  成本决策", "长期稳定运行，才有理由承担 RI 的期限",
             "先确认业务质量和流量，再比较常驻成本；Jev 按实际输入用量付费。",
             "计算费外另有 EBS、网络、应用和运维。区域 RI 不预留容量；这里没有执行 RI 购买。",
             cost_notes + "\n与可停机按需比较，RI年度运行小时交点=RI总额/(按需小时价×期限小时)。"
             "1年无预付约65.1%，1年全预付60.8%，3年全预付40.0%；这是可关机小时占比，不是GPU利用率。")
    utilization = [(r["label"], r["annualized_break_even_running_fraction"]) for r in chosen_cost[1:]]
    table(s, ["RI 方式", "相对可停机按需的运行时长交点"],
          [[name, f'{value * 100:.1f}% 的期限小时'] for name, value in utilization],
          [4.6, 7.35], row_h=.58, size=17)
    L.text(s, "运行时长交点 ≠ GPU 利用率；前者比较“开机多久”，后者影响“每次决策成本”。", .86, 4.88, 11.7, .48, 16, L.MUTED)
    card(s, .78, 5.49, 5.75, .93, "稳定常驻且质量达标", "RI 降低固定费；仍需规划冗余与模型维护。", L.TEAL, size=13.5)
    card(s, 6.76, 5.49, 5.80, .93, "流量少或波动大", "按量 API 避免长期闲置成本与预付款。", PURPLE, size=13.5)

    s = base(prs, "09  结论", "用业务结果选模型，再用流量决定部署方式",
             "简单语义选择可以先试；证据、完成条件、多约束和金额计算需要更严格验收。",
             "这些结论来自当前小样本；性能测试不替代质量验证，价格交点不替代容量测试。",
             validation_notes + "\n最终综合：分类与局部选择接近，Jev本批复杂证据/多步完成更好但也有金融和路由错误；"
             "Decider支持可控自托管，vLLM提升并发，RI降固定成本。GPT用于开放式生成和复杂异常，具体任务仍须回归。",
             (8, 9, 32, 35, 49, 50, 51))
    panel(s, "先放入明确任务", "分类 / 工具选择 / 语义分流\n\n固定候选、缺证据可拒绝。\n以独立业务样本验收。", .76, w=3.79, h=3.79, icon_name="list-filter", color=L.TEAL, size=17)
    panel(s, "复杂流程看完成度", "证据核验 / 浏览器 / 金融\n\nJev 本批表现更稳，仍有错。\n确定性条件由程序复核。", 4.80, w=3.79, h=3.79, icon_name="clipboard-check", color=PURPLE, size=17)
    panel(s, "规模确定后算成本", "Decider 自托管 / Jev 按量\n\nvLLM 改善并发，RI 降固定费。\n同时计入运维与冗余。", 8.84, w=3.79, h=3.79, icon_name="banknote", color=L.BLUE, size=17)

    s = base(prs, "09  演示与资料", "打开演示，查看真实过程和原始记录",
             "演示展示输入、原始选择、实际动作和结果；仓库保留验证脚本、截图与数据。",
             "PPT 内含逐页讲稿与来源；v9 保留，本版新增验证与 RI 后将 52 页重排为 32 页。",
             f"Demo：{DEMO}\n主项目：{REPO}\nvLLM：{VREPO}\n验证：{REPO}/blob/main/validation/README.md"
             "\n官方介绍：https://strandsagents.com/blog/introducing-strands-decider/"
             "\n评测：https://benchmarkheaven.com/jev-models/api\n价格与条款见sources。", (52,))
    links = [
        ("在线演示", DEMO),
        ("主项目", REPO),
        ("vLLM 适配与公共 API", VREPO),
        ("场景验证结果与截图", REPO + "/blob/main/validation/README.md"),
        ("Strands Decider 官方介绍", "https://strandsagents.com/blog/introducing-strands-decider/"),
        ("JevBench 评测结果与方法", "https://benchmarkheaven.com/jev-models/api"),
    ]
    for i, (title, url) in enumerate(links):
        yy = 2.25 + i * .66
        L.text(s, title, .8, yy, 3.2, .34, 16, L.TEAL, True)
        hyperlink(s, url, url, 4.05, yy, 8.45, 12.5 if len(url) > 75 else 14)

    assert len(prs.slides) == 32
    pptx = OUT / f"{NAME}.pptx"
    prs.save(pptx)
    assert digest(BASE) == base_hash
    # Copy compact machine-readable facts, avoiding external mutable sources.
    snapshots = {
        "payment-warm.json": PAY, "app-routing.json": ROUTE,
        "validation-cross-scenario.json": VALID, "browser-finance.json": BROWSER,
        "policy-routing-original.json": POLICY, "policy-routing-natural.json": NATURAL,
        "jevbench.json": BENCH, "vllm-concurrency.json": f"{VLLM}/sweep-1-128.json",
        "vllm-batches.json": f"{VLLM}/batch-summary.json", "vllm-jev.json": f"{VLLM}/jev-summary.json",
    }
    snapshot_manifest = []
    for name, src in snapshots.items():
        dst = SOURCES / name
        shutil.copy2(ROOT / src, dst)
        snapshot_manifest.append({"file": str(dst.relative_to(ROOT)), "source": src, "sha256": digest(dst)})
    manifest = {
        "created_at": datetime.now(UTC).isoformat(), "base_file": str(BASE.relative_to(ROOT)),
        "base_sha256": base_hash, "pptx_file": str(pptx.relative_to(ROOT)), "pptx_sha256": digest(pptx),
        "pages": 32, "previous_pages": 52, "new_model_calls": 0, "aws_mutations": 0,
        "slides": SLIDES, "evidence": EVIDENCE, "snapshots": snapshot_manifest, "images": IMAGES,
        "price_queried_at": cost["queried_at"],
        "deduplication": [
            "模型机制与BERT/Jev/GPT接口说明合并，避免重复介绍",
            "JevBench的三页合并成一页，保留方法解释与链接",
            "支付逐场景重复截图和结论合并，保留总表、退款截图与一个闭环",
            "主应用路由保留实现、总表、检索证据，合并代码和语义解释",
            "优化建议合并为实测改进与后续路线两页",
            "vLLM部署与测试方法压缩，并发与batch共用一页",
            "成本加入RI条款、价格和交点，最终选型仅保留一页",
        ],
    }
    (OUT / "build-manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    notes = ["# v10 讲稿与来源", "", "32 页；2026-10-15。各页备注同样内嵌在 PPTX，可在演讲者视图查看。",
             "", "价格为新查询；性能和质量为原始保存实验。本次没有新增模型调用或购买 RI。", ""]
    for meta in SLIDES:
        notes.extend([f'## {meta["page"]:02} · {meta["title"]}', "", meta["notes"], ""])
    (OUT / "讲稿与来源.md").write_text("\n".join(notes))
    print(json.dumps({"pptx": str(pptx), "pages": len(prs.slides), "evidence_files": len(EVIDENCE)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
