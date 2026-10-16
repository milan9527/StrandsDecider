"""Audit real-call evidence, paired inputs, original choices, slides and PDF."""

import hashlib
import json
import math
import statistics
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path

import pymupdf as fitz
from PIL import Image, ImageChops, ImageDraw
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/presentations/v12"
PROVIDERS = ("decider", "jev", "bedrock")
PROFILES = ("baseline", "optimized")
STAGES = ("intake", "verify", "approve", "execute", "close")


def load(path):
    return json.loads(path.read_text())


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normal(text):
    return "".join(text.split())


def body(slide):
    # The lower explanatory row is intentionally removed from every v12 slide.
    return [s.text for s in slide.shapes if s.has_text_frame and s.top < Inches(6.5)]


def images(slide):
    return [hashlib.sha256(s.image.blob).hexdigest() for s in slide.shapes if s.shape_type == 13]


def check_capture():
    data = load(OUT / "sources/payment-live.json")
    stats = load(OUT / "sources/payment-comparison.json")
    capture = load(OUT / "capture-manifest.json")
    assert digest(OUT / "sources/payment-live.json") == capture["evidence_sha256"]
    assert data["complete"] and not data["page_errors"] and not data["gpu_writes"]
    assert data["gpu"] == {"phase": "ready", "provider": "ec2", "instance_type": "g6.xlarge"}
    warmups = [r for r in data["http_decisions"] if r["phase"] == "warmup"]
    calls = [r for r in data["http_decisions"] if r["phase"] == "formal"]
    assert len(warmups) == 6 and len(calls) == 120
    keys = Counter((r["request"]["snapshot"]["scenario"], r["request"]["input_profile"],
                    r["request"]["provider"], r["request"]["stage"]) for r in calls)
    assert len(keys) == 120 and set(keys.values()) == {1}
    for r in [*warmups, *calls]:
        assert r["status"] == 200 and r["response"]["source"] == "live"
    for case, runs in data["cases"].items():
        assert runs["baseline"]["snapshot"] == runs["optimized"]["snapshot"]
        for profile in PROFILES:
            run = runs[profile]
            for stage in STAGES:
                rr = [run["results"][p][stage] for p in PROVIDERS]
                assert all(r["model_input"] == rr[0]["model_input"] for r in rr)
                assert len({r["input_hash"] for r in rr}) == 1
                state = json.loads(rr[0]["model_input"]["state"])
                # Do not mistake the independently returned reference for input.
                assert "reference_choice" not in json.dumps(rr[0]["model_input"])
                baseline = runs["baseline"]["results"]["decider"][stage]["model_input"]
                original_state = json.loads(baseline["state"])
                assert all(original_state[k] == v for k, v in state.items())
                assert list(rr[0]["model_input"]["questions"]["action"]["criteria"]) == list(
                    baseline["questions"]["action"]["criteria"])
                for provider, result in zip(PROVIDERS, rr, strict=True):
                    http = next(r for r in calls if (
                        r["request"]["snapshot"]["scenario"] == case
                        and r["request"]["input_profile"] == profile
                        and r["request"]["provider"] == provider and r["request"]["stage"] == stage))
                    assert http["request"]["snapshot"] == run["snapshot"]
                    assert http["response"] == {k: v for k, v in result.items() if k != "client_ms"}
                    audit = next(a for a in run["executions"][provider]["audit"] if a["stage"] == stage)
                    assert audit["requested_action"] == result["choice"]
                    assert audit["external_action_executed"] is False
                    assert result["matches_reference"] == (result["choice"] == result["reference_choice"])
            for provider in PROVIDERS:
                rows = list(run["results"][provider].values())
                saved = stats["cases"][case][profile][provider]
                assert len(rows) == saved["calls"] == 5
                assert saved["matches"] == sum(r["matches_reference"] for r in rows)
                assert math.isclose(saved["total_backend_ms"], sum(r["timings"]["backend_total_ms"] for r in rows))
                assert saved["business_outcome"] == run["businessOutcomes"][provider]
    for profile, correct in (("baseline", (14, 18, 20)), ("optimized", (18, 20, 20))):
        for provider, expected in zip(PROVIDERS, correct, strict=True):
            rows = [r for case in data["cases"].values() for r in case[profile]["results"][provider].values()]
            saved = stats["aggregate"][profile][provider]
            assert saved["matches"] == sum(r["matches_reference"] for r in rows) == expected
            assert saved["calls"] == len(rows) == 20
            assert math.isclose(saved["mean_backend_ms"], statistics.mean(r["timings"]["backend_total_ms"] for r in rows))
    pending = data["cases"]["pending"]
    a, b = [json.loads(pending[pr]["results"]["decider"]["approve"]["model_input"]["state"]) for pr in PROFILES]
    assert (len(a), len(b)) == (18, 12)
    assert set(a) - set(b) == {"refund_approved", "refund_destination_verified", "settled",
                              "settlement_minutes", "settlement_overdue", "settlement_sla_minutes"}
    for p in PROVIDERS:
        assert pending["baseline"]["results"][p]["approve"]["choice"] == ("approved" if p == "bedrock" else "await_approval")
        assert pending["optimized"]["results"][p]["approve"]["choice"] == "approved"
        assert pending["baseline"]["executions"][p]["snapshot"]["review_pending"] == (p != "bedrock")
        assert pending["optimized"]["executions"][p]["snapshot"]["review_pending"] is False
        for pr in PROFILES:
            assert pending[pr]["businessOutcomes"][p]["fundsComplete"] is False
            assert pending[pr]["businessOutcomes"][p]["complete"] is False
            assert not any(a["ledger_entry"] for a in pending[pr]["executions"][p]["audit"])
    refund = data["cases"]["refund"]["optimized"]
    assert refund["results"]["decider"]["approve"]["choice"] == "approved"
    audit = next(a for a in refund["executions"]["decider"]["audit"] if a["stage"] == "approve")
    assert audit["requested_action"] == "approved" and audit["action"] == "verify_destination" and audit["safety_adjusted"]
    for shot in capture["screenshots"]:
        assert digest(ROOT / shot["file"]) == shot["sha256"]
        assert shot["data_replaced"] is False
    return {"formal_calls": 120, "warmups_excluded": 6, "provider_inputs_equal": True,
            "baseline_optimized_facts_preserved": True, "candidate_ids_unchanged": True,
            "reference_not_in_model_input": True, "raw_choices_matched_to_execution_audit": True,
            "metrics_recomputed": True, "pending_state_fields": "18 -> 12",
            "raw_screenshots_verified": 10, "external_funds_actions": 0, "gpu_lifecycle_calls": 0}


def main():
    manifest = load(OUT / "build-manifest.json")
    pptx = ROOT / manifest["pptx_file"]
    prs = Presentation(pptx)
    base = Presentation(ROOT / manifest["base_file"])
    pdf = fitz.open(pptx.with_suffix(".pdf"))
    assert len(prs.slides) == len(pdf) == manifest["pages"] == 35
    assert digest(ROOT / manifest["base_file"]) == manifest["base_sha256"]
    assert digest(pptx) == manifest["pptx_sha256"]
    for item in manifest["evidence"]:
        assert digest(ROOT / item["path"]) == item["sha256"]
    for item in manifest["images"]:
        dest = ROOT / item["output"]
        assert digest(dest) == item["sha256"]
        with Image.open(ROOT / item["source"]) as original, Image.open(dest) as rendered:
            crop = original.crop(item["crop"]) if item["crop"] else original
            assert crop.size == rendered.size and ImageChops.difference(crop, rendered).getbbox() is None
    missing, bounds, preserved, thumbs = [], [], [], []
    previews = OUT / "previews"
    previews.mkdir(exist_ok=True)
    for i, (slide, page, meta) in enumerate(zip(prs.slides, pdf, manifest["slides"], strict=True), 1):
        rendered = normal(page.get_text())
        assert not any(s.has_text_frame and s.text.strip() and Inches(6.5) <= s.top < Inches(7.03)
                       for s in slide.shapes), i
        assert meta["removed_lower_explanation"] in slide.notes_slide.notes_text_frame.text
        assert any(s.has_text_frame and s.text == f"{i:02} / 35" for s in slide.shapes)
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip() and normal(shape.text) not in rendered:
                missing.append({"page": i, "text": shape.text})
            if min(shape.left, shape.top) < 0 or shape.left + shape.width > prs.slide_width + 12 or shape.top + shape.height > prs.slide_height + 12:
                bounds.append({"page": i, "shape": shape.name})
        assert slide.notes_slide.notes_text_frame.text == meta["notes"]
        if meta["source"] == "v11":
            original = base.slides[meta["v11_page"] - 1]
            assert images(slide) == images(original)
            if not meta["body_modified"]:
                assert body(slide) == body(original)
                preserved.append(i)
        path = previews / f"page-{i:02}.png"
        page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(path)
        with Image.open(path) as im:
            im.thumbnail((384, 216))
            tile = Image.new("RGB", (404, 248), "#e7edf2")
            tile.paste(im, ((404 - im.width) // 2, 6))
            ImageDraw.Draw(tile).text((12, 226), f"{i:02}", fill="#203448")
            thumbs.append(tile)
    for name, indexes in (("slides-overview", range(35)), ("payment-overview", range(7, 13))):
        cols = 3 if name == "payment-overview" else 4
        canvas = Image.new("RGB", (cols * 404, math.ceil(len(indexes) / cols) * 248), "#e7edf2")
        for n, index in enumerate(indexes):
            canvas.paste(thumbs[index], ((n % cols) * 404, (n // cols) * 248))
        canvas.save(OUT / f"{name}.png")
    report = {
        "audited_at": datetime.now(UTC).isoformat(), "pages": 35,
        "v11_file_unchanged": True, "preserved_body_pages": preserved,
        "lower_explanations_removed_pages": 35,
        "bottom_footer_and_page_numbers_preserved": True,
        "all_retained_images_preserved": True, "crop_pixels_identical_to_capture": True,
        "rendered_text_missing": missing, "shape_bounds_issues": bounds,
        "capture_checks": check_capture(),
        "pptx_sha256": digest(pptx), "pdf_sha256": digest(pptx.with_suffix(".pdf")),
    }
    (OUT / "render-audit.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    assert not missing and not bounds


if __name__ == "__main__":
    main()
