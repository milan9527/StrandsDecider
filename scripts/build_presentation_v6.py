"""Add two observed input-optimization slides to the latest EC2 presentation."""

import copy
import hashlib
import io
import json
import re
import shutil
from datetime import UTC, datetime
from pathlib import Path

import cairosvg
import presentation_layout as L
from build_presentation_v5 import replace_text
from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/presentations/v5"
OUT = ROOT / "docs/presentations/v6"
BASE = SOURCE / "Strands-Decider-模型原理与支付业务演示-v5-EC2部署版-25页-20261011.pptx"
EVIDENCE = ROOT / "docs/results/optimized-2026-10-11.json"
NAME = "Strands-Decider-模型原理与支付业务演示-v6-优化流程实测版-27页-20261011"
DEMO = "https://d3i7ogfx5ml7dc.cloudfront.net/#crypto"
REPO = "https://github.com/milan9527/StrandsDecider"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def icon(slide, name, x, y, color=L.TEAL, size=0.35):
    svg = (ROOT / "docs/presentations/v3/icons" / f"{name}.svg").read_text()
    bitmap = cairosvg.svg2png(
        bytestring=svg.replace("currentColor", f"#{color}").encode(),
        output_width=128,
        output_height=128,
    )
    shape = slide.shapes.add_picture(
        io.BytesIO(bitmap), Inches(x), Inches(y), Inches(size), Inches(size)
    )
    shape.name = "ICON: Lucide " + name


def new_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = L.rgb(L.BG)
    L.text(slide, "STRANDS DECIDER   /   05 / 优化流程实测", 0.6, 0.27, 12, 0.25, 10, L.TEAL, True)
    L.text(slide, title, 0.6, 0.79, 12.1, 0.69, 30, L.INK, True)
    L.text(slide, subtitle, 0.62, 1.56, 12, 0.43, 16.5, L.MUTED)
    L.line(slide, [(0.62, 7.04), (12.68, 7.04)], L.LINE, 0.6)
    return slide


def prepare_images():
    assets = OUT / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    original = assets / "optimized-pending-live.png"
    if not original.exists():
        shutil.copyfile(ROOT / "artifacts/optimized-pending-live.png", original)
    with Image.open(original) as image:
        assert image.size == (1600, 1100)
        crops = {
            "optimized-pending-timings.png": (264, 498, 1559, 650),
        }
        for name, box in crops.items():
            image.crop(box).save(assets / name)
    selector_capture = json.loads((assets / "selector-capture.json").read_text())
    assert selector_capture["sha256"] == digest(assets / "optimized-selector.png")
    return {
        "capture_script": "scripts/check_optimized_crypto_browser.py",
        "mode": "live",
        "scenario": "pending",
        "profile": "optimized",
        "capture_url": DEMO,
        "original_image": str(original.relative_to(ROOT)),
        "original_sha256": digest(original),
        "crops": [
            {
                "file": str((assets / name).relative_to(ROOT)),
                "source_box_pixels": list(box),
                "sha256": digest(assets / name),
            }
            for name, box in crops.items()
        ],
        "editing": "Rectangular crops only; no values or UI text replaced.",
        "selector_screenshot": {
            "file": str((assets / "optimized-selector.png").relative_to(ROOT)),
            **selector_capture,
        },
    }


def case_for(data, profile, name):
    return next(case for case in data["profiles"][profile]["cases"] if case["case"] == name)


def method_slide(prs, data):
    title = "优化示例：让退款审批只作用于退款"
    slide = new_slide(prs, title, "USDC 收款 · 同一业务证据，调整输入字段与候选条件")
    L.text(slide, "真实界面入口", 0.67, 2.31, 1.7, 0.3, 14, L.MUTED, True)
    L.picture(slide, OUT / "assets/optimized-selector.png", 2.4, 2.17, 3.25, 0.76)
    icon(slide, "list-filter", 6.04, 2.35)
    L.text(slide, "精简字段 → 明确适用场景 → 重新判断", 6.53, 2.35, 6.1, 0.38, 17, L.TEAL, True)
    for x, profile, name, color, fill in [
        (0.65, "baseline", "固定流程", L.ORANGE, L.PALE_ORANGE),
        (6.79, "optimized", "优化流程", L.TEAL, L.PALE_TEAL),
    ]:
        case = case_for(data, profile, "pending")
        step = next(s for s in case["steps"] if s["stage"] == "approve")
        L.rect(slide, x, 3.16, 5.84, 2.89, fill, L.LINE, True)
        icon(slide, "layers" if profile == "baseline" else "list-filter", x + 0.2, 3.38, color)
        L.text(slide, name, x + 0.7, 3.35, 4.9, 0.39, 21, color, True)
        body = (
            "收款事实 + refund_approved=false\n退款字段仍出现在普通收款输入中"
            if profile == "baseline"
            else "收款事实，移除无关退款字段\n候选明确：非退款不要求退款审批"
        )
        L.text(slide, body, x + 0.22, 3.98, 5.4, 0.79, 16, L.INK)
        L.line(slide, [(x + 0.22, 4.9), (x + 5.62, 4.9)], L.LINE, 0.8)
        for offset, provider in [(0, "decider"), (0.47, "bedrock")]:
            result = step[provider]
            label = "Decider" if provider == "decider" else "GPT"
            L.text(
                slide,
                f"{label}：{result['label']}",
                x + 0.22,
                5.08 + offset,
                4.0,
                0.32,
                17,
                color if provider == "decider" else L.BLUE,
                True,
            )
            L.text(
                slide,
                f"{result['timings']['backend_total_ms']:,.1f} ms",
                x + 4.23,
                5.12 + offset,
                1.36,
                0.25,
                12,
                L.MUTED,
                align=PP_ALIGN.RIGHT,
            )
    L.text(
        slide,
        "本次优化改变输入与条件表述；候选 ID、模型权重和执行校验保持一致。",
        0.65,
        6.29,
        12,
        0.34,
        15,
        L.INK,
        True,
    )
    L.text(
        slide,
        "确认数仍为 3/12：两侧资金继续挂起。审批选择改善，不表示已经到账。",
        0.65,
        6.75,
        12,
        0.25,
        12.5,
        L.MUTED,
    )
    slide.notes_slide.notes_text_frame.text = (
        "展示现有真实测试，不是新编造的模型输出，也未重新训练或微调模型。\n"
        "演示操作：支付流程对比 → USDC未到账 → 固定流程 → 开始对比；保持证据不变，"
        "再选择优化流程。可用对比记录切换两次运行，查看审批节点及实际模型输入。\n"
        "普通收款输入原本含refund_approved=false。优化后只保留本场景相关字段，"
        "候选明确退款审批只适用于refund。候选ID集合完全不变，两模型使用同一优化输入。"
        "程序没有把参考答案输入模型，也没有替换模型原始选择。\n"
        "本次固定流程Decider选择await_approval，75.722ms；优化后选择approved，80.217ms。"
        "该单步没有提速。GPT两次均选择approved，分别1080.104ms和820.332ms。"
        "这是2026-10-11 EC2批次，不能与第14页2026-10-10历史记录的87.406ms混淆。\n"
        "观察支持本次优化后选择改善，但不能从输出确认模型内部错误原因。"
        "输入字段与候选表述同时修改，因此也不能单独归因于删除某个字段。\n"
        "确认数仍为3/12；两侧实际资金动作均挂起，未生成资金账本，不是业务完成。"
        "优化流程是固定五步流程，不自动推进模拟事件；自主闭环仍是独立测试选项。\n"
        "来源：docs/results/optimized-2026-10-11.json，pending/approve；"
        "入口截图为登录现网后选择优化流程的高分辨率元素截图，未启动决策；"
        "采集时间、选择器与哈希见capture-manifest.json。\n"
        "配置：EC2 g6.xlarge / NVIDIA L4，Decider v19；Bedrock GPT-6.1 Sol / reasoning low。"
        "API 20261011-optimized1，前端optimized1。耗时为backend_total_ms，含通信与排队。"
    )
    return {"title": title, "section": "05 / 优化流程实测 · 方法", "dark": False}


def results_slide(prs, data):
    title = "四场景实测：参考一致 14/20 → 18/20"
    slide = new_slide(prs, title, "固定与优化各 20 步 × 两模型 = 80 次真实调用，全部成功")
    x, y = 0.65, 2.2
    widths = [2.05, 1.62, 1.62, 1.78]
    headers = ["业务场景", "Decider 固定", "Decider 优化", "GPT 固定 / 优化"]
    rows = []
    for name, label in [
        ("pending", "USDC 未到账"),
        ("underpaid", "USDT 少付"),
        ("refund", "稳定币退款"),
        ("settlement", "商户结算"),
    ]:

        def count(profile, provider, case_name=name):
            return sum(
                s[provider]["matches_reference"]
                for s in case_for(data, profile, case_name)["steps"]
            )

        rows.append(
            [
                label,
                f"{count('baseline', 'decider')}/5",
                f"{count('optimized', 'decider')}/5",
                f"{count('baseline', 'bedrock')}/5 → {count('optimized', 'bedrock')}/5",
            ]
        )
    for index, cells in enumerate([headers, *rows]):
        xx = x
        for col, (value, width) in enumerate(zip(cells, widths, strict=True)):
            fill = L.NAVY if index == 0 else (L.WHITE if index % 2 else "EDF2F6")
            L.rect(slide, xx, y + index * 0.49, width - 0.035, 0.46, fill)
            L.text(
                slide,
                value,
                xx + 0.09,
                y + index * 0.49 + 0.11,
                width - 0.2,
                0.25,
                12.2 if index == 0 else 15,
                L.WHITE if index == 0 else L.INK if col == 0 else L.BLUE if col == 3 else L.TEAL,
                index == 0 or col > 0,
            )
            xx += width
    icon(slide, "timer", 8.16, 2.24)
    L.text(slide, "累计后端耗时 · 固定 → 优化", 8.64, 2.23, 4.0, 0.34, 16, L.INK, True)
    for i, (provider, label, color) in enumerate(
        [("decider", "Decider", L.TEAL), ("bedrock", "GPT", L.BLUE)]
    ):
        before = data["profiles"]["baseline"]["summary"][provider]["backend_total_ms"] / 1000
        after = data["profiles"]["optimized"]["summary"][provider]["backend_total_ms"] / 1000
        L.text(slide, label, 8.18, 2.87 + i * 0.57, 1.21, 0.3, 15, color, True)
        L.text(
            slide,
            f"{before:.3f} → {after:.3f} s",
            9.52,
            2.87 + i * 0.57,
            3.06,
            0.32,
            17,
            color,
            True,
        )
    L.text(slide, "Decider 保护调整：4 次 → 1 次", 8.18, 4.08, 4.4, 0.28, 14, L.INK, True)
    L.text(slide, "仍有差异：少付核验、退款审批", 8.18, 4.49, 4.4, 0.28, 13, L.ORANGE)
    L.text(
        slide,
        "真实界面 · USDC 优化流程的五步累计耗时（单场景）",
        0.67,
        4.92,
        12,
        0.24,
        11.5,
        L.MUTED,
    )
    L.picture(
        slide, OUT / "assets/optimized-pending-timings.png", 0.65, 5.25, 12.02, 1.39, frame=True
    )
    L.text(
        slide,
        "参考一致不是通用准确率；每配置仅采集一次，耗时受调用顺序与缓存影响。",
        0.65,
        6.77,
        12,
        0.25,
        12.5,
        L.MUTED,
    )
    slide.notes_slide.notes_text_frame.text = (
        "本页独立呈现2026-10-11 EC2 g6.xlarge常驻环境的优化实验，不与之前SageMaker图表合并。\n"
        "四场景×五步骤×两输入配置×两提供方=80次真实模型调用，全部成功。"
        "每场景先固定流程后优化流程；两者初始业务证据相同，字段和候选描述不同。"
        "同一种配置内部两模型收到相同业务输入，但不同配置不能算同输入配对。\n"
        "Decider参考一致：USDC 4→5、少付4→4、退款2→4、结算4→5，合计14/20→18/20。"
        "GPT两配置均20/20。参考来自演示策略，不是行业统一答案或通用准确率。\n"
        "Decider执行保护调整4→1。优化后少付核验误选await_confirmations，"
        "退款审批过早选择approved，后端仍改为验证退款地址。不能隐藏剩余错误，"
        "也不能把系统保护记成模型答对。四个初始场景资金条件均未补齐，业务仍待跟进。\n"
        "20次后端累计：Decider 2059.338→1577.990ms；GPT 18694.280→23174.244ms。"
        "包含通信与排队，不是并发墙钟时间或GPU纯推理。没有随机化、重复采样或置信区间，"
        "顺序和缓存可能影响速度；不能保证优化一定提速，GPT本批次反而更慢。\n"
        "下方截图是USDC优化流程的真实界面，仅展示该场景5步累计："
        "Decider390.9ms、GPT9.77s、比值25.0×，不是上方四场景合计。\n"
        "来源：docs/results/optimized-2026-10-11.json；截图来源与裁剪坐标见capture-manifest.json。"
        "真实采集已在上一轮完成；本次编辑PPT没有重新调用模型。"
    )
    return {"title": title, "section": "05 / 优化流程实测 · 结果", "dark": False}


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    before_hash = digest(BASE)
    data = json.loads(EVIDENCE.read_text())
    capture = prepare_images()
    prs = Presentation(BASE)
    metadata = copy.deepcopy(json.loads((SOURCE / "build-manifest.json").read_text())["slides"])
    assert len(prs.slides) == 25
    new_metadata = [method_slide(prs, data), results_slide(prs, data)]
    # Insert the two new slides before the three closing slides.
    ids = prs.slides._sldIdLst
    added = list(ids)[-2:]
    for sid in added:
        ids.remove(sid)
    for index, sid in enumerate(added, 22):
        ids.insert(index, sid)
    metadata[22:22] = new_metadata
    metadata[0]["section"] = "决策模型与支付业务 / v6 · 优化实测 · 27 页"
    metadata[25]["section"] = "06 / 结论"
    metadata[26]["section"] = "06 / 演示与复现"

    for number, slide in enumerate(prs.slides, 1):
        numbered = False
        for shape in list(slide.shapes):
            if not shape.has_text_frame:
                continue
            if re.fullmatch(r"\d{2} / 25", shape.text):
                replace_text(shape, shape.text, f"{number:02} / 27")
                numbered = True
            replacements = {
                "v5 · EC2 部署版 · 25 页": "v6 · 优化实测 · 27 页",
                "05 / 结论": "06 / 结论",
                "05 / 演示与复现": "06 / 演示与复现",
            }
            if number == 26:
                replacements |= {
                    "历史实测": "本次优化",
                    "Decider 决策更快": "优化后参考一致 18/20",
                    "选择差异与失败同样可见": "仍有错误，执行校验不可省",
                }
            for old, new in replacements.items():
                if old in shape.text:
                    replace_text(shape, old, new)
        if not numbered:
            L.text(
                slide,
                f"{number:02} / 27",
                11.76,
                7.14,
                0.92,
                0.23,
                10,
                L.MUTED,
                align=PP_ALIGN.RIGHT,
            )
        if number not in (23, 24):
            slide.notes_slide.notes_text_frame.text += (
                "\n\nv6补充：第23–24页新增2026-10-11 EC2优化流程的独立实测，"
                "输入字段与候选描述优化，未修改模型权重。原2026-10-10数据仍保留原采集口径。"
                "新增页不能用于把旧数据改标为EC2测速；优化后参考一致18/20也不是通用准确率。"
                f"来源：{EVIDENCE.relative_to(ROOT)}。Demo：{DEMO}；Repo：{REPO}。"
            )
        for effect in slide._element.xpath(".//a:effectRef"):
            effect.set("idx", "0")
        visible = "\n".join(s.text for s in slide.shapes if s.has_text_frame)
        assert "历史实测：2026-10-10 · 迁移前的托管部署" not in visible
        for shape in slide.shapes:
            assert min(shape.left, shape.top) >= 0, (number, shape.name)
            assert shape.left + shape.width <= prs.slide_width + 10, (number, shape.name)
            assert shape.top + shape.height <= prs.slide_height + 10, (number, shape.name)
    assert len(prs.slides) == 27
    prs.core_properties.title = "Strands Decider & Jev · 优化流程实测版 · 27页"
    prs.core_properties.subject = "EC2架构、决策对比与四场景输入优化实测"
    prs.core_properties.comments = (
        "v6 adds two slides using the observed 2026-10-11 EC2 optimization batch. "
        "Earlier 2026-10-10 measurements retain their original environment in notes."
    )
    target = OUT / f"{NAME}.pptx"
    prs.save(target)
    for path in (SOURCE / "assets").glob("*-qr.png"):
        shutil.copyfile(path, OUT / "assets" / path.name)
    shutil.copyfile(SOURCE / "source-manifest.json", OUT / "source-manifest.json")
    (OUT / "讲稿与来源.md").write_text(
        "# Strands Decider · 优化流程实测版 v6 · 27页讲稿\n\n"
        + "\n\n".join(
            f"## {i:02}. {meta['title']}\n\n{slide.notes_slide.notes_text_frame.text}"
            for i, (slide, meta) in enumerate(zip(prs.slides, metadata, strict=True), 1)
        )
        + "\n"
    )
    capture["observed_at"] = case_for(data, "optimized", "pending")["captured_at"]
    capture["evidence_sha256"] = digest(EVIDENCE)
    (OUT / "capture-manifest.json").write_text(
        json.dumps(capture, ensure_ascii=False, indent=2) + "\n"
    )
    manifest = {
        "generated_at": datetime.now(UTC).isoformat(),
        "slides": metadata,
        "base_file": str(BASE.relative_to(ROOT)),
        "base_sha256": before_hash,
        "pptx_sha256": digest(target),
        "added_slides": [23, 24],
        "previous_to_current": {str(i): i if i <= 22 else i + 2 for i in range(1, 26)},
        "evidence_file": str(EVIDENCE.relative_to(ROOT)),
        "evidence_sha256": digest(EVIDENCE),
        "new_model_calls_during_presentation_edit": False,
        "observed_model_calls_in_source_batch": 80,
        "model_weights_changed": False,
        "original_measurement_values_retained": True,
        "demo_url": DEMO,
        "repo_url": REPO,
    }
    (OUT / "build-manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n"
    )
    assert digest(BASE) == before_hash
    print(f"Created {target.relative_to(ROOT)} — 27 slides, two observed optimization slides.")


if __name__ == "__main__":
    main()
