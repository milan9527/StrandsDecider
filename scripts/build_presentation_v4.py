"""Reorganize the original narrative around decision models, evidence and conclusions."""

import argparse
import copy
import hashlib
import io
import json
from datetime import UTC, datetime
from pathlib import Path

import cairosvg
import presentation_layout as L
import qrcode
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/presentations/v4"
V3 = ROOT / "docs/presentations/v3"
V2 = ROOT / "docs/presentations/v2"
ORIGINAL = ROOT / "artifacts/presentation/Strands-Decider-模型原理与支付业务演示.pptx"
V3_DECK = V3 / "Strands-Decider-模型原理与支付业务演示-v3-25页-20261010.pptx"
NAME = "Strands-Decider-模型原理与支付业务演示-v4-决策模型与支付对比-25页-20261010"
BLOG = "https://strandsagents.com/blog/introducing-strands-decider/"
JEV = "https://docs.typesafe.ai/concepts/system-one"
JEV_LAUNCH = "https://typesafe.ai/blog/introducing-system-one-models-and-jev"
DEMO = "https://d3i7ogfx5ml7dc.cloudfront.net/#crypto"
REPO = "https://github.com/milan9527/StrandsDecider"
TEAL, BLUE, NAVY, PURPLE = "087F72", "4669C4", "132639", "8055A4"
INK, MUTED, WHITE, LINE = L.INK, L.MUTED, L.WHITE, L.LINE


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--base",
        type=Path,
        default=ORIGINAL,
        help="The original 25-slide deck provided by the project owner.",
    )
    args = parser.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "assets").mkdir(exist_ok=True)
    prs, v3 = Presentation(args.base), Presentation(V3_DECK)
    v3_meta = json.loads((V3 / "build-manifest.json").read_text())["slides"]
    loops = json.loads((V2 / "loop-observations.json").read_text())
    bench_path = ROOT / "docs/results/benchmark-2026-10-10.json"
    bench = json.loads(bench_path.read_text())
    for sid in list(prs.slides._sldIdLst):
        prs.part.drop_rel(sid.rId)
        prs.slides._sldIdLst.remove(sid)
    metadata, image_cache = [], {}
    sources = (
        "\n\n来源与口径：\n"
        f"原始文稿：{args.base.name}；SHA256 {digest(args.base)}\n"
        f"Strands 官方介绍：{BLOG}\n"
        "官方 v19：https://huggingface.co/StrandsAgents/strands-decider-2B-hobson-v19\n"
        "判断实现：https://github.com/strands-labs/strands-decider/tree/"
        "63d24ae286e50105fba0bd1db15b0e243aea4650/src/strands_decider\n"
        f"Jev / System One 官方文档：{JEV}\n"
        f"Jev 发布介绍：{JEV_LAUNCH}\n"
        "BERT：https://aclanthology.org/N19-1423/；"
        "NLI 候选评分：https://aclanthology.org/D19-1404/\n"
        f"Demo：{DEMO}\nRepo：{REPO}\n"
        "固定输入证据：docs/results/benchmark-2026-10-10.json\n"
        "闭环证据：docs/presentations/v2/loop-observations.json\n"
        "截图来源：docs/presentations/v2/capture-manifest.json；"
        "裁剪来源：docs/presentations/v3/build-manifest.json\n"
        "数据：2026-10-10 12:26 UTC 起，crypto10 / API 20261010-loop1；"
        "本次仅编辑演示文稿，没有重新调用模型。\n"
        "Decider：官方 v19 / SageMaker ml.g6.xlarge / L4；"
        "GPT：Bedrock us.openai.gpt-6.1-sol / reasoning low。\n"
        "时间口径：backend_total_ms 包含服务通信与排队；"
        "不含外部事件等待，非浏览器全程用时。Jev、BERT 未参与本项目实测。\n"
        "业务与资金是模拟案例；客服回复来自应用模板。"
    )

    def text(s, value, x, y, w, h=0.4, size=18, color=INK, bold=False, align=None):
        return L.text(s, str(value), x, y, w, h, size, color, bold, align)

    def icon(s, name, x, y, size=0.45, color=TEAL):
        key = (name, color)
        if key not in image_cache:
            svg = (V3 / "icons" / f"{name}.svg").read_text().replace("currentColor", f"#{color}")
            image_cache[key] = cairosvg.svg2png(
                bytestring=svg.encode(), output_width=256, output_height=256
            )
        sh = s.shapes.add_picture(
            io.BytesIO(image_cache[key]), Inches(x), Inches(y), Inches(size), Inches(size)
        )
        sh.name = "ICON: Lucide " + name

    def new(title, section, subtitle="", notes="", dark=False, original_pages=()):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = L.rgb(NAVY if dark else L.BG)
        text(
            s,
            f"STRANDS DECIDER   /   {section}",
            0.6,
            0.27,
            12,
            0.25,
            10,
            "86D7C6" if dark else TEAL,
            True,
        )
        text(s, title, 0.6, 0.79, 12.1, 0.69, 30, WHITE if dark else INK, True)
        if subtitle:
            text(s, subtitle, 0.62, 1.56, 12, 0.56, 15, "B8C9DB" if dark else MUTED)
        s.notes_slide.notes_text_frame.text = notes + sources
        metadata.append(
            {
                "title": title,
                "section": section,
                "dark": dark,
                "original_narrative_pages": list(original_pages),
                "reused_v3_slide": None,
            }
        )
        return s

    def keep(number, original_pages=()):
        source = v3.slides[number - 1]
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        meta = v3_meta[number - 1]
        s.background.fill.fore_color.rgb = L.rgb(NAVY if meta["dark"] else L.BG)
        remap = {}
        for rel in source.part.rels.values():
            if rel.reltype == RT.IMAGE:
                _, rid = s.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
                remap[rel.rId] = rid
        for shape in source.shapes:
            if shape.top >= Inches(7.03):
                continue
            node = copy.deepcopy(shape._element)
            for child in node.iter():
                for key, value in list(child.attrib.items()):
                    if value in remap and key.startswith(
                        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
                    ):
                        child.set(key, remap[value])
            s.shapes._spTree.insert_element_before(node, "p:extLst")
        s.notes_slide.notes_text_frame.text = (
            source.notes_slide.notes_text_frame.text.split("\n\n来源与口径：")[0] + sources
        )
        metadata.append(
            {
                "title": meta["title"],
                "section": meta["section"],
                "dark": meta["dark"],
                "original_narrative_pages": list(original_pages),
                "reused_v3_slide": number,
            }
        )
        return s

    def note(s, value, dark=False):
        text(s, value, 0.62, 6.65, 12.05, 0.31, 11.5, "B8C9DB" if dark else MUTED)

    def pill(s, value, x, y, w, color=TEAL, fill=L.PALE_TEAL):
        L.rect(s, x, y, w, 0.35, fill, radius=True)
        text(s, value, x + 0.05, y + 0.055, w - 0.1, 0.25, 11, color, True, PP_ALIGN.CENTER)

    def card(s, x, y, w, h, title, body="", glyph=None, color=TEAL, fill=WHITE):
        L.rect(s, x, y, w, h, fill, LINE, True)
        if glyph:
            icon(s, glyph, x + 0.22, y + 0.22, 0.43, color)
        title_x = 0.82 if glyph else 0.23
        text(s, title, x + title_x, y + 0.23, w - title_x - 0.20, 0.43, 19, color, True)
        if body:
            text(s, body, x + 0.23, y + 0.93, w - 0.46, h - 1.02, 18)

    def arrow(s, points, color=MUTED):
        L.line(s, points, color, 1.8)
        ln = s.shapes[-1]._element.spPr.find(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}ln"
        )
        tip = OxmlElement("a:tailEnd")
        tip.set("type", "triangle")
        ln.append(tip)

    def flow(s, items, x=0.65, y=3.0, w=12.0, color=TEAL, height=1.6):
        gap = 0.31
        width = (w - gap * (len(items) - 1)) / len(items)
        for i, (title, body, glyph) in enumerate(items):
            xx = x + i * (width + gap)
            L.rect(s, xx, y, width, height, WHITE, LINE, True)
            icon(s, glyph, xx + 0.17, y + 0.16, 0.38, color)
            text(s, title, xx + 0.67, y + 0.17, width - 0.8, 0.4, 17, color, True)
            text(s, body, xx + 0.17, y + 0.75, width - 0.34, height - 0.82, 15)
            if i < len(items) - 1:
                arrow(
                    s,
                    [
                        (xx + width + 0.03, y + height / 2),
                        (xx + width + gap - 0.04, y + height / 2),
                    ],
                    color,
                )

    def link(s, label, url, x, y, w, color=BLUE, size=14):
        sh = text(s, label, x, y, w, 0.46, size, color, True)
        sh.text_frame.paragraphs[0].runs[0].hyperlink.address = url
        return sh

    # 01 — Recover the original central theme.
    s = new(
        "让小模型做好高频决策",
        "决策模型与支付业务 / v4 · 25 页",
        "Strands Decider & Jev：从模型范式，到 Decider × GPT 的真实演示",
        "参照原始25页文稿的模型目的、原理、场景、实测、结论主线重新组织。"
        "开头加入System One/Jev背景；闭环仅保留两页。仍为25页，BERT一页，无GPU提示。"
        "Jev只作模型类别与产品介绍，不参与本项目性能比较。",
        dark=True,
        original_pages=(1,),
    )
    for x, number, title, body, glyph in [
        (0.7, "01", "认识决策模型", "Strands Decider / Jev\n面向软件的有限判断", "list-filter"),
        (4.87, "02", "走进支付演示", "同一业务输入\n选择、速度与执行结果", "workflow"),
        (9.04, "03", "形成应用结论", "哪些节点适合使用\n怎样评估真实收益", "target"),
    ]:
        text(s, number, x, 2.58, 1.1, 0.65, 35, "567B99", True)
        icon(s, glyph, x, 3.63, 0.48, "86D7C6")
        text(s, title, x + 0.67, 3.65, 3.10, 0.47, 23, WHITE, True)
        text(s, body, x, 4.51, 3.6, 1.1, 19, "DCE6F0")
    note(s, "模型介绍 → 工作原理 → 支付对比 → 实测结论 · 2026.10.10", True)

    # 02 — Motivation, before terminology.
    s = new(
        "软件流程，需要大量小而明确的判断",
        "01 / 为什么需要决策模型",
        "选哪个队列？调用哪个工具？证据是否足够继续？",
        "从原版第二页的高频重复、结构化输出和嵌入流程展开。这里的决策模型"
        "处理语义判断，并不代替确定性的金额运算、权限校验或业务状态机。",
        original_pages=(2,),
    )
    flow(
        s,
        [
            ("接收请求", "客户消息 / 当前状态", "receipt-text"),
            ("选择路径", "队列 / 工具 / 下一步", "route"),
            ("检查结果", "是否完成 / 是否升级", "clipboard-check"),
        ],
        y=2.44,
        height=1.62,
    )
    card(
        s,
        0.65,
        4.61,
        5.86,
        1.64,
        "这些判断的共同点",
        "候选明确、频繁调用、结果被代码消费",
        "list-filter",
    )
    card(
        s,
        6.80,
        4.61,
        5.86,
        1.64,
        "决策模型的目标",
        "用受约束的输出，降低判断开销",
        "timer",
        BLUE,
        L.PALE_BLUE,
    )
    note(s, "自然语言理解与路径选择交给模型；金额、权限与执行前提由代码核验。")

    # 03 — The decision-model interface.
    s = new(
        "决策模型：状态输入，类型化答案输出",
        "01 / Decision Models · System One",
        "把一个复杂流程拆成明确问题，再在应用代码中组合答案。",
        "Strands官方把这种模型称为decision models或system one models。"
        "TypeSafe文档定义System One为对state返回typed answers与probabilities的模型。"
        "Choice、Noul、Score在两方材料中均有介绍。此页不展示虚构模型运行数字。"
        "输出结构受约束不代表判断一定正确；校准需在一组预测上评估，不能保证单条答案。",
        original_pages=(2, 4),
    )
    pill(s, "共享输入：业务状态 + 问题 + 候选或判断条件", 1.44, 2.22, 10.4)
    for x, title, question, answer, glyph in [
        (0.65, "Choice · 多选一", "该交给哪个客服队列？", "支付 / 退款 / 商户运营", "route"),
        (4.78, "Noul · 是 / 否", "是否表达了退款诉求？", "返回 P(true)", "circle-check"),
        (
            8.91,
            "Score · 有序评分",
            "这条请求有多紧急？",
            "在定义好的等级上评分",
            "chart-no-axes-combined",
        ),
    ]:
        card(s, x, 2.93, 3.78, 2.72, title, f"{question}\n\n{answer}", glyph, TEAL, WHITE)
    note(s, "输出类型受约束 ≠ 业务判断必然正确；概率与置信度需要结合业务数据验证。")

    # 04 — Jev and Decider are related directions, not the same product.
    s = new(
        "Jev 与 Strands Decider：同一方向，两种实现",
        "01 / 代表性模型",
        "共同强调快速、结构化的判断，以及可供程序使用的概率信息。",
        "Jev是TypeSafe的旗舰System One模型，通过官方API/SDK调用。官方强调"
        "RLCD（Reinforcement Learning for Calibrated Decisions）、typed primitives与"
        "并行问题评估。Strands Decider是Strands Labs的开源2B实验模型，公开代码和权重，"
        "本项目使用v19自行部署。不能因此断言二者内部架构相同、API完全兼容、"
        "或Decider和Jev有相同性能。没有采用Jev营销页面的速度/成本倍数。",
        original_pages=(2, 3),
    )
    card(
        s,
        0.65,
        2.50,
        5.85,
        3.30,
        "Jev / TypeSafe",
        "System One 模型 · 官方 API / SDK\n强调校准决策训练（RLCD）\n类型化输出与多问题并行评估",
        "brain-circuit",
        PURPLE,
        "F3EEFA",
    )
    card(
        s,
        6.80,
        2.50,
        5.85,
        3.30,
        "Strands Decider / Strands Labs",
        "开源 2B 决策模型 · 代码与权重公开\n轻量模型主体 + Pointer Head\n面向实验，可本地或自行部署",
        "list-filter",
        TEAL,
        L.PALE_TEAL,
    )
    note(s, "Jev 为产品背景介绍；本项目实测的是 Strands Decider 与 GPT，没有接入 Jev。")

    # 05 — Restore the original architecture explanation.
    s = new(
        "Strands Decider：用候选评分头完成选择",
        "02 / 模型原理",
        "官方 v19：Qwen3.5-2B-Base + rank-16 LoRA + Pointer Head",
        "沿用原版第三页的核心机制。模型移除语言模型词表生成头，使用答案位置"
        "和候选位置的隐藏表示，经归一化、投影和点积得到候选分数，再计算概率。"
        "候选从请求提供；choice取最高概率项。支持同一state的多个问题共享计算。"
        "单次请求评分的优势不等于整个多步业务只需一次模型调用。",
        original_pages=(3,),
    )
    flow(
        s,
        [
            ("输入业务", "state / question\ncriteria / 候选", "receipt-text"),
            ("模型主体", "Qwen3.5-2B\nLoRA 任务适配", "layers"),
            ("候选评分", "Pointer Head\n候选位置 ↔ 答案位置", "list-filter"),
            ("输出决策", "选择 / 分布\n概率 / 置信信息", "circle-check"),
        ],
        y=2.57,
        height=1.84,
    )
    pill(s, "对当前候选直接评分，减少逐 token 文本生成", 1.44, 5.03, 10.4)
    note(s, "适合明确、聚焦的判断；开放式写作与复杂长链推理，需要不同模型能力。")

    # 06–07 — Keep concise comparisons.
    keep(3, original_pages=(3, 5))
    keep(5, original_pages=(5,))

    # 08 — Applicability, matching the original flow into the demo.
    s = new(
        "适用场景：候选清晰，结果可检查",
        "02 / 从模型走向应用",
        "从 Agent 路由与守卫，到客服、支付与运营判断。",
        "恢复原版第六页的六类应用，依据Strands官方博客。支付只是其中一个"
        "易于观察的演示载体，并非Decider或Jev专门面向金融。模型不负责执行资金操作。",
        original_pages=(6,),
    )
    for i, (title, body, glyph) in enumerate(
        [
            ("模型与工具路由", "选择模型、工具或业务队列", "route"),
            ("工具调用检查", "参数有无依据，是否需澄清", "shield-check"),
            ("客服分流与升级", "识别诉求，跟进或转人工", "headset"),
            ("策略分类与审核", "将自然语言映射到策略", "list-filter"),
            ("RAG 与内容评估", "判断相关性、质量与覆盖", "scan-eye"),
            ("记忆与上下文选择", "决定保留或检索哪些片段", "layers"),
        ]
    ):
        card(s, 0.65 + (i % 3) * 4.13, 2.35 + (i // 3) * 1.88, 3.78, 1.64, title, body, glyph)
    note(s, "接下来用支付业务，把「选择 → 执行 → 结果」放到同一条可观察流程中。")

    # 09–10 — Project and scenarios.
    keep(6, original_pages=(7, 8))
    keep(7, original_pages=(10, 13))

    # 11 — Focus the experiment on three questions.
    s = new(
        "比较方法：同一业务输入，独立执行",
        "03 / 支付对比设计",
        "候选与条件相同；保留两模型的原始选择、耗时和实际业务结果。",
        "参照原版第11页。固定输入每一步双侧输入指纹一致，同步请求；收齐方案"
        "后独立执行。后续反馈模式各侧状态独立，不能把不同输入下的时延混为配对基准。"
        "成功时延含通信排队，不等于浏览器全程时间。当前最新基准19组成功配对，另有1次GPT超时。",
        original_pages=(11,),
    )
    pill(s, "同一业务快照 + 同一问题 + 同一候选条件", 1.44, 2.18, 10.4)
    card(
        s,
        0.65,
        2.94,
        5.85,
        2.14,
        "Decider / SageMaker",
        "返回选择、分布与置信信息\n记录后端决策耗时，独立执行",
        "list-filter",
        TEAL,
        L.PALE_TEAL,
    )
    card(
        s,
        6.80,
        2.94,
        5.85,
        2.14,
        "GPT / Bedrock",
        "返回候选 ID JSON，校验结果\n记录后端决策耗时，独立执行",
        "code-xml",
        BLUE,
        L.PALE_BLUE,
    )
    for x, title, glyph in [
        (0.85, "选择有何差异", "route"),
        (5.03, "每步与累计多快", "timer"),
        (9.18, "业务是否完成", "circle-check"),
    ]:
        icon(s, glyph, x, 5.67, 0.39)
        text(s, title, x + 0.57, 5.68, 3.04, 0.4, 18, INK, True)
    note(s, "速度主统计使用成功同输入配对；反馈执行结果单独观察。")

    # 12–16 — Five consecutive evidence slides tell one comparison story.
    keep(9, original_pages=(12,))
    keep(10, original_pages=(14,))
    keep(11, original_pages=(15, 21))
    keep(12, original_pages=(13,))
    keep(16, original_pages=(16,))

    # 17 — Compress the loop to one explanation and one result page.
    s = new(
        "继续推进：证据更新后，再判断与执行",
        "03 / 反馈执行 · 1 / 2",
        "条件不足时等待；把最新事实和执行结果交给下一次决策。",
        "闭环从v3六页收紧为两页。这里承接原版17至19页推进模拟交易的故事，"
        "反映当前自主反馈模式：每侧独立读状态、判断、执行、反馈，条件缺失时等待事件。"
        "确认数、地址验证和审批来自模拟器，模型不能制造外部证据。无效动作可有限重判，"
        "同节点连续两次无进展转人工；最多20决策/32tick。后续输入可能不同。",
        original_pages=(17, 18, 19),
    )
    flow(
        s,
        [
            ("读取状态", "事实 + 上轮反馈", "receipt-text"),
            ("选择动作", "Decider 或 GPT", "route"),
            ("实际执行", "更新订单与账本", "workflow"),
            ("检查结果", "完成 / 等待 / 再判", "circle-check"),
        ],
        y=2.39,
        height=1.49,
    )
    arrow(s, [(11.14, 3.94), (11.14, 4.37), (1.79, 4.37), (1.79, 3.94)], TEAL)
    L.rect(s, 4.62, 4.14, 4.19, 0.40, L.BG)
    text(s, "最新证据 + 执行反馈", 4.64, 4.20, 4.12, 0.32, 14, TEAL, True, PP_ALIGN.CENTER)
    card(s, 0.65, 4.95, 5.85, 1.38, "USDC 未到账", "模拟确认数：3 / 12 → 12 / 12", color=TEAL)
    card(s, 6.80, 4.95, 5.85, 1.38, "稳定币退款", "模拟验证地址 → 模拟审批通过", color=BLUE)
    note(s, "同一套机制，各侧独立推进；外部事件由模拟器更新。")

    s = new(
        "执行结果：两组业务，双方均完成",
        "03 / 反馈执行 · 2 / 2",
        "完成标准：目标订单、对应账本与工单结案。",
        "USDC每侧6次决策、1事件，Decider570.140ms/GPT7268.536ms；"
        "退款每侧9次决策、2事件，Decider897.332ms/GPT10541.95ms。"
        "下方截图为同次真实退款最终结果。双方都complete=true，但Decider在过程中有"
        "不同原始选择及执行保护调整。成功结果不能消除这些差异。后续状态可能不同，"
        "这一页不是同输入配对基准。",
        original_pages=(18, 19),
    )
    columns = [
        (2.50, "场景"),
        (2.37, "Decider"),
        (2.37, "GPT"),
        (2.35, "每侧决策数"),
        (2.43, "业务结果"),
    ]
    x = 0.65
    for width, label in columns:
        L.rect(s, x, 2.27, width - 0.04, 0.49, NAVY)
        text(s, label, x + 0.12, 2.38, width - 0.24, 0.30, 13, WHITE, True)
        x += width
    for i, (case, label) in enumerate([("pending", "USDC 未到账"), ("refund", "稳定币退款")]):
        data = loops[case]["providers"]
        elapsed = {
            p: sum(d["timings"]["backend_total_ms"] for d in data[p]["decisions"]) / 1000
            for p in ("decider", "bedrock")
        }
        row = [
            label,
            f"{elapsed['decider']:.3f} s",
            f"{elapsed['bedrock']:.3f} s",
            f"{len(data['decider']['decisions'])} 次",
            "双方完成",
        ]
        x, y = 0.65, 2.83 + i * 0.70
        for j, ((width, _), value) in enumerate(zip(columns, row, strict=True)):
            L.rect(s, x, y, width - 0.04, 0.64, WHITE if i == 0 else "ECF1F6")
            text(
                s,
                value,
                x + 0.12,
                y + 0.17,
                width - 0.24,
                0.35,
                17,
                TEAL if j == 1 else BLUE if j == 2 else INK,
                j in (1, 2),
            )
            x += width
    text(s, "退款案例 · 实际执行结果", 0.68, 4.27, 11.97, 0.27, 12, TEAL, True)
    L.picture(s, V2 / "screenshots/refund-loop-final.png", 0.68, 4.58, 11.97, 1.56, frame=True)
    note(s, "后端累计决策耗时，不含外部事件等待；两侧后续输入可能不同，不计入同输入速度基准。")

    # 19–22 — Results retain the real successful-pair basis and failure disclosure.
    keep(13, original_pages=(20,))
    keep(15, original_pages=(20,))
    keep(14, original_pages=(20, 22))
    s = new(
        "本轮对比：速度、规则选择与实际执行",
        "04 / 如何理解测试结果",
        "四个案例 × 五个步骤，每侧调用 20 次；各项数字衡量不同的事情。",
        "本轮固定输入测试每侧尝试20次。Decider成功返回20次，GPT成功返回19次，"
        "其商户结算execute请求出现ReadTimeoutError。成功返回是获得有效选择，"
        "并非业务完成。速度与规则符合度只比较双方都成功的19组同输入请求；"
        "GPT失败对应的Decider成功请求也排除，保证双方使用相同记录。"
        "19组累计后端耗时为Decider1809.597ms、GPT22068.204ms，未包含失败等待。"
        "与开发者业务参考一致的次数为13/19与19/19，这不是通用准确率。"
        "例如普通收款无风险标记，无需退款人工审批，预设规则对应approved；"
        "该次Decider选择await_approval，GPT选择approved。"
        "商户结算中Decider处理五步后仍待跟进；GPT虽取得其他四步选择，"
        "但缺少完整方案，因此该侧本次未进入沙盒执行。"
        "本页分别解释调用成功、规则选择、实际执行；所有结论仅描述这一轮观测。",
        original_pages=(20, 21, 22),
    )
    totals = bench["totals"]
    paired = totals["paired_steps"]
    attempts = sum(len(case["steps"]) for case in bench["cases"])
    successful = {
        provider: attempts - totals["failed_requests"][provider]
        for provider in ("decider", "bedrock")
    }
    columns = [(2.35, "比较什么"), (2.25, "Decider"), (2.25, "GPT"), (5.17, "怎么理解")]
    x, y = 0.65, 2.15
    for width, label in columns:
        L.rect(s, x, y, width - 0.04, 0.47, NAVY)
        text(s, label, x + 0.12, y + 0.10, width - 0.24, 0.30, 14, WHITE, True)
        x += width
    rows = [
        (
            "累计决策耗时",
            f"{totals['decider']['backend_ms'] / 1000:.3f} 秒",
            f"{totals['bedrock']['backend_ms'] / 1000:.3f} 秒",
            "19 组同输入成功配对的累计时间\n失败等待未计入这两个数字",
            0.68,
        ),
        (
            "调用是否成功",
            f"{successful['decider']} / {attempts} 次",
            f"{successful['bedrock']} / {attempts} 次",
            "GPT 有 1 次读取超时\n返回了选择，不代表业务已经完成",
            0.68,
        ),
        (
            "与业务规则一致",
            f"{totals['decider']['reference_matches']} / {paired} 次",
            f"{totals['bedrock']['reference_matches']} / {paired} 次",
            "选择与开发者预设规则一致的次数\n这不是通用准确率",
            0.68,
        ),
        (
            "商户结算执行",
            "处理完 5 步\n结算仍待跟进",
            "方案未执行\n处置步骤超时",
            "GPT 其余 4 步虽返回选择\n但缺少完整方案，本次未进入沙盒执行",
            0.83,
        ),
    ]
    y += 0.53
    for i, row in enumerate(rows):
        x, height = 0.65, row[-1]
        for j, ((width, _), value) in enumerate(zip(columns, row[:4], strict=True)):
            fill = WHITE if i % 2 == 0 else "ECF1F6"
            L.rect(s, x, y, width - 0.04, height - 0.035, fill)
            text(
                s,
                value,
                x + 0.12,
                y + 0.09,
                width - 0.24,
                height - 0.12,
                13.5 if j == 3 else 15.5 if i == 3 else 17,
                TEAL if j == 1 else BLUE if j == 2 else INK,
                j in (1, 2),
            )
            x += width
        y += height
    L.rect(s, 0.65, 5.66, 12.0, 0.42, L.PALE_TEAL, radius=True)
    text(
        s,
        "为什么比 19 组？20 组尝试 − 1 组 GPT 超时；双方使用同一组有效记录。",
        0.80,
        5.75,
        11.70,
        0.27,
        13,
        TEAL,
        True,
    )
    text(
        s,
        "规则例：普通收款无风险标记，应选「检查通过」；Decider 那次选「等待人工审批」，GPT 选「检查通过」。",
        0.67,
        6.20,
        11.98,
        0.33,
        12,
        INK,
    )
    note(s, "仅描述本轮测试；决策速度、规则符合度和业务完成是不同指标。")

    # 23 — Where to use the model.
    keep(24, original_pages=(6, 24))

    # 24 — Summarize the entire narrative, rather than another loop explanation.
    s = new(
        "总结：用合适的模型，完成明确的判断",
        "05 / 结论",
        "从 Jev 与 Decider 的模型方向，回到本项目的真实观察。",
        "原始24页文稿以模型价值、项目启示、下一步评估结尾。本版总结模型方向、"
        "本次实测及应用落地三层：决策模型面向有限判断，Decider在本次19组成功同输入"
        "配对中累计耗时低于GPT，但原始选择有差异。外部证据和执行校验仍由业务系统"
        "负责；不能把Jev官方营销结果或BERT推断混为本项目实测。下一步应按真实业务分布"
        "重复评估准确性、校准、吞吐与整体成本。",
        dark=True,
        original_pages=(24,),
    )
    for x, title, body, glyph in [
        (0.70, "模型方向", "有限输出，聚焦判断\nJev 与 Decider 的共同方向", "list-filter"),
        (4.87, "本次实测", "Decider 决策更快\n选择差异与失败同样可见", "timer"),
        (9.04, "应用落地", "模型给建议，系统做校验\n以真实业务结果检验收益", "shield-check"),
    ]:
        icon(s, glyph, x, 2.77, 0.58, "86D7C6")
        text(s, title, x, 3.72, 3.56, 0.55, 25, WHITE, True)
        text(s, body, x, 4.56, 3.56, 1.18, 19, "DCE6F0")
    note(s, "下一步：业务数据、重复采样、概率校准、吞吐与整体成本。", True)

    # 25 — Explicit, clickable links and scannable QR codes.
    s = new(
        "打开 Demo，查看代码与实测记录",
        "05 / 演示与复现",
        "从同一业务输入开始，观察两模型的选择、速度与执行结果。",
        "结尾给出用户要求的Demo和Repo完整可点击地址，二维码内容与地址一致。"
        "Demo需要现有用户名登录；GitHub仓库为私有仓库，需要访问权限。"
        "不在文稿中嵌入密码或GitHub令牌。官方资料和证据来源完整保存在各页备注及讲稿。",
        original_pages=(23, 25),
    )
    for yy, label, desc, url, glyph, filename in [
        (2.30, "在线 Demo", "加密支付业务台 · 使用现有账号登录", DEMO, "workflow", "demo-qr.png"),
        (
            4.34,
            "GitHub Repo",
            "架构、代码、演示资料与实测记录 · 需仓库权限",
            REPO,
            "code-xml",
            "repo-qr.png",
        ),
    ]:
        L.rect(s, 0.65, yy, 12.0, 1.74, WHITE, LINE, True)
        icon(s, glyph, 0.92, yy + 0.23, 0.46)
        text(s, label, 1.56, yy + 0.21, 8.18, 0.45, 22, TEAL, True)
        link(s, url, url, 0.94, yy + 0.82, 9.72, size=16)
        text(s, desc, 0.94, yy + 1.34, 9.72, 0.29, 12, MUTED)
        qr = qrcode.QRCode(
            version=None, error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=12, border=4
        )
        qr.add_data(url)
        qr.make(fit=True)
        path = OUT / "assets" / filename
        qr.make_image(fill_color=f"#{NAVY}", back_color="white").save(path)
        sh = s.shapes.add_picture(
            str(path), Inches(10.77), Inches(yy + 0.12), Inches(1.50), Inches(1.50)
        )
        sh.click_action.hyperlink.address = url
        sh.name = f"QR: {label}"
    note(s, "官方来源：Strands Decider 博客 / TypeSafe System One 文档 · 详见逐页讲稿。")

    assert len(prs.slides) == 25
    for i, (s, meta) in enumerate(zip(prs.slides, metadata, strict=True), 1):
        color = "8CA6BF" if meta["dark"] else MUTED
        L.line(s, [(0.62, 7.04), (12.68, 7.04)], "365168" if meta["dark"] else LINE, 0.6)
        text(s, "决策模型 → 支付对比 → 应用结论", 0.64, 7.16, 9.9, 0.21, 8.5, color)
        text(s, f"{i:02} / 25", 11.76, 7.14, 0.92, 0.23, 10, color, align=PP_ALIGN.RIGHT)
        for effect in s._element.xpath(".//a:effectRef"):
            effect.set("idx", "0")
        for sh in s.shapes:
            assert sh.left >= -10 and sh.top >= -10, (i, sh.name)
            assert sh.left + sh.width <= prs.slide_width + 10, (i, sh.name, "width")
            assert sh.top + sh.height <= prs.slide_height + 10, (i, sh.name, "height")
        assert "GPU" not in "\n".join(sh.text for sh in s.shapes if sh.has_text_frame)
    prs.core_properties.title = "Strands Decider & Jev · 决策模型与支付业务演示 · 25页"
    prs.core_properties.subject = "决策模型介绍、Decider与GPT实测、执行结果与应用结论"
    prs.core_properties.comments = (
        "Reorganized from the original 25-slide deck. Jev is background only, not benchmarked. "
        "Two feedback slides. Dated production evidence; clickable demo and repository links."
    )
    target = OUT / f"{NAME}.pptx"
    prs.save(target)
    (OUT / "讲稿与来源.md").write_text(
        "# 决策模型与支付业务演示 v4 · 25 页讲稿\n\n"
        + "\n\n".join(
            f"## {i:02}. {m['title']}\n\n{s.notes_slide.notes_text_frame.text}"
            for i, (s, m) in enumerate(zip(prs.slides, metadata, strict=True), 1)
        )
        + "\n"
    )
    (OUT / "build-manifest.json").write_text(
        json.dumps(
            {
                "generated_at": datetime.now(UTC).isoformat(),
                "slides": metadata,
                "base_file": args.base.name,
                "base_sha256": digest(args.base),
                "reused_file": V3_DECK.relative_to(ROOT).as_posix(),
                "reused_sha256": digest(V3_DECK),
                "pptx_sha256": digest(target),
                "benchmark_file": bench_path.relative_to(ROOT).as_posix(),
                "benchmark_sha256": digest(bench_path),
                "loop_sha256": digest(V2 / "loop-observations.json"),
                "source_capture_manifest": "docs/presentations/v2/capture-manifest.json",
                "source_crop_manifest": "docs/presentations/v3/build-manifest.json",
                "capture_date": "2026-10-10",
                "new_model_calls": False,
                "benchmark_models": ["Strands Decider v19", "Bedrock GPT-6.1 Sol"],
                "jev_benchmarked": False,
                "bert_benchmarked": False,
                "feedback_slides": [17, 18],
                "bert_slides": [7],
                "successful_pairs": bench["totals"]["paired_steps"],
                "failed_requests": bench["totals"]["failed_requests"],
                "demo_url": DEMO,
                "repo_url": REPO,
                "qr_files": {p.name: digest(p) for p in (OUT / "assets").glob("*-qr.png")},
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n"
    )
    print(f"Created {target.relative_to(ROOT)} — {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
