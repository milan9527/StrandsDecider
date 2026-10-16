"""Add a reproducible, real-request payment comparison to the v11 deck."""

import json
import statistics
from datetime import UTC, datetime
from pathlib import Path

import build_presentation_v10 as V
import presentation_layout as L
from build_presentation_v5 import replace_text
from build_presentation_v11 import update_text
from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = V.ROOT
OUT = ROOT / "docs/presentations/v12"
BASE = next((ROOT / "docs/presentations/v11").glob("*.pptx"))
NAME = "Strands-Decider-v12-支付流程真实请求与优化对照-35页-20261015"
PROVIDERS = ("decider", "jev", "bedrock")
PROFILES = ("baseline", "optimized")
STAGES = ("intake", "verify", "approve", "execute", "close")
LABELS = {"decider": "Decider", "jev": "Jev", "bedrock": "GPT"}
CASES = {"pending": "USDC 待确认", "underpaid": "USDT 少付", "refund": "重复付款退款", "settlement": "商户结算"}
COLORS = {"decider": L.TEAL, "jev": V.PURPLE, "bedrock": L.BLUE}


def save(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def results(data, case, profile, stage):
    return {p: data["cases"][case][profile]["results"][p][stage] for p in PROVIDERS}


def metrics(data):
    summary = {
        "source": "sources/payment-live.json",
        "timing": "Application backend_total_ms, including queue/provider call; not pure GPU time or browser RTT.",
        "method": "One run per case/profile/provider; 6 warmups excluded. Baseline before optimized; no randomized repeats.",
        "cases": {}, "aggregate": {},
    }
    for case, runs in data["cases"].items():
        summary["cases"][case] = {}
        for profile in PROFILES:
            summary["cases"][case][profile] = {}
            for provider in PROVIDERS:
                rows = list(runs[profile]["results"][provider].values())
                summary["cases"][case][profile][provider] = {
                    "matches": sum(r["matches_reference"] is True for r in rows),
                    "calls": len(rows),
                    "total_backend_ms": sum(r["timings"]["backend_total_ms"] for r in rows),
                    "mean_backend_ms": statistics.mean(r["timings"]["backend_total_ms"] for r in rows),
                    "safety_adjustments": sum(a["safety_adjusted"] for a in runs[profile]["executions"][provider]["audit"]),
                    "business_outcome": runs[profile]["businessOutcomes"][provider],
                }
    for profile in PROFILES:
        summary["aggregate"][profile] = {}
        for provider in PROVIDERS:
            rows = [r for case in data["cases"].values() for r in case[profile]["results"][provider].values()]
            summary["aggregate"][profile][provider] = {
                "matches": sum(r["matches_reference"] is True for r in rows),
                "calls": len(rows),
                "mean_backend_ms": statistics.mean(r["timings"]["backend_total_ms"] for r in rows),
                "mean_browser_rtt_ms": statistics.mean(r["client_ms"] for r in rows),
                "models": sorted({r["model"] for r in rows}),
            }
    return summary


def code(s, value, x, y, w, h, size=14, color=L.INK):
    shape = L.text(s, value, x, y, w, h, size, color)
    for p in shape.text_frame.paragraphs:
        p.space_after = Pt(1)
        p.line_spacing = 1.10
    return shape


def request_notes(data, case, profiles=PROFILES, stages=("approve",)):
    output = []
    for profile in profiles:
        run = data["cases"][case][profile]
        for stage in stages:
            rr = results(data, case, profile, stage)
            payload = {"model_input": rr["decider"]["model_input"]}
            responses = {p: {k: r[k] for k in (
                "provider", "model", "input_hash", "choice", "reference_choice",
                "matches_reference", "timings", "client_ms",
            )} for p, r in rr.items()}
            executions = {p: next(a for a in run["executions"][p]["audit"] if a["stage"] == stage) for p in PROVIDERS}
            output.append(f"\n{case}/{profile}/{stage}\n" + json.dumps({
                **payload, "application_responses": responses, "sandbox_executions": executions,
            }, ensure_ascii=False, indent=2))
    return "\n".join(output)


METHOD = (
    "2026-10-15 03:53–03:54 UTC，真实线上演示同一采集批次。"
    "四场景×两种输入×三模型×五步=120次正式调用，另有6次预热且不计入正式统计。"
    "每个场景先baseline后optimized，同一步三提供方并行，不同节点顺序请求；每组合只跑一次，没有随机顺序或置信区间。"
    "这是输入优化复现实验，不是稳定延迟基准。图中耗时取应用返回timings.backend_total_ms，"
    "包含应用排队和提供方调用；不是纯GPU前向时间，也不是浏览器client_ms。"
    "三模型收到相同model_input（hash相同），实际供应商API封装不同。GPT模型us.openai.gpt-6.1-sol、low effort；"
    "Decider strands-decider-2B-hobson-v19运行在现有g6.xlarge原版服务；Jev响应模型jev-1.13.0。"
    "接口返回节选中的choice来自应用保存的真实模型选择；不是声称展示完整供应商原生HTTP响应。"
    "原始HTTP请求正文、应用响应、完整模型输入、前端执行日志和最终状态均在sources/payment-live.json。"
    "演示参考由业务代码独立给出，不作为答案放入模型输入。模型权重未修改，不是fine-tuning。"
    "业务执行全部为浏览器沙盒，未执行外部资金动作；所有120次正式调用均成功。"
)


def main():
    (OUT / "assets").mkdir(exist_ok=True)
    V.OUT, V.ASSETS, V.SOURCES = OUT, OUT / "assets", OUT / "sources"
    V.EVIDENCE.clear()
    V.SLIDES.clear()
    V.IMAGES.clear()
    data = V.load(V.record(OUT / "sources/payment-live.json"))
    capture = V.load(V.record(OUT / "capture-manifest.json"))
    assert data["complete"]
    summary = metrics(data)
    save(OUT / "sources/payment-comparison.json", summary)
    V.record(OUT / "sources/payment-comparison.json")
    old = V.load(V.record(ROOT / "docs/presentations/v11/build-manifest.json"))
    before_hash = V.digest(BASE)
    prs = Presentation(BASE)
    originals = list(prs.slides)
    ids = list(prs.slides._sldIdLst)
    order, metas = [], []

    def keep(page, changed=False):
        order.append(ids[page - 1])
        meta = dict(old["slides"][page - 1])
        meta.update(page=len(order), source="v11", v11_page=page, body_modified=changed)
        metas.append(meta)
        return originals[page - 1]

    def new(title, subtitle, notes, foot=None):
        s = V.base(prs, "03  支付流程 · 实际请求与结果", title, subtitle,
                   foot or "2026-10-15 实测 · 耗时为应用后端耗时 · 输入与接口返回均为节选，完整请求和执行记录见讲稿与来源。",
                   METHOD + "\n" + notes)
        order.append(prs.slides._sldIdLst[-1])
        meta = dict(V.SLIDES[-1])
        meta.update(page=len(order), source="new")
        metas.append(meta)
        return s

    for page in range(1, 8):
        keep(page, page in (1, 6))
    update_text(originals[0], "v11 · 2026-10-15", "v12 · 2026-10-15")
    update_text(originals[5], "独立 vLLM API 见第 25 页", "独立 vLLM API 见第 28 页")

    # 8: actual baseline input and raw choice excerpts for every provider.
    rr = results(data, "pending", "baseline", "approve")
    mi = rr["decider"]["model_input"]
    state = json.loads(mi["state"])
    s = new("原始固定流程：同一收款审批，三个模型怎样选择",
            "用户：“我已经支付了 129 USDC，钱包显示转账成功，为什么订单还是待支付？”",
            "本页展示实际model_input.state与criteria的节选，以及应用返回的choice节选。"
            "基线approved选项已经包含其他场景无需退款审批的说明；问题是await_approval等其他候选"
            "未同样明确Only refund，且收款state中混入退款字段，不能说原来完全没有非退款豁免规则。"
            + request_notes(data, "pending", ("baseline",)))
    V.card(s, .68, 2.19, 5.57, 2.69, "实际 state · 节选", "", size=14)
    snippet = {
        "case_type": state["case_type"], "risk_flag": state["risk_flag"],
        "received_amount": state["received_amount"], "confirmations": state["confirmations"],
        "required_confirmations": state["required_confirmations"],
        "refund_approved": state["refund_approved"],
        "refund_destination_verified": state["refund_destination_verified"],
    }
    # Combine short properties so the closing brace stays inside the card.
    raw = json.dumps(snippet, indent=2)
    raw = raw.replace('"pending",\n  "risk_flag"', '"pending", "risk_flag"')
    raw = raw.replace('"confirmations": 3,\n  "required_confirmations"',
                      '"confirmations": 3, "required_confirmations"')
    code(s, raw, .87, 2.75, 5.14, 2.05, 13.4)
    V.card(s, 6.48, 2.19, 6.18, 2.69, "实际候选条件 · 节选", "", L.BLUE)
    code(s, 'await_approval:\n"' + mi["questions"]["action"]["criteria"]["await_approval"] + '"',
         6.68, 2.76, 5.78, 1.02, 14)
    code(s, 'approved（末句）:\n"Other cases do not need refund approval."',
         6.68, 3.91, 5.78, .64, 14)
    L.text(s, "接口返回的 choice · 与参考 approved 对照", .72, 5.01, 11.8, .31, 13, L.MUTED)
    for i, p in enumerate(PROVIDERS):
        r = rr[p]
        V.card(s, .68 + i * 4.07, 5.43, 3.84, .97,
               f'{LABELS[p]} · {r["timings"]["backend_total_ms"]:.1f} ms',
               json.dumps({"choice": r["choice"]}), COLORS[p], 12.3)

    # 9: full fixed five-step results and real sandbox state.
    run = data["cases"]["pending"]["baseline"]
    s = new("五步都执行了，业务结果仍要单独验收",
            "固定流程 · 每格为真实 choice / 后端毫秒；输入保持同一业务快照，按阶段更换问题与候选。",
            "这不是自主闭环：五个模型请求使用冻结的初始事实，每步问题和候选变化，"
            "模型选择再分别由各提供方的前端沙盒执行。表格是原始模型选择，"
            "不是执行器纠正后的动作。原始模式Decider与Jev多产生一个人工复核阻塞，"
            "GPT避免了这个额外阻塞，但所有模型仍需等待链上3/12确认。"
            + request_notes(data, "pending", ("baseline",), STAGES))
    rows = []
    for stage, name in zip(STAGES, ("客服受理", "支付核验", "风控审批", "资金处理", "客户回复"), strict=True):
        rows.append([name, *[
            f'{run["results"][p][stage]["choice"]} / {run["results"][p][stage]["timings"]["backend_total_ms"]:.1f}'
            for p in PROVIDERS]])
    V.table(s, ["阶段", "Decider", "Jev", "GPT"], rows,
            [1.70, 3.42, 3.42, 3.42], y=2.18, row_h=.49, size=13.2, highlights=(2,))
    totals = [summary["cases"]["pending"]["baseline"][p]["total_backend_ms"] for p in PROVIDERS]
    L.text(s, "五步耗时相加：" + "；".join(f"{LABELS[p]} {t:.1f} ms" for p, t in zip(PROVIDERS, totals, strict=True)),
           .84, 5.25, 11.6, .35, 14, L.MUTED)
    V.card(s, .68, 5.78, 5.86, .75, "Decider / Jev：额外等待人工审批", "", L.ORANGE)
    V.card(s, 6.80, 5.78, 5.86, .75, "三者最终：未入账、未结案；确认数仍为 3/12", "", L.TEAL)
    # Avoid an unused blank body in these deliberately title-only cards.
    for shape in s.shapes:
        if shape.has_text_frame and shape.text.startswith("三者最终："):
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(15.8)

    # 10: show actual changed criteria, rather than a generic tuning recipe.
    optimized = results(data, "pending", "optimized", "approve")["decider"]["model_input"]
    scoped = optimized["questions"]["action"]
    s = new("怎么优化：删去不适用字段，把适用范围写进每个候选",
            "同一笔 129 USDC、同样 3/12 确认、相同 4 个候选 ID；调整的是模型输入与条件表述。",
            "代码为app/crypto_profiles.py中的scoped_state和SCOPED_INSTRUCTIONS。"
            "pending的state由18字段变为12字段，删除refund_approved、refund_destination_verified、"
            "settled、settlement_minutes、settlement_overdue、settlement_sla_minutes。"
            "保留的字段和值全部不变。criteria和instructions同时重写，"
            "本实验不能单独归因于删字段或某一句话，也没有改变模型权重。"
            "全部4个候选ID及其顺序不变，未向模型输入reference_choice。"
            + request_notes(data, "pending"))
    V.card(s, .68, 2.19, 5.86, 3.43, "原始输入 · baseline", "", L.ORANGE)
    V.card(s, 6.79, 2.19, 5.86, 3.43, "优化输入 · optimized", "", L.TEAL)
    code(s, '"refund_approved": false\n"refund_destination_verified": true',
         .88, 2.78, 5.45, .68, 13.6)
    L.text(s, "收款请求里也带着退款字段。", .88, 3.48, 5.4, .34, 14, L.MUTED)
    code(s, 'await_approval:\n"' + mi["questions"]["action"]["criteria"]["await_approval"] + '"',
         .88, 4.02, 5.42, 1.26, 14)
    L.text(s, "移除无关退款与结算字段：18 → 12", 6.99, 2.78, 5.42, .39, 16, L.TEAL, True)
    code(s, 'await_approval:\n"' + scoped["criteria"]["await_approval"] + '"',
         6.99, 3.31, 5.43, 1.49, 13.7)
    L.text(s, "每个候选都明确业务类型、前提与优先级。", 6.99, 5.02, 5.42, .34, 14, L.MUTED)
    V.callout(s, "再明确阶段边界：通过审批，不等于资金已到账；执行步骤仍检查金额与链上确认。", y=5.94, size=15.5)

    # 11: unaltered pixels from the actual comparison UI, before and after.
    s = new("同一审批的实际变化：两次“等待人工审批”变为“通过”",
            "Decider / Jev：review_pending 由 true 变为 false；GPT 两次都通过。资金状态仍然未完成。",
            "上下两图分别是同次采集中pending-baseline-approve与pending-optimized-approve的真实界面裁剪。"
            "只裁剪边界，不改界面数据。两者保留三提供方名字、耗时、原始choice以及实际执行文字。"
            "底层执行状态Decider和Jev review_pending:true→false，GPT false→false。"
            "三者execute仍返回hold，fundsComplete:false、complete:false、无账本条目。"
            "这次改善是去掉不必要的人工审批，没有把链上3/12确认当成满足入账条件。"
            + request_notes(data, "pending"))
    before = V.image("docs/presentations/v12/screenshots/pending-baseline-approve.png",
                     "pending-baseline-outputs.png", (0, 254, 1865, 555))
    after = V.image("docs/presentations/v12/screenshots/pending-optimized-approve.png",
                    "pending-optimized-outputs.png", (0, 293, 1865, 594))
    L.text(s, "优化前 · 固定流程", .74, 2.11, 11.8, .3, 15, L.ORANGE, True)
    L.picture(s, before, .68, 2.44, 11.98, 1.80, frame=True)
    L.text(s, "优化后 · 仅改输入与候选条件", .74, 4.32, 11.8, .3, 15, L.TEAL, True)
    L.picture(s, after, .68, 4.66, 11.98, 1.80, frame=True)

    # 12: score changes plus actual per-case five-step latency, no stale batch.
    s = new("四场景一起重测：选择更一致，仍有两项 Decider 错误",
            "每格上行：与演示参考一致次数，原始 → 优化；下行：五步后端耗时之和，原始 → 优化。",
            "一致性按原始choice计算，执行器修正不计为模型答对。"
            "Decider14/20→18/20，Jev18/20→20/20，GPT20/20→20/20。"
            "Decider少付核验由request_topup变成await_confirmations，属于退化；"
            "同时少付执行由credit_order改为hold，场景总分仍4/5。"
            "退款审批await_approval→approved仍错，参考verify_destination；"
            "退款execute和close从refund_payment/resolve改为hold/follow_up，故2/5→4/5。"
            "商户结算close从resolve→follow_up，4/5→5/5。"
            "优化后四场景仍未完成资金业务，是冻结业务事实的固定流程试验，不能以四场景全部未结案推断模型全错。"
            "\n逐场景精确统计：\n" + json.dumps(summary, ensure_ascii=False, indent=2),
            "4 场景 × 2 种输入 × 3 模型 × 5 步 = 120 次成功调用；另有 6 次预热。一次配对观测，不据此声称稳定加速。")
    rows = []
    for case, label in CASES.items():
        cells = []
        for p in PROVIDERS:
            a, b = [summary["cases"][case][profile][p] for profile in PROFILES]
            cells.append(f'{a["matches"]}/5 → {b["matches"]}/5\n{a["total_backend_ms"]:.0f} → {b["total_backend_ms"]:.0f} ms')
        rows.append([label, *cells])
    V.table(s, ["场景", "Decider", "Jev", "GPT"], rows, [2.17, 3.26, 3.26, 3.27],
            row_h=.66, size=14)
    for shape in s.shapes:
        if shape.has_text_frame and "\n" in shape.text:
            for paragraph in shape.text_frame.paragraphs:
                paragraph.space_after = Pt(0)
                paragraph.line_spacing = 1.0
    L.text(s, "合计：Decider 14/20 → 18/20；Jev 18/20 → 20/20；GPT 20/20 → 20/20。",
           .84, 5.69, 11.64, .4, 16, L.TEAL, True)
    V.callout(s, "仍需处理：少付时误选“等待确认”；退款地址未验证却选择“批准”。优化并非每个节点都改善。", y=6.16, size=13.6,
              color=L.ORANGE, fill=L.PALE_ORANGE)

    # 13: real failure is preserved, including raw model output and protection.
    s = new("退款仍会选错：原始返回与执行结果必须分开看",
            "实际输入：退款 250 USDC；refund_destination_verified=false；refund_approved=false。",
            "图为优化模式退款审批的真实界面，去掉顶部轮次栏而保留客户请求、输入摘要、三模型返回与实际执行。"
            "原始Decider await_approval错误；优化后approved仍错误。"
            "Jev与GPT在两个profile中都返回verify_destination。"
            "优化后Decider原始返回approved，执行器记录requested_action:approved、action:verify_destination、"
            "safety_adjusted:true，原因是退款地址未验证。该修正不能计入模型正确率。"
            "三者后续资金步骤均hold，回复follow_up，未退款、未结案。"
            + request_notes(data, "refund"))
    shot = V.image("docs/presentations/v12/screenshots/refund-optimized-approve.png",
                   "refund-optimized-detail.png", (0, 67, 1865, 720))
    L.picture(s, shot, .68, 2.20, 11.98, 4.23, frame=True)

    # Existing loop, routing, validation, vLLM and costs are preserved.
    for page in range(10, 22):
        keep(page)
    # Old generic optimization slide (22) is replaced by the preceding evidence.
    for page in range(23, 32):
        keep(page, page == 31)
    means = summary["aggregate"]["optimized"]
    update_text(originals[30],
                "同一支付批次：Decider 102.2 ms / 18/20；Jev 416.5 ms / 20/20；GPT 1116.7 ms / 20/20。",
                f'优化后支付：Decider {means["decider"]["mean_backend_ms"]:.1f} ms / 18/20；'
                f'Jev {means["jev"]["mean_backend_ms"]:.1f} ms / 20/20；GPT {means["bedrock"]["mean_backend_ms"]:.1f} ms / 20/20。')
    originals[30].notes_slide.notes_text_frame.text = (
        originals[30].notes_slide.notes_text_frame.text.replace("平均1116.7ms", f'平均{means["bedrock"]["mean_backend_ms"]:.1f}ms')
        + "\nv12支付指标更新为2026-10-15正式采集的optimized输入，每模型20个节点；其余验证、vLLM与成本数据保持各自原始口径。\n" + METHOD)
    keep(32, True)
    update_text(originals[31],
                "PPT 内含逐页讲稿与来源；v11 聚焦 Jev 场景结论、按需与 1 年 RI 成本交叉图。",
                "PPT 内含逐页讲稿与来源；v12 新增三模型支付实际请求、优化前后结果与真实界面。")

    used = {node.rId for node in order}
    for node in list(prs.slides._sldIdLst):
        prs.slides._sldIdLst.remove(node)
        if node.rId not in used:
            prs.part.drop_rel(node.rId)
    for node in order:
        prs.slides._sldIdLst.append(node)
    assert len(prs.slides) == 35
    for i, (slide, meta) in enumerate(zip(prs.slides, metas, strict=True), 1):
        # Keep provenance in speaker notes, not in a second visible footer.
        explanations = [
            shape for shape in slide.shapes
            if shape.has_text_frame and Inches(6.5) <= shape.top < Inches(7.03)
        ]
        assert len(explanations) == 1, (i, len(explanations))
        meta["removed_lower_explanation"] = explanations[0].text
        slide.notes_slide.notes_text_frame.text += (
            "\n原页底说明（仅保留于讲稿）：\n" + explanations[0].text
        )
        for shape in explanations:
            shape._element.getparent().remove(shape._element)
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.endswith(" / 32") and shape.top > Inches(7):
                replace_text(shape, shape.text, f"{i:02} / 35")
        meta["notes"] = slide.notes_slide.notes_text_frame.text
    prs.core_properties.title = "Strands Decider × Jev × GPT：支付流程实际请求与优化对照"
    prs.core_properties.subject = "v12 / 35 pages / 120 real calls, baseline versus scoped input"
    output = OUT / f"{NAME}.pptx"
    prs.save(output)
    assert V.digest(BASE) == before_hash
    save(OUT / "build-manifest.json", {
        "created_at": datetime.now(UTC).isoformat(),
        "base_file": str(BASE.relative_to(ROOT)), "base_sha256": before_hash,
        "pptx_file": str(output.relative_to(ROOT)), "pptx_sha256": V.digest(output),
        "pages": 35, "slides": metas, "evidence": V.EVIDENCE, "images": V.IMAGES,
        "new_pages": [s["page"] for s in metas if s["source"] == "new"],
        "new_formal_model_calls": 120, "new_warmup_model_calls": 6,
        "aws_resource_mutations": 0, "external_funds_actions": 0,
        "lower_explanations_removed_from_all_slides": True,
        "capture": capture,
    })
    notes = ["# v12 讲稿与来源", "", "35 页；v11 保留。新增 120 次正式调用与 6 次预热。", ""]
    for meta in metas:
        notes.extend([f'## {meta["page"]:02} · {meta["title"]}', "", meta["notes"], ""])
    (OUT / "讲稿与来源.md").write_text("\n".join(notes))
    details = ["# 支付流程：实际请求、返回与执行结果", "", METHOD, ""]
    for case, label in CASES.items():
        details += [f"## {label}", "", "```text", request_notes(data, case, stages=STAGES), "```", ""]
    (OUT / "请求与结果对照.md").write_text("\n".join(details))
    print(json.dumps({"pptx": str(output), "pages": 35, "new_pages": list(range(8, 14)),
                      "metrics": summary["aggregate"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
