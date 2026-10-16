"""Package the Chinese AWS article with editable Word, HTML and sourced figures."""

import base64
import hashlib
import html
import json
import re
import shutil
from datetime import UTC, datetime
from pathlib import Path

import cairosvg
import markdown
from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/blogs/strands-decider"
NAME = "Strands-Decider-AWS实践"
ASSETS = OUT / "assets"
FONT = "Noto Sans CJK SC"


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def architecture():
    parts = [
        '<svg xmlns="http://www.w3.org/2000/svg" width="1600" height="1000" viewBox="0 0 1600 1000">',
        '<defs><marker id="arrow" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0,0 L8,4 L0,8 Z" fill="#64788a"/></marker></defs>',
        '<rect width="1600" height="1000" fill="white"/>',
        f'<g font-family="{FONT},sans-serif" fill="#203448">',
    ]

    def txt(x, y, text, size=24, anchor="middle", color="#203448", weight="normal"):
        parts.append(f'<text x="{x}" y="{y}" font-size="{size}" text-anchor="{anchor}" fill="{color}" font-weight="{weight}">{html.escape(text)}</text>')

    def rect(x, y, w, h, stroke, fill="white", dashed=False):
        dash = 'stroke-dasharray="8 5"' if dashed else ""
        parts.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="7" fill="{fill}" stroke="{stroke}" stroke-width="2" {dash}/>')

    def arrow(points, color="#64788a"):
        value = " ".join(f"{x},{y}" for x, y in points)
        parts.append(f'<polyline points="{value}" fill="none" stroke="{color}" stroke-width="2.5" marker-end="url(#arrow)"/>')

    def node(name, x, y, title, detail):
        data = base64.b64encode((ROOT / f"docs/diagrams/aws-icons/{name}.svg").read_bytes()).decode()
        parts.append(f'<image x="{x-32}" y="{y-48}" width="64" height="64" href="data:image/svg+xml;base64,{data}"/>')
        txt(x, y + 50, title, 23, weight="bold")
        txt(x, y + 83, detail, 19, color="#64788a")

    rect(265, 55, 1310, 905, "#64788a")
    txt(300, 92, "AWS Cloud", 24, "start", weight="bold")
    rect(290, 298, 1260, 635, "#17aeb7", dashed=True)
    txt(320, 333, "AWS Region · us-east-1", 24, "start", weight="bold")
    rect(320, 383, 810, 410, "#8b5cf6", "#fbfaff")
    txt(346, 420, "应用 VPC · 私有子网", 24, "start", weight="bold")

    arrow([(178, 176), (448, 176)])
    txt(315, 157, "HTTPS", 18)
    arrow([(584, 176), (1338, 176), (1338, 330)])
    txt(994, 157, "OAC 签名读取静态资源", 20)
    arrow([(520, 263), (520, 280), (850, 280), (850, 442), (444, 442), (444, 506)])
    txt(868, 369, "/api/* · VPC Origin", 18, "start")
    arrow([(508, 554), (644, 554)])
    arrow([(774, 554), (904, 554)], "#087f72")
    arrow([(752, 516), (806, 516), (806, 457), (1220, 457), (1220, 530), (1274, 530)], "#4669c4")
    arrow([(752, 580), (804, 580), (804, 734), (1220, 734), (1220, 680), (1274, 680)], "#8055a4")
    arrow([(710, 649), (710, 837), (1274, 837)])
    arrow([(1274, 706), (1200, 706), (1200, 897), (120, 897), (120, 880)], "#8055a4")

    node("client", 112, 176, "浏览器", "登录与场景对比")
    node("cloudfront", 520, 176, "Amazon CloudFront", "静态站点与 API 入口")
    node("s3", 1338, 380, "Amazon S3", "静态前端 · 禁止公共访问")
    node("alb", 444, 554, "内部 ALB", "API 请求转发")
    node("fargate", 710, 554, "ECS Fargate", "FastAPI · 业务后端")
    node("ec2", 972, 554, "Amazon EC2", "v19 · g6.xlarge / L4")
    txt(972, 675, "私网 HTTP · GPU 常驻", 18, color="#087f72")
    node("bedrock", 1338, 530, "Amazon Bedrock", "GPT · Converse / low")
    node("lambda", 1338, 680, "AWS Lambda", "Jev 转发 · 密钥存 Secrets Manager")
    node("dynamodb", 1338, 828, "Amazon DynamoDB", "账号与会话")
    rect(12, 752, 216, 128, "#8055a4", "#faf6fd")
    txt(120, 795, "Typesafe Jev API", 21, weight="bold")
    txt(120, 825, "AWS 账户外", 19, color="#64788a")
    txt(120, 854, "HTTPS / System One", 17, color="#8055a4")
    txt(330, 980, "示意图省略私有网络端点和 GPU 管理路径；模型路由模块使用同一业务后端。", 18, "start", color="#64788a")
    parts.extend(["</g>", "</svg>"])
    svg = ASSETS / "01-aws-architecture.svg"
    svg.write_text("\n".join(parts))
    cairosvg.svg2png(url=str(svg), write_to=str(svg.with_suffix(".png")), output_width=2400)


def decision_flow():
    parts = [
        '<svg xmlns="http://www.w3.org/2000/svg" width="1600" height="600" viewBox="0 0 1600 600">',
        '<defs><marker id="arrow" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0,0 L8,4 L0,8 Z" fill="#527083"/></marker></defs>',
        '<rect width="1600" height="600" fill="white"/>',
        f'<g font-family="{FONT},sans-serif" fill="#203448">',
    ]

    def text(x, y, content, size=25, color="#203448", weight="normal"):
        parts.append(f'<text x="{x}" y="{y}" font-size="{size}" text-anchor="middle" fill="{color}" font-weight="{weight}">{html.escape(content)}</text>')

    def box(x, y, title, lines, fill="#f0f7fa", stroke="#aac6d4"):
        parts.append(f'<rect x="{x}" y="{y}" width="328" height="158" rx="12" fill="{fill}" stroke="{stroke}" stroke-width="2"/>')
        text(x + 164, y + 43, title, 28, weight="bold")
        for i, line in enumerate(lines):
            text(x + 164, y + 85 + i * 36, line, 23)

    def arrow(points):
        path = " ".join(f"{x},{y}" for x, y in points)
        parts.append(f'<polyline points="{path}" fill="none" stroke="#527083" stroke-width="3" marker-end="url(#arrow)"/>')

    text(800, 57, "一次前向计算：读取上下文，对候选评分", 29, weight="bold")
    box(30, 110, "业务输入", ["事实、问题", "候选 ID 与适用条件"])
    box(430, 110, "Qwen3.5-2B 主体", ["rank-16 LoRA 适配", "取得各位置隐藏表示"])
    box(830, 110, "Pointer Head", ["候选位置与答案位置", "评分 → 温度校准"])
    box(1230, 110, "结构化结果", ["候选 ID、概率", "原始结果留存"], "#edf8f3", "#80bda6")
    for x in [358, 758, 1158]:
        arrow([(x + 5, 189), (x + 62, 189)])
    box(1230, 365, "业务执行器", ["核验金额、地址与权限", "执行实际允许的动作"], "#fff5e8", "#dbb37b")
    box(430, 365, "更新业务状态", ["保存执行结果", "下一轮读取最新事实"], "#edf8f3", "#80bda6")
    arrow([(1394, 273), (1394, 360)])
    arrow([(1225, 444), (766, 444)])
    text(994, 422, "执行结果与未完成原因", 23)
    arrow([(425, 444), (194, 444), (194, 276)])
    text(293, 417, "反馈", 23)
    text(800, 571, "独立问题可以共享上下文；依赖执行结果的问题需要分轮处理。", 23, color="#64788a")
    parts.extend(["</g>", "</svg>"])
    svg = ASSETS / "05-decision-flow.svg"
    svg.write_text("\n".join(parts))
    cairosvg.svg2png(url=str(svg), write_to=str(svg.with_suffix(".png")), output_width=2400)


def set_font(run, size=None, bold=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def inline(paragraph, node, size=None):
    for child in node.children:
        if isinstance(child, NavigableString):
            run = paragraph.add_run(str(child))
            set_font(run, size)
        elif child.name == "a":
            rel = paragraph.part.relate_to(
                child.get("href"), "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                is_external=True,
            )
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), rel)
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "176EAE")
            properties.append(color)
            run.append(properties)
            text = OxmlElement("w:t")
            text.text = child.get_text()
            run.append(text)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
        elif child.name == "br":
            paragraph.add_run().add_break()
        else:
            run = paragraph.add_run(child.get_text())
            set_font(run, size, child.name in {"strong", "b"})
            if child.name == "code":
                run.font.color.rgb = RGBColor.from_string("087F72")


def build_docx(body):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.74)
    section.left_margin = section.right_margin = Inches(.72)
    for name, size in [("Normal", 10.5), ("Title", 23), ("Heading 1", 16), ("Heading 2", 13)]:
        style = doc.styles[name]
        style.font.name = FONT
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.25
    doc.styles["Normal"].paragraph_format.widow_control = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    soup = BeautifulSoup(body, "html.parser")
    images, tables = 0, 0
    for node in soup.children:
        if not getattr(node, "name", None):
            continue
        if node.name in {"h1", "h2", "h3"}:
            doc.add_heading(node.get_text(), 0 if node.name == "h1" else int(node.name[1]) - 1)
        elif node.name == "p":
            image = node.find("img")
            if image:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str(OUT / image["src"]), width=Inches(6.80))
                images += 1
            else:
                p = doc.add_paragraph()
                inline(p, node)
                if node.get_text().startswith("图 "):
                    p.paragraph_format.space_after = Pt(12)
                    for run in p.runs:
                        run.font.size = Pt(9)
                        run.font.color.rgb = RGBColor.from_string("64788A")
        elif node.name == "pre":
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.0
            for i, line in enumerate(node.get_text().splitlines()):
                if i:
                    p.add_run().add_break()
                set_font(p.add_run(line), 9)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "F0F4F7")
            p._p.get_or_add_pPr().append(shade)
        elif node.name == "table":
            rows = node.find_all("tr")
            table = doc.add_table(rows=len(rows), cols=len(rows[0].find_all(["td", "th"])))
            table.style = "Table Grid"
            for ri, row in enumerate(rows):
                for cell, value in zip(table.rows[ri].cells, row.find_all(["td", "th"]), strict=True):
                    inline(cell.paragraphs[0], value, 8.5)
                    cell.paragraphs[0].paragraph_format.space_after = Pt(4)
                    if ri == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
                props = table.rows[ri]._tr.get_or_add_trPr()
                props.append(OxmlElement("w:cantSplit"))
                if ri == 0:
                    props.append(OxmlElement("w:tblHeader"))
            doc.add_paragraph()
            tables += 1
        elif node.name in {"ol", "ul"}:
            for item in node.find_all("li", recursive=False):
                p = doc.add_paragraph(style="List Number" if node.name == "ol" else "List Bullet")
                inline(p, item)
    doc.core_properties.title = soup.h1.get_text()
    doc.core_properties.subject = "Strands Decider 原理、AWS 部署、支付与模型路由、vLLM 并发和成本"
    doc.save(OUT / f"{NAME}.docx")
    return {"embedded_figures": images, "tables": tables}


def main():
    ASSETS.mkdir(parents=True, exist_ok=True)
    architecture()
    decision_flow()
    figures = [
        ("02-refund-decision.png", "docs/presentations/v8/screenshots/refund-decision.png"),
        ("03-routing-retrieval.png", "docs/presentations/v9/assets/routing-retrieval-results.png"),
        ("04-original-vllm-rps.png", "docs/presentations/v9/assets/original-vllm-rps.png"),
    ]
    provenance = [{"file": "assets/01-aws-architecture.png",
                   "source": "assets/01-aws-architecture.svg",
                   "source_sha256": digest(ASSETS / "01-aws-architecture.svg"),
                   "sha256": digest(ASSETS / "01-aws-architecture.png"),
                   "transformation": "Original simplified architecture, using repository AWS service icons"},
                  {"file": "assets/05-decision-flow.png",
                   "source": "assets/05-decision-flow.svg",
                   "source_sha256": digest(ASSETS / "05-decision-flow.svg"),
                   "sha256": digest(ASSETS / "05-decision-flow.png"),
                   "reference": "https://strandsagents.com/blog/introducing-strands-decider/",
                   "transformation": "Original explanatory diagram; no model output or measurement fabricated"}]
    for target, source in figures:
        shutil.copyfile(ROOT / source, ASSETS / target)
        provenance.append({"file": "assets/" + target, "source": source,
                           "source_sha256": digest(ROOT / source), "sha256": digest(ASSETS / target),
                           "transformation": "unchanged copy"})
    (OUT / "figure-sources.json").write_text(json.dumps({
        "figures": provenance, "routing_crop_source": "docs/presentations/v9/build-manifest.json",
        "aws_icon_attribution": "docs/diagrams/aws-icons/README.md",
    }, ensure_ascii=False, indent=2) + "\n")

    source = (OUT / f"{NAME}.md").read_text()
    body = markdown.markdown(source, extensions=["tables", "fenced_code"])
    stats = build_docx(body)
    embedded = BeautifulSoup(body, "html.parser")
    for img in embedded.find_all("img"):
        data = base64.b64encode((OUT / img["src"]).read_bytes()).decode()
        img["src"] = "data:image/png;base64," + data
        paragraph = img.parent
        caption = paragraph.find_next_sibling()
        if caption and caption.name == "p" and caption.get_text().startswith("图 "):
            figure = embedded.new_tag("figure")
            paragraph.insert_before(figure)
            figure.append(img.extract())
            caption.name = "figcaption"
            figure.append(caption.extract())
            paragraph.decompose()
    css = """
    :root{color-scheme:light}body{margin:0;background:#edf1f5;color:#243447;
    font-family:"Noto Sans CJK SC","Microsoft YaHei",sans-serif}
    article{max-width:1000px;margin:32px auto;padding:48px 58px;background:white}
    h1{font-size:30px;line-height:1.55;margin:0 0 28px}h2{font-size:24px;margin:42px 0 16px}
    h3{font-size:20px;line-height:1.6;margin:28px 0 12px}
    p,li{font-size:17px;line-height:1.95;margin:15px 0}a{color:#176eae}
    img{display:block;max-width:100%;height:auto;margin:28px auto 8px}
    figure{margin:26px 0}figure img{margin:0 auto 8px}
    figcaption{font-size:14px;line-height:1.7;color:#64788a}
    pre{padding:18px 22px;background:#f0f4f7;overflow-wrap:anywhere;white-space:pre-wrap;font-size:14px}
    table{border-collapse:collapse;width:100%;margin:22px 0;font-size:14px}
    th,td{border:1px solid #dce4ea;padding:12px 10px;vertical-align:top}
    th{background:#edf4f4;text-align:left}code{color:#087f72}
    @media(max-width:700px){article{margin:0;padding:24px 18px}h1{font-size:25px}
    table{font-size:11px}th,td{padding:8px 5px}}
    @media print{@page{size:A4;margin:17mm}body{background:white}article{padding:0;margin:0}
    h1{font-size:23px}h2{font-size:18px;break-after:avoid}
    h3{font-size:14px;break-after:avoid}p,li{font-size:12px;line-height:1.75;orphans:3;widows:3}
    table{font-size:10px}tr,img,pre,figure{break-inside:avoid}th,td{padding:7px}
    figcaption{font-size:10px}}
    """
    (OUT / f"{NAME}.html").write_text(
        '<!doctype html><html lang="zh-CN"><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1">'
        f"<title>{html.escape(embedded.h1.get_text())}</title><style>{css}</style>"
        f"<body><article>{embedded}</article></body></html>"
    )
    counted = source.split("## 参考资料")[0]
    counted = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", counted)
    counted = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", counted)
    counted = re.sub(r"```.*?```", "", counted, flags=re.S)
    words = len(re.findall(r"[\u4e00-\u9fff]|[A-Za-z0-9]+(?:[.\-/][A-Za-z0-9]+)*", counted))
    cjk = len(re.findall(r"[\u4e00-\u9fff]", counted))
    assert 6000 <= words <= 6500, words
    manifest = {
        "built_at": datetime.now(UTC).isoformat(), "word_count": words,
        "chinese_character_count": cjk,
        "revision": "6000-word-rewrite",
        "count_method": "Chinese characters plus English/numeric words; title, captions and tables included; code and references excluded",
        **stats,
        "files": {p.name: digest(p) for p in OUT.glob(NAME + ".*") if p.suffix != ".pdf"},
        "new_model_calls": 0, "screenshots_rewritten": False,
        "style_references": [
            "https://aws.amazon.com/cn/blogs/china/kiro-quick-deploy-agent-deploy-amazon-bedrock-agentcore/",
            "https://aws.amazon.com/cn/blogs/china/amazon-bedrock-agentcore-file-system-session-storage-amazon-efs-s3-files/",
            "https://aws.amazon.com/cn/blogs/china/self-learning-evolvable-agents-for-cultural-tourism-info-extraction-with-agentcore/",
        ],
    }
    (OUT / "build-manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"word_count": words, **stats, "output": str(OUT.relative_to(ROOT))}, ensure_ascii=False))


if __name__ == "__main__":
    main()
