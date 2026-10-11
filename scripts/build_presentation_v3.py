"""Build the concise 25-slide Decider/GPT comparison from the v2 evidence."""

import hashlib
import io
import json
from datetime import UTC, datetime
from pathlib import Path

import cairosvg
import presentation_layout as L
from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/presentations/v3"
V2 = ROOT / "docs/presentations/v2"
BASE = V2 / "Strands-Decider-模型原理与支付业务演示-v2-20261010.pptx"
NAME = "Strands-Decider-模型原理与支付业务演示-v3-25页-20261010"
BLOG = "https://strandsagents.com/blog/introducing-strands-decider/"
TEAL, BLUE, NAVY = "087F72", "4669C4", "132639"
INK, MUTED, WHITE, LINE = L.INK, L.MUTED, L.WHITE, L.LINE
PALE_T, PALE_B, PALE_O = L.PALE_TEAL, L.PALE_BLUE, L.PALE_ORANGE
ORANGE = L.ORANGE
STAGES = {
    "intake": "客服受理",
    "verify": "交易核验",
    "approve": "风控审批",
    "execute": "业务处置",
    "close": "回复结案",
}
CASES = {
    "pending": "USDC 未到账",
    "underpaid": "USDT 少付",
    "refund": "稳定币退款",
    "settlement": "商户结算",
}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "screenshots").mkdir(exist_ok=True)
    bench = json.loads((ROOT / "docs/results/benchmark-2026-10-10.json").read_text())
    loops = json.loads((V2 / "loop-observations.json").read_text())
    cases = {c["summary"]["case"]: c for c in bench["cases"]}
    totals = bench["totals"]
    prs = Presentation(BASE)
    # Keep the source's page format and theme; rebuild the narrative and all layouts.
    for sid in list(prs.slides._sldIdLst):
        prs.part.drop_rel(sid.rId)
        prs.slides._sldIdLst.remove(sid)
    metadata, crops = [], []
    icon_cache = {}
    sources = (
        "\n\n来源与口径：\n"
        f"官方博客：{BLOG}\n"
        "模型：https://huggingface.co/StrandsAgents/strands-decider-2B-hobson-v19\n"
        "判断实现：https://github.com/strands-labs/strands-decider/tree/"
        "63d24ae286e50105fba0bd1db15b0e243aea4650/src/strands_decider\n"
        "BERT：https://aclanthology.org/N19-1423/；"
        "NLI 候选评分：https://aclanthology.org/D19-1404/\n"
        "固定输入：docs/results/benchmark-2026-10-10.json\n"
        "闭环：docs/presentations/v2/loop-observations.json\n"
        "截图：docs/presentations/v2/capture-manifest.json\n"
        "数据采集：2026-10-10 12:26 UTC 起，crypto10 / API 20261010-loop1。\n"
        "Decider：官方 v19 / SageMaker ml.g6.xlarge / L4；"
        "GPT：Bedrock us.openai.gpt-6.1-sol / reasoning low。\n"
        "累计决策耗时 = backend_total_ms 求和，包含服务通信与排队；"
        "不含人工/外部事件等待，非浏览器全程用时。\n"
        "所有资金与客户场景为模拟数据；客服回复由应用模板生成。"
    )

    def text(s, value, x, y, w, h=0.4, size=18, color=INK, bold=False, align=None):
        return L.text(s, str(value), x, y, w, h, size, color, bold, align)

    def icon(s, name, x, y, size=0.45, color=TEAL, aws=False):
        key = (name, color, aws)
        if key not in icon_cache:
            folder = ROOT / "docs/diagrams/aws-icons" if aws else OUT / "icons"
            svg = (folder / f"{name}.svg").read_text()
            if not aws:
                svg = svg.replace("currentColor", f"#{color}")
            icon_cache[key] = cairosvg.svg2png(
                bytestring=svg.encode(), output_width=256, output_height=256
            )
        shape = s.shapes.add_picture(
            io.BytesIO(icon_cache[key]), Inches(x), Inches(y), Inches(size), Inches(size)
        )
        shape.name = f"ICON: {'AWS ' if aws else 'Lucide '}{name}"
        return shape

    def slide(
        title,
        section="模型与原理",
        subtitle="",
        notes="",
        dark=False,
        evidence="2026-10-10 实测与项目实现",
    ):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = L.rgb(NAVY if dark else L.BG)
        text(
            s,
            f"STRANDS DECIDER   /   {section}",
            0.6,
            0.27,
            11.9,
            0.25,
            10,
            "86D7C6" if dark else TEAL,
            True,
        )
        text(s, title, 0.6, 0.79, 12.1, 0.69, 30, WHITE if dark else INK, True)
        if subtitle:
            text(s, subtitle, 0.62, 1.56, 12, 0.53, 15, "B8C9DB" if dark else MUTED)
        s.notes_slide.notes_text_frame.text = notes + sources
        metadata.append({"title": title, "section": section, "dark": dark, "evidence": evidence})
        return s

    def note(s, value, dark=False):
        text(s, value, 0.62, 6.65, 12.05, 0.31, 11.5, "B8C9DB" if dark else MUTED)

    def pill(s, value, x, y, w, color=TEAL, fill=PALE_T):
        L.rect(s, x, y, w, 0.35, fill, radius=True)
        text(s, value, x + 0.05, y + 0.055, w - 0.1, 0.25, 11, color, True, PP_ALIGN.CENTER)

    def card(s, x, y, w, h, title, body="", glyph=None, color=TEAL, fill=WHITE):
        L.rect(s, x, y, w, h, fill, LINE, True)
        if glyph:
            icon(s, glyph, x + 0.22, y + 0.22, 0.43, color)
            text(s, title, x + 0.82, y + 0.23, w - 1.02, 0.40, 19, color, True)
        else:
            text(s, title, x + 0.23, y + 0.23, w - 0.46, 0.4, 19, color, True)
        if body:
            text(s, body, x + 0.23, y + 0.93, w - 0.46, h - 1.02, 18)

    def arrow(s, x1, y1, x2, y2, color=MUTED):
        # DrawingML tailEnd places the arrowhead at the destination.
        L.line(s, [(x1, y1), (x2, y2)], color, 1.8)
        connector = s.shapes[-1]
        from pptx.oxml.xmlchemy import OxmlElement

        ln = connector._element.spPr.find(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}ln"
        )
        tip = OxmlElement("a:tailEnd")
        tip.set("type", "triangle")
        ln.append(tip)

    def flow(s, items, x=0.65, y=3.0, w=12.0, color=TEAL, height=1.45):
        gap = 0.31
        width = (w - gap * (len(items) - 1)) / len(items)
        for i, (heading, desc, glyph) in enumerate(items):
            xx = x + i * (width + gap)
            L.rect(s, xx, y, width, height, WHITE, LINE, True)
            icon(s, glyph, xx + 0.17, y + 0.16, 0.38, color)
            text(s, heading, xx + 0.67, y + 0.17, width - 0.8, 0.4, 17, color, True)
            text(s, desc, xx + 0.17, y + 0.75, width - 0.34, 0.5, 15)
            if i < len(items) - 1:
                arrow(
                    s,
                    xx + width + 0.03,
                    y + height / 2,
                    xx + width + gap - 0.04,
                    y + height / 2,
                    color,
                )

    def metric(s, value, label, x, y, w, color=TEAL, dark=False):
        text(s, value, x, y, w, 0.75, 39, color, True)
        text(s, label, x + 0.02, y + 0.89, w, 0.57, 14, "B8C9DB" if dark else MUTED)

    def table(s, headers, rows, widths, y=2.26, row_h=0.63, size=17, highlights=()):
        x = 0.63
        for h, w in zip(headers, widths, strict=True):
            L.rect(s, x, y, w - 0.03, 0.50, NAVY)
            text(s, h, x + 0.13, y + 0.10, w - 0.26, 0.33, 13, WHITE, True)
            x += w
        for i, row in enumerate(rows):
            yy, x = y + 0.56 + row_h * i, 0.63
            for j, (val, w) in enumerate(zip(row, widths, strict=True)):
                fill = PALE_O if i in highlights else (WHITE if i % 2 == 0 else "ECF1F6")
                L.rect(s, x, yy, w - 0.03, row_h - 0.04, fill)
                color = TEAL if j == 1 else BLUE if j == 2 else INK
                text(s, val, x + 0.13, yy + 0.13, w - 0.26, row_h - 0.14, size, color, j in (1, 2))
                x += w

    def shot(s, name, x, y, w, h, bounds=None):
        source = V2 / "screenshots" / name
        path = source
        if bounds:
            path = OUT / "screenshots" / (source.stem + "-crop.png")
            with Image.open(source) as im:
                im.crop(bounds).save(path)
            entry = {
                "file": path.relative_to(OUT).as_posix(),
                "source": source.relative_to(ROOT).as_posix(),
                "source_sha256": digest(source),
                "crop_pixels": list(bounds),
                "sha256": digest(path),
                "data_replaced": False,
            }
            if entry not in crops:
                crops.append(entry)
        return L.picture(s, path, x, y, w, h, frame=True)

    def loop_total(case, provider):
        return sum(
            d["timings"]["backend_total_ms"]
            for d in loops[case]["providers"][provider]["decisions"]
        )

    def pairbars(s, labels, dvalues, gvalues, x=2.8, y=2.7, w=7.8, gap=0.8, units="ms"):
        maxval = max(gvalues + dvalues) * 1.05
        pill(s, "Decider", x, y - 0.56, 1.1)
        pill(s, "GPT", x + 1.28, y - 0.56, 0.9, BLUE, PALE_B)
        for i, (label, d, g) in enumerate(zip(labels, dvalues, gvalues, strict=True)):
            yy = y + gap * i
            text(s, label, 0.68, yy + 0.03, x - 0.9, 0.6, 15, INK, True)
            for j, (val, color) in enumerate(((d, TEAL), (g, BLUE))):
                bar_w = w * val / maxval
                L.rect(s, x, yy + j * 0.26, w, 0.17, "E6EBF1", radius=True)
                L.rect(s, x, yy + j * 0.26, bar_w, 0.17, color, radius=True)
                text(
                    s,
                    f"{val:,.3f} s" if units == "s" else f"{val:,.1f} {units}",
                    x + w + 0.16,
                    yy - 0.045 + j * 0.26,
                    1.5,
                    0.28,
                    12,
                    color,
                    True,
                )
        return maxval

    # 01 — Cover.
    s = slide(
        "Strands Decider × GPT",
        "支付业务对比 / v3 · 25 页",
        "从下一步选择，到执行反馈，再到业务结果",
        dark=True,
        notes="本版基于34页v2重新编排为25页。模型基础精简，BERT仅一页，主体是决策过程与速度。截图和数据复用同日已经完成的真实采集，不发起新模型调用。",
    )
    for x, name, desc, glyph, color in [
        (0.65, "Decider", "对候选动作直接评分", "list-filter", "86D7C6"),
        (7.0, "GPT-6.1 Sol", "生成结构化动作选择", "code-xml", "A6B9FF"),
    ]:
        L.rect(s, x, 2.6, 5.68, 2.18, "1D354C", "365168", True)
        icon(s, glyph, x + 0.28, 2.91, 0.6, color)
        text(s, name, x + 1.09, 2.92, 4.2, 0.53, 28, WHITE, True)
        text(s, desc, x + 0.29, 4.03, 5.0, 0.43, 20, color)
    text(s, "VS", 6.30, 3.30, 0.66, 0.5, 20, "B8C9DB", True, PP_ALIGN.CENTER)
    for x, label, glyph in [
        (1.1, "如何选择", "route"),
        (5.25, "决策多快", "timer"),
        (9.38, "是否完成", "circle-check"),
    ]:
        icon(s, glyph, x, 5.57, 0.38, "86D7C6")
        text(s, label, x + 0.57, 5.57, 2.0, 0.4, 19, WHITE, True)
    note(s, "2026.10.10 真实模型采集 · 加密支付模拟业务 · 原版与 v2 保留", True)

    # 02 — Purpose.
    s = slide(
        "把业务中的「下一步选择」交给模型",
        subtitle="应用定义候选与条件，模型选择动作，执行器落实业务。",
        notes="Decider面向有限候选的高频决策。候选名称和criteria由应用提供，并非模型临时发明。模型没有资金权限。",
    )
    flow(
        s,
        [
            ("业务输入", "客户消息 + 订单事实", "receipt-text"),
            ("模型选择", "等待 / 补款 / 退款", "route"),
            ("实际执行", "更新订单与模拟账本", "clipboard-check"),
        ],
        y=2.55,
        height=1.65,
    )
    card(
        s,
        0.65,
        4.72,
        5.85,
        1.52,
        "Decider",
        "专注候选评分与结构化决策",
        "list-filter",
        TEAL,
        PALE_T,
    )
    card(
        s,
        6.80,
        4.72,
        5.85,
        1.52,
        "GPT",
        "通用生成模型，也可选择候选",
        "brain-circuit",
        BLUE,
        PALE_B,
    )
    note(s, "本项目比较的是动作选择；客服回复来自同一套应用模板。")

    # 03 — Mechanism.
    s = slide(
        "同样的候选，走两条决策路径",
        subtitle="Decider 评分后取最高项；GPT 逐 token 生成候选 ID。",
        notes="官方v19基于Qwen3.5-2B，配合LoRA和Pointer Head。答案位置与候选位置隐藏表示经投影、归一化和点积计算logits，按类型温度计算概率。GPT在本项目输出JSON，再由应用解析验证。两者都接受候选描述；GPT也可受限选择，不能泛化为必然输出长文。",
    )
    pill(s, "共同输入：业务事实 + 问题 + 候选条件", 2.38, 2.13, 8.5)
    text(s, "DECIDER", 0.68, 2.95, 1.7, 0.4, 17, TEAL, True)
    flow(
        s,
        [
            ("编码输入", "业务与候选表示", "layers"),
            ("Pointer Head", "对候选计算分数", "list-filter"),
            ("最高分候选", "choice + 概率分布", "circle-check"),
        ],
        x=2.35,
        y=2.73,
        w=10.3,
        color=TEAL,
    )
    text(s, "GPT", 0.68, 4.92, 1.7, 0.4, 17, BLUE, True)
    flow(
        s,
        [
            ("读取提示", "同一业务与候选", "receipt-text"),
            ("生成 token", "结构化选择 JSON", "code-xml"),
            ("解析与校验", "有效候选 ID", "shield-check"),
        ],
        x=2.35,
        y=4.68,
        w=10.3,
        color=BLUE,
    )
    note(s, "两条路径的输出都交给应用执行器；模型本身不执行转账。")

    # 04 — Visible decision process.
    s = slide(
        "可观察的决策过程：输入 → 输出 → 执行",
        subtitle="原始选择、执行调整和业务状态，分别展示。",
        notes="这张图表示应用可观察的数据流，不是模型内部思维链。Decider概率与confidence不是业务正确概率，GPT的自然语言解释也不是执行凭据。应用独立核验前置条件，原始输出保留。",
    )
    flow(
        s,
        [
            ("共享事实", "同一快照 / 候选", "receipt-text"),
            ("原始选择", "Decider  /  GPT", "arrow-left-right"),
            ("执行检查", "金额 / 地址 / 审批", "shield-check"),
            ("实际结果", "订单 / 账本 / 工单", "check-check"),
        ],
        y=2.68,
        height=1.7,
    )
    card(s, 0.65, 4.90, 5.84, 1.34, "看选择", "模型建议是否适合当前条件", color=TEAL)
    card(s, 6.79, 4.90, 5.84, 1.34, "看执行", "实际做了什么，是否推进业务", color=BLUE)
    note(s, "概率表示候选偏好；不展示或推测模型隐藏思维链。")

    # 05 — BERT, one slide only.
    s = slide(
        "BERT、Decider、GPT：适配方式一页看懂",
        subtitle="区别在输出头、候选表达与任务范围。",
        notes="BERT仅在此页概念比较，未部署或测速。常见BERT+固定分类头需要领域标签与监督适配。基于BERT的NLI或文本对模型也能按动态候选评分，因此动态候选不是Decider独有。",
    )
    for x, title, body, glyph, color, fill in [
        (
            0.65,
            "BERT 分类器",
            "固定标签 → 分类头\n标签稳定、领域标注\n新增类别通常需适配",
            "layers",
            MUTED,
            WHITE,
        ),
        (
            4.78,
            "Decider",
            "请求候选 → 共享评分\n条件用自然语言描述\n专注有限动作选择",
            "list-filter",
            TEAL,
            PALE_T,
        ),
        (
            8.91,
            "GPT",
            "提示候选 → 生成 JSON\n支持开放式生成与规划\n也可做受限动作选择",
            "brain-circuit",
            BLUE,
            PALE_B,
        ),
    ]:
        card(s, x, 2.48, 3.78, 3.2, title, body, glyph, color, fill)
    note(s, "BERT/NLI 也可做动态候选评分；本项目未对 BERT 进行速度或质量实测。")

    # 06 — Architecture, official service icons.
    s = slide(
        "一套业务后端，对接两种模型",
        "项目架构",
        "AWS us-east-1 · 私有前端源站 · 相同业务执行逻辑",
        notes="CloudFront通过OAC访问私有S3，四项Block Public Access开启。API路径经VPC Origin到内网ALB再到ECS Fargate。后端直接调用SageMaker、Bedrock、DynamoDB。图中省略网络端点节点，认证是自建用户名密码而非Cognito。",
    )

    def service(name, title, desc, x, y, w=2.42, h=1.28, fill=WHITE):
        L.rect(s, x, y, w, h, fill, LINE, True)
        icon(s, name, x + 0.17, y + 0.2, 0.57, aws=True)
        text(s, title, x + 0.85, y + 0.19, w - 1.0, 0.35, 16, INK, True)
        text(s, desc, x + 0.17, y + 0.85, w - 0.34, 0.3, 12, MUTED)

    service("cloudfront", "CloudFront", "页面入口 / API 路由", 0.7, 3.52)
    service("s3", "Amazon S3", "OAC · 禁止公开访问", 0.7, 2.02)
    arrow(s, 1.9, 3.45, 1.9, 3.34)
    L.rect(s, 3.62, 2.44, 4.35, 3.73, None, LINE, True)
    text(s, "VPC / 私有业务服务", 3.83, 2.63, 3.9, 0.32, 13, MUTED, True)
    service("alb", "内网 ALB", "CloudFront VPC Origin", 3.96, 3.03, 3.64)
    service("ecs", "ECS Fargate", "FastAPI · 决策 / 执行 / 反馈", 3.96, 4.56, 3.64)
    L.line(s, [(3.17, 4.17), (3.44, 4.17), (3.44, 3.67)], MUTED, 1.8)
    arrow(s, 3.44, 3.67, 3.90, 3.67)
    arrow(s, 5.8, 4.34, 5.8, 4.50)
    service("sagemaker", "SageMaker", "Decider v19", 9.14, 2.18, 3.5, fill=PALE_T)
    service("bedrock", "Bedrock", "GPT-6.1 Sol", 9.14, 3.72, 3.5, fill=PALE_B)
    service("dynamodb", "DynamoDB", "用户 / 会话", 9.14, 5.26, 3.5)
    for y in (2.82, 4.36, 5.90):
        L.line(s, [(7.65, 5.19), (8.46, 5.19), (8.46, y)], MUTED, 1.5)
        arrow(s, 8.46, y, 9.06, y)
    note(s, "后端调用相应服务；订单、工单与模拟账本用于演示，不广播链上交易。")

    # 07 — Scenarios.
    s = slide(
        "四个支付场景，贯穿五个业务节点",
        "演示设计",
        "同一笔业务，分别交给 Decider 和 GPT。",
        notes="四个固定案例覆盖129USDC未到账、100USDT订单少付0.20、250USDC退款和商户结算超时。五步都是模型选择候选，执行器决定实际状态。确认门槛12、SLA60分钟是演示参数。",
    )
    for x, heading, body, glyph in [
        (0.65, "未到账", "确认数 3 / 12", "wallet"),
        (3.75, "少付", "100 → 99.80", "banknote"),
        (6.85, "退款", "地址与审批待补齐", "repeat-2"),
        (9.95, "结算", "到账状态与时效", "receipt-text"),
    ]:
        card(s, x, 2.43, 2.72, 1.87, heading, body, glyph)
    flow(
        s,
        [
            (name, f"节点 {i + 1}", glyph)
            for i, (name, glyph) in enumerate(
                zip(
                    STAGES.values(),
                    ["headset", "search", "shield-check", "workflow", "circle-check"],
                    strict=True,
                )
            )
        ],
        y=4.75,
        height=1.43,
    )
    note(s, "客户消息、事实快照与候选条件均可追溯。")

    # 08 — Method.
    s = slide(
        "两种对比：决策速度与业务完成",
        "演示设计",
        "固定输入用于逐步比较；闭环用于观察执行反馈。",
        notes="固定模式每一步两模型同输入，并发请求；仅均成功且输入指纹一致的请求进入速度配对。闭环各侧独立推进，后续输入可能不同，因此只比较各自运行结果、决策数和累计耗时。每个时延包含通信排队。",
    )
    card(
        s,
        0.65,
        2.40,
        5.85,
        3.52,
        "固定输入对比",
        "相同快照 → 各自 5 个选择\n逐步比较选择与耗时\n收齐方案后执行",
        "arrow-left-right",
        TEAL,
        PALE_T,
    )
    card(
        s,
        6.8,
        2.40,
        5.85,
        3.52,
        "自主反馈闭环",
        "各自判断 → 执行 → 读取结果\n等待事件后重新判断\n比较业务完成与总决策耗时",
        "repeat-2",
        BLUE,
        PALE_B,
    )
    note(s, "以下时间为后端累计决策耗时，含通信与排队；不等于浏览器全程等待。")

    # 09 — Actual input and output.
    s = slide(
        "真实界面：同一输入，输出与耗时并排看",
        "固定输入 / 过程",
        "USDC 未到账 · 客服受理：两模型都选择「支付客服」。",
        notes="这是v2采集的真实固定输入第1轮截图。Decider91.831ms，GPT1169.203ms。输入指纹一致，原始模型选择和实际执行分别展示。",
    )
    shot(s, "fixed-turn-1.png", 0.67, 2.16, 12.0, 4.27, bounds=(0, 65, 1400, 482))
    note(s, "线上原始截图 · 2026-10-10 · 每轮保留输入、原始选择、实际动作和耗时。")

    # 10 — Five steps with each time.
    pending = cases["pending"]
    s = slide(
        "五步对照：耗时差异贯穿每次选择",
        "固定输入 / USDC",
        "确认数 3 / 12，固定事实保持不变；审批节点出现选择差异。",
        notes="表中为真实5个同输入配对请求。橙色行显示Decider等待人工审批，而GPT规则检查通过。后续动作是固定输入独立选择，不表示读取了上一行执行结果。",
    )
    rows = [
        [
            STAGES[st["stage"]],
            f"{st['decider']['label']} · {st['decider']['timings']['backend_total_ms']:.1f} ms",
            f"{st['bedrock']['label']} · {st['bedrock']['timings']['backend_total_ms'] / 1000:.2f} s",
        ]
        for st in pending["steps"]
    ]
    table(
        s, ["节点", "Decider", "GPT"], rows, [1.72, 5.2, 5.15], row_h=0.61, size=16, highlights=(2,)
    )
    pill(s, "累计 0.458 s", 2.42, 6.02, 4.35)
    pill(s, "累计 7.792 s", 7.65, 6.02, 4.35, BLUE, PALE_B)
    note(s, "累计差约 17.00×；这是该案例 5 次成功同输入决策的单次观测。")

    # 11 — Decision evidence.
    s = slide(
        "同一审批问题，为什么选择不一样？",
        "固定输入 / 决策差异",
        "观察候选评分与输出；不把评分解释成程序规则。",
        notes="pending.approve的候选概率：risk_review0.162、verify_destination0.0678、await_approval0.5854、approved0.1849。Decider取最高项await_approval；GPT生成approved。参考来自开发者策略reference()。这些可见数据不能证明模型内部为什么偏好该项。",
    )
    pill(s, "同一输入：无风险标记 / 收款场景 / 审批候选条件", 0.65, 2.17, 12.02)
    card(
        s, 0.65, 2.89, 5.84, 3.18, "Decider · 87.4 ms", glyph="list-filter", color=TEAL, fill=PALE_T
    )
    for i, (label, val) in enumerate(
        [
            ("等待人工审批", 0.5854),
            ("规则检查通过", 0.1849),
            ("风险复核", 0.1620),
            ("验证退款地址", 0.0678),
        ]
    ):
        yy = 3.75 + i * 0.48
        text(s, label, 0.93, yy, 2.08, 0.3, 14, INK, i == 0)
        L.rect(s, 3.07, yy + 0.08, 2.08 * val, 0.14, TEAL)
        text(s, f"{val:.1%}", 5.25, yy, 0.93, 0.3, 13, TEAL, True)
    card(
        s,
        6.79,
        2.89,
        5.84,
        3.18,
        "GPT · 3,771.8 ms",
        '输出候选 ID\n{"choice": "approved"}\n→ 规则检查通过',
        "code-xml",
        BLUE,
        PALE_B,
    )
    note(s, "演示参考为「规则检查通过」；参考符合度不是通用准确率，概率也不是正确率。")

    # 12 — Actual fixed result.
    s = slide(
        "处理 5/5 步，资金仍在等待确认",
        "固定输入 / 结果",
        "两模型最终都挂起并跟进；确认条件尚未满足。",
        notes="两个固定计划都处理5步，但此时没有资金账本且工单仍待跟进，因此不是业务完成。截图为原始fixed-final。客服回复为应用模板。",
    )
    flow(
        s,
        [
            ("完成选择", "各返回 5 个动作", "list-filter"),
            ("挂起资金", "确认数仍为 3 / 12", "shield-check"),
            ("持续跟进", "工单未结案", "headset"),
        ],
        y=2.30,
        height=1.45,
    )
    shot(s, "fixed-final.png", 0.66, 4.04, 12.0, 1.63)
    pill(s, "完成标准：目标订单 + 对应账本 + 工单结案", 1.47, 5.97, 10.35)
    note(s, "下一步需要外部证据更新；反复调用模型不会增加链上确认数。")

    # 13 — Main result.
    dtotal, gtotal = (totals[p]["backend_ms"] for p in ("decider", "bedrock"))
    n = totals["paired_steps"]
    s = slide(
        "19 组成功配对，累计决策耗时相差 12.20×",
        "速度对比 / 总览",
        "同一批业务输入与候选 · 每侧 20 次尝试，其中 19 组成功配对",
        notes="计入成功同输入配对的Decider1809.597ms、GPT22068.204ms，比值12.195093。均值95.242ms和1161.484ms。GPT另有一次ReadTimeoutError，失败不进入成功配对总时延，不能将该比值描述为全程用户体验或稳定性能保证。",
        dark=True,
    )
    metric(s, f"{dtotal / 1000:.3f} s", "Decider · 19 次累计", 0.7, 2.46, 5.2, "86D7C6", True)
    metric(s, f"{gtotal / 1000:.3f} s", "GPT · 19 次累计", 7.05, 2.46, 5.3, "A6B9FF", True)
    for yy, val, color, label in [
        (4.6, dtotal, "86D7C6", "Decider"),
        (5.27, gtotal, "A6B9FF", "GPT"),
    ]:
        text(s, label, 0.72, yy - 0.07, 1.5, 0.4, 16, WHITE, True)
        L.rect(s, 2.29, yy, 9.45, 0.24, "294258", radius=True)
        L.rect(s, 2.29, yy, 9.45 * val / gtotal, 0.24, color, radius=True)
    note(s, "单次观测；另有 1 次 GPT 读取超时，未计入成功耗时。不是全程用时或稳定性能保证。", True)

    # 14 — By stage.
    s = slide(
        "按节点比较：每次判断平均花多久",
        "速度对比 / 业务节点",
        "每个节点仅统计成功同输入配对，横轴为平均毫秒数。",
        notes="按stage分别对paired=true的backend_total_ms取均值。客服、核验、审批、结案各4对，业务处置3对。没有跨不同请求数量比较累计值。",
    )
    dvals, gvals, labels = [], [], []
    for stage, label in STAGES.items():
        pairs = [
            st for c in bench["cases"] for st in c["steps"] if st["stage"] == stage and st["paired"]
        ]
        labels.append(f"{label} · n={len(pairs)}")
        for provider, values in (("decider", dvals), ("bedrock", gvals)):
            values.append(
                sum(st[provider]["timings"]["backend_total_ms"] for st in pairs) / len(pairs)
            )
    pairbars(s, labels, dvals, gvals, x=2.8, y=2.77, w=7.9, gap=0.68)
    note(s, "业务处置仅 3 组成功配对；商户结算的 GPT 处置请求超时。条形均从 0 起。")

    # 15 — By scenario.
    s = slide(
        "按场景比较：相同输入下的累计耗时",
        "速度对比 / 四类场景",
        "同一场景中，两侧使用完全相同的成功配对子集。",
        notes="四场景分别是5、5、5、4组配对。图中每一对条形比较同一子集。结算Decider还有一个成功处置请求，但对应GPT超时，因此也排除该Decider请求，保证配对口径。",
    )
    summaries = [c["summary"] for c in bench["cases"]]
    pairbars(
        s,
        [f"{CASES[c['case']]}\n{c['paired_steps']} 组配对" for c in summaries],
        [c["decider_total_ms"] / 1000 for c in summaries],
        [c["bedrock_total_ms"] / 1000 for c in summaries],
        x=2.8,
        y=2.86,
        w=7.9,
        gap=0.83,
        units="s",
    )
    note(s, "结算仅含 4 对成功请求，不能当作完整五步运行耗时；跨场景比较需留意请求数。")

    # 16 — Refund fixed process.
    s = slide(
        "退款差异：模型想执行，条件是否允许？",
        "固定输入 / 退款",
        "地址未验证、审批待完成：两模型建议不同，实际资金动作都挂起。",
        notes="固定退款Decider在execute选择refund_payment，145.857ms，应用改为hold；GPT直接选择hold，1297.061ms。Decider还在close选择resolve，应用改为follow_up。此页截图为真实第4轮。",
    )
    shot(s, "refund-execute.png", 0.66, 2.18, 12.0, 4.19, bounds=(0, 65, 1400, 520))
    note(s, "Decider：提交模拟退款 → 应用挂起；GPT：直接挂起。执行保护不能计作模型判断正确。")

    # 17 — Feedback diagram.
    s = slide(
        "闭环把执行结果带回下一次判断",
        "自主闭环 / 原理",
        "两侧各自维护业务状态，逐轮判断并执行。",
        notes="自主闭环每次只处理当前节点，读取实际执行反馈。条件不足等待模拟证据更新；无效动作可有限重判，同一节点连续两次无进展转人工，预算20次决策/32tick。两侧后续输入可能不同。",
    )
    for yy, label, color in [(2.54, "DECIDER", TEAL), (4.51, "GPT", BLUE)]:
        text(s, label, 0.68, yy + 0.35, 1.5, 0.4, 17, color, True)
        flow(
            s,
            [
                ("读当前状态", "业务事实 + 上轮反馈", "receipt-text"),
                ("选择并执行", "候选 → 执行结果", "workflow"),
                ("判断去向", "完成 / 等待 / 重判", "route"),
            ],
            x=2.24,
            y=yy,
            w=10.39,
            color=color,
            height=1.35,
        )
        L.line(
            s,
            [(10.98, yy + 1.41), (10.98, yy + 1.69), (3.90, yy + 1.69)],
            color,
            1.6,
        )
        arrow(s, 3.90, yy + 1.69, 3.90, yy + 1.39, color)
        text(s, "证据更新 + 执行反馈", 5.95, yy + 1.41, 3.27, 0.28, 11, color)
    note(s, "外部事件由模拟器推进；模型不会自行制造链上确认、地址验证或审批事实。")

    # 18 — USDC loop route.
    s = slide(
        "USDC 闭环：一次证据更新，六次模型判断",
        "自主闭环 / 决策过程",
        "两侧都完成业务，但第 2 次核验选择不同。",
        notes="真实pending闭环每侧6次模型调用和1次确认事件。Decider第2次选择settlement_verified，执行后事实校验仍缺确认，因此等待事件，没有被该标签误判为业务成功。确认3变12后再核验到付款成功。",
    )
    dp = loops["pending"]["providers"]["decider"]["decisions"]
    gp = loops["pending"]["providers"]["bedrock"]["decisions"]
    rows = [
        [f"{i + 1} · {STAGES[d['stage']]}", d["label"], g["label"]]
        for i, (d, g) in enumerate(zip(dp, gp, strict=True))
    ]
    table(
        s,
        ["轮次 / 节点", "Decider", "GPT"],
        rows,
        [2.43, 4.83, 4.81],
        y=2.14,
        row_h=0.53,
        size=16,
        highlights=(1,),
    )
    pill(s, "第 2 → 3 次之间：模拟确认数 3 / 12 → 12 / 12", 1.57, 6.09, 10.0, ORANGE, PALE_O)
    note(s, "本表按调用顺序展示两次独立运行；不假定后续请求的输入相同。")

    # 19 — USDC observed feedback screenshot.
    s = slide(
        "真实反馈：选择不同，执行都等待证据",
        "自主闭环 / 实际过程",
        "Decider「结算已完成」与 GPT「等待链上确认」之后，都读到条件未满足。",
        notes="截图为pending-loop第2轮原始界面。为了投影可读，裁去轮次栏、模板回复与底部日志，保留输入、选择、实际执行与反馈；不改写结果。Decider90.031ms，GPT1193.33ms。",
    )
    shot(s, "pending-loop-feedback.png", 0.66, 2.15, 12.0, 4.34, bounds=(0, 65, 1400, 610))
    note(s, "等待事件 → 确认数达标 → 重新核验；最终由订单与账本确认完成。")

    # 20 — Refund route with actual adjustments.
    s = slide(
        "退款闭环：两次事件补齐执行前提",
        "自主闭环 / 决策过程",
        "共 9 次决策；下表聚焦 3 次审批与最后 2 次动作。",
        notes="退款第3次Deciderapproved实际改verify_destination；第5次approved改await_approval；第7次approved才实际通过。GPT分别选择verify_destination、await_approval、approved。事件分别为模拟验证地址和模拟审批通过。中间第4、6次是重新核验到账。",
    )
    rows = [
        ["3 · 地址未验证", "检查通过 → 验证地址", "验证退款地址"],
        ["5 · 待人工审批", "检查通过 → 等待审批", "等待人工审批"],
        ["7 · 前提已补齐", "规则检查通过", "规则检查通过"],
        ["8 · 业务处置", "提交模拟退款", "提交模拟退款"],
        ["9 · 回复结案", "回复并结案", "回复并结案"],
    ]
    table(
        s,
        ["调用 / 当前条件", "Decider 原始 → 实际", "GPT 原始与实际"],
        rows,
        [2.55, 4.75, 4.77],
        y=2.19,
        row_h=0.60,
        size=16,
        highlights=(0, 1),
    )
    pill(s, "事件 1：验证地址", 1.43, 6.01, 4.84, ORANGE, PALE_O)
    pill(s, "事件 2：审批通过", 7.00, 6.01, 4.84, ORANGE, PALE_O)
    note(s, "箭头左边保留模型选择，右边为执行器实际动作；事件由模拟器推进。")

    # 21 — Actual completed refund.
    s = slide(
        "退款结果：订单、账本与结案全部确认",
        "自主闭环 / 结果",
        "两模型均完成模拟退款；总耗时与最终状态同时可见。",
        notes="图为同次真实refund-loop的累计耗时和最终结果截图。Decider9次决策897.332ms，GPT9次10541.95ms，各2次模拟事件。截图中用时为决策耗时，非全程用户耗时。",
    )
    shot(s, "refund-loop-summary.png", 0.66, 2.26, 12.0, 1.91)
    shot(s, "refund-loop-final.png", 0.66, 4.41, 12.0, 1.56)
    note(s, "完成不是因为「已走完节点」，而是模拟订单、对应账本与工单结案均已确认。")

    # 22 — Loop summary.
    s = slide(
        "两组闭环：业务都完成，累计决策时间不同",
        "自主闭环 / 对比结果",
        "闭环后续输入可能不同；这里比较运行结果，不计为同输入速度基准。",
        notes="pending每侧6次决策、1事件，Decider570.140ms/GPT7268.536ms。refund每侧9次、2事件，Decider897.332ms/GPT10541.95ms。双方complete=true且fundsComplete=true。数值按观测JSON实时求和。",
    )
    rows = []
    for case, count, event_count in [("pending", 6, 1), ("refund", 9, 2)]:
        rows.append(
            [
                CASES[case],
                f"{loop_total(case, 'decider') / 1000:.3f} s",
                f"{loop_total(case, 'bedrock') / 1000:.3f} s",
                f"{count} 次 / {event_count} 个",
                "双方完成",
            ]
        )
    table(
        s,
        ["闭环场景", "Decider", "GPT", "每侧决策 / 事件", "结果"],
        rows,
        [2.48, 2.16, 2.16, 3.10, 2.17],
        row_h=0.91,
        size=18,
    )
    flow(
        s,
        [
            ("同一完成标准", "订单 + 对应账本", "receipt-text"),
            ("同一结束要求", "工单已结案", "circle-check"),
            ("独立决策过程", "反馈可能不同", "arrow-left-right"),
        ],
        y=4.88,
        height=1.37,
    )
    note(s, "以上仅为两组真实运行；不代表未来任务保证完成，也不包含现实事件等待时间。")

    # 23 — Contrast quality and failure.
    s = slide(
        "速度之外，同时看选择差异与调用失败",
        "对比解读",
        "快慢、参考符合度和业务完成，分别衡量。",
        notes="固定配对Decider13/19符合开发者参考、GPT19/19；不是准确率。20次尝试Decider返回20个有效选择，GPT19个，GPT结算处置ReadTimeoutError导致其方案不完整，未执行该侧沙盒。成功速度不能掩盖失败。",
    )
    table(
        s,
        ["固定输入观测", "Decider", "GPT"],
        [
            ["成功返回 / 尝试", "20 / 20", "19 / 20"],
            ["成功配对参考符合", "13 / 19", "19 / 19"],
            ["商户结算执行", "处理 5 步，继续跟进", "处置超时，方案未执行"],
        ],
        [3.65, 4.21, 4.21],
        y=2.30,
        row_h=0.82,
        size=18,
    )
    card(
        s,
        0.65,
        5.58,
        12.0,
        0.65,
        "单次样本：Decider 决策更快；GPT 在本批次更符合演示策略。",
        color=INK,
        fill=PALE_O,
    )
    note(s, "开发者参考不是通用准确率；19 组成功延迟不包含那次失败等待。")

    # 24 — Application.
    s = slide(
        "在支付流程中，按任务分工组合使用",
        "适用方式",
        "以本次观测为起点，继续按真实业务分布评测。",
        notes="本页是面向未来的组合建议，不是本项目已实现的路由功能。高频受限选择可评估Decider；复杂语义或开放式内容可使用GPT；金额权限账本最终校验由确定性代码承担。低confidence可辅助分流但需要业务校准。",
    )
    flow(
        s,
        [
            ("业务请求", "消息 + 当前事实", "receipt-text"),
            ("任务分流", "范围 / 复杂度 / 风险", "route"),
            ("模型建议", "Decider 或 GPT", "brain-circuit"),
            ("执行系统", "校验与落实动作", "shield-check"),
        ],
        y=2.34,
        height=1.61,
    )
    card(
        s,
        0.65,
        4.56,
        5.85,
        1.69,
        "Decider",
        "队列路由、受限候选、高频下一步选择",
        "list-filter",
        TEAL,
        PALE_T,
    )
    card(
        s,
        6.80,
        4.56,
        5.85,
        1.69,
        "GPT",
        "复杂语义、开放式内容、异常分析",
        "brain-circuit",
        BLUE,
        PALE_B,
    )
    note(s, "组合方案为应用建议；当前演示以两模型并排对比为主。")

    # 25 — Close.
    s = slide(
        "看清选择，量出速度，核实结果",
        "结论",
        dark=True,
        subtitle="本次支付演示的三点观察",
        notes="25页版本结论：Decider在本次19组成功同输入配对中累计时延约为GPT的1/12.20，但原始选择和GPT有差异，也依赖相同执行保护。两组闭环双方均完成；执行反馈和证据是完成业务的必要组成。BERT未实测。所有来源见逐页讲稿。",
    )
    for x, number, title, body, glyph in [
        (0.69, "01", "速度", "19 组成功配对\n累计耗时比 12.20×", "timer"),
        (4.86, "02", "决策", "原始选择有差异\n执行调整单独保留", "route"),
        (9.03, "03", "结果", "两组反馈闭环\n两模型均完成业务", "circle-check"),
    ]:
        text(s, number, x, 2.48, 1.4, 0.73, 37, "567B99", True)
        icon(s, glyph, x, 3.50, 0.55, "86D7C6")
        text(s, title, x + 0.80, 3.52, 2.5, 0.52, 26, WHITE, True)
        text(s, body, x, 4.47, 3.59, 1.05, 20, "DCE6F0")
    note(s, "来源：Strands 官方博客 / v19 模型实现 / 2026-10-10 实测与截图；详见逐页讲稿。", True)

    assert len(prs.slides) == 25
    for i, (s, meta) in enumerate(zip(prs.slides, metadata, strict=True), 1):
        footer_color = "8CA6BF" if meta["dark"] else MUTED
        L.line(s, [(0.62, 7.04), (12.68, 7.04)], "365168" if meta["dark"] else LINE, 0.6)
        text(s, "DECIDER × GPT  /  支付业务演示", 0.64, 7.16, 9.9, 0.21, 8.5, footer_color)
        text(s, f"{i:02} / 25", 11.76, 7.14, 0.92, 0.23, 10, footer_color, align=PP_ALIGN.RIGHT)
        for effect in s._element.xpath(".//a:effectRef"):
            effect.set("idx", "0")
        for shape in s.shapes:
            assert shape.left >= -10 and shape.top >= -10, (i, shape.name)
            assert shape.left + shape.width <= prs.slide_width + 10, (i, shape.name)
            assert shape.top + shape.height <= prs.slide_height + 10, (i, shape.name)
        assert "GPU" not in "\n".join(sh.text for sh in s.shapes if sh.has_text_frame)
    prs.core_properties.title = "Strands Decider × GPT · 支付业务演示 · 25页精简版"
    prs.core_properties.subject = "决策过程、速度与业务结果对比"
    prs.core_properties.comments = (
        "Revised from v2; 25 slides, one BERT comparison; dated real screenshots and observations."
    )
    target = OUT / f"{NAME}.pptx"
    prs.save(target)
    (OUT / "讲稿与来源.md").write_text(
        "# Strands Decider × GPT · 25 页讲稿\n\n"
        + "\n\n".join(
            f"## {i:02}. {meta['title']}\n\n{s.notes_slide.notes_text_frame.text}"
            for i, (s, meta) in enumerate(zip(prs.slides, metadata, strict=True), 1)
        )
        + "\n"
    )
    (OUT / "build-manifest.json").write_text(
        json.dumps(
            {
                "generated_at": datetime.now(UTC).isoformat(),
                "slides": metadata,
                "base_file": BASE.relative_to(ROOT).as_posix(),
                "base_sha256": digest(BASE),
                "pptx_sha256": digest(target),
                "benchmark_file": "docs/results/benchmark-2026-10-10.json",
                "benchmark_sha256": digest(ROOT / "docs/results/benchmark-2026-10-10.json"),
                "loop_sha256": digest(V2 / "loop-observations.json"),
                "capture_date": "2026-10-10",
                "new_model_calls": False,
                "source_capture_manifest": "docs/presentations/v2/capture-manifest.json",
                "crops": crops,
                "icon_sources": ["Lucide 0.468.0 (ISC)", "AWS official architecture icons"],
                "successful_pairs": n,
                "failed_requests": totals["failed_requests"],
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n"
    )
    print(f"Created {target.relative_to(ROOT)} — {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
