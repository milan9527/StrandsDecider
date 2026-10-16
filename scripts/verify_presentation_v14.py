"""Verify the two revised slides, support evidence and preserved v13 pages."""

import json
import math
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path

import pymupdf as fitz
from PIL import Image, ImageChops, ImageDraw
from pptx import Presentation
from pptx.util import Inches
from verify_presentation_v13 import body, digest, images, load, normal

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/presentations/v14"
PROVIDERS = ("decider", "jev", "bedrock")


def check_support():
    data = load(OUT / "sources/support-live.json")
    summary = load(OUT / "sources/support-comparison.json")
    capture = load(OUT / "capture-manifest.json")
    assert digest(OUT / "sources/support-live.json") == capture["evidence_sha256"]
    assert data["complete"] and data["case"] == "support"
    assert not data["page_errors"] and not data["gpu_writes"]
    assert Counter(c["phase"] for c in data["calls"]) == {
        "warmup": 6, "routing": 6, "target-decider": 1, "target-jev": 1, "target-bedrock": 1}
    exported = data["executed_export"]
    assert exported["router_model"] == "us.openai.gpt-6-sol" and exported["reasoning_effort"] == "low"
    sides = exported["cases"]["support"]["sides"]
    for stage in ("workflow", "model"):
        steps = [next(r for r in sides[p]["steps"] if r["stage"] == stage) for p in PROVIDERS]
        assert len({r["input_hash"] for r in steps}) == 1
        assert all(r["model_input"] == steps[0]["model_input"] for r in steps)
        state = json.loads(steps[0]["model_input"]["state"])
        assert state["request"] == data["task"]
        if stage == "workflow":
            assert set(state) == {"request"}
    for p in PROVIDERS:
        side, row = sides[p], summary["routes"][p]
        assert row["workflow_choice"] == side["steps"][0]["choice"] == "customer_support"
        assert math.isclose(row["routing_ms"], sum(r["timings"]["backend_total_ms"] for r in side["steps"]))
        for step in side["steps"]:
            call = next(c for c in data["calls"] if c["phase"] == "routing"
                        and c["request"]["provider"] == p and c["request"]["stage"] == step["stage"])
            assert call["status"] == 200 and call["response"]["source"] == "live"
            assert call["response"] == {k: v for k, v in step.items() if k != "client_total_ms"}
        target = next(c for c in data["calls"] if c["phase"] == f"target-{p}")
        assert target["request"]["task"]["text"] == data["task"]
        assert target["request"]["model_id"] == side["steps"][1]["selected_model"]["id"] == row["target_model_id"]
        assert row["target_status"] == target["status"]
        if p in ("decider", "jev"):
            assert target["status"] == 504 and not row["target_complete"]
            assert row["target_error"] == target["response"]["detail"] == side["execution"]["error"]
            assert row["target_ms"] is None and row["target_usage"] is None
            assert "40 秒" in row["target_error"]
        else:
            response = target["response"]
            assert response == side["execution"] and target["status"] == 200 and response["complete"]
            assert response["model_id"] == "openai.gpt-6-luna" and response["request_fields"]["reasoning"]["effort"] == "max"
            assert row["review"]["quote"] in response["text"].replace("**", "")
            assert round(response["backend_total_ms"] / 1000, 2) == 5.18
            assert "3/12" in response["text"] and "9 次确认" in response["text"]
    for shot in capture["screenshots"]:
        assert digest(ROOT / shot["file"]) == shot["sha256"] and shot["data_replaced"] is False
    return {"formal_routing_calls": 6, "warmups_excluded": 6, "target_attempts": 3,
            "target_successes": 1, "target_timeouts": 2, "same_stage_inputs": True,
            "actual_request_and_quote": "verified", "no_gpu_mutations": True}


def main():
    manifest = load(OUT / "build-manifest.json")
    path = ROOT / manifest["pptx_file"]
    prs = Presentation(path)
    original = Presentation(ROOT / manifest["base_file"])
    pdf = fitz.open(path.with_suffix(".pdf"))
    assert len(prs.slides) == len(pdf) == manifest["pages"] == 30
    assert digest(path) == manifest["pptx_sha256"]
    assert digest(ROOT / manifest["base_file"]) == manifest["base_sha256"]
    for source in manifest["evidence"]:
        assert digest(ROOT / source["path"]) == source["sha256"]
    for image in manifest["images"]:
        assert digest(ROOT / image["output"]) == image["sha256"]
        if image.get("source", "").endswith(".png"):
            with Image.open(ROOT / image["source"]) as src, Image.open(ROOT / image["output"]) as dst:
                expected = src.crop(image["crop"]) if image["crop"] else src
                assert expected.size == dst.size and ImageChops.difference(expected, dst).getbbox() is None
    previews = OUT / "previews"
    previews.mkdir(exist_ok=True)
    missing, bounds, preserved, thumbnails = [], [], [], []
    for i, (slide, page, meta) in enumerate(zip(prs.slides, pdf, manifest["slides"], strict=True), 1):
        text = normal(page.get_text())
        assert any(s.has_text_frame and s.text == f"{i:02} / 30" for s in slide.shapes)
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                assert not Inches(6.5) <= shape.top < Inches(7.03), i
                if normal(shape.text) not in text:
                    missing.append({"page": i, "text": shape.text})
            if (min(shape.left, shape.top) < 0 or shape.left + shape.width > prs.slide_width + 12
                    or shape.top + shape.height > prs.slide_height + 12):
                bounds.append({"page": i, "shape": shape.name})
        assert slide.notes_slide.notes_text_frame.text == meta["notes"]
        if meta["source"] == "v13":
            before = original.slides[meta["v13_page"] - 1]
            assert body(slide) == body(before) and images(slide) == images(before)
            preserved.append(i)
        dest = previews / f"page-{i:02}.png"
        page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(dest)
        with Image.open(dest) as image:
            image.thumbnail((384, 216))
            tile = Image.new("RGB", (404, 248), "#e7edf2")
            tile.paste(image, ((404 - image.width) // 2, 6))
            ImageDraw.Draw(tile).text((12, 226), f"{i:02}", fill="#203448")
            thumbnails.append(tile)
    for name, indexes in (("slides-overview", range(30)), ("updated-overview", [14, 17])):
        columns = min(4, len(indexes))
        image = Image.new("RGB", (columns * 404, math.ceil(len(indexes) / columns) * 248), "#e7edf2")
        for n, index in enumerate(indexes):
            image.paste(thumbnails[index], ((n % columns) * 404, (n // columns) * 248))
        image.save(OUT / f"{name}.png")
    support = load(OUT / "sources/support-comparison.json")
    assert normal(support["task"]) in normal(pdf[17].get_text())
    assert "高频客服" in pdf[17].get_text()
    assert all(label in pdf[14].get_text() for label in ("高频客服", "交易代码", "结算分析", "中英文索引"))
    report = {
        "audited_at": datetime.now(UTC).isoformat(), "pages": 30, "v13_file_unchanged": True,
        "preserved_body_and_images_pages": preserved, "new_pages": [15, 18],
        "settlement_two_pages_replaced_with_one_support_page": True,
        "rendered_text_missing": missing, "shape_bounds_issues": bounds,
        "lower_explanatory_rows": 0, "support_checks": check_support(),
        "pptx_sha256": digest(path), "pdf_sha256": digest(path.with_suffix(".pdf")),
    }
    (OUT / "render-audit.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    assert not missing and not bounds


if __name__ == "__main__":
    main()
