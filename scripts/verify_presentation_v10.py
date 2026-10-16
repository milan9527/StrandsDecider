"""Check v10's rendered pages, unchanged source deck and auditable cost data."""

import hashlib
import json
import math
import re
import unicodedata
from datetime import UTC, datetime
from pathlib import Path

import pymupdf as fitz
from PIL import Image, ImageDraw
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/presentations/v10"


def load(path):
    return json.loads(Path(path).read_text())


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def normal(value):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", value)).replace("\u00ad", "")


def check_prices():
    raw = load(OUT / "sources/aws-g6.2xlarge-pricing.json")
    calculated = load(OUT / "sources/cost-calculation.json")
    product = raw["products"][0]
    assert product["product"]["attributes"]["instanceType"] == "g6.2xlarge"
    assert product["product"]["attributes"]["regionCode"] == "us-east-1"
    od = next(iter(next(iter(product["terms"]["OnDemand"].values()))["priceDimensions"].values()))
    hourly = float(od["pricePerUnit"]["USD"])
    assert hourly == .9776
    terms = {(int(t["termAttributes"]["LeaseContractLength"][0]), t["termAttributes"]["PurchaseOption"]): t
             for t in product["terms"]["Reserved"].values() if t["termAttributes"]["OfferingClass"] == "standard"}
    assert len(terms) == 6
    for row in calculated["rows"]:
        if row["years"]:
            term = terms[(row["years"], row["payment"])]
            dimensions = {d["unit"]: float(d["pricePerUnit"]["USD"]) for d in term["priceDimensions"].values()}
            assert row["upfront_usd"] == dimensions.get("Quantity", 0)
            assert row["recurring_hourly_usd"] == dimensions.get("Hrs", 0)
            monthly = row["upfront_usd"] / (12 * row["years"]) + row["recurring_hourly_usd"] * 730
        else:
            monthly = hourly * 730
        assert math.isclose(row["monthly_amortized_usd"], monthly)
        assert math.isclose(row["savings"], 1 - monthly / (730 * hourly), abs_tol=1e-12)
        for label, unit in [
            ("payment_decisions", calculated["jev_payment_usd_per_million_decisions"]),
            ("two_step_routes", calculated["jev_route_usd_per_million_routes"]),
        ]:
            assert math.isclose(row[f"break_even_monthly_{label}"], monthly / unit * 1e6)
    old = load(ROOT / "docs/presentations/v9/sources/cost-calculation.json")
    assert old["measured_jev_input_tokens"] / old["measured_jev_payment_calls"] == calculated["jev_payment_mean_input_tokens"]
    assert math.isclose(calculated["jev_payment_mean_input_tokens"] * calculated["jev_input_per_million_usd"],
                        calculated["jev_payment_usd_per_million_decisions"])
    assert math.isclose(calculated["jev_route_mean_input_tokens"] * calculated["jev_input_per_million_usd"],
                        calculated["jev_route_usd_per_million_routes"])
    return {"standard_ri_offers_verified": len(terms), "on_demand_hourly_usd": hourly,
            "cost_formula_checks": "passed", "measured_token_cost_checks": "passed"}


def main():
    manifest = load(OUT / "build-manifest.json")
    pptx = ROOT / manifest["pptx_file"]
    prs = Presentation(pptx)
    pdf = fitz.open(pptx.with_suffix(".pdf"))
    assert len(prs.slides) == len(pdf) == manifest["pages"] == 32
    assert digest(ROOT / manifest["base_file"]) == manifest["base_sha256"]
    assert digest(pptx) == manifest["pptx_sha256"]
    for source in manifest["evidence"]:
        assert digest(ROOT / source["path"]) == source["sha256"], source["path"]
    for source in manifest["snapshots"]:
        assert digest(ROOT / source["file"]) == source["sha256"]
        assert digest(ROOT / source["source"]) == source["sha256"]
    for source in manifest["images"]:
        assert digest(ROOT / source["output"]) == source["sha256"]
    prices = check_prices()
    previews = OUT / "previews"
    previews.mkdir(exist_ok=True)
    missing, bounds, notes, tiles = [], [], [], []
    for i, (slide, page, meta) in enumerate(zip(prs.slides, pdf, manifest["slides"], strict=True), 1):
        rendered = normal(page.get_text())
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip() and normal(shape.text) not in rendered:
                missing.append({"page": i, "text": shape.text})
            if (min(shape.left, shape.top) < 0 or shape.left + shape.width > prs.slide_width + 12
                    or shape.top + shape.height > prs.slide_height + 12):
                bounds.append({"page": i, "shape": shape.name})
        assert slide.notes_slide.notes_text_frame.text == meta["notes"]
        notes.append(i)
        preview = previews / f"page-{i:02}.png"
        page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(preview)
        with Image.open(preview) as im:
            im.thumbnail((384, 216))
            tile = Image.new("RGB", (404, 248), "#e7edf2")
            tile.paste(im, ((404 - im.width) // 2, 6))
            ImageDraw.Draw(tile).text((12, 226), f"{i:02}", fill="#203448")
            tiles.append(tile)
    for name, start, end in [("slides-overview", 0, 32), ("validation-overview", 13, 20), ("cost-overview", 27, 30)]:
        cols = 4 if end - start > 3 else 3
        canvas = Image.new("RGB", (cols * 404, ((end - start + cols - 1) // cols) * 248), "#e7edf2")
        for i, tile in enumerate(tiles[start:end]):
            canvas.paste(tile, ((i % cols) * 404, (i // cols) * 248))
        canvas.save(OUT / f"{name}.png")
    audit = {
        "audited_at": datetime.now(UTC).isoformat(), "pages": 32, "base_v9_unchanged": True,
        "source_hashes_verified": len(manifest["evidence"]), "snapshots_verified": len(manifest["snapshots"]),
        "image_hashes_verified": len(manifest["images"]), "speaker_notes_pages_verified": notes,
        "shape_bounds_issues": bounds, "rendered_text_missing": missing, "cost_checks": prices,
        "pptx_sha256": digest(pptx), "pdf_sha256": digest(pptx.with_suffix(".pdf")),
        "new_model_calls": 0, "aws_resource_mutations": 0,
    }
    (OUT / "render-audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(audit, ensure_ascii=False, indent=2))
    assert not missing
    assert not bounds


if __name__ == "__main__":
    main()
