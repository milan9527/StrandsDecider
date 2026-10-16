"""Verify retained slides, rendered content, Jev usage and crossover formulas."""

import hashlib
import json
import math
import statistics
from datetime import UTC, datetime
from pathlib import Path

import pymupdf as fitz
import verify_presentation_v10 as previous
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/presentations/v11"
load, digest, normal = previous.load, previous.digest, previous.normal


def check_metrics():
    saved = load(OUT / "sources/jev-scenario-metrics.json")
    records = {
        name: [json.loads(line) for line in (ROOT / path).read_text().splitlines() if line.strip()]
        for name, path in saved["sources"].items()
    }
    groups = {
        "skills": [p for row in records["skills"] for p in row["providers"].values()],
        "browser": [r for r in records["browser_finance"] if r["experiment"] == "browser"],
        "finance_original": [r for r in records["browser_finance"] if r["experiment"] == "finance" and r["variant"] == "original_order"],
        "policy_routing": records["routing"],
    }
    for name, rows in groups.items():
        for provider in ("decider", "jev"):
            selected = [r for r in rows if r["provider"] == provider]
            stats = saved["groups"][name][provider]
            times = sorted(r["elapsed_ms"] if "elapsed_ms" in r else r["client_elapsed_ms"] for r in selected)
            assert stats["calls"] == len(selected)
            assert stats["median_rtt_ms"] == statistics.median(times)
            assert stats["p95_rtt_ms"] == times[math.ceil(len(times) * .95) - 1]
            if provider == "jev":
                tokens = sum(r["response"]["usage"]["input_tokens"] for r in selected)
                assert stats["input_tokens"] == tokens
                assert math.isclose(stats["batch_list_price_usd"], tokens * .042 / 1e6)
                assert math.isclose(stats["usd_per_1000_api_calls"], tokens / len(selected) * .042 / 1000)
    browser = saved["browser_jev_complete_tasks"]
    assert math.isclose(browser["usd_per_1000_observed_tasks"],
                        saved["groups"]["browser"]["jev"]["batch_list_price_usd"] / 10 * 1000)
    # Verify displayed business scores independently of the slide literals.
    correct = {p: sum(j["matches_reference"] is True for r in records["skills"]
                      for j in r["providers"][p]["judgments"].values()) for p in ("decider", "jev")}
    assert correct == {"decider": 50, "jev": 65}
    bf = load(ROOT / "validation/browser_finance/runs/2026-10-13-live/scored-summary.json")
    assert (bf["browser"]["jev"]["success"], bf["browser"]["decider"]["success"]) == (10, 4)
    for p, expected in [("jev", 28), ("decider", 22)]:
        assert sum(v["original_order_correct"] for v in bf["finance"][p]["categories"].values()) == expected
    policy = load(ROOT / "validation/model_routing/runs/2026-10-13-live/summary.json")
    assert (policy["jev"]["correct"], policy["decider"]["correct"]) == (46, 25)
    return {"scenarios": 4, "providers": 2, "latency_usage_cost_and_business_scores": "passed"}


def check_crossovers():
    costs = load(OUT / "sources/cost-calculation.json")
    plots = load(OUT / "sources/crossover-calculation.json")
    assert plots["ri_years"] == 1 and plots["ri_payment"] == "No Upfront"
    ri = next(r for r in costs["rows"] if r["years"] == 1 and r["payment"] == "No Upfront")
    od = next(r for r in costs["rows"] if r["years"] == 0)
    for key, plot in plots["plots"].items():
        unit = costs["jev_payment_usd_per_million_decisions"] if key == "payment" else costs["jev_route_usd_per_million_routes"]
        assert plot["on_demand_monthly_usd"] == od["monthly_amortized_usd"]
        assert plot["one_year_ri_monthly_usd"] == ri["monthly_amortized_usd"]
        assert math.isclose(plot["on_demand_millions"] * unit, od["monthly_amortized_usd"])
        assert math.isclose(plot["one_year_ri_millions"] * unit, ri["monthly_amortized_usd"])
        for ext in ("png", "svg"):
            assert digest(ROOT / plot[ext]) == plot[f"{ext}_sha256"]
    previous.OUT = OUT
    previous.check_prices()
    return {"payment_and_routing_intersections": "passed", "ri": "1-year Standard / No Upfront"}


def body(slide):
    return [shape.text for shape in slide.shapes if shape.has_text_frame and shape.top < Inches(7.03)]


def images(slide):
    return [hashlib.sha256(shape.image.blob).hexdigest() for shape in slide.shapes if shape.shape_type == 13]


def main():
    manifest = load(OUT / "build-manifest.json")
    pptx = ROOT / manifest["pptx_file"]
    prs = Presentation(pptx)
    base = Presentation(ROOT / manifest["base_file"])
    pdf = fitz.open(pptx.with_suffix(".pdf"))
    assert len(prs.slides) == len(pdf) == 32
    assert digest(ROOT / manifest["base_file"]) == manifest["base_sha256"]
    assert digest(pptx) == manifest["pptx_sha256"]
    for source in manifest["evidence"]:
        assert digest(ROOT / source["path"]) == source["sha256"]
    for source in manifest["snapshots"]:
        assert digest(ROOT / source["file"]) == digest(ROOT / source["source"]) == source["sha256"]
    missing, bounds, preserved, thumbnails = [], [], [], []
    previews = OUT / "previews"
    previews.mkdir(exist_ok=True)
    for i, (slide, page, meta) in enumerate(zip(prs.slides, pdf, manifest["slides"], strict=True), 1):
        text = normal(page.get_text())
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip() and normal(shape.text) not in text:
                missing.append({"page": i, "text": shape.text})
            if (min(shape.left, shape.top) < 0 or shape.left + shape.width > prs.slide_width + 12
                    or shape.top + shape.height > prs.slide_height + 12):
                bounds.append({"page": i, "shape": shape.name})
        assert slide.notes_slide.notes_text_frame.text == meta["notes"]
        assert "3年" not in normal(" ".join(body(slide))), i
        if meta["source"] == "v10":
            before = base.slides[meta["v10_page"] - 1]
            assert images(slide) == images(before)
            if not meta["body_modified"]:
                assert body(slide) == body(before)
                preserved.append(i)
        path = previews / f"page-{i:02}.png"
        page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(path)
        with Image.open(path) as im:
            im.thumbnail((384, 216))
            tile = Image.new("RGB", (404, 248), "#e7edf2")
            tile.paste(im, ((404 - im.width) // 2, 6))
            ImageDraw.Draw(tile).text((12, 226), f"{i:02}", fill="#203448")
            thumbnails.append(tile)
    for name, indexes in [
        ("slides-overview", list(range(32))),
        ("updated-slides-overview", [19, 20, 28, 29, 30]),
        ("cost-overview", [28, 29]),
    ]:
        cols = min(4, len(indexes))
        canvas = Image.new("RGB", (cols * 404, math.ceil(len(indexes) / cols) * 248), "#e7edf2")
        for n, index in enumerate(indexes):
            canvas.paste(thumbnails[index], ((n % cols) * 404, (n // cols) * 248))
        canvas.save(OUT / f"{name}.png")
    report = {
        "audited_at": datetime.now(UTC).isoformat(), "pages": 32, "v10_file_unchanged": True,
        "preserved_body_pages": preserved, "all_retained_slide_images_preserved": True,
        "shape_bounds_issues": bounds, "rendered_text_missing": missing,
        "scenario_checks": check_metrics(), "crossover_checks": check_crossovers(),
        "pptx_sha256": digest(pptx), "pdf_sha256": digest(pptx.with_suffix(".pdf")),
        "new_model_calls": 0, "aws_resource_mutations": 0,
    }
    (OUT / "render-audit.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    assert not missing
    assert not bounds


if __name__ == "__main__":
    main()
